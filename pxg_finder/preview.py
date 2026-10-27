"""Предпросмотр файла справа от списка: текст, таблица (Excel/CSV), Word, изображение, PDF. Большие и сетевые файлы читаются ограниченно."""
from __future__ import annotations

import csv
import io
import os
import re
import zipfile
from typing import Any, Dict, List

TEXT_BYTES = 64 * 1024
TABLE_ROWS = 60
TABLE_COLS = 24
MAX_OFFICE = 60 * (1 << 20)  # Excel/Word крупнее не открываем для предпросмотра
IMAGE_EXT = {"png", "jpg", "jpeg", "gif", "bmp", "webp", "svg", "ico"}
TEXT_EXT = {"txt", "log", "md", "json", "xml", "yaml", "yml", "ini", "cfg", "dat", "inc", "data", "py", "sql", "sh", "bat",
            "csv", "tsv", "html", "htm", "css", "js", "ts", "grdecl", "include", "schedule"}


def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return ("%.6g" % v)
    return str(v)[:200]


def _trim(rows: List[List[str]]) -> List[List[str]]:
    """Пустые хвосты строк и пустые строки внизу не показываем."""
    width = max((max((i + 1 for i, c in enumerate(r) if c), default=0) for r in rows), default=0)
    rows = [r[:width] for r in rows]
    while rows and not any(rows[-1]):
        rows.pop()
    return rows


def _table_xlsx(path: str, sheet: str) -> Dict[str, Any]:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        names = list(wb.sheetnames)
        ws = wb[sheet] if sheet in names else wb[names[0]]
        rows = []  # type: List[List[str]]
        for i, row in enumerate(ws.iter_rows(max_row=TABLE_ROWS, max_col=TABLE_COLS, values_only=True)):
            rows.append([_cell(v) for v in row])
        rows = _trim(rows)
        return {"kind": "table", "sheets": names, "sheet": ws.title, "rows": rows, "more": (ws.max_row or 0) > TABLE_ROWS}
    finally:
        wb.close()


def _table_other(path: str, ext: str) -> Dict[str, Any]:
    from python_calamine import CalamineWorkbook
    wb = CalamineWorkbook.from_path(path)
    names = list(wb.sheet_names)
    data = wb.get_sheet_by_name(names[0]).to_python()
    rows = _trim([[_cell(v) for v in r[:TABLE_COLS]] for r in data[:TABLE_ROWS]])
    return {"kind": "table", "sheets": names, "sheet": names[0], "rows": rows, "more": len(data) > TABLE_ROWS}


def _docx_text(path: str) -> str:
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8", errors="replace")
    xml = re.sub(r"</w:p>", "\n", xml)
    xml = re.sub(r"<w:tab/>", "\t", xml)
    text = re.sub(r"<[^>]+>", "", xml)
    for a, b in (("&amp;", "&"), ("&lt;", "<"), ("&gt;", ">"), ("&quot;", '"'), ("&apos;", "'")):
        text = text.replace(a, b)
    return text


def make(path: str, sheet: str = "") -> Dict[str, Any]:
    path = os.path.normpath(path)
    if os.path.isdir(path):
        try:
            names = sorted(os.listdir(path), key=str.lower)
        except OSError as e:
            return {"kind": "error", "error": str(e)}
        return {"kind": "folder", "count": len(names), "names": names[:80]}
    if not os.path.isfile(path):
        return {"kind": "error", "error": "Файл не найден."}
    ext = os.path.splitext(path)[1].lstrip(".").lower()
    try:
        size = os.path.getsize(path)
    except OSError as e:
        return {"kind": "error", "error": str(e)}
    base = {"size": size, "ext": ext}
    try:
        if ext in IMAGE_EXT:
            return {**base, "kind": "image"}
        if ext == "pdf":
            return {**base, "kind": "pdf"}
        if ext in ("xlsx", "xlsm", "xltx"):
            if size > MAX_OFFICE:
                return {**base, "kind": "none", "note": "Файл больше 60 МБ: предпросмотр отключён."}
            return {**base, **_table_xlsx(path, sheet)}
        if ext in ("xls", "xlsb", "ods"):
            if size > MAX_OFFICE:
                return {**base, "kind": "none", "note": "Файл больше 60 МБ: предпросмотр отключён."}
            return {**base, **_table_other(path, ext)}
        if ext in ("csv", "tsv"):
            with open(path, "rb") as f:
                text = _decode(f.read(TEXT_BYTES))
            dialect = "\t" if ext == "tsv" or text.count("\t") > text.count(";") + text.count(",") else (";" if text.count(";") >= text.count(",") else ",")
            rows = [[c[:200] for c in r[:TABLE_COLS]] for _, r in zip(range(TABLE_ROWS), csv.reader(io.StringIO(text), delimiter=dialect))]
            return {**base, "kind": "table", "sheets": [], "sheet": "", "rows": rows, "more": size > TEXT_BYTES}
        if ext == "docx":
            if size > MAX_OFFICE:
                return {**base, "kind": "none", "note": "Файл больше 60 МБ: предпросмотр отключён."}
            t = _docx_text(path)
            return {**base, "kind": "text", "text": t[:20000], "truncated": len(t) > 20000}
        with open(path, "rb") as f:
            raw = f.read(TEXT_BYTES)
        if ext in TEXT_EXT or b"\x00" not in raw[:2048]:
            return {**base, "kind": "text", "text": _decode(raw), "truncated": size > TEXT_BYTES}
        return {**base, "kind": "none", "note": "Для этого типа файла предпросмотра нет."}
    except Exception as e:  # noqa: BLE001 - повреждённый файл не должен ронять окно
        return {**base, "kind": "error", "error": "%s: %s" % (type(e).__name__, e)}
