"""Генератор тестовых исходников для модулей расходов («создание_БД_расходов», «нулевой_расход»).

Структура строится по тому, что ожидают скрипты:

    <корень>/Отбор/<ГГГГ-ГГГГ>/Результаты работы скважин отбор <ГГГГ-ГГГГ>гг/<файл>.xlsx
    <корень>/Закачка/<ГГГГ>/Результаты работы скважин в закачку <ГГГГ>гг/<файл>.xlsx

В книге по листу на месяц (отбор: октябрь–апрель, закачка: апрель–октябрь). На листе: заголовок «№№ скв.» с датами по дням,
ниже скважины с часовым расходом газа и строка «Итого»; через пустые строки вторая таблица, время работы (0–24 ч), тоже с «Итого».
Рядом создаётся файл периодов («ДД.ММ.ГГГГ prod|inj|none»).

Данные детерминированы (seed). В нескольких ячейках намеренно нарушено соответствие «расход / время»: они записаны в
манифест, по нему тесты проверяют, что модули находят ровно эти случаи.
"""
from __future__ import annotations

import calendar
import random
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

MONTHS = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь",
          "Декабрь"]
PRODUCTION_MONTHS = [10, 11, 12, 1, 2, 3, 4]
INJECTION_MONTHS = [4, 5, 6, 7, 8, 9, 10]

# период действует с указанной даты до следующей строки
DEFAULT_PERIODS: List[Tuple[str, str]] = [
    ("01.10.2023", "prod"),
    ("01.12.2023", "none"),
    ("05.12.2023", "prod"),
    ("01.04.2024", "inj"),
    ("01.10.2024", "prod"),
]


@dataclass
class Manifest:
    root: Path
    periods_file: Path
    wells: List[str]
    # (отбор|закачка, дата, скважина, вид) — вид: «Время есть, расход 0» или «Расход есть, время 0»
    planted: List[Tuple[str, date, str, str]] = field(default_factory=list)
    # те же случаи, но только в периодах, которые не отсекает файл периодов (для «создание_БД_расходов»)
    planted_after_periods: List[Tuple[str, date, str, str]] = field(default_factory=list)
    # число строк «скважина × день» по каждому виду данных до фильтрации по периодам
    rows: Dict[str, int] = field(default_factory=dict)
    # то же после фильтрации по файлу периодов (как в скрипте)
    rows_after_periods: Dict[str, int] = field(default_factory=dict)
    neutral_rows: int = 0
    files: List[Path] = field(default_factory=list)


def _period_of(day: date, periods: List[Tuple[date, str]]) -> Optional[str]:
    current = None
    for start, kind in periods:
        if day >= start:
            current = kind
        else:
            break
    return current


def _write_book(path: Path, months: List[Tuple[int, int]], wells: List[str], rng: random.Random,
                kind: str, planted: List[Tuple[str, date, str, str]], plant_per_sheet: int) -> int:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    total = 0
    for year, month in months:
        ws = wb.create_sheet(MONTHS[month - 1])
        days = calendar.monthrange(year, month)[1]
        ws.cell(row=1, column=2, value="Суточные показатели, %s %d г." % (MONTHS[month - 1], year))
        header_row = 3
        ws.cell(row=header_row, column=2, value="№№ скв.")
        for d in range(days):
            ws.cell(row=header_row, column=3 + d, value=datetime(year, month, d + 1))
        n = len(wells)
        gas, hours = {}, {}
        for w in wells:
            for d in range(days):
                if rng.random() < 0.12:           # простой: расход и время нулевые
                    gas[(w, d)], hours[(w, d)] = 0, 0
                else:
                    gas[(w, d)] = round(rng.uniform(800, 4200), 1)
                    hours[(w, d)] = rng.choice([24, 24, 24, 22, 20, 18.5, 12])
        # намеренные несоответствия
        cells = [(w, d) for w in wells for d in range(days)]
        for w, d in rng.sample(cells, plant_per_sheet):
            if rng.random() < 0.5:
                gas[(w, d)] = 0
                hours[(w, d)] = rng.choice([6, 12, 24])
                what = "Время есть, расход 0"
            else:
                gas[(w, d)] = round(rng.uniform(800, 4200), 1)
                hours[(w, d)] = 0
                what = "Расход есть, время 0"
            planted.append((kind, date(year, month, d + 1), w, what))
        # таблица газа
        for i, w in enumerate(wells):
            row = header_row + 1 + i
            ws.cell(row=row, column=2, value=w)
            for d in range(days):
                ws.cell(row=row, column=3 + d, value=gas[(w, d)])
        total_row = header_row + 1 + n
        ws.cell(row=total_row, column=2, value="Итого")
        for d in range(days):
            ws.cell(row=total_row, column=3 + d, value=round(sum(gas[(w, d)] for w in wells), 1))
        # таблица времени работы (после пустых строк)
        title_row = total_row + 3
        ws.cell(row=title_row, column=2, value="Время работы, ч")
        for i, w in enumerate(wells):
            row = title_row + 1 + i
            ws.cell(row=row, column=2, value=w)
            for d in range(days):
                ws.cell(row=row, column=3 + d, value=hours[(w, d)])
        ws.cell(row=title_row + 1 + n, column=2, value="Итого")
        for d in range(days):
            ws.cell(row=title_row + 1 + len(wells), column=3 + d, value=round(sum(hours[(w, d)] for w in wells), 1))
        total += n * days
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return total


def make_flow_tree(root, *, wells: Optional[List[str]] = None, season: str = "2023-2024", injection_year: str = "2024",
                   plant_per_sheet: int = 2, seed: int = 20240101,
                   periods: Optional[List[Tuple[str, str]]] = None) -> Manifest:
    """Создаёт дерево папок и книг с расходами и файл периодов; возвращает манифест с ожидаемыми результатами."""
    root = Path(root)
    wells = wells or ["101", "102", "103", "104", "105"]
    rng = random.Random(seed)
    periods = periods or DEFAULT_PERIODS
    parsed = [(datetime.strptime(d, "%d.%m.%Y").date(), k) for d, k in periods]
    periods_file = root / "периоды.txt"
    root.mkdir(parents=True, exist_ok=True)
    periods_file.write_text("".join("%s %s\n" % p for p in periods), encoding="utf-8")
    manifest = Manifest(root=root, periods_file=periods_file, wells=wells)

    y1, y2 = [int(x) for x in season.split("-")]
    prod_months = [(y1 if m >= 10 else y2, m) for m in PRODUCTION_MONTHS]
    inj_months = [(int(injection_year), m) for m in INJECTION_MONTHS]
    jobs = [
        ("отбор", prod_months,
         root / "Отбор" / season / ("Результаты работы скважин отбор %sгг" % season), "Отбор_%s.xlsx" % season),
        ("закачка", inj_months,
         root / "Закачка" / injection_year / ("Результаты работы скважин в закачку %sгг" % injection_year),
         "Закачка_%s.xlsx" % injection_year),
    ]
    for kind, months, folder, name in jobs:
        path = folder / name
        manifest.rows[kind] = _write_book(path, months, wells, rng, kind, manifest.planted, plant_per_sheet)
        manifest.files.append(path)
        keep = {"отбор": ("prod", "none"), "закачка": ("inj", "none")}[kind]
        after = neutral = 0
        for year, month in months:
            for d in range(calendar.monthrange(year, month)[1]):
                p = _period_of(date(year, month, d + 1), parsed)
                if p is None or p in keep:
                    after += len(wells)
                    if p == "none":
                        neutral += len(wells)
        manifest.rows_after_periods[kind] = after
        manifest.neutral_rows += neutral
    # несоответствия вне подходящих периодов «создание_БД_расходов» отфильтрует вместе со строками
    keep_map = {"отбор": ("prod", "none"), "закачка": ("inj", "none")}
    manifest.planted_after_periods = [x for x in manifest.planted
                        if _period_of(x[1], parsed) is None or _period_of(x[1], parsed) in keep_map[x[0]]]
    return manifest


if __name__ == "__main__":  # python -m pxg_base.testdata.flows <папка>
    import sys
    target = sys.argv[1] if len(sys.argv) > 1 else "тестовые_расходы"
    m = make_flow_tree(target)
    print("Создано в %s: файлов %d, строк отбор/закачка %s, намеренных несоответствий %d"
          % (m.root, len(m.files), m.rows, len(m.planted)))
