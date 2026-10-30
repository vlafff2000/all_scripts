"""База ГДИ из PDF-отчётов «Мониторинг ГДИС»: папка (или ZIP) с отчётами → книга в формате базы ГДИ (лист «Лист1»).

Одно исследование — блок строк: в первой строке данные скважины и исследования (№ п/п, скважина, ГСП, дата, метод,
способ, a, b, DPmax, Qmax, статика, Туст, Рпл), в каждой — режим (№, дебит, Руст, Рзаб, DP). Столбцы и порядок те же, что в
образце базы ГДИ, который читает Газовый Атлас.

Откуда берутся значения: a, b и Рпл — «Результаты метода»; дебит, Рзаб (Pман на ИП) и dP — таблица «P(Q) на ИП»; Руст и
Туст — «Данные на режимах». % КРУ берётся из отчёта, если он есть (исследование в шлейф). Чего в отчёте нет (ГСП), остаётся пустым; ГСП и пустые a, b, Рпл жёлтые и перечислены на листе «Не распознано»:
итоговое слово за инженером, исправления вносятся прямо в «Лист1».

Параметры: --folder (папка с отчётами или ZIP), --db (прежняя база для дозаписи), --groups (таблица «скважина → ГСП»),
--output. Строки прежней базы не перезаписываются: ручные
правки инженера сохраняются, повторное исследование той же скважины за ту же дату пропускается и указано в «Не распознано».
"""
from __future__ import annotations

import os
import re
import tempfile
from datetime import datetime
from typing import Dict, List, Optional, Tuple

import pandas as pd

from pxg_core import гдис_карта as kart
from pxg_core import гдис_pdf as pdf
from pxg_core import параметры
from pxg_core.parallel import map_jobs

COLUMNS = ["№ п/п", "№скв", "ГСП", "дата", "метод исслед.", "способ исслед.", "a", "b", "DPmax кгс/см2",
           "Qmax тыс.м3/сут", "Рст. тр, кгс/см2", "Рст. затр, кгс/см2", "Тст. тр, \n˚С", "Тст. затр, ˚С", "№ режима",
           "Qгаза тыс.м3/сут", "% КРУ", "Руст. тр, кгс/см2", "Руст. затр, кгс/см2", "Туст. ˚С", "Рпл, \nна ИП кгс/см2",
           "Рзаб, кгс/см2", "DP кгс/см2", "примечание", "метод"]
SHEET = "Лист1"
METHODS = {"ИД": "стац"}          # «метод исследования» отчёта → запись в базе
NUMERIC = ("a", "b", "DPmax кгс/см2", "Qmax тыс.м3/сут", "Рпл, \nна ИП кгс/см2")
UNRESOLVED_COLUMNS = ["Файл", "Скважина", "Дата", "Поле", "Что не так", "Как поправить"]
HOW = {
    "ГСП": "вписать ГСП в «Лист1» или указать таблицу «скважина → ГСП» (--groups) и запустить заново",
}


IN_FLOWLINE, IN_ATMOSPHERE = "в шлейф", "в атмосферу"


def delivery_method(sposob, method_text, kru, dshb) -> Tuple[str, str]:
    """Метод исследования по способу вывода газа: в атмосферу (ДИКТ, диаметры шайбы) или в шлейф (% раскрытия КРУ).

    `kru` и `dshb` — значения по режимам. Возвращает (метод, замечание): без признаков или при двух сразу метод пуст,
    решает инженер (замечание попадает на лист «Не распознано»)."""
    dikt = bool(re.search(r"дикт", "%s %s" % (sposob or "", method_text or ""), re.I))
    washer = any(to_number(v) is not None for v in dshb)
    krus = any(to_number(v) is not None for v in kru)
    if (dikt or washer) and krus:
        return "", "в отчёте и ДИКТ/шайба, и %% КРУ: метод не выбран, проставьте «%s» или «%s» вручную" % (IN_ATMOSPHERE, IN_FLOWLINE)
    if dikt or washer:
        return IN_ATMOSPHERE, ""
    if krus:
        return IN_FLOWLINE, ""
    return "", "нет ни ДИКТ и диаметров шайбы, ни % КРУ: метод (в шлейф / в атмосферу) не определён"


def to_number(v):
    try:
        return None if v is None or str(v).strip() == "" else float(str(v).replace(",", "."))
    except ValueError:
        return None


def num(x, nd: int = 10):
    return None if x is None else round(float(x), nd)


def rows_of_study(study: kart.Study, n: int, group) -> Tuple[List[list], List[tuple]]:
    """Строки блока исследования и список замечаний (поле, текст) к нему."""
    notes: List[tuple] = []
    sposob = study.sposob
    method = METHODS.get(study.method_raw.strip())
    if method is None:
        method = study.method_raw.strip().lower()
        notes.append(("метод исслед.", "метод «%s» не сопоставлен с принятым в базе (ИД → «стац»): записан как есть"
                      % study.method_raw))
    flow_method, why = delivery_method(sposob, study.method_raw, [r.cells.get("kru") for r in study.modes],
                                       [r.cells.get("dshb") for r in study.modes])
    if why:
        notes.append(("метод", why))
    modes = [r for r in study.modes if r.q is not None]
    used = [r for r in modes if r.in_calc] or modes
    top = max(used, key=lambda r: r.q) if used else None
    static = study.static
    t_first = next((r.cells.get("tu") for r in study.modes if r.cells.get("tu") is not None), None)
    if group in (None, ""):
        notes.append(("ГСП", "в отчёте группы нет, в таблице «скважина → ГСП» скважина не найдена"))
    out = []
    for i, r in enumerate(study.modes):
        dp = r.dp
        if dp is None and study.p_res is not None and r.p_ip is not None:
            dp = study.p_res - r.p_ip
        row = [None] * len(COLUMNS)
        if i == 0:
            row[0:14] = [n, int(study.well) if study.well.isdigit() else study.well, group, study.started.date() if study.started else None,
                         method, sposob, num(study.a), num(study.b), num(top.dp if top else None, 4),
                         num(top.q if top else None, 2), num(static.cells.get("p_tube") if static else None),
                         num(static.cells.get("p_ann") if static else None), num(static.cells.get("tu") if static else None),
                         num(static.cells.get("tu") if static else None)]
            row[19] = num(t_first)
            row[20] = num(study.p_res)
        row[14] = r.n
        row[15] = num(r.q, 2)
        kru = r.cells.get("kru")
        row[16] = num(kru) if isinstance(kru, (int, float)) else (kru or None)
        row[17] = num(r.cells.get("p_tube"))
        row[18] = num(r.cells.get("p_ann"))
        row[21] = num(r.p_ip)
        row[22] = num(dp, 4)
        if not r.in_calc:
            row[23] = "режим не вошёл в расчёт a и b"
        if i == 0:
            row[24] = flow_method or None
        out.append(row)
    return out, notes


def load_group_map(path: str) -> Dict[str, object]:
    """Таблица «скважина → ГСП»: на любом листе столбцы с «скв» и «гсп» в названии (в т. ч. прежняя база ГДИ)."""
    result: Dict[str, object] = {}
    if not path:
        return result
    for _name, df in pd.read_excel(path, sheet_name=None).items():
        well = next((c for c in df.columns if "скв" in str(c).lower()), None)
        grp = next((c for c in df.columns if str(c).strip().lower() in ("гсп", "gsp", "группа")
                    or ("гсп" in str(c).lower() and not re.search(r"т\s*гсп|р\s*гсп|давл|темп", str(c).lower()))), None)
        if well is None or grp is None:
            continue
        for w, g in zip(df[well], df[grp]):
            if pd.isna(w) or pd.isna(g) or str(g).strip() in ("", "-"):
                continue
            key = re.sub(r"\D", "", str(w).split(".")[0]) or str(w).strip()
            result.setdefault(str(int(key)) if key.isdigit() else key, group_value(g))
    return result


def group_value(g):
    """ГСП как число, если это «1», «ГСП-1», «ГСП 1»; иначе как есть."""
    m = re.fullmatch(r"\s*(?:гсп)?[\s-]*(\d+)(?:\.0)?\s*", str(g), flags=re.I)
    return int(m.group(1)) if m else str(g).strip()


def read_db(path: str) -> Optional[pd.DataFrame]:
    if not path:
        return None
    df = pd.read_excel(path, sheet_name=SHEET if SHEET in pd.ExcelFile(path).sheet_names else 0)
    df = df.iloc[:, :len(COLUMNS)]
    same = [re.sub(r"\s+", "", str(a)) == re.sub(r"\s+", "", b) for a, b in zip(df.columns, COLUMNS)]
    if not all(same) or df.shape[1] < len(COLUMNS) - 2:
        raise ValueError("прежняя база не в формате образца (столбцы «№ п/п», «№скв», «ГСП», «дата», … «примечание»): %s" % path)
    df.columns = COLUMNS[:df.shape[1]]
    df = df.reindex(columns=COLUMNS)
    return fill_method(df)


def fill_method(df: pd.DataFrame) -> pd.DataFrame:
    """Прежняя база без «метод»: по блоку исследования. ДИКТ в способе — «в атмосферу» (в % КРУ у таких строк стоят
    диаметры шайбы), иначе есть % КРУ — «в шлейф»; без признаков пусто."""
    df = df.copy()
    df["метод"] = df["метод"].astype(object)
    if df["метод"].notna().any() or df.empty:
        return df
    starts = list(df.index[df["№скв"].notna()]) + [len(df)]
    for a, b in zip(starts, starts[1:]):
        block = df.iloc[a:b]
        dikt = bool(re.search(r"дикт", " ".join(str(v) for v in block.iloc[0][["метод исслед.", "способ исслед."]] if not pd.isna(v)), re.I))
        if dikt:
            df.iloc[a, df.columns.get_loc("метод")] = IN_ATMOSPHERE
        elif block["% КРУ"].map(to_number).notna().any():
            df.iloc[a, df.columns.get_loc("метод")] = IN_FLOWLINE
    return df


def db_keys(df: Optional[pd.DataFrame]) -> set:
    if df is None or df.empty:
        return set()
    head = df[df["№скв"].notna()]
    return {(str(int(float(w))) if re.fullmatch(r"\d+(\.0)?", str(w)) else str(w), pd.Timestamp(d).date())
            for w, d in zip(head["№скв"], head["дата"]) if not pd.isna(d)}


def build(studies: List[kart.Study], groups: Dict[str, object], start: int = 1,
          taken: Optional[set] = None) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Таблица базы, замечания, режимы подробно. `taken` — исследования, уже лежащие в прежней базе."""
    taken = set(taken or ())
    wide: List[list] = []
    bad: List[dict] = []
    detail: List[dict] = []
    n = start
    ordered = sorted(studies, key=lambda s: (int(s.well) if s.well.isdigit() else 10**9, s.started or datetime.min, s.file))
    seen: Dict[tuple, str] = {}

    def add(s: kart.Study, field: str, text: str, how: str = ""):
        bad.append({"Файл": s.file, "Скважина": s.well or s.well_raw, "Дата": s.started.date() if s.started else None,
                    "Поле": field, "Что не так": text, "Как поправить": how or HOW.get(field, "проверить по PDF и поправить в «Лист1»")})

    for s in ordered:
        if not s.well or s.started is None or not s.modes:
            for f, t in s.issues:
                add(s, f, t)
            add(s, "файл", "исследование не попало в базу: не хватает номера скважины, даты или режимов",
                "прочитать PDF вручную и добавить строки в «Лист1»")
            continue
        key = (s.well, s.started.date())
        if key in taken:
            add(s, "дубликат", "исследование этой скважины за %s уже есть в прежней базе или в этой пачке (%s): пропущено"
                % (s.started.strftime("%d.%m.%Y"), seen.get(key, "прежняя база")), "если нужна замена — удалить старый блок строк из базы и запустить заново")
            continue
        taken.add(key)
        seen[key] = s.file
        rows, notes = rows_of_study(s, n, groups.get(s.well))
        wide.extend(rows)
        n += 1
        for f, t in list(s.issues) + notes:
            add(s, f, t)
        for r in s.modes:
            detail.append({"Файл": s.file, "Скважина": s.well, "Дата начала": s.started, "Режим": r.n,
                           "Dшб, мм": r.cells.get("dshb"), "Дебит в таблице режимов": r.cells.get("q_table"), "Дебит P(Q), тыс.м3/сут": r.q,
                           "Pтруб": r.cells.get("p_tube"), "Pзатруб": r.cells.get("p_ann"), "Ту": r.cells.get("tu"),
                           "Рзаб": r.p_ip, "dP": r.dp, "В расчёт": "да" if r.in_calc else "нет",
                           "Перфорация, м": "; ".join("%g–%g" % p for p in s.perforations),
                           "Метод обработки (a, b)": s.method_title, "Рпл": s.p_res, "Qабс.своб": s.q_aof, "Qсвоб": s.q_free,
                           "Шаблон": s.template})
    return (pd.DataFrame(wide, columns=COLUMNS), pd.DataFrame(bad, columns=UNRESOLVED_COLUMNS), pd.DataFrame(detail))


def summary(files: int, studies: List[kart.Study], new_rows: pd.DataFrame, bad: pd.DataFrame, kept: int) -> pd.DataFrame:
    wells = int(new_rows["№скв"].nunique()) if not new_rows.empty else 0
    rows = [["PDF-файлов обработано", files], ["Исследований добавлено", int(new_rows["№скв"].notna().sum())],
            ["Скважин", wells], ["Строк-режимов", len(new_rows)], ["Исследований в прежней базе (сохранены)", kept],
            ["Замечаний (лист «Не распознано»)", len(bad)],
            ["Столбец «% КРУ»", "берётся из отчёта по режимам, если он там есть (исследование в шлейф); иначе пуст (диаметр шайбы Dшб — на листе «Режимы (подробно)»)"],
            ["Столбец «метод»", "«в атмосферу» — в отчёте ДИКТ или диаметры шайбы; «в шлейф» — есть % КРУ; нет признаков или оба сразу — пусто (замечание)"],
            ["Допущение: способ исследования", "берётся из отчёта («ДИКТ» + число); если его нет — пусто"],
            ["Допущение: Тст. затр", "равна Тст. тр: в отчёте одна температура статики Ту"],
            ["Допущение: Qmax и DPmax", "дебит и dP режима с наибольшим дебитом среди вошедших в расчёт"],
            ["Допущение: Рзаб", "давление на ИП из таблицы P(Q) на ИП, как в образце базы; другие величины — на листе «Шаблоны»"],
            ["Единицы", "давления в кгс/см², дебиты в тыс.м³/сут, a и b в (кгс/см²)²/(тыс.м³/сут)[²]: пересчёт по единицам отчёта"]]
    return pd.DataFrame(rows, columns=["Показатель", "Значение"])


def read_studies(scans: List[pdf.Scan], mapping_file: str = "", method_pref: str = "4"
                 ) -> Tuple[List[kart.Study], pd.DataFrame]:
    """Отчёты → исследования по сопоставлению своего шаблона (сохранённое/из файла поверх подсказки) и лист «Шаблоны»."""
    groups: Dict[str, List[pdf.Scan]] = {}
    for sc in scans:
        groups.setdefault(sc.fingerprint if not sc.error else "ошибка", []).append(sc)
    studies: List[kart.Study] = []
    info: List[dict] = []
    for n, (fp, members) in enumerate(sorted(groups.items(), key=lambda kv: -len(kv[1])), 1):
        if fp == "ошибка":
            studies += [kart.build_study(sc, {}, "") for sc in members]
            continue
        rep = max(members, key=lambda sc: len(sc.items))
        mapping, source = kart.effective(rep, kart.load_saved(fp, mapping_file), method_pref)
        studies += [kart.build_study(sc, mapping, fp) for sc in members]
        for t in kart.TARGETS:
            info.append({"Шаблон": fp, "Файлов": len(members), "Программа": rep.program, "Поле базы": t.label,
                         "Источник в отчёте": kart.describe(mapping.get(t.key)), "Откуда": source[t.key],
                         "Пример": (kart.find_item(rep.items, mapping[t.key]).sample[:60] if mapping.get(t.key)
                                    and kart.find_item(rep.items, mapping[t.key]) else "")})
    return studies, pd.DataFrame(info, columns=["Шаблон", "Файлов", "Программа", "Поле базы", "Источник в отчёте", "Откуда", "Пример"])


def write(path: str, wide: pd.DataFrame, bad: pd.DataFrame, detail: pd.DataFrame, info: pd.DataFrame,
          templates: Optional[pd.DataFrame] = None) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    with pd.ExcelWriter(path, engine="openpyxl", date_format="DD.MM.YYYY", datetime_format="DD.MM.YYYY HH:MM") as xw:
        wide.to_excel(xw, sheet_name=SHEET, index=False)
        bad.to_excel(xw, sheet_name="Не распознано", index=False)
        detail.to_excel(xw, sheet_name="Режимы (подробно)", index=False)
        info.to_excel(xw, sheet_name="Сводка", index=False)
        if templates is not None and not templates.empty:
            templates.to_excel(xw, sheet_name="Шаблоны", index=False)
            for i, w in enumerate((12, 8, 24, 36, 60, 12, 50)):
                xw.sheets["Шаблоны"].column_dimensions[chr(65 + i)].width = w
            for c in xw.sheets["Шаблоны"][1]:
                c.font = Font(bold=True)
        ws = xw.sheets[SHEET]
        yellow = PatternFill("solid", fgColor="FFF2CC")
        first = {c: i + 1 for i, c in enumerate(COLUMNS)}
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 2).value is not None:        # первая строка блока: поля исследования
                for col in ("ГСП", "a", "b", "Рпл, \nна ИП кгс/см2"):
                    cell = ws.cell(r, first[col])
                    if cell.value is None:
                        cell.fill = yellow
            for col in ("Qгаза тыс.м3/сут", "Рзаб, кгс/см2"):
                cell = ws.cell(r, first[col])
                if cell.value is None and ws.cell(r, first["№ режима"]).value is not None:
                    cell.fill = yellow
            ws.cell(r, first["дата"]).number_format = "dd/mm/yyyy"
        for c in ws[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.freeze_panes = "C2"
        for name, widths in (("Не распознано", (10, 10, 12, 16, 70, 60)), ("Сводка", (44, 80))):
            sheet = xw.sheets[name]
            for i, w in enumerate(widths):
                sheet.column_dimensions[chr(65 + i)].width = w
            for c in sheet[1]:
                c.font = Font(bold=True)
        for c in xw.sheets["Режимы (подробно)"][1]:
            c.font = Font(bold=True)


def main(argv=None):
    args = параметры.parse(
        "База ГДИ из PDF-отчётов «Мониторинг ГДИС»", argv,
        folder="папка с PDF-отчётами или ZIP-архив", db="прежняя база ГДИ для дозаписи, необязательно",
        groups="таблица «скважина → ГСП» (Excel), необязательно",
        mapping="файл сопоставления шапок с полями базы (JSON из мастера), необязательно", method="метод обработки для a и b (по умолчанию 4)",
        output="имя файла результата")
    src = параметры.clean_path(args.folder or параметры.ask("Папка с PDF-отчётами или ZIP"))
    if not src or not (os.path.isdir(src) or os.path.isfile(src)):
        print("❌ Папка или архив не найдены: %s" % src)
        return 1
    tmp = None
    folder = src
    if os.path.isfile(src):
        if not src.lower().endswith(".zip"):
            print("❌ Ожидалась папка или ZIP: %s" % src)
            return 1
        tmp = tempfile.TemporaryDirectory()
        folder = pdf.unpack_zip(src, tmp.name)
    files = pdf.find_reports(folder)
    if not files:
        print("⚠️ PDF-отчётов не найдено: %s" % src)
        return 0
    db_path = параметры.ask_path("Прежняя база ГДИ для дозаписи (Enter — создать новую)", args.db if args.db is not None else "")
    groups_path = параметры.ask_path("Таблица «скважина → ГСП» (Enter — пропустить)", args.groups if args.groups is not None else "")

    print("Чтение отчётов: %d…" % len(files))
    scans = map_jobs(pdf.scan_report, files, 150.0)       # отчёты независимы: по ядрам, результат тот же, что подряд
    studies, templates = read_studies(scans, параметры.clean_path(args.mapping or ""), (args.method or "").strip() or "4")
    if len(templates["Шаблон"].unique()) > 1 if not templates.empty else False:
        print("ℹ️ Найдено шаблонов отчёта: %d. Сопоставление по каждому — на листе «Шаблоны»; изменить можно в мастере «Сопоставление шапок»." % templates["Шаблон"].nunique())
    try:
        old = read_db(db_path)
    except ValueError as exc:
        print("❌ %s" % exc)
        return 1
    groups = load_group_map(groups_path)
    if old is not None:
        groups = {**load_group_map(db_path), **groups}
        print("📚 Прежняя база: %d строк" % len(old))
    start = int(pd.to_numeric(old["№ п/п"], errors="coerce").max() or 0) + 1 if old is not None and not old.empty else 1
    new, bad, detail = build(studies, groups, start=start, taken=db_keys(old))
    if new.empty and old is None:
        print("⚠️ Ни одно исследование не удалось прочитать. Причины — в «Не распознано».")
    wide = pd.concat([old, new], ignore_index=True) if old is not None and not old.empty else new
    kept = int(old["№скв"].notna().sum()) if old is not None else 0
    out = (args.output or "").strip() or "БД_ГДИ_из_отчётов.xlsx"
    if not out.lower().endswith(".xlsx"):
        out += ".xlsx"
    write(out, wide, bad, detail, summary(len(files), studies, new, bad, kept), templates)
    print("✅ База ГДИ: %s (%d исследований добавлено, %d скважин)" % (out, int(new["№скв"].notna().sum()), new["№скв"].nunique()))
    if len(bad):
        print("⚠️ Замечаний: %d — смотрите лист «Не распознано» (пустые и спорные значения в «Лист1» выделены жёлтым)" % len(bad))
    if tmp:
        tmp.cleanup()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
