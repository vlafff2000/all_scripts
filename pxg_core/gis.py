"""Разбор книги «База ГИС-контроль»: лист = скважина, слева полная таблица интервалов, справа выжимка (её не читаем).

Что делает разбор (этапы А2–А4 плана «ГИС и ГВК»):
  * находит шапку листа («С к в а ж и н а №…», «Алт.=», строку «Дата»), блоки замеров по датам, интервалы, метки ГВК,
    примечания и статусы («не охвачена», «прихват» и т. п.) — ничего не теряет, а непонятное кладёт в «Проверку»;
  * приводит даты к единому виду (месяц словами, сокращения, пробелы, «Март1983г», опечатки) и ставит сезон:
    сент.–дек. — осенний контроль (после закачки), март–май — весенний (после отбора), остальное — без сезона;
  * считает по формулам листа: hi, hэф (Кг > порога), hг эф, Кг ср.взв, этаж, ГВК а.о. = −(«до» − Алт);
  * определяет пачку по заливке столбца «от» и легенде «цвет → пачка» (своя на каждый лист, правится вручную).

Слово за инженером: легенда пачек, ошибки и предупреждения — подсказки с причиной. Любой выбор можно задать вручную
(параметр `legends`), без подтверждения нераспознанный цвет остаётся без пачки и попадает в «Проверку».
Значения, сохранённые в Excel, служат эталоном: `parity()` сверяет с ними пересчёт (допуск 1e-6).
"""
from __future__ import annotations

import datetime as dt
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple

import pandas as pd

KG_MIN = 0.25  # hэф = hi только при Кг > порога (одна константа на весь расчёт)

VERH, NIZ, PERE, POKR = "Верхняя пачка", "Нижняя пачка", "Перемычка", "Покрышка"
PACKS = (POKR, VERH, PERE, NIZ)

# Легенда по умолчанию (раздел 1а плана): ключ цвета → пачка. «нет» = без заливки или белая.
DEFAULT_LEGEND = {"нет": VERH, "99CCFF": NIZ, "FFCC99": PERE, "FFFF99": POKR}

# Какая пачка подразумевается меткой ГВК (подсказка для легенды): «ГВК**» — перемычка, «ГВК*» — покрышка.
LABEL_PACK = {"ГВК": VERH, "ГВК-1": VERH, "ГВК-2": NIZ, "ГВК*": POKR, "ГВК**": PERE}

SEASON_AUTUMN, SEASON_SPRING, SEASON_NONE = "осень", "весна", "без сезона"
UNKNOWN_PACK = ""  # пачка не определена (цвет не в легенде)

_GVK_RE = re.compile(r"^ГВК\s*(-\s*\d)?\s*(\*{0,2})$", re.I)
_TOP_RE = re.compile(r"^кровля\s*(\d*)$", re.I)
_DATE_RE = re.compile(r"([А-Яа-яЁё]{2,})\.?\s*(\d{4})")
_YEAR_RE = re.compile(r"(?<!\d)((?:19|20)\d{2})\s*г")
_WELL_RE = re.compile(r"№\s*(\d+)")
_SUM_RE = re.compile(r"SUM\(\s*[A-Z](\d+)(?::[A-Z](\d+))?\s*\)", re.I)


# ----------------------------------------------------------------------------------------------- даты

def _month_of(token: str) -> Optional[int]:
    """Номер месяца по слову: полному, сокращённому («Дек.», «Сент.») или с опечаткой («Аперель»)."""
    t = token.lower()
    two = {"ян": 1, "фе": 2, "ап": 4, "ав": 8, "се": 9, "ок": 10, "но": 11, "де": 12}
    if t[:2] in two:
        return two[t[:2]]
    if t[:2] == "ма" and len(t) >= 3:
        return 3 if t[2] == "р" else 5 if t[2] in "йя" else None
    if t[:2] == "ию" and len(t) >= 3:
        return 6 if t[2] == "н" else 7 if t[2] == "л" else None
    return None


def season_of(month: Optional[int]) -> str:
    if month in (9, 10, 11, 12):
        return SEASON_AUTUMN
    if month in (3, 4, 5):
        return SEASON_SPRING
    return SEASON_NONE


@dataclass
class ParsedDate:
    """Дата замера. date — 1-е число месяца; exact — точный день, если в ячейке настоящая дата."""
    date: Optional[dt.date] = None
    exact: Optional[dt.date] = None
    season: str = SEASON_NONE
    quality: str = ""        # «», «только год», «нет даты»
    note: str = ""           # хвост текста: «спец. исследования» и т. п.


def parse_date(value) -> ParsedDate:
    """Дата из ячейки столбца A. Нераспознанное → quality «нет даты» (весь текст уходит в note)."""
    if isinstance(value, (dt.datetime, dt.date)):
        d = value.date() if isinstance(value, dt.datetime) else value
        first = d.replace(day=1)
        return ParsedDate(first, d, season_of(d.month))
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    for m in _DATE_RE.finditer(text):
        month = _month_of(m.group(1)) if len(m.group(1)) <= 9 else None
        if month:
            first = dt.date(int(m.group(2)), month, 1)
            tail = (text[:m.start()] + " " + text[m.end():]).strip(" ,.;")
            tail = re.sub(r"^г{1,2}\.?\s*", "", tail)
            return ParsedDate(first, None, season_of(month), "", tail)
    y = _YEAR_RE.search(text)
    if y and len(text) < 20:
        return ParsedDate(dt.date(int(y.group(1)), 1, 1), None, SEASON_NONE, "только год")
    return ParsedDate(None, None, SEASON_NONE, "нет даты", text)


# ---------------------------------------------------------------------------------------------- ячейки

def _num(v) -> Tuple[Optional[float], str]:
    """(число, текст). Число из «811,5*» тоже читается (текст сохраняется); «не ясно» → (None, «не ясно»)."""
    if v is None:
        return None, ""
    if isinstance(v, bool):
        return None, str(v)
    if isinstance(v, (int, float)):
        return float(v), ""
    s = str(v).strip()
    if not s:
        return None, ""
    t = s.replace(",", ".").replace(" ", "")
    star = t.endswith("*")
    try:
        return float(t.rstrip("*")), (s if star else "")
    except ValueError:
        return None, s


def _str(v) -> str:
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


def color_key(cell) -> str:
    """Ключ заливки ячейки: «нет» (без заливки или белая), RGB без прозрачности, «theme9/0.4», «idx5»."""
    f = cell.fill
    if f is None or f.fill_type != "solid":
        return "нет"
    c = f.fgColor
    if c.type == "rgb" and isinstance(c.rgb, str):
        rgb = c.rgb.upper()
        return "нет" if rgb in ("FFFFFFFF", "00FFFFFF", "00000000", "FFFFFF") else rgb[-6:]
    if c.type == "theme":
        return f"theme{c.theme}/{round(c.tint or 0, 2)}"
    if c.type == "indexed":
        return "нет" if c.indexed in (64, 65) else f"idx{c.indexed}"
    return "нет"


# ---------------------------------------------------------------------------------------------- модель

@dataclass
class Issue:
    level: str   # error | warn | note
    code: str
    sheet: str
    row: int
    message: str
    value: str = ""


@dataclass
class Row:
    sheet: str
    well: str
    xl_row: int
    block: int
    top_label: str = ""       # «кровля 1»
    gvk_label: str = ""       # «ГВК», «ГВК-1», «ГВК*» …
    label_raw: str = ""       # всё, что стоит в столбце B
    frm: Optional[float] = None
    to: Optional[float] = None
    kp: Optional[float] = None
    kg: Optional[float] = None
    kg_text: str = ""         # Кг не число: «не ясно», «30-80», «возможно газонасыщен»
    status: str = ""          # текст вместо интервала: «не охвачена», «газ … не отмечается»
    note: str = ""
    color: str = "нет"
    pack: str = UNKNOWN_PACK
    hi: Optional[float] = None
    hef: float = 0.0
    kg_hef: float = 0.0       # Кгi·hэф (столбец I листа)
    gvk_abs: Optional[float] = None
    gvk_extra: bool = False   # второй ГВК в той же пачке того же замера: не используется, ошибка проверки
    xl: Dict[str, object] = field(default_factory=dict)   # сохранённые значения Excel (эталон)
    xf: Dict[str, str] = field(default_factory=dict)      # формулы Excel


@dataclass
class Block:
    sheet: str
    well: str
    block: int
    xl_row: int
    date_text: str
    date: Optional[dt.date]
    exact: Optional[dt.date]
    season: str
    quality: str
    note: str
    rows: List[Row] = field(default_factory=list)


@dataclass
class LegendEntry:
    sheet: str
    color: str
    pack: str
    source: str     # по умолчанию | по меткам ГВК | вручную | нет в легенде
    status: str     # ок | проверить | не распознан
    reason: str
    rows: int
    labels: str


@dataclass
class SheetResult:
    sheet: str
    well: str
    header_row: Optional[int]
    alt: Optional[float]
    alt_source: str
    hints: List[str]
    blocks: List[Block]
    rows: List[Row]
    legend: List[LegendEntry]
    issues: List[Issue]
    alt_sheet: Optional[float] = None    # альтитуда по подписи «Алт.=» на листе (даже если взята другая)
    perf: Optional[Tuple[float, float]] = None   # (верх, низ) перфорации, а.о., из файла «Глубины_перфораций»
    tab_color: str = ""                  # цвет ярлыка листа (RGB без прозрачности) — атрибут скважины

    @property
    def data_rows(self) -> List[Row]:
        return [r for r in self.rows if r.frm is not None or r.to is not None]


# ---------------------------------------------------------------------------------------------- статусы

DONE = "выполнено"
# Статус наблюдения по тексту примечания: порядок важен (первое совпадение для одной строки текста).
STATUS_RULES = [
    (re.compile(r"не\s+были\s+выполнен|не\s+выполнен", re.I), "не выполнено"),
    (re.compile(r"не\s*охвач|не\s*охв\.|охвачен[аы]?\s+не", re.I), "пласт не охвачен"),
    (re.compile(r"обводнен", re.I), "обводнена"),
    (re.compile(r"остановк", re.I), "остановка"),
    (re.compile(r"прихват", re.I), "прихват"),
    (re.compile(r"газ\s+в\s+пласте.*не\s+отмеча", re.I), "газ не отмечается"),
    (re.compile(r"работы\s+проведены\s+на\s+скважине", re.I), "работы на другой скважине"),
    (re.compile(r"заглушен|ремонт", re.I), "скважина в ремонте"),
    (re.compile(r"спец\.?\s*исследован", re.I), "спец. исследования"),
    (re.compile(r"переаттестац", re.I), "переаттестация"),
]


def status_of(text: str) -> str:
    """Код статуса по тексту примечания; пусто, если текст ничего не говорит о ходе исследования."""
    for rx, code in STATUS_RULES:
        if rx.search(text or ""):
            return code
    return ""


def block_statuses(b: "Block") -> Tuple[List[str], str]:
    """(коды статусов, весь текст) замера: примечания в дате, тексты вместо интервалов, метки «остановка», Кг-слова."""
    texts = [b.note or ""]
    for r in b.rows:
        texts += [r.status, r.note, r.kg_text if status_of(r.kg_text) else "", r.label_raw if not r.gvk_label and not r.top_label else ""]
    codes: List[str] = []
    for t in texts:
        c = status_of(t)
        if c and c not in codes:
            codes.append(c)
    has_data = any(r.hi is not None for r in b.rows)
    if has_data:
        codes.insert(0, DONE)
    return codes, "; ".join(dict.fromkeys(t for t in texts if t))


# Цвета ярлыков листов: смысл, который назвал инженер (остальные цвета читаются, но смысл не придумывается).
TAB_MEANING = {"00FF00": "газонасыщенная перемычка между двумя пачками (выборка для Костикова)"}


def tab_color(ws) -> str:
    c = ws.sheet_properties.tabColor
    if c is None:
        return ""
    if c.type == "rgb" and isinstance(c.rgb, str):
        return c.rgb.upper()[-6:]
    return f"theme{c.theme}" if c.type == "theme" else f"idx{c.indexed}"


def read_bottom(path: str, sheet: str = "Тек.забой") -> pd.DataFrame:
    """Лист «Тек.забой»: максимальная глубина прохождения прибора по скважинам и датам.

    Строка заголовка — «№ скв.» и номера скважин по столбцам; первая строка без даты — текущий забой скважины, ниже даты
    замеров. Результат длинный: Скважина, Дата, Глубина, Вид («текущий забой» / «прохождение прибора»).
    """
    import openpyxl
    wv = openpyxl.load_workbook(path, data_only=True)
    cols = ["Скважина", "Дата", "Глубина, м", "Вид"]
    if sheet not in wv.sheetnames:
        return pd.DataFrame(columns=cols)
    ws = wv[sheet]
    head = next((r for r in range(1, 10) if "скв" in _str(ws.cell(r, 1).value).lower()), None)
    if head is None:
        return pd.DataFrame(columns=cols)
    wells = {}
    for c in range(2, ws.max_column + 1):
        v, _ = _num(ws.cell(head, c).value)
        if v is not None:
            wells[c] = str(int(v))
    out = []
    for r in range(head + 1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        d = parse_date(a) if a is not None else None
        for c, w in wells.items():
            depth, _ = _num(ws.cell(r, c).value)
            if depth is None:
                continue
            if a is None:
                out.append([w, None, depth, "текущий забой"])
            elif d is not None and d.date is not None:
                out.append([w, d.exact or d.date, depth, "прохождение прибора"])
    return pd.DataFrame(out, columns=cols)


# ----------------------------------------------------------------------------------------------- разбор

def _find_header(ws) -> Optional[int]:
    for r in range(1, 15):
        if _str(ws.cell(r, 1).value).lower() == "дата":
            return r
    return None


def _find_alt(ws, wv, header: int, rows: List[Row]) -> Tuple[Optional[float], str]:
    """Альтитуда: число справа от «Алт.=» в шапке; иначе ячейка, на которую ссылаются формулы ГВК ($M$3)."""
    for r in range(1, header + 1):
        for c in range(1, ws.max_column + 1):
            if _str(wv.cell(r, c).value).startswith("Алт"):
                for cc in range(c + 1, c + 4):
                    v, _ = _num(wv.cell(r, cc).value)
                    if v is not None:
                        return v, f"«Алт.=» {ws.cell(r, cc).coordinate}"
    for row in rows:
        for f in row.xf.values():
            m = re.search(r"\$([A-Z]+)\$(\d+)", f or "")
            if m:
                v, _ = _num(wv[f"{m.group(1)}{m.group(2)}"].value)
                if v is not None:
                    return v, f"ячейка {m.group(1)}{m.group(2)} из формулы ГВК (подписи «Алт.» нет)"
    return None, ""


def _hints(ws, header: int) -> List[str]:
    """Пачки, названные в заголовках правой таблицы (подсказка для легенды; сама таблица не читается)."""
    found = []
    for row in ws.iter_rows(min_row=1, max_row=header, min_col=15):
        for c in row:
            s = _str(c.value).lower()
            for key, pack in (("покрышк", POKR), ("перемычк", PERE), ("верхняя", VERH), ("нижняя", NIZ)):
                if s and key in s and pack not in found:
                    found.append(pack)
    return found


def _classify_b(row: Row, b_cached) -> None:
    s = _str(b_cached)
    row.label_raw = s
    if not s:
        return
    if _GVK_RE.match(s):
        row.gvk_label = re.sub(r"\s+", "", s.upper())
    elif _TOP_RE.match(s):
        row.top_label = s
    else:
        row.note = s


def _fix_shifted_dates(blocks: List[Block], issues: List[Issue]) -> None:
    """Дата записана строкой ниже первого интервала замера (бывает после вставки строк): Excel считал замер по диапазону,
    который начинается выше строки с датой (SUM(I24:I27) при дате в строке 25). Строки между началом диапазона и датой
    переходят в этот замер; в «Проверку» идёт заметка."""
    for prev, b in zip(blocks, blocks[1:]):
        if not b.rows:
            continue
        head = b.rows[0]
        m = _SUM_RE.search(head.xf.get("K", "") or head.xf.get("H", ""))
        start = int(m.group(1)) if m else head.xl_row
        moved = [r for r in prev.rows if start <= r.xl_row < head.xl_row]
        if not moved or len(moved) == len(prev.rows):
            continue
        prev.rows = [r for r in prev.rows if r not in moved]
        for r in moved:
            r.block = b.block
        b.rows = moved + b.rows
        issues.append(Issue("note", "date_shifted", b.sheet, head.xl_row,
                            f"Дата стоит ниже первого интервала замера: строки {moved[0].xl_row}–{moved[-1].xl_row} "
                            f"отнесены к замеру {b.date_text.strip()} (так же считает формула листа)", b.date_text))


def parse_sheet(ws, wv, legend: Optional[Dict[str, str]] = None, alt: Optional[float] = None,
                kg_min: float = KG_MIN, file_alts: Optional[Dict[str, float]] = None,
                perfs: Optional[Dict[str, Tuple[float, float]]] = None) -> SheetResult:
    """Разбор одного листа. ws — лист с формулами, wv — тот же лист со значениями; legend — ручные правки {цвет: пачка}.

    Альтитуда: вручную (alt) главнее файла (file_alts: {скважина: Алт}), файл главнее подписи на листе. При расхождении
    с листом больше 0,05 м в проверку идут оба значения. perfs: {скважина: (верх, низ)} перфорации из того же файла.
    """
    name = ws.title
    issues: List[Issue] = []
    header = _find_header(ws)
    m = None
    for r in range(1, (header or 5) + 1):
        for c in range(1, 10):
            m = m or _WELL_RE.search(_str(wv.cell(r, c).value))
    well = m.group(1) if m else name
    if header is None:
        issues.append(Issue("error", "no_header", name, 0, "Не найдена строка шапки «Дата»: лист пропущен"))
        return SheetResult(name, well, None, alt, "", [], [], [], [], issues)

    blocks: List[Block] = []
    rows: List[Row] = []
    cur: Optional[Block] = None
    for r in range(header + 1, ws.max_row + 1):
        a, b, cv, dv = (wv.cell(r, c).value for c in (1, 2, 3, 4))
        if all(wv.cell(r, c).value in (None, "") for c in (1, 2, 3, 4, 6, 7)):
            continue    # пустая строка: в ней остались только формулы-заготовки
        a_s = _str(a)
        if a_s or isinstance(a, (dt.datetime, dt.date)):
            pd_ = parse_date(a)
            if pd_.quality == "нет даты" and cur is not None:
                cur.note = (cur.note + "; " if cur.note else "") + a_s    # примечание внутри блока даты
            else:
                cur = Block(name, well, len(blocks) + 1, r, a_s or str(a), pd_.date, pd_.exact, pd_.season,
                            pd_.quality, pd_.note)
                blocks.append(cur)
                if pd_.quality == "нет даты":
                    issues.append(Issue("warn", "date_unparsed", name, r, "Дата не распознана: блок без даты", a_s))
                elif pd_.quality == "только год":
                    issues.append(Issue("warn", "date_year_only", name, r, "В дате нет месяца: взято 1 января", a_s))
                elif pd_.season == SEASON_NONE:
                    issues.append(Issue("note", "season_none", name, r,
                                        "Месяц вне сезонов (сент.–дек., март–май): замер без сезона, выбор вручную", a_s))
        if cur is None:
            cur = Block(name, well, 1, r, "", None, None, SEASON_NONE, "нет даты", "")
            blocks.append(cur)
            issues.append(Issue("warn", "rows_without_date", name, r, "Строки до первой даты: блок без даты"))
        row = Row(name, well, r, cur.block)
        for col, key in ((5, "E"), (8, "H"), (9, "I"), (10, "J"), (11, "K"), (12, "L"), (13, "M"), (14, "N")):
            row.xl[key] = wv.cell(r, col).value
            fv = ws.cell(r, col).value
            if isinstance(fv, str) and fv.startswith("="):
                row.xf[key] = fv
        _classify_b(row, b)
        row.frm, t1 = _num(cv)
        row.to, t2 = _num(dv)
        row.kp, _ = _num(wv.cell(r, 6).value)
        row.kg, row.kg_text = _num(wv.cell(r, 7).value)
        if row.frm is None or row.to is None:
            row.status = _str(cv) if row.frm is None else _str(dv)
        if row.frm is not None and row.to is not None:
            for t in (t1, t2):
                if t:
                    row.note = (row.note + "; " if row.note else "") + f"в ячейке интервала «{t}»"
                    issues.append(Issue("warn", "interval_text", name, r,
                                        f"В интервале текст «{t}»: число прочитано, проверьте", t))
            row.hi = row.to - row.frm
            if row.hi < 0:
                issues.append(Issue("error", "interval_reversed", name, r, "«До» меньше «от»", f"{row.frm}–{row.to}"))
            if row.kg_text:
                issues.append(Issue("warn", "kg_text", name, r,
                                    f"Кг не число («{row.kg_text}»): интервал не входит в hэф", row.kg_text))
            row.hef = row.hi if (row.kg is not None and row.kg > kg_min) else 0.0
            row.kg_hef = (row.kg or 0.0) * row.hef
            row.color = color_key(ws.cell(r, 3))
        elif row.frm is not None or row.to is not None:
            row.color = color_key(ws.cell(r, 3))
            issues.append(Issue("warn", "interval_incomplete", name, r, "Интервал неполный", row.status))
        rows.append(row)
        cur.rows.append(row)

    _fix_shifted_dates(blocks, issues)
    sheet_alt, sheet_src = _find_alt(ws, wv, header, rows)
    file_alt = (file_alts or {}).get(well)
    if alt is not None:
        the_alt, alt_src = alt, "задана вручную"
    elif file_alt is not None:
        the_alt, alt_src = file_alt, "файл альтитуд"
        if sheet_alt is not None and abs(sheet_alt - file_alt) > 0.05:
            issues.append(Issue("warn", "alt_mismatch", name, 0,
                                f"Альтитуда: на листе {sheet_alt:g}, в файле {file_alt:g}; взята из файла", f"{sheet_alt:g} / {file_alt:g}"))
    else:
        the_alt, alt_src = sheet_alt, sheet_src
    if the_alt is None:
        has_gvk = any(x.gvk_label for x in rows)
        issues.append(Issue("error" if has_gvk else "warn", "no_alt", name, 0,
                            "Нет альтитуды: ГВК а.о. не посчитан" if has_gvk else "Нет альтитуды"))
    hints = _hints(ws, header)

    legend_rows = _build_legend(name, rows, hints, legend or {}, issues)
    pack_of = {e.color: e.pack for e in legend_rows}
    for row in rows:
        if row.frm is not None or row.to is not None:
            row.pack = pack_of.get(row.color, UNKNOWN_PACK)

    for b in blocks:
        _gvk_per_pack(b, the_alt, issues)
        codes, text = block_statuses(b)
        if not codes:
            issues.append(Issue("warn", "block_empty", name, b.xl_row,
                                "Замер без интервалов и без понятного статуса: текст сохранён как примечание", text))
    res = SheetResult(name, well, header, the_alt, alt_src, hints, blocks, rows, legend_rows, issues, sheet_alt,
                      (perfs or {}).get(well), tab_color(ws))
    if res.tab_color:
        issues.append(Issue("note", "tab_color", name, 0,
                            "Цвет ярлыка: " + (TAB_MEANING.get(res.tab_color) or f"{res.tab_color} (смысл не задан)"), res.tab_color))
    return res


def _build_legend(sheet: str, rows: List[Row], hints: List[str], manual: Dict[str, str],
                  issues: List[Issue]) -> List[LegendEntry]:
    """Легенда листа «цвет → пачка» с причиной выбора. Ручной выбор главнее подсказки."""
    colors = Counter(r.color for r in rows if r.frm is not None or r.to is not None)
    evidence: Dict[str, Counter] = defaultdict(Counter)
    for r in rows:
        if r.gvk_label in LABEL_PACK and (r.frm is not None or r.to is not None):
            evidence[r.color][LABEL_PACK[r.gvk_label]] += 1
    out = []
    for color, n in sorted(colors.items(), key=lambda kv: (kv[0] != "нет", -kv[1])):
        labels = ", ".join(sorted({r.gvk_label for r in rows if r.color == color and r.gvk_label}))
        default = DEFAULT_LEGEND.get(color)
        ev = evidence.get(color, Counter())
        ev_pack = next(iter(ev)) if len(ev) == 1 else None
        if color in manual:
            pack, source, status, reason = manual[color], "вручную", "ок", "выбор инженера"
        elif default is None and ev_pack:
            pack, source, status = ev_pack, "по меткам ГВК", "проверить"
            reason = f"цвета нет в легенде; метки ГВК ({labels}) указывают на «{ev_pack}»"
        elif default is None:
            pack, source, status = UNKNOWN_PACK, "нет в легенде", "не распознан"
            reason = "цвета нет в легенде, меток ГВК на нём нет: выберите пачку вручную"
        elif ev_pack and ev_pack != default and color != "нет":
            pack, source, status = ev_pack, "по меткам ГВК", "проверить"
            reason = f"по умолчанию «{default}», но метки ГВК ({labels}) указывают на «{ev_pack}»"
        else:
            pack, source, status, reason = default, "по умолчанию", "ок", "легенда по умолчанию"
            if color == "нет" and ev_pack and ev_pack != default:
                status, reason = "проверить", f"без заливки, но метки ГВК ({labels}) указывают на «{ev_pack}»"
        if hints and pack and pack not in hints and status == "ок" and source != "вручную":
            status = "проверить"
            reason += f"; в заголовках справа такой пачки нет ({', '.join(hints)})"
        out.append(LegendEntry(sheet, color, pack, source, status, reason, n, labels))
    free = [h for h in hints if h not in {e.pack for e in out}]
    if len(free) == 1:      # заголовки назвали пачку, которой нет ни у одного цвета: подсказка для строк «проверить»
        for e in out:
            if e.status != "ок" and e.pack != free[0]:
                e.reason += f"; в заголовках есть «{free[0]}», пока ни за одним цветом не закреплена"
    for e in out:
        if e.status != "ок":
            issues.append(Issue("warn" if e.pack else "error", "legend_" + ("check" if e.pack else "unknown"), sheet, 0,
                                f"Заливка {e.color}: {e.reason}", e.pack or "пачка не определена"))
    return out


def _gvk_per_pack(block: Block, alt: Optional[float], issues: List[Issue]) -> None:
    """ГВК а.о. и правило «один ГВК на пачку»: второй в той же пачке того же замера — ошибка и не используется."""
    seen: Dict[str, Row] = {}
    for r in block.rows:
        if not r.gvk_label or r.to is None:
            continue
        key = r.pack or f"цвет {r.color}"
        if key in seen:
            r.gvk_extra = True
            issues.append(Issue("error", "gvk_twice", r.sheet, r.xl_row,
                                f"Второй ГВК ({r.gvk_label}) в пачке «{key}» одного замера (первый — {seen[key].gvk_label}, "
                                f"строка {seen[key].xl_row}): не используется", r.gvk_label))
            continue
        seen[key] = r
        if alt is not None:
            r.gvk_abs = -(r.to - alt)


# --------------------------------------------------------------------------------------------- итоги

def _aggregate(rows: List[Row]) -> dict:
    """hг эф, Кг ср.взв, этаж и ГВК по группе строк одного замера (пачка или весь замер)."""
    num = [r for r in rows if r.frm is not None and r.to is not None]
    hg = sum(r.hef for r in num)
    gvk = next((r for r in rows if r.gvk_label and not r.gvk_extra and r.to is not None), None)
    return {
        "Строк": len(num),
        "Кровля, м": num[0].frm if num else None,
        "Подошва, м": num[-1].to if num else None,
        "hг эф, м": hg if num else None,
        "Кг ср.взв": (sum(r.kg_hef for r in num) / hg) if hg > 0 else None,
        "Этаж, м": (num[-1].to - num[0].frm) if num else None,
        "ГВК": gvk.gvk_label if gvk else "",
        "ГВК глубина, м": gvk.to if gvk else None,
        "ГВК а.о., м": gvk.gvk_abs if gvk else None,
        "Лишних ГВК": sum(1 for r in rows if r.gvk_extra),
    }


def to_frames(results: Dict[str, SheetResult]) -> Dict[str, pd.DataFrame]:
    """Таблицы разбора: строки, замеры, замеры по пачкам, легенды, проверка. Название пачки — в каждой строке."""
    rows, blocks, packs, legend, issues = [], [], [], [], []
    for res in results.values():
        for b in res.blocks:
            for r in b.rows:
                has = r.hi is not None
                rows.append({
                    "Скважина": res.well, "Лист": res.sheet, "Строка листа": r.xl_row, "Дата": b.date,
                    "Дата (точная)": b.exact, "Сезон": b.season, "Пачка": r.pack, "Цвет": r.color, "Метка": r.label_raw,
                    "ГВК": r.gvk_label, "От, м": r.frm, "До, м": r.to, "hi, м": r.hi, "Кп": r.kp, "Кг": r.kg,
                    "hэф, м": r.hef if has else None, "Кг·hэф": r.kg_hef if has else None, "ГВК а.о., м": r.gvk_abs,
                    "ГВК лишний": r.gvk_extra, "Статус": r.status, "Кг (текст)": r.kg_text, "Примечание": r.note,
                })
            data = [x for x in b.rows if x.frm is not None or x.to is not None]
            allp = _aggregate(b.rows)
            notes = [b.note] + [r.status for r in b.rows if r.status and r.frm is None and r.to is None]
            codes, full_text = block_statuses(b)
            blocks.append({
                "Скважина": res.well, "Лист": res.sheet, "Замер": b.block, "Строка листа": b.xl_row,
                "Дата в файле": b.date_text, "Дата": b.date, "Дата (точная)": b.exact, "Сезон": b.season,
                "Качество даты": b.quality, "Примечание": "; ".join(x for x in notes if x),
                "Статус": ", ".join(codes) if codes else "нет данных", "Текст статуса": full_text,
                "Строк интервалов": len(data), "Пачек": len({x.pack or x.color for x in data}),
                "hг эф (весь замер), м": allp["hг эф, м"], "Кг ср.взв (весь замер)": allp["Кг ср.взв"],
                "Этаж (весь замер), м": allp["Этаж, м"],
            })
            groups: Dict[str, List[Row]] = defaultdict(list)
            for x in data:
                groups[x.pack or f"цвет {x.color}"].append(x)
            for grp in groups.values():
                packs.append({"Скважина": res.well, "Лист": res.sheet, "Замер": b.block, "Дата": b.date, "Сезон": b.season,
                              "Пачка": grp[0].pack, "Цвет": grp[0].color, "Альтитуда, м": res.alt,
                              "Альтитуда по листу, м": res.alt_sheet,
                              "Перфорация верх, м": res.perf[0] if res.perf else None,
                              "Перфорация низ, м": res.perf[1] if res.perf else None, **_aggregate(grp)})
        legend += [{"Лист": e.sheet, "Скважина": res.well, "Цвет": e.color, "Пачка": e.pack, "Источник": e.source,
                    "Статус": e.status, "Причина": e.reason, "Строк": e.rows, "Метки ГВК": e.labels} for e in res.legend]
        issues += [{"Уровень": i.level, "Код": i.code, "Скважина": res.well, "Лист": i.sheet, "Строка": i.row,
                    "Сообщение": i.message, "Значение": i.value} for i in res.issues]
    sheets = [{"Лист": r.sheet, "Скважина": r.well, "Альтитуда, м": r.alt, "Источник альтитуды": r.alt_source,
               "Цвет ярлыка": r.tab_color, "Смысл цвета": TAB_MEANING.get(r.tab_color, "смысл не задан" if r.tab_color else ""),
               "Замеров": len(r.blocks), "Строк интервалов": len(r.data_rows)} for r in results.values()]
    return {"Строки": pd.DataFrame(rows), "Замеры": pd.DataFrame(blocks), "Пачки": pd.DataFrame(packs),
            "Листы": pd.DataFrame(sheets), "Легенда": pd.DataFrame(legend), "Проверка": pd.DataFrame(issues)}


def load_book(path: str, legends: Optional[Dict[str, Dict[str, str]]] = None, alts: Optional[Dict[str, float]] = None,
              sheets: Optional[List[str]] = None, kg_min: float = KG_MIN, file_alts: Optional[Dict[str, float]] = None,
              perfs: Optional[Dict[str, Tuple[float, float]]] = None) -> Dict[str, SheetResult]:
    """Все листы-скважины книги. Служебные листы («Тек.забой», «примеч») и листы без шапки «Дата» пропускаются."""
    import openpyxl
    wb = openpyxl.load_workbook(path)
    wv = openpyxl.load_workbook(path, data_only=True)
    out: Dict[str, SheetResult] = {}
    for name in sheets or wb.sheetnames:
        ws = wb[name]
        if _find_header(ws) is None:
            continue
        out[name] = parse_sheet(ws, wv[name], (legends or {}).get(name), (alts or {}).get(name), kg_min, file_alts, perfs)
    return out


# ------------------------------------------------------------------------------------ сверка с Excel

_PLUS_RE = re.compile(r"^=\s*J(\d+)(?:\s*\+\s*J(\d+))+\s*$", re.I)
_DIFF_RE = re.compile(r"^=\s*D(\d+)\s*-\s*C(\d+)\s*$", re.I)
_GVK_F_RE = re.compile(r'IF\(\s*B\d+\s*=\s*"([^"]+)"\s*,\s*\(\s*D\d+\s*-\s*\$[A-Z]+\$\d+\s*\)\s*\*\s*\(?-1\)?', re.I)
_THR_RE = re.compile(r"IF\(\s*G\d+\s*<=\s*(\d+(?:[.,]\d+)?)\s*,", re.I)
_COLS = ["Лист", "Скважина", "Строка", "Столбец", "Excel", "Пересчёт", "Вердикт", "Пояснение"]


def sheet_threshold(row: Row) -> Optional[float]:
    """Порог Кг, записанный в формуле hэф этой строки (если отличается от общего — инженер правил формулу вручную)."""
    m = _THR_RE.search(row.xf.get("J", ""))
    return float(m.group(1).replace(",", ".")) if m else None


def _why(res: SheetResult, byrow: Dict[int, Row], r: Row, key: str, f: str) -> str:
    """Причина расхождения, если она видна в самой книге: правка формулы вручную, пустые ячейки, текст вместо Кг."""
    refs = {int(x) for x in re.findall(r"[A-Z]\$?(\d+)", re.sub(r"\$[A-Z]+\$\d+", "", f))}
    if key in ("E", "I", "J") and refs and refs != {r.xl_row}:
        return "формула ссылается на другую строку: " + ", ".join(str(x) for x in sorted(refs))
    if key in ("H", "K", "L") and len(refs) >= 1:
        lo, hi = min(refs | {r.xl_row}), max(refs | {r.xl_row})
        rng = [byrow[i] for i in range(lo, hi + 1) if i in byrow]
        bare = [x.xl_row for x in rng if x.hi is not None and "J" not in x.xf]
        if bare:
            return "в диапазоне строки без формулы hэф: " + ", ".join(str(x) for x in bare[:5])
        txt = [x.xl_row for x in rng if x.kg_text]
        if txt:
            return "в диапазоне Кг текстом: " + ", ".join(str(x) for x in txt[:5])
        if key == "K" and any(not (isinstance(x.xl.get("J"), (int, float))) for x in rng if x.hi is not None):
            return "в диапазоне hэф не число (ошибка или текст в Excel)"
    if key in ("I", "J", "K") and r.kg_text:
        return f"Кг текстом («{r.kg_text}»): Excel считает такой интервал эффективным, мы — нет"
    n = r.xl_row
    std = {"I": (f"=G{n}*J{n}", f"=J{n}*G{n}"), "J": ()}
    if key == "I" and f.replace(" ", "") not in std["I"]:
        return "формула строки нестандартная (правлена вручную)"
    if key == "J" and not _THR_RE.search(f):
        return "формула строки нестандартная (правлена вручную)"
    return "Excel считал по другому диапазону или формуле (правлена вручную)"


def parity(results: Dict[str, SheetResult], tol: float = 1e-6) -> pd.DataFrame:
    """Пересчёт по формулам листа против сохранённых значений Excel. Одна строка = одна ячейка Excel с формулой.

    Диапазоны берутся из самой формулы Excel (SUM(I4:I8), =D8-C4), поэтому проверяется формула, а не наша группировка.
    Вердикт: «совпало»; «порог листа» (в формуле строки другой порог Кг, чем 0,25: с ним результат совпал, пересчёт с общим
    порогом нет); «расхождение»; «Excel: ошибка» (в ячейке #VALUE! и т. п.); «Excel: пусто/текст».
    """
    out = []
    for res in results.values():
        byrow = {r.xl_row: r for r in res.rows}

        def hef(r, own):
            thr = (sheet_threshold(r) if own else None) or KG_MIN
            return r.hi if (r.hi is not None and r.kg is not None and r.kg > thr) else 0.0

        def ssum(lo, hi, own, weighted=False):
            tot = 0.0
            for i in range(lo, hi + 1):
                r = byrow.get(i)
                if r is not None and r.hi is not None:
                    h = hef(r, own)
                    tot += h * (r.kg or 0.0) if weighted else h
            return tot

        for r in res.rows:
            for key, f in r.xf.items():
                excel = r.xl.get(key)
                note = ""
                variants = {}    # own=False: общий порог, own=True: порог из формулы листа
                for own in (False, True):
                    mine = None
                    if key == "E":
                        mine = r.hi
                    elif key == "J":
                        mine = hef(r, own) if r.hi is not None else None
                    elif key == "I":
                        mine = hef(r, own) * (r.kg or 0.0) if r.hi is not None else None
                    elif key == "K":
                        m = _SUM_RE.search(f)
                        if m:
                            mine = ssum(int(m.group(1)), int(m.group(2) or m.group(1)), own)
                        elif _PLUS_RE.match(f):
                            mine = sum(ssum(int(i), int(i), own) for i in re.findall(r"J(\d+)", f, re.I))
                        elif re.fullmatch(r"=\s*J\d+\s*", f):
                            mine = hef(r, own) if r.hi is not None else None
                        elif re.fullmatch(r"=\s*E\d+\s*", f):
                            mine = r.hi
                    elif key == "H":
                        m = _SUM_RE.search(f)
                        if m:
                            lo, hi = int(m.group(1)), int(m.group(2) or m.group(1))
                            hg = ssum(lo, hi, own)
                            mine = ssum(lo, hi, own, True) / hg if hg > 0 else None
                            note = "hг эф = 0: у Excel деление на 0" if hg <= 0 else ""
                        elif re.fullmatch(r"=\s*G\d+\s*", f):
                            mine = r.kg
                            note = "Excel берёт Кг строки как есть"
                        elif re.fullmatch(r"=\s*I\d+\s*/\s*K\d+\s*", f):
                            hg = hef(r, own)
                            mine = hg * (r.kg or 0.0) / hg if hg > 0 else None
                    elif key == "L":
                        m = _DIFF_RE.match(f)
                        if m:
                            d, c = byrow.get(int(m.group(1))), byrow.get(int(m.group(2)))
                            if d is not None and c is not None and d.to is not None and c.frm is not None:
                                mine = d.to - c.frm
                        elif re.fullmatch(r"=\s*[EK]\d+\s*", f):
                            mine = r.hi if f.strip()[1:2].upper() == "E" else (hef(r, own) if r.hi is not None else None)
                    elif key in ("M", "N"):
                        m = _GVK_F_RE.search(f)
                        if m and res.alt is not None and r.to is not None and r.gvk_label == m.group(1).replace(" ", ""):
                            mine = -(r.to - res.alt)
                    variants[own] = mine
                mine = variants[False]
                num = isinstance(excel, (int, float)) and not isinstance(excel, bool)
                empty_row = r.frm is None and r.to is None
                if mine is None and not num:
                    continue   # обе стороны пусты: нечего сверять
                close = lambda x: x is not None and num and abs(x - excel) <= tol * max(1.0, abs(excel))
                if not num:
                    verdict = "Excel: ошибка" if isinstance(excel, str) and excel.startswith("#") else "Excel: пусто/текст"
                    note = note or f"в ячейке «{excel}»"
                elif close(mine):
                    verdict = "совпало"
                elif close(variants[True]):
                    verdict, note = "порог листа", "в формуле листа порог Кг не 0,25"
                elif mine is None and empty_row and excel == 0:
                    verdict, note = "совпало", "пустая строка: Excel даёт 0"
                else:
                    verdict = "расхождение"
                    note = _why(res, byrow, r, key, f) or note or ("у нас нет значения" if mine is None else "")
                out.append({"Лист": res.sheet, "Скважина": res.well, "Строка": r.xl_row, "Столбец": key, "Excel": excel,
                            "Пересчёт": mine, "Вердикт": verdict, "Пояснение": note})
    return pd.DataFrame(out, columns=_COLS)


def report(results: Dict[str, SheetResult]) -> pd.DataFrame:
    """«Сколько распознано» по листам: замеры, строки интервалов, даты, пачки, ГВК, замечания."""
    rows = []
    for res in results.values():
        data = res.data_rows
        rows.append({
            "Лист": res.sheet, "Скважина": res.well, "Альтитуда": res.alt, "Замеров": len(res.blocks),
            "Дат распознано": sum(1 for b in res.blocks if b.date), "Без даты": sum(1 for b in res.blocks if not b.date),
            "Строк интервалов": len(data), "Пачек": len({r.pack or r.color for r in data}),
            "Пачки": ", ".join(sorted({r.pack for r in data if r.pack})),
            "ГВК": sum(1 for r in data if r.gvk_label and not r.gvk_extra),
            "Цветов в «Проверку»": sum(1 for e in res.legend if e.status != "ок"),
            "Ошибок": sum(1 for i in res.issues if i.level == "error"),
            "Предупреждений": sum(1 for i in res.issues if i.level == "warn"),
        })
    return pd.DataFrame(rows)


def keyed(d: dict) -> dict:
    """{номер скважины: значение} из gsp_maps.inputs.load_perforations → ключи-строки, как номера скважин в разборе."""
    return {str(int(k)): v for k, v in d.items()}


def main(argv=None) -> None:
    import argparse
    import os
    ap = argparse.ArgumentParser(description="Разбор «Базы ГИС-контроль»: отчёт и сверка с сохранёнными значениями Excel")
    ap.add_argument("xlsx")
    ap.add_argument("--perf", help="«Глубины_перфораций.xlsx»: листы «перфорации» и «альтитуды»")
    ap.add_argument("--out", help="папка для выгрузки таблиц разбора в Excel")
    a = ap.parse_args(argv)
    file_alts = perfs = None
    if a.perf:
        from gsp_maps.inputs import load_perforations
        depths, alt = load_perforations(a.perf)
        file_alts, perfs = keyed(alt), keyed(depths)
    res = load_book(a.xlsx, file_alts=file_alts, perfs=perfs)
    rep, par = report(res), parity(res)
    pd.set_option("display.width", 250, "display.max_columns", 30, "display.max_rows", 200)
    print(rep.to_string(index=False))
    print("\nСверка с Excel:", par["Вердикт"].value_counts().to_dict())
    if a.out:
        os.makedirs(a.out, exist_ok=True)
        with pd.ExcelWriter(os.path.join(a.out, "ГИС_разбор.xlsx")) as xw:
            for k, v in {**to_frames(res), "Сверка с Excel": par, "Отчёт": rep}.items():
                v.to_excel(xw, sheet_name=k[:31], index=False)


if __name__ == "__main__":
    main()
