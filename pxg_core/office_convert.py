"""Конвертация между форматами Office: Excel, Word, презентации.

Основной движок — LibreOffice (soffice --headless). Без него умеем на чистом Python только .xls → .xlsx (xlrd + openpyxl).
Исходные файлы не меняются; готовый результат повторно не пересоздаётся.
"""
from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from typing import Dict, List, Optional, Tuple

FAMILIES = {
    "excel": ("xls", "xlsx", "xlsm", "ods", "csv"),
    "word": ("doc", "docx", "odt", "rtf"),
    "slides": ("ppt", "pptx", "odp"),
}
# целевой формат -> семейство (pdf годится для всех)
TARGETS = {"xlsx": "excel", "xls": "excel", "ods": "excel", "csv": "excel",
           "docx": "word", "doc": "word", "odt": "word", "rtf": "word",
           "pptx": "slides", "ppt": "slides", "odp": "slides", "pdf": "any"}
SOURCES = tuple(e for f in FAMILIES.values() for e in f) + ("pdf",)
PDF_FILTER = {"word": "writer_pdf_import", "slides": "impress_pdf_import"}


def family_of(ext: str) -> Optional[str]:
    for fam, exts in FAMILIES.items():
        if ext in exts:
            return fam
    return "pdf" if ext == "pdf" else None


def pair_supported(src: str, dst: str) -> Tuple[bool, str]:
    """Можно ли пару (расширения без точки) вообще ждать от LibreOffice; второй элемент — причина отказа."""
    sf, tf = family_of(src), TARGETS.get(dst)
    if sf is None:
        return False, "формат .%s не поддерживается" % src
    if tf is None:
        return False, "целевой формат .%s не поддерживается" % dst
    if src == dst:
        return False, "уже в формате .%s" % dst
    if sf == "pdf":
        return (tf in ("word", "slides"), "PDF можно превратить только в документ Word или презентацию")
    if tf == "any" or tf == sf:
        return True, ""
    return False, "из %s в %s конвертировать нельзя" % ({"excel": "Excel", "word": "Word", "slides": "презентации"}[sf],
                                                         {"excel": "Excel", "word": "Word", "slides": "презентацию"}[tf])


def find_soffice() -> Optional[str]:
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    cands: List[str] = []
    for env in ("ProgramFiles", "ProgramFiles(x86)", "ProgramW6432"):
        base = os.environ.get(env)
        if base:
            cands += [os.path.join(base, d, "program", "soffice.exe") for d in ("LibreOffice", "LibreOffice 7", "LibreOffice 24")]
    cands += ["/opt/libreoffice/program/soffice", "/usr/lib64/libreoffice/program/soffice", "/usr/lib/libreoffice/program/soffice",
              "/opt/libreoffice7.6/program/soffice", "/opt/libreoffice24.8/program/soffice",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"]
    try:
        for d in sorted(os.listdir("/opt")):
            if d.lower().startswith("libreoffice"):
                cands.append(os.path.join("/opt", d, "program", "soffice"))
    except OSError:
        pass
    for c in cands:
        if os.path.isfile(c):
            return c
    return None


def collect(paths: List[str], recursive: bool = True) -> List[Tuple[str, str]]:
    """Файлы из списка путей и папок: (полный путь, корень для сохранения подпапок). Временные «~$…» пропускаются."""
    out: List[Tuple[str, str]] = []
    seen = set()
    for p in paths:
        p = p.strip().strip('"')
        if os.path.isfile(p):
            files = [(p, os.path.dirname(p))]
        elif os.path.isdir(p):
            files = []
            for root, dirs, names in os.walk(p):
                dirs.sort()
                for n in sorted(names):
                    files.append((os.path.join(root, n), p))
                if not recursive:
                    break
        else:
            continue
        for f, base in files:
            ext = os.path.splitext(f)[1].lower().lstrip(".")
            if ext in SOURCES and not os.path.basename(f).startswith(("~$", ".~lock")) and f not in seen:
                seen.add(f)
                out.append((f, base))
    return out


def xls_to_xlsx_python(src: str, dst: str) -> str:
    """Запасной путь: значения и числовые форматы; формулы, диаграммы и рисунки .xls не переносятся. Возвращает заметку о потерях."""
    import xlrd
    from openpyxl import Workbook
    book = xlrd.open_workbook(src, formatting_info=True)
    wb = Workbook()
    wb.remove(wb.active)
    for sh in book.sheets():
        ws = wb.create_sheet(sh.name[:31] or "Лист")
        for r in range(sh.nrows):
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                t = cell.ctype
                if t in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    continue
                v = cell.value
                if t == xlrd.XL_CELL_DATE:
                    try:
                        v = xlrd.xldate.xldate_as_datetime(v, book.datemode)
                    except Exception:
                        pass
                elif t == xlrd.XL_CELL_BOOLEAN:
                    v = bool(v)
                elif t == xlrd.XL_CELL_ERROR:
                    v = xlrd.error_text_from_code.get(v, "#ERR")
                out = ws.cell(row=r + 1, column=c + 1, value=v)
                try:
                    fmt = book.format_map[book.xf_list[cell.xf_index].format_key].format_str
                    if fmt and fmt != "General":
                        out.number_format = fmt
                except Exception:
                    pass
        for rng in sh.merged_cells:
            ws.merge_cells(start_row=rng[0] + 1, end_row=rng[1], start_column=rng[2] + 1, end_column=rng[3])
    if not wb.sheetnames:
        wb.create_sheet("Лист1")
    wb.save(dst)
    return "формулы заменены значениями, рисунки и диаграммы не перенесены"


def _soffice_convert(soffice: str, src: str, dst_ext: str, out_dir: str, timeout: int = 300) -> str:
    """Конвертирует src в out_dir, возвращает путь результата; бросает RuntimeError."""
    profile = tempfile.mkdtemp(prefix="lo_profile_")
    try:
        cmd = [soffice, "-env:UserInstallation=" + _file_url(profile), "--headless", "--norestore", "--nolockcheck"]
        src_ext = os.path.splitext(src)[1].lower().lstrip(".")
        fam = TARGETS.get(dst_ext)
        if src_ext == "pdf" and fam in PDF_FILTER:
            cmd += ["--infilter=" + PDF_FILTER[fam]]
        target = dst_ext
        if dst_ext == "csv":
            target = "csv:Text - txt - csv (StarCalc):44,34,76,1"   # UTF-8, запятая, кавычки, первый лист
        cmd += ["--convert-to", target, "--outdir", out_dir, src]
        try:
            res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=timeout)
        except subprocess.TimeoutExpired:
            raise RuntimeError("LibreOffice не ответил за %d с" % timeout)
        produced = os.path.join(out_dir, os.path.splitext(os.path.basename(src))[0] + "." + dst_ext)
        if res.returncode != 0 or not os.path.isfile(produced):
            msg = (res.stderr or res.stdout or b"").decode("utf-8", "replace").strip().splitlines()
            raise RuntimeError("LibreOffice не смог сконвертировать" + (": " + msg[-1] if msg else ""))
        return produced
    finally:
        shutil.rmtree(profile, ignore_errors=True)


def _file_url(path: str) -> str:
    p = os.path.abspath(path).replace("\\", "/")
    return "file:///" + p.lstrip("/")


def convert_file(src: str, dst_ext: str, dst: str, soffice: Optional[str]) -> Dict[str, str]:
    """Один файл. Результат: status = ok | skip | error, engine, note."""
    src_ext = os.path.splitext(src)[1].lower().lstrip(".")
    ok, why = pair_supported(src_ext, dst_ext)
    if not ok:
        return {"status": "skip" if src_ext == dst_ext else "error", "engine": "", "note": why}
    if os.path.exists(dst):
        return {"status": "skip", "engine": "", "note": "уже есть: " + os.path.basename(dst)}
    os.makedirs(os.path.dirname(dst) or ".", exist_ok=True)
    tmp = tempfile.mkdtemp(prefix="conv_")
    try:
        if soffice:
            try:
                produced = _soffice_convert(soffice, src, dst_ext, tmp)
                shutil.move(produced, dst)
                note = "оформление и формулы по возможностям LibreOffice"
                if src_ext == "pdf":
                    note = "из PDF: вёрстка и таблицы сохраняются приблизительно"
                elif src_ext in ("xls", "doc", "ppt") or dst_ext in ("xls", "doc", "ppt", "csv"):
                    note = "возможны мелкие отличия оформления" + ("; csv: только первый лист" if dst_ext == "csv" else "")
                return {"status": "ok", "engine": "LibreOffice", "note": note}
            except RuntimeError as e:
                if not (src_ext == "xls" and dst_ext == "xlsx"):
                    return {"status": "error", "engine": "LibreOffice", "note": str(e)}
                lo_err = str(e)
        else:
            lo_err = ""
        if src_ext == "xls" and dst_ext == "xlsx":
            try:
                part = os.path.join(tmp, "out.xlsx")
                note = xls_to_xlsx_python(src, part)
                shutil.move(part, dst)
                return {"status": "ok", "engine": "xlrd+openpyxl", "note": note + (" (LibreOffice: %s)" % lo_err if lo_err else "")}
            except Exception as e:
                return {"status": "error", "engine": "xlrd+openpyxl", "note": "не удалось прочитать файл: %s" % e}
        return {"status": "error", "engine": "", "note": "нужен LibreOffice (бесплатный, libreoffice.org): без него эта пара форматов не поддерживается"}
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def dest_path(src: str, base: str, dst_ext: str, out_root: Optional[str]) -> str:
    stem = os.path.splitext(os.path.basename(src))[0] + "." + dst_ext
    if out_root:
        rel = os.path.relpath(os.path.dirname(src), base)
        return os.path.join(out_root, "" if rel == "." else rel, stem)
    return os.path.join(os.path.dirname(src), stem)
