"""Шахматки отбора и закачки («Показатели × сутки», блоки по скважинам) и книга воды: свод в единую базу и проверка формул Qгаз.пром.

Формулы проверяются, если файл .xlsx или если для .xls есть LibreOffice (soffice); без него .xls читается только по значениям.
"""
from __future__ import annotations

import argparse
import collections
import datetime as dt
import math
import os
import re
import shutil
import subprocess
import tempfile
from typing import Dict, List, Optional, Tuple

import openpyxl
import pandas as pd
from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl

MONTHS = ["январь", "февраль", "март", "апрель", "май", "июнь", "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь"]
REF = re.compile(r"(\$?)([A-Z]{1,3})(\$?)(\d+)")
Issue = Dict[str, object]


def to_xlsx(path: str, workdir: str, use_soffice: bool = True) -> Optional[str]:
    """Путь к .xlsx для чтения формул; для .xls — через LibreOffice, если он есть (иначе None)."""
    if path.lower().endswith((".xlsx", ".xlsm")):
        return path
    exe = (shutil.which("soffice") or shutil.which("libreoffice")) if use_soffice else None
    if not exe:
        return None
    src = os.path.join(workdir, "src" + os.path.splitext(path)[1])
    shutil.copy(path, src)
    try:
        subprocess.run([exe, "--headless", "--convert-to", "xlsx", "--outdir", workdir, src],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300, check=False)
    except (OSError, subprocess.SubprocessError):
        return None
    out = os.path.join(workdir, "src.xlsx")
    return out if os.path.exists(out) else None


def _text(v) -> str:
    return "" if v is None else re.sub(r"\s+", " ", str(v)).strip()


def _num(v) -> Optional[float]:
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return None if (isinstance(v, float) and math.isnan(v)) else float(v)
    try:
        return float(str(v).replace(",", ".").strip())
    except ValueError:
        return None


def _use_calamine() -> bool:
    if os.environ.get("PXG_EXCEL_ENGINE", "").lower() == "openpyxl":
        return False
    try:
        import python_calamine  # noqa: F401
        return True
    except ImportError:
        return False


def _cut(rows, max_row: int, max_col: int):
    out = []
    for r in list(rows)[:max_row]:
        r = list(r)[:max_col]
        out.append([None if (v == "" or v is None) else v for v in r])
    return out


def _values_calamine(path: str, max_row: int, max_col: int):
    from python_calamine import CalamineWorkbook
    wb = CalamineWorkbook.from_path(path)
    name = wb.sheet_names[0]
    return _cut(wb.get_sheet_by_name(name).to_python(skip_empty_area=False), max_row, max_col), name


def _values_plain(path: str, max_row: int, max_col: int):
    if path.lower().endswith((".xlsx", ".xlsm")):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return _cut(ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True), max_row, max_col), ws.title
    import xlrd
    sh = xlrd.open_workbook(path).sheet_by_index(0)
    return _cut(([sh.cell_value(r, c) for c in range(min(sh.ncols, max_col))] for r in range(min(sh.nrows, max_row))), max_row, max_col), sh.name


def _values(path: str, max_row: int, max_col: int):
    """Значения первого листа: calamine (быстро, и .xls, и .xlsx), иначе openpyxl в режиме чтения / xlrd.

    calamine на некоторых .xls падает паникой Rust (PanicException - не Exception): ловим и читаем запасным движком."""
    if _use_calamine():
        try:
            return _values_calamine(path, max_row, max_col)
        except (KeyboardInterrupt, SystemExit):
            raise
        except BaseException:  # noqa: B036  включая pyo3_runtime.PanicException
            pass
    return _values_plain(path, max_row, max_col)


def _formulas(path: str, max_row: int, max_col: int):
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.worksheets[0]
    return _cut(ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True), max_row, max_col)


class Shakh:
    """Шахматка «показатели × сутки» с блоками по скважинам (одинаковый шаблон у отбора и закачки)."""

    def __init__(self, path: str, kind: str, start_year: int, workdir: str, max_row: int = 2600, max_col: int = 260,
                 xls_formulas: bool = False):
        self.kind, self.start_year = kind, start_year
        self.label = os.path.basename(path)
        self.issues: List[Issue] = []
        self.src = to_xlsx(path, workdir, xls_formulas)
        self.has_formulas = self.src is not None
        if self.has_formulas:
            self.V, self.sheet = _values(self.src if not _use_calamine() else path, max_row, max_col)
            self.F = _formulas(self.src, max_row, max_col)
        else:
            self.V, self.sheet = _values(path, max_row, max_col)
            self.F = self.V
            self.issue("предупреждение", "Формулы недоступны", os.path.basename(path),
                       "Файл .xls прочитан по значениям: формулы не проверялись (сохраните его как .xlsx или включите «Проверять формулы в .xls» – нужен LibreOffice, дольше минуты)")
        self._find_layout()

    def f(self, r, c):  # 1-based
        return self.F[r - 1][c - 1] if r - 1 < len(self.F) and c - 1 < len(self.F[r - 1]) else None

    def v(self, r, c):
        return self.V[r - 1][c - 1] if r - 1 < len(self.V) and c - 1 < len(self.V[r - 1]) else None

    def _col_has_data(self, c) -> str:
        n = sum(1 for _s, _w, labs in self.blocks for r in labs.values() if self.v(r, c) not in (None, "")) if hasattr(self, "blocks") else 0
        return "" if not n else " (в нём есть данные: %d ячеек)" % n

    def issue(self, level, kind, where, text):
        self.issues.append({"Файл": self.kind, "Уровень": level, "Вид": kind, "Где": where, "Описание": text})

    def _find_layout(self):
        hdr = next(r for r in range(1, 12) if _text(self.v(r, 1)) == "№№")
        self.month_row, self.day_row = hdr, hdr + 1
        ncol = len(self.V[0])
        zone_hdr = None
        for r in range(1, 12):
            for c in range(150, ncol + 1):
                if _text(self.v(r, c)) == "Зоны":
                    zone_hdr = (r, c)
        self.zone_hdr = zone_hdr
        last = (zone_hdr[1] - 1) if zone_hdr else ncol
        # столбцы: сутки и итоги
        self.day_cols: List[Tuple[int, dt.date, str]] = []
        self.total_cols: List[Tuple[int, str]] = []
        cur: Optional[dt.date] = None
        label_runs: List[Tuple[int, int, str, str]] = []   # (с, по, подпись, дата по ряду)
        for c in range(4, last + 1):
            m, d = _text(self.v(self.month_row, c)), self.v(self.day_row, c)
            dn = _num(d)
            if m.lower().startswith("итог"):
                self.total_cols.append((c, "%s %s" % (m, _text(d)))); continue
            if dn is None or not m:
                continue
            day = int(dn)
            where = "%s%d" % (gl(c), self.day_row)
            if cur is None:
                mi = MONTHS.index(m.lower()) + 1 if m.lower() in MONTHS else 1
                cur = dt.date(self.start_year, mi, day)
            else:
                nxt = cur + dt.timedelta(days=1)
                if day == nxt.day:
                    cur = nxt
                elif day == 1 and cur.day >= 28:
                    self.issue("предупреждение", "Пропуск суток", where,
                               "Между %s и 1-м числом следующего месяца нет столбцов (пропущены сутки)" % cur.isoformat())
                    cur = (cur.replace(day=1) + dt.timedelta(days=32)).replace(day=1)
                elif day > cur.day:
                    try:
                        new = cur.replace(day=day)
                        self.issue("предупреждение", "Пропуск суток", where, "Нет столбцов между %s и %s" % (cur.isoformat(), new.isoformat()))
                        cur = new
                    except ValueError:
                        self.issue("ошибка", "Лишний столбец суток", where,
                                   "Номер дня %d после %s не существует; столбец пропущен%s" % (
                                       day, cur.isoformat(), self._col_has_data(c)))
                        continue
                else:
                    self.issue("ошибка", "Нумерация суток", where, "Номер дня %d не продолжает ряд после %s; столбец пропущен%s" % (
                        day, cur.isoformat(), self._col_has_data(c)))
                    continue
            if m.lower() in MONTHS and MONTHS.index(m.lower()) + 1 != cur.month:
                if label_runs and label_runs[-1][1] == c - 1 and label_runs[-1][2] == m:
                    label_runs[-1] = (label_runs[-1][0], c, m, label_runs[-1][3])
                else:
                    label_runs.append((c, c, m, cur.isoformat()))
            self.day_cols.append((c, cur, m))
        for c0, c1, m, ds in label_runs:
            self.issue("предупреждение", "Подпись месяца", "%s%d:%s%d" % (gl(c0), self.month_row, gl(c1), self.month_row),
                       "Подпись «%s» не совпадает с датой по ряду суток (с %s); дата взята по ряду суток" % (m, ds))
        # блоки скважин
        self.blocks: List[Tuple[int, str, Dict[str, int]]] = []   # (стартовая строка, скважина, {показатель: строка})
        r = 1
        nrow = len(self.V)
        while r <= nrow:
            if _text(self.v(r, 2)) == "Состояние":
                well = _text(self.v(r, 1))
                try:
                    well = str(int(float(well)))
                except ValueError:
                    pass
                labs: Dict[str, int] = {}
                rr = r
                while rr <= nrow and _text(self.v(rr, 2)) and (rr == r or _text(self.v(rr, 2)) != "Состояние"):
                    lab = _text(self.v(rr, 2))
                    labs[lab] = rr
                    rr += 1
                if any(b[1] == well for b in self.blocks):
                    self.issue("предупреждение", "Повтор скважины", "строка %d" % r,
                               "Скважина «%s» уже есть выше; блок загружен как «%s (дубль, строка %d)»" % (well, well, r))
                    well = "%s (дубль, строка %d)" % (well, r)
                self.blocks.append((r, well, labs))
                r = rr
            else:
                r += 1
        self.first_block_row = self.blocks[0][0]
        self.last_block_row = max(max(b[2].values()) for b in self.blocks)
        names = [b[1] for b in self.blocks]

    # --- чтение данных
    def wells_table(self) -> pd.DataFrame:
        recs = []
        for start, well, labs in self.blocks:
            for c, date, _m in self.day_cols:
                rec = {"Скважина": well, "Дата": date, "Строка блока": start}
                nonempty = False
                for lab, r in labs.items():
                    val = self.v(r, c)
                    if val is not None and val != "":
                        nonempty = True
                    rec[lab] = val
                if nonempty:
                    recs.append(rec)
        return pd.DataFrame(recs)

    def totals_table(self) -> pd.DataFrame:
        recs = []
        for start, well, labs in self.blocks:
            for c, name in self.total_cols:
                for lab, r in labs.items():
                    val = self.v(r, c)
                    if val is not None and val != "":
                        recs.append({"Скважина": well, "Итог": name, "Показатель": lab, "Значение": val})
        return pd.DataFrame(recs)

    def object_rows(self) -> Dict[str, int]:
        """Строки ПХГ-уровня: выше первого блока (подпись в столбце B) и расчётные строки ниже последнего блока."""
        rows: Dict[str, int] = {}
        for r in range(self.day_row + 1, self.first_block_row):
            lab = _text(self.v(r, 2))
            if lab:
                rows.setdefault(lab, r)
        for r in range(self.last_block_row + 1, len(self.V) + 1):
            lab = _text(self.v(r, 2))
            if lab:
                rows.setdefault("(внизу) " + lab, r)
        return rows

    def object_table(self) -> pd.DataFrame:
        rows = self.object_rows()
        recs = []
        for c, date, _m in self.day_cols:
            rec: Dict[str, object] = {"Дата": date}
            nonempty = False
            for lab, r in rows.items():
                val = self.v(r, c)
                rec[lab] = val
                if val not in (None, ""):
                    nonempty = True
            if nonempty:
                recs.append(rec)
        return pd.DataFrame(recs)

    # --- зоны
    def zones(self) -> pd.DataFrame:
        if not self.zone_hdr:
            return pd.DataFrame(columns=["Скважина", "Группа"])
        hr, zc = self.zone_hdr
        recs = []
        r = hr + 3
        while r <= len(self.V):
            w = _num(self.v(r, zc + 2))
            if w is None:
                if r > hr + 3 and _text(self.v(r, zc + 2)) == "" and _text(self.v(r + 1, zc + 2)) == "":
                    break
                r += 1
                continue
            recs.append({"Скважина": str(int(w)), "Группа": re.sub(r"\s+", " ", _text(self.v(r, zc))), "Строка": r})
            r += 1
        return pd.DataFrame(recs)

    def zones_from_formulas(self, zones: pd.DataFrame) -> pd.DataFrame:
        """Второй источник групп: формулы SUM(...) в таблице «по зонам» (имя зоны слева, ссылки на строки со скважинами)."""
        hr, zc = self.zone_hdr
        sum_col = zc + 10   # столбец «По зонам, тыс.м3»
        row2well = {int(r): w for w, r in zip(zones["Скважина"], zones["Строка"])}
        recs = []
        for r in range(hr + 1, hr + 40):
            name = re.sub(r"\s+", " ", _text(self.v(r, zc + 9))).strip()
            fml = _text(self.f(r, sum_col))
            if not name or name.startswith("Итого") or not fml.startswith("=SUM("):
                continue
            for m in REF.finditer(fml):
                rr = int(m.group(4))
                if rr in row2well:
                    recs.append({"Скважина": row2well[rr], "Группа (по формуле)": name})
        return pd.DataFrame(recs)

    # --- проверки формул
    def _leaves(self, r, c, label_rows, depth=0):
        fm = self.f(r, c)
        if not (isinstance(fm, str) and fm.startswith("=")) or depth > 5:
            return []
        out = []
        seen = set()
        for m in REF.finditer(fm):
            rr = int(m.group(4))
            if (m.group(2), rr) in seen:
                continue
            seen.add((m.group(2), rr))
            if rr in label_rows:
                out.append(label_rows[rr])
            elif ci(m.group(2)) == c and rr != r:
                out.extend(self._leaves(rr, c, label_rows, depth + 1))
        return out

    def check_sum_coverage(self, k_row: Optional[int]):
        """Коэффициент пересчёта K: из каких скважин складывается знаменатель (по цепочке формул)."""
        label_rows = {r: (well, lab) for _s, well, labs in self.blocks for lab, r in labs.items()}
        all_wells = [b[1] for b in self.blocks]
        res = collections.defaultdict(list)
        self.k_missing = {}
        for c, d, _m in self.day_cols:
            leaves = self._leaves(k_row, c, label_rows) if k_row else []
            if not leaves:
                continue
            lab = collections.Counter(l for _w, l in leaves).most_common(1)[0][0]
            got = collections.Counter(w for w, l in leaves if l == lab)
            key = (lab, tuple(w for w in all_wells if w not in got), tuple(sorted(w for w, n in got.items() if n > 1)))
            self.k_missing[d] = set(key[1])
            res[key].append(d)
        for (lab, missing, dup), days in res.items():
            if missing or dup:
                self.issue("ошибка", "K считается не по всем скважинам", "строка K %s, суток: %d (%s…%s)" % (k_row, len(days), min(days), max(days)),
                           "В знаменателе K (показатель «%s») нет скважин: %s%s" % (lab, ", ".join(missing) or "нет",
                                                                                     ("; учтены дважды: " + ", ".join(dup)) if dup else ""))

    def check_formula_patterns(self, labels: List[str]):
        """Однотипные ячейки блоков должны иметь одну и ту же формулу (в относительных ссылках)."""
        pats: Dict[str, Dict[str, List[str]]] = collections.defaultdict(lambda: collections.defaultdict(list))
        obj_rows = {r for r in range(self.last_block_row + 1, len(self.F) + 1)}
        for start, well, labs in self.blocks:
            for lab in labels:
                r = labs.get(lab)
                if not r:
                    continue
                for c, _d, _m in self.day_cols:
                    fm = self.f(r, c)
                    if fm is None:
                        continue
                    if isinstance(fm, str) and fm.startswith("="):
                        def sub(m, r=r, c=c):
                            cc, rr = ci(m.group(2)), int(m.group(4))
                            if rr in obj_rows:
                                return "c[%+d]<строка ПХГ %d>" % (cc - c, rr)
                            return ("C%d" % cc if m.group(1) else "c[%+d]" % (cc - c)) + ("R%d" % rr if m.group(3) else "r[%+d]" % (rr - r))
                        key = REF.sub(sub, fm)
                    else:
                        key = "КОНСТАНТА" if _num(fm) is not None else "ТЕКСТ"
                    pats[lab][key].append("%s%d" % (gl(c), r))
        for lab, d in pats.items():
            ordered = sorted(d.items(), key=lambda x: -len(x[1]))
            # строки ПХГ-ссылок у каждой скважины свои: сводим к виду «ссылка на строку ПХГ»
            merged: Dict[str, List[str]] = collections.defaultdict(list)
            for k, cells in ordered:
                merged[re.sub(r"<строка ПХГ \d+>", "<строка ПХГ>", k)].extend(cells)
            main = max(merged.items(), key=lambda x: len(x[1]))[0]
            for k, cells in merged.items():
                if k == main or k.startswith("=SUM(") and main.startswith("=SUM("):
                    continue
                if len(cells) > 400:
                    continue
                self.issue("ошибка" if k != "КОНСТАНТА" else "предупреждение", "Формула «%s» отличается от основной" % lab,
                           ", ".join(cells[:6]) + (" и ещё %d" % (len(cells) - 6) if len(cells) > 6 else ""),
                           "%d яч.: %s (основная: %s)" % (len(cells), k[:120], main[:120]))
            # ссылка на строку K у разных скважин должна быть одна и та же строка ПХГ
            rows_by = collections.defaultdict(list)
            for k, cells in ordered:
                for m in re.finditer(r"<строка ПХГ (\d+)>", k):
                    rows_by[m.group(1)].extend(cells)
            if len(rows_by) > 1:
                top = max(rows_by.items(), key=lambda x: len(x[1]))[0]
                for row, cells in rows_by.items():
                    if row != top:
                        self.issue("ошибка", "Сдвинутая ссылка на коэффициент пересчёта «%s»" % lab,
                                   ", ".join(cells[:8]) + (" и ещё %d" % (len(cells) - 8) if len(cells) > 8 else ""),
                                   "ссылка на строку %s вместо основной строки %s" % (row, top))


# ---------------------------------------------------------------- пересчёт Qгаз.пром

def _k_row(sh: Shakh, label: str) -> Optional[int]:
    cnt = collections.Counter()
    if not sh.has_formulas:
        return None
    for start, well, labs in sh.blocks:
        r = labs.get(label)
        if not r:
            continue
        for c, _d, _m in sh.day_cols[::10]:
            fm = sh.f(r, c)
            if isinstance(fm, str) and fm.startswith("="):
                for m in REF.finditer(fm):
                    if int(m.group(4)) > sh.last_block_row:
                        cnt[int(m.group(4))] += 1
    return cnt.most_common(1)[0][0] if cnt else None


def recompute(sh: Shakh, df: pd.DataFrame) -> pd.DataFrame:
    """Qгаз.пром = (расчётный дебит скважины) × K, K = Qотч по ПХГ / сумма расчётных дебитов по ВСЕМ скважинам."""
    otbor = sh.kind == "отбор"
    obj = sh.object_rows()
    q_obj_row = obj["Qгаза отч"] if otbor else obj["Qгаза сут"]
    target = "q газа пром" if otbor else "Qгаза пром."
    krow = _k_row(sh, target)
    sh.k_row = krow
    if sh.has_formulas:
        sh.check_sum_coverage(krow)
    base = "qгаза (расч)" if otbor else "q газа пром"
    n = lambda x: _num(x) or 0.0
    if otbor:
        df["q расч (пересчёт)"] = [0.0141 * n(d) ** 2 * n(p) * math.sqrt(1 + 0.002253 * n(p)) / 24 * n(t)
                                    for d, p, t in zip(df["Диам. штуцера"], df["Ршт. блок"], df["t работы"])]
    else:
        df["q расч (пересчёт)"] = [n(q) / 24 * n(t) for q, t in zip(df["q (шум)"], df["t работы"])]
    sums = df.groupby("Дата")["q расч (пересчёт)"].sum()
    qobj = {}
    for c, date, _m in sh.day_cols:
        qobj[date] = n(sh.v(q_obj_row, c))
    df["Qотч по ПХГ"] = df["Дата"].map(qobj)
    df["K (пересчёт)"] = [(q / sums[d]) if sums[d] > 0 else 0.0 for q, d in zip(df["Qотч по ПХГ"], df["Дата"])]
    df["Qгаз.пром (пересчёт)"] = df["q расч (пересчёт)"] * df["K (пересчёт)"]
    date2col = {d: c for c, d, _m in sh.day_cols}
    if krow:
        df["K (в файле)"] = [sh.v(krow, date2col[d]) for d in df["Дата"]]
    df["Qгаз.пром (в файле)"] = pd.to_numeric(df[target], errors="coerce")
    diff = (df["Qгаз.пром (в файле)"].fillna(0) - df["Qгаз.пром (пересчёт)"])
    df["Расхождение"] = diff
    rel = diff.abs() / df["Qгаз.пром (пересчёт)"].abs().replace(0, float("nan"))
    bad = (diff.abs() > 0.01) & ((rel > 0.001) | rel.isna())
    # причина
    starts = {w: (s, labs) for s, w, labs in sh.blocks}
    base_f_all = pd.to_numeric(df[base], errors="coerce").fillna(0)
    isf = (lambda fm: isinstance(fm, str) and fm.startswith("=")) if sh.has_formulas else (lambda fm: True)
    own = []   # (по строке df) причина самой ячейки либо ""
    for is_b, w, d, bf, bc in zip(bad, df["Скважина"], df["Дата"], base_f_all, df["q расч (пересчёт)"]):
        labs = starts[w][1]
        col = date2col[d]
        fm_t, fm_b = sh.f(labs[target], col), sh.f(labs[base], col)
        reason = ""
        if is_b and not isf(fm_t):
            reason = "в ячейке число, а не формула"
        elif abs(bf - bc) > max(0.01, 0.001 * abs(bc)) and not isf(fm_b):
            reason = "в расчётном дебите число, а не формула"
        elif abs(bf - bc) > max(0.01, 0.001 * abs(bc)):
            reason = "расчётный дебит в файле не сходится с диаметром/давлением/часами"
        elif is_b and sh.has_formulas and isf(fm_t):
            refs = [int(m.group(4)) for m in REF.finditer(fm_t) if int(m.group(4)) > sh.last_block_row]
            if krow and refs and krow not in refs:
                reason = "ссылка на К указывает на другую строку"
        own.append(reason)
    df["_own"] = own
    day_const = df[df["_own"].isin(["в расчётном дебите число, а не формула", "расчётный дебит в файле не сходится с диаметром/давлением/часами"])] \
        .groupby("Дата")["Скважина"].agg(lambda x: ", ".join(map(str, x)))
    day_t = df[df["_own"] == "в ячейке число, а не формула"].groupby("Дата")["Скважина"].agg(lambda x: ", ".join(map(str, x)))
    qpos = df[df["q расч (пересчёт)"] > 0].groupby("Дата")["Скважина"].agg(set).to_dict()
    reasons = []
    for is_b, o_r, d in zip(bad, df["_own"], df["Дата"]):
        if not is_b:
            reasons.append(o_r)       # собственная ошибка без влияния на сумму тоже отмечается
            continue
        if o_r:
            reasons.append(o_r)
            continue
        parts = []
        if d in day_const.index:
            parts.append("в расчётных дебитах дня числа/ошибки у скв. " + day_const[d][:80])
        if d in day_t.index:
            parts.append("в Qгаз.пром дня числа вместо формул у скв. " + day_t[d][:80])
        miss = sorted(w for w in getattr(sh, "k_missing", {}).get(d, set()) if w in qpos.get(d, set()))
        if miss:
            parts.append("в сумме K не учтены скв. " + ", ".join(miss))
        reasons.append("К дня искажён: " + "; ".join(parts) if parts else "К в файле отличается от K по полному списку скважин")
    df["Причина расхождения"] = reasons
    df["Расхождение > допуска"] = bad
    df.drop(columns=["_own"], inplace=True)
    # сводка по замечаниям: одна запись на вид причины
    cat = df["Причина расхождения"].map(lambda x: x.split(":")[0])
    for why, part in df[bad | (df["Причина расхождения"] != "")].groupby(cat[bad | (df["Причина расхождения"] != "")]):
        wells_l = sorted(set(part["Скважина"]), key=str)
        ex = part.assign(a=part["Расхождение"].abs()).sort_values("a", ascending=False).iloc[0]
        sh.issue("ошибка", "Qгаз.пром в файле: " + why, "суток: %d, скважин: %d" % (part["Дата"].nunique(), len(wells_l)),
                 "ячеек %d; скважины: %s%s; наибольшее расхождение %.1f тыс.м3 (скв. %s, %s: файл %.1f, пересчёт %.1f)" % (
                     len(part), ", ".join(map(str, wells_l[:25])), " …" if len(wells_l) > 25 else "", abs(ex["Расхождение"]),
                     ex["Скважина"], ex["Дата"], ex["Qгаз.пром (в файле)"], ex["Qгаз.пром (пересчёт)"]))
    if krow:
        kf, kc = df.drop_duplicates("Дата").set_index("Дата")["K (в файле)"], df.drop_duplicates("Дата").set_index("Дата")["K (пересчёт)"]
        kd = ((pd.to_numeric(kf, errors="coerce").fillna(0) - kc).abs() > 1e-3)
        if kd.any():
            sh.issue("ошибка", "K в файле ≠ K по всем скважинам", "строка %s, дней: %d" % (krow, int(kd.sum())),
                     "Расхождения K из-за неполных сумм; пример: %s файл=%.4f пересчёт=%.4f" % (
                         kd.index[kd][0], float(pd.to_numeric(kf, errors="coerce").loc[kd.index[kd][0]] or 0), float(kc.loc[kd.index[kd][0]])))
    return df


# ---------------------------------------------------------------- вода

def read_water(path: str) -> Tuple[pd.DataFrame, pd.DataFrame]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Данные"]
    ncol = ws.max_column
    heads = [(c, _text(ws.cell(1, c).value)) for c in range(1, ncol + 1) if _text(ws.cell(1, c).value)]
    recs = []
    for i, (c0, title) in enumerate(heads):
        c1 = (heads[i + 1][0] - 1) if i + 1 < len(heads) else ncol
        names = {c: _text(ws.cell(2, c).value) for c in range(c0, c1 + 1) if _text(ws.cell(2, c).value)}
        dcol = next(c for c, nm in names.items() if nm.lower() == "дата")
        for r in range(3, ws.max_row + 1):
            d = ws.cell(r, dcol).value
            if not isinstance(d, (dt.datetime, dt.date)):
                continue
            rec = {"Сезон": title.replace("Факт", "").replace("гг.", "").strip(), "Дата": d.date() if isinstance(d, dt.datetime) else d}
            for c, nm in names.items():
                if c != dcol:
                    rec[nm] = ws.cell(r, c).value
            recs.append(rec)
    season = pd.DataFrame(recs)
    wpd = wb["Water per day"]
    rows = [(a, b) for a, b in wpd.iter_rows(min_row=1, max_col=2, values_only=True) if isinstance(a, (dt.datetime, dt.date))]
    daily = pd.DataFrame(rows, columns=["Дата", "Вынос пластовой жидкости, м3 (лист Water per day)"])
    daily["Дата"] = daily["Дата"].dt.date
    return season, daily


# ---------------------------------------------------------------- сборка базы

def _infer_year(path: str, default: int) -> int:
    m = re.search(r"(20\d\d)", os.path.basename(path))
    return int(m.group(1)) if m else default


def _norm_group(x) -> str:
    return str(x).replace(" ", "").upper()


def _label_cells(path: str, limit: int = 80) -> List[str]:
    """Подписи показателей (столбец B) первых строк первого листа — для опознания вида файла."""
    if path.lower().endswith((".xlsx", ".xlsm")):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return [_text(r[1]) if len(r) > 1 else "" for r in ws.iter_rows(min_row=1, max_row=limit, max_col=2, values_only=True)]
    import xlrd
    sh = xlrd.open_workbook(path).sheet_by_index(0)
    return [_text(sh.cell_value(r, 1)) if sh.ncols > 1 else "" for r in range(min(sh.nrows, limit))]


def classify(path: str) -> str:
    """'отбор' | 'закачка' | 'вода' | '' по содержимому файла (имя файла - только запасной признак)."""
    low = os.path.basename(path).lower()
    try:
        if path.lower().endswith((".xlsx", ".xlsm")):
            names = openpyxl.load_workbook(path, read_only=True).sheetnames
            if "Water per day" in names or ("Данные" in names and "вода" in low):
                return "вода"
        labels = " | ".join(_label_cells(path, 150)).lower()
    except Exception:
        return ""
    if "состояние" in labels and "t работы" in labels:
        if "диам. штуцера" in labels or "qгаза (расч)" in labels:
            return "отбор"
        if "q (шум)" in labels or "шабл" in labels:
            return "закачка"
        return "закачка" if "зак" in low and "отбор" not in low else "отбор"
    return ""


def expand_paths(items: List[str]) -> List[str]:
    """Файлы Excel из списка путей: файл, папка (все книги внутри) или маска."""
    import glob
    out: List[str] = []
    for it in items:
        it = (it or "").strip().strip('"')
        if not it:
            continue
        found = glob.glob(it) if any(ch in it for ch in "*?[") else [it]
        for f in found:
            if os.path.isdir(f):
                out += [os.path.join(f, n) for n in sorted(os.listdir(f))]
            else:
                out.append(f)
    seen, res = set(), []
    for f in out:
        if f.lower().endswith((".xlsx", ".xlsm", ".xls")) and not os.path.basename(f).startswith("~$") and f not in seen:
            seen.add(f)
            res.append(f)
    return res


def build(paths: List[str], out: str = "База_шахматка.xlsx", otbor_year: Optional[int] = None, zak_year: Optional[int] = None,
          xls_formulas: bool = False, log=print) -> Dict[str, object]:
    """Сводит файлы (шахматки отбора и закачки по сезонам, книги воды) в книгу Excel и базу SQLite рядом (расширение .sqlite).

    Вид каждого файла определяется по содержимому. Год первых суток шахматки берётся из имени файла (иначе из otbor_year / zak_year).
    """
    import sqlite3
    tmp = tempfile.mkdtemp()
    files = expand_paths(paths)
    parts: Dict[str, List[Tuple[Shakh, pd.DataFrame]]] = {"отбор": [], "закачка": []}
    water: List[Tuple[pd.DataFrame, pd.DataFrame]] = []
    unknown: List[str] = []
    for n, path in enumerate(files):
        kind = classify(path)
        name = os.path.basename(path)
        if not kind:
            unknown.append(name)
            log("⚠️ %s: вид файла не опознан (ни шахматка, ни книга воды), пропущен" % name)
            continue
        if kind == "вода":
            water.append(read_water(path))
            log("💧 %s: вода" % name)
            continue
        wd = os.path.join(tmp, str(n))
        os.makedirs(wd)
        year = (otbor_year if kind == "отбор" else zak_year) or _infer_year(path, dt.date.today().year)
        sh = Shakh(path, kind, year, wd, xls_formulas=xls_formulas)
        sh.label = name
        df = recompute(sh, sh.wells_table())
        if sh.has_formulas:
            sh.check_formula_patterns(["qгаза (расч)", "q газа пром"] if kind == "отбор" else ["q газа пром", "Qгаза пром.", "Qгаза отч."])
        df.insert(0, "Файл", name)
        parts[kind].append((sh, df))
        log("📄 %s (%s): %d скважин, %d суток (%s…%s)" % (name, kind, len(sh.blocks), len(sh.day_cols),
                                                       sh.day_cols[0][1] if sh.day_cols else "", sh.day_cols[-1][1] if sh.day_cols else ""))
    all_sh = [sh for lst in parts.values() for sh, _d in lst]
    # группы: таблица зон последнего файла, где скважина встречается
    zrows = []
    for sh, _d in parts["отбор"] + parts["закачка"]:
        if not sh.zone_hdr:
            continue
        zn = sh.zones()
        zf = sh.zones_from_formulas(zn) if sh.has_formulas and len(zn) else pd.DataFrame(columns=["Скважина", "Группа (по формуле)"])
        zn = zn.merge(zf.groupby("Скважина")["Группа (по формуле)"].agg(lambda x: ", ".join(sorted(set(x)))).reset_index(),
                      on="Скважина", how="left") if len(zf) else zn.assign(**{"Группа (по формуле)": None})
        zrows.append(zn.assign(Файл=sh.label))
    zn2 = pd.concat(zrows).drop_duplicates("Скважина", keep="last").reset_index(drop=True) if zrows else \
        pd.DataFrame(columns=["Скважина", "Группа", "Группа (по формуле)", "Группа итог", "Источник группы"])
    if len(zn2):
        canon = {_norm_group(part): part for g in zn2["Группа (по формуле)"].dropna() for part in str(g).split(", ")}

        def final(row):
            if isinstance(row["Группа (по формуле)"], str) and row["Группа (по формуле)"]:
                return row["Группа (по формуле)"], "формулы в таблице зон"
            return canon.get(_norm_group(row["Группа"]), row["Группа"]), "подпись в столбце зон"
        zn2[["Группа итог", "Источник группы"]] = zn2.apply(lambda r: pd.Series(final(r)), axis=1)
    gmap = dict(zip(zn2["Скважина"], zn2["Группа итог"])) if len(zn2) else {}
    res: Dict[str, pd.DataFrame] = {}
    for kind, lst in parts.items():
        if not lst:
            continue
        df = pd.concat([d for _s, d in lst], ignore_index=True)
        df.insert(2, "Группа", [gmap.get(str(w).split(" ")[0], "") for w in df["Скважина"]])
        df.rename(columns={"Строка блока": "Строка в исходнике"}, inplace=True)
        dup = df.duplicated(["Скважина", "Дата"], keep=False)
        if dup.any():
            lst[0][0].issue("предупреждение", "Одни и те же сутки в нескольких файлах", ", ".join(sorted(set(df.loc[dup, "Файл"]))),
                            "Строк с повтором «скважина + дата»: %d; оставлена запись из последнего файла" % int(dup.sum()))
            df = df.drop_duplicates(["Скважина", "Дата"], keep="last").reset_index(drop=True)
        miss = sorted(set(df.loc[df["Группа"] == "", "Скважина"]), key=str)
        if miss and gmap:
            lst[0][0].issue("предупреждение", "Нет группы", ", ".join(map(str, miss)), "Скважин без группы в таблице зон: %d" % len(miss))
        res[kind] = df
    main_cols = ["Файл", "Скважина", "Группа", "Дата", "Состояние", "t работы", "Qгаз.пром (пересчёт)", "Qгаз.пром (в файле)",
                 "Расхождение", "Причина расхождения"]

    def lead(df):
        cols = [c for c in main_cols if c in df.columns]
        return df[cols + [c for c in df.columns if c not in cols]].drop(columns=["Расхождение > допуска"])

    def objt(lst):
        frames = []
        for sh, df in lst:
            t = sh.object_table()
            spec = {"Σ расчётных дебитов": ("q расч (пересчёт)", "sum"), "Σ Qгаз.пром (пересчёт)": ("Qгаз.пром (пересчёт)", "sum"),
                    "Σ Qгаз.пром (в файле)": ("Qгаз.пром (в файле)", "sum"), "K (пересчёт)": ("K (пересчёт)", "first")}
            if "K (в файле)" in df.columns:
                spec["K (в файле)"] = ("K (в файле)", "first")
            frames.append(t.merge(df.groupby("Дата").agg(**spec).reset_index(), on="Дата", how="left").assign(Файл=sh.label))
        o = pd.concat(frames, ignore_index=True)
        return o[["Файл"] + [c for c in o.columns if c != "Файл"]]

    sheets: Dict[str, pd.DataFrame] = {}
    for kind, title in (("отбор", "Отбор"), ("закачка", "Закачка")):
        if kind in res:
            sheets[title + "_по_скважинам"] = lead(res[kind])
            sheets[title + "_ПХГ_по_суткам"] = objt(parts[kind])
    tot = []
    for kind, lst in parts.items():
        for sh, _d in lst:
            t = sh.totals_table()
            if len(t):
                tot.append(t.assign(Режим=kind, Файл=sh.label))
    totals = pd.concat(tot)[["Файл", "Режим", "Скважина", "Итог", "Показатель", "Значение"]] if tot else pd.DataFrame()
    wells = []
    all_w = sorted({w for df in res.values() for w in df["Скважина"]}, key=lambda x: (len(str(x).split(" ")[0]), str(x)))
    for w in all_w:
        rec = {"Скважина": w, "Группа": gmap.get(str(w).split(" ")[0], "")}
        for kind, df in res.items():
            part = df[df["Скважина"] == w]
            rec["Дней данных, %s" % kind] = len(part)
            rec["Дней работы (t>0), %s" % kind] = int((pd.to_numeric(part["t работы"], errors="coerce") > 0).sum())
        wells.append(rec)
    wells = pd.DataFrame(wells)
    season = daily = None
    if water:
        season = pd.concat([w[0] for w in water], ignore_index=True).drop_duplicates(["Сезон", "Дата"], keep="last")
        daily = pd.concat([w[1] for w in water], ignore_index=True).drop_duplicates("Дата", keep="last").sort_values("Дата")
    issues = pd.DataFrame([dict(i, Файл=getattr(sh, "label", i["Файл"])) for sh in all_sh for i in sh.issues])
    if unknown:
        issues = pd.concat([issues, pd.DataFrame([{"Файл": n, "Уровень": "предупреждение", "Вид": "Файл не опознан", "Где": "",
                                                    "Описание": "Не шахматка отбора/закачки и не книга воды: пропущен"} for n in unknown])], ignore_index=True)
    if len(issues):
        issues = issues.sort_values(["Файл", "Уровень", "Вид"]).reset_index(drop=True)
    stat = []
    for kind, df in res.items():
        d = df.groupby("Дата").agg(p=("Qгаз.пром (пересчёт)", "sum"), f=("Qгаз.пром (в файле)", "sum"))
        stat.append([kind, len(parts[kind]), str(df["Дата"].min()), str(df["Дата"].max()), len(d), round(float(d.p.sum())),
                     round(float(d.f.sum())), int(df["Причина расхождения"].ne("").sum()),
                     int(df.loc[df["Причина расхождения"] != "", "Скважина"].nunique())])
    summary = pd.DataFrame(stat, columns=["Режим", "Файлов", "Первая дата", "Последняя дата", "Суток", "Σ Qгаз.пром пересчёт, тыс.м3",
                                          "Σ Qгаз.пром в файле, тыс.м3", "Ячеек с расхождением", "Скважин затронуто"])
    flist = [[k, sh.label, len(sh.blocks), len(sh.day_cols), sh.k_row] for k, lst in parts.items() for sh, _d in lst]
    flist += [["вода", os.path.basename(p_), "", "", ""] for p_ in files if classify(p_) == "вода"] if water else []
    files_df = pd.DataFrame(flist, columns=["Данные", "Файл", "Блоков скважин", "Суточных столбцов", "Строка К"])
    rules = pd.DataFrame({"Правило": [
        "Берём только Qгаз.пром. (основной показатель); остальные показатели шахматок сохранены рядом, ничего не отброшено.",
        "Отбор: q расч = 0,0141·d²·P·√(1+0,002253·P)/24·t (d – диаметр штуцера, мм; P – Ршт блока; t – часы). Закачка: q = q(шум)/24·t.",
        "K = Qотч по ПХГ (за сутки) / сумма расчётных дебитов ВСЕХ скважин. Qгаз.пром скважины = q расч × K. Каждый файл (сезон) считается отдельно.",
        "Дата столбца определяется рядом номеров суток, а не подписью месяца (подписи в файлах бывают неверными). Столбцы «Итого» читаются в отдельный лист.",
        "Группы берутся из формул в таблице зон; если скважины в формулах нет, берётся подпись зоны. Если таблицы зон в нескольких файлах разные, берётся из последнего.",
        "Вода: каждому сезону свой набор столбцов (подпись сезона в первой строке); значения формул берутся из сохранённых результатов."]})
    if os.path.dirname(out):
        os.makedirs(os.path.dirname(out), exist_ok=True)
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Сводка", index=False)
        rules.to_excel(xw, sheet_name="Сводка", index=False, startrow=len(summary) + 3)
        files_df.to_excel(xw, sheet_name="Файлы", index=False)
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
        if len(totals):
            totals.to_excel(xw, sheet_name="Итоги_месяцев", index=False)
        if len(zn2):
            zn2[["Скважина", "Группа итог", "Источник группы", "Группа (по формуле)", "Группа", "Файл"]].rename(
                columns={"Группа итог": "Группа", "Группа": "Подпись зоны в таблице", "Файл": "Файл таблицы зон"}).to_excel(
                xw, sheet_name="Группы", index=False)
        wells.to_excel(xw, sheet_name="Скважины", index=False)
        if season is not None:
            season.to_excel(xw, sheet_name="Вода_по_сезонам", index=False)
            daily.to_excel(xw, sheet_name="Вода_по_суткам", index=False)
        (issues if len(issues) else pd.DataFrame({"Заметка": ["Замечаний к формулам нет"]})).to_excel(xw, sheet_name="Контроль_формул", index=False)
        for ws in xw.book.worksheets:
            if ws.title != "Сводка":
                ws.freeze_panes = "A2"
            for col in ws.columns:
                w = max((len(str(c.value)) for c in list(col)[:60] if c.value is not None), default=8)
                ws.column_dimensions[col[0].column_letter].width = min(max(10, w + 2), 60)
    dbp = os.path.splitext(out)[0] + ".sqlite"
    if os.path.exists(dbp):
        os.remove(dbp)
    con = sqlite3.connect(dbp)
    tables = dict(sheets)
    for name, df in (("Итоги_месяцев", totals), ("Группы", zn2), ("Скважины", wells), ("Вода_по_сезонам", season),
                     ("Вода_по_суткам", daily), ("Контроль_формул", issues)):
        if df is not None and len(df):
            tables[name] = df
    for name, df in tables.items():
        d = df.copy()
        seen_c: Dict[str, int] = {}
        newc = []
        for c in d.columns:
            k = str(c).lower()
            seen_c[k] = seen_c.get(k, 0) + 1
            newc.append(c if seen_c[k] == 1 else "%s (%d)" % (c, seen_c[k]))
        d.columns = newc
        for c in d.columns:
            if d[c].dtype == object:
                d[c] = d[c].map(lambda x: x.isoformat() if isinstance(x, (dt.date, dt.datetime)) else
                                (x if x is None or isinstance(x, (int, float)) else str(x)))
        d.to_sql(name, con, index=False, if_exists="replace")
    con.close()
    return {"out": out, "sqlite": dbp, "summary": summary, "issues": issues, "wells": len(wells), "files": len(files) - len(unknown)}
