"""Выгрузка общей таблицы (TableView) в XLSX/CSV: ``POST /api/table/export`` {table, format, name}.

Таблица — как в контракте Атласа: ``columns`` (key, label, unit, decimals, kind) и ``rows`` по столбцам
(``rows[столбец][строка]``). Подключается в любом приложении: ``*table_export.routes()`` в списке маршрутов Starlette.
"""
from __future__ import annotations

import csv
import io
from datetime import datetime
from typing import Any
from urllib.parse import quote

from starlette.requests import Request
from starlette.responses import JSONResponse, Response
from starlette.routing import Route

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _label(col: dict) -> str:
    unit = col.get("unit") or ""
    return f'{col.get("label", col.get("key", ""))}, {unit}' if unit else str(col.get("label", col.get("key", "")))


def _date(v: Any):
    if not isinstance(v, str) or not v:
        return v
    try:
        return datetime.fromisoformat(v.replace("Z", "+00:00")).replace(tzinfo=None)
    except ValueError:
        return v


def _rows(table: dict) -> list[list[Any]]:
    cols = table.get("columns") or []
    data = table.get("rows") or []
    n = max((len(c) for c in data), default=0)
    out = []
    for i in range(n):
        row = []
        for k, col in enumerate(cols):
            v = data[k][i] if k < len(data) and i < len(data[k]) else None
            row.append(_date(v) if col.get("kind") == "date" else v)
        out.append(row)
    return out


def to_csv(table: dict) -> bytes:
    """CSV для Excel на русской локали: разделитель «;», запятая в числах, BOM."""
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    w.writerow([_label(c) for c in table.get("columns") or []])
    for row in _rows(table):
        w.writerow(["" if v is None else v.strftime("%d.%m.%Y") if isinstance(v, datetime) else str(v).replace(".", ",") if isinstance(v, float) else v for v in row])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(table: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = (str(table.get("title") or "Таблица"))[:31].translate({ord(c): "_" for c in "[]:*?/\\"}) or "Таблица"
    cols = table.get("columns") or []
    ws.append([_label(c) for c in cols])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in _rows(table):
        ws.append(row)
    for k, col in enumerate(cols, start=1):
        dec = col.get("decimals")
        fmt = {"date": "DD.MM.YYYY", "number": "0" if dec == 0 else "0." + "0" * dec if isinstance(dec, int) and dec > 0 else None}.get(col.get("kind"))
        width = max([len(_label(col))] + [len(str(v)) for v in (table.get("rows") or [[]] * k)[k - 1][:200] if v is not None] or [8])
        ws.column_dimensions[get_column_letter(k)].width = min(max(width + 2, 8), 48)
        if fmt:
            for r in range(2, ws.max_row + 1):
                ws.cell(r, k).number_format = fmt
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


async def export(request: Request):
    try:
        body = await request.json()
    except ValueError:
        body = None
    table = body.get("table") if isinstance(body, dict) else None
    if not isinstance(table, dict) or not isinstance(table.get("columns"), list):
        return JSONResponse({"error": "Ожидается JSON с полем table (columns, rows)."}, status_code=400)
    fmt = body.get("format", "xlsx")
    if fmt not in ("xlsx", "csv"):
        return JSONResponse({"error": "Формат: xlsx или csv."}, status_code=400)
    name = (str(body.get("name") or table.get("title") or "таблица")) + "." + fmt
    data = to_xlsx(table) if fmt == "xlsx" else to_csv(table)
    return Response(data, media_type=XLSX if fmt == "xlsx" else "text/csv; charset=utf-8",
                    headers={"Content-Disposition": "attachment; filename*=UTF-8''" + quote(name)})


def routes() -> list:
    return [Route("/api/table/export", export, methods=["POST"])]
