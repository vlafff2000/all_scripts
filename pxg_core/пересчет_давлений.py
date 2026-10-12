"""Пересчёт пластовых давлений на верх перфораций (логика Excel-файла «ФАЙЛ_СО_ВСЕМИ_ДАВЛЕНИЯМИ»).

Вход: строки замеров (скважина, дата, Рпл приведённое, Руст, избыточное давление, уровень жидкости), справочники
«Альтитуды» (Z, м) и «Перфорации» (абсолютная отметка верха перфораций). Глубина приведения 670 м одна для всех скважин.

Плотность воды ρ (г/см³), как в столбце I Excel:
  есть уровень E:               ρ = 10·D / (E + 670 + Z)
  нет уровня, есть избыточное G: ρ = 10·(D − G) / (670 + Z)
  есть Руст (газовая скважина):  ρ не считается
Давление на верх перфораций J (столбец J Excel):
  скважина из списка исключений или есть Руст:  J = D (без пересчёта)
  есть уровень и ρ:                             J = (E + Z + отметка)·ρ/10
  есть избыточное давление и ρ:                 J = (Z + отметка)·ρ/10 + G
K = J·0,980665 (кгс/см² в бар), плюс 1,01325 по флажку «+1 атм».
Excel при отсутствии отметки перфорации молча считал отметку нулём; здесь такая строка остаётся без J и помечается.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from typing import Dict, Iterable, List, Optional, Tuple

import pandas as pd

DEFAULT_EXCLUDED = (8, 11, 110, 122, 447)
DEPTH = 670.0
KGF_TO_BAR = 0.980665
ATM = 1.01325

# Способ расчёта (столбец «Способ расчёта»)
M_EXCLUDED = "исключённая скважина, без пересчёта"
M_GAS = "есть Руст (газовая), без пересчёта"
M_LEVEL = "по уровню жидкости"
M_EXCESS = "по избыточному давлению"
M_NO_ALT = "нет альтитуды"
M_NO_PERF = "нет отметки перфорации"
M_NO_DATA = "нет данных"
M_ZERO = "расчёт невозможен (деление на ноль)"


def _num(x) -> Optional[float]:
    if x is None or isinstance(x, bool):
        return None
    try:
        v = float(x)
    except (TypeError, ValueError):
        return None
    return None if v != v else v


@dataclass
class Refs:
    altitude: Dict[int, float] = field(default_factory=dict)
    perf_mark: Dict[int, float] = field(default_factory=dict)


def read_refs(path: str) -> Refs:
    """Альтитуды и абсолютные отметки верха перфораций из книги со справочниками (листы по названию «Альтитуд…», «Перфорац…»)."""
    import openpyxl
    refs = Refs()
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        for ws in wb.worksheets:
            low = ws.title.lower()
            if "альтитуд" in low:
                for row in ws.iter_rows(min_row=2, values_only=True):
                    w, z = _num(row[0]), _num(row[1]) if len(row) > 1 else None
                    if w is not None and z is not None:
                        refs.altitude[int(w)] = z
            elif "перфорац" in low:
                rows = list(ws.iter_rows(values_only=True))
                head = [str(c or "").lower() for c in (rows[0] if rows else ())]
                mark_col = next((i for i, h in enumerate(head) if "абсолютн" in h), None)
                pair: Dict[int, float] = {}
                for row in rows[1:]:
                    w = _num(row[0]) if row else None
                    if w is None:
                        continue
                    if mark_col is not None and mark_col < len(row) and _num(row[mark_col]) is not None:
                        refs.perf_mark[int(w)] = _num(row[mark_col])
                    elif len(row) > 1 and _num(row[1]) is not None and all(c in (None, "") for c in row[2:]):
                        pair[int(w)] = _num(row[1])   # нижняя таблица «скважина — отметка»
                for w, m in pair.items():
                    refs.perf_mark.setdefault(w, m)
    finally:
        wb.close()
    return refs


def read_average(path: str) -> pd.DataFrame:
    """Дата и «Р среднее по ПХГ» из книги с исходниками по эксплуатационным скважинам (как в Excel: значение вводится вручную, не считается).

    Ищется лист, в первых строках которого есть столбец со словом «среднее»; даты — в первом столбце.
    """
    import datetime as dt
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            col = None
            for r in rows[:5]:
                col = next((i for i, c in enumerate(r) if "средн" in str(c or "").lower()), None)
                if col is not None:
                    break
            if col is None:
                continue
            out = [(r[0], _num(r[col])) for r in rows if r and isinstance(r[0], (dt.datetime, dt.date)) and len(r) > col and _num(r[col]) is not None]
            if out:
                df = pd.DataFrame(out, columns=["Дата", "Среднее пластовое, кгс/см2"])
                df["Дата"] = pd.to_datetime(df["Дата"])
                return df.sort_values("Дата").reset_index(drop=True)
    finally:
        wb.close()
    return pd.DataFrame(columns=["Дата", "Среднее пластовое, кгс/см2"])


def level_value(raw, sign: str) -> Optional[float]:
    """Уровень жидкости E, м: «auto» — всегда ниже устья (−|Нст|), «as_is» — как в отчёте."""
    v = _num(raw)
    if v is None:
        return None
    return -abs(v) if sign == "auto" else v


def recalc_row(well: int, d, f, g, e, refs: Refs, excluded: Iterable[int] = DEFAULT_EXCLUDED, depth: float = DEPTH,
               target: Optional[float] = None):
    """Одна строка: (ρ, J, способ). d — Рпл приведённое, f — Руст, g — избыточное, e — уровень (м, обычно отрицательный).

    target — одна глубина для всех скважин (абсолютная отметка, м; знак как у отметки перфорации); без неё берётся верх перфораций скважины."""
    d, f, g, e = _num(d), _num(f), _num(g), _num(e)
    if well in set(excluded):
        return None, d, M_EXCLUDED if d is not None else M_NO_DATA
    if d is None or (e is None and f is None and g is None):
        return None, None, M_NO_DATA
    if f is not None and f != 0 and d != 0:
        return None, d, M_GAS
    z, mark = refs.altitude.get(well), (refs.perf_mark.get(well) if target is None else target)
    if z is None:
        return None, None, M_NO_ALT
    if mark is not None:
        # Абсолютная отметка ниже уровня моря записывается со знаком минус (−600); в формуле нужна глубина ниже
        # уровня моря: знак убирается, как у уровня жидкости. Положительное число — та же глубина (как раньше).
        mark = abs(mark)
    if e is not None:
        den = e + depth + z
        rho = 10 * d / den if den else None
    elif g is not None:
        rho = 10 * (d - g) / (depth + z)
    else:
        return None, None, M_NO_DATA
    if rho is None:
        return None, None, M_ZERO
    if mark is None:
        return rho, None, M_NO_PERF
    if e is not None:
        return rho, (e + z + mark) * rho / 10, M_LEVEL
    return rho, (z + mark) * rho / 10 + g, M_EXCESS


RHO_TOLERANCE = 0.005    # г/см3: расхождение плотности воды из файла и расчётной
ES_TOLERANCE = 0.003     # расхождение ES из файла и Рпл/Руст


def target_label(target: Optional[float]) -> str:
    return "верх перфораций" if target is None else "%g м" % abs(target)


def recalc(db: pd.DataFrame, refs: Refs, excluded: Iterable[int] = DEFAULT_EXCLUDED, level_sign: str = "auto",
           add_atm: bool = False, depth: float = DEPTH, target: Optional[float] = None) -> pd.DataFrame:
    """Таблица замеров (столбцы «База давлений по замерам») -> таблица пересчёта.

    target — одна глубина пересчёта для всех скважин (абсолютная отметка, м), иначе верх перфораций каждой скважины.
    Если в базе есть столбцы «L», «ES», «Плотность воды», рядом с расчётом добавляются эти значения из файла и отметка расхождения.
    """
    excluded = tuple(excluded)
    out: List[list] = []
    has_file = any(c in db.columns for c in ("L", "ES", "Плотность воды"))
    if has_file:
        db = db.rename(columns={"Плотность воды": "Плотность_воды"})
        for c in ("L", "ES", "Плотность_воды"):
            if c not in db.columns:
                db[c] = None
    for r in db.itertuples(index=False):
        row = r._asdict() if hasattr(r, "_asdict") else {}
        well = int(row["Скважина"])
        level = level_value(row.get("Уровень_жидкости"), level_sign)
        rho, j, how = recalc_row(well, row.get("Рпл"), row.get("Руст"), row.get("Избыточное_давление"), level, refs, excluded, depth, target)
        k = None if j is None else j * KGF_TO_BAR + (ATM if add_atm else 0.0)
        rec = [well, row.get("Дата"), row.get("Горизонт"), row.get("Категория"), _num(row.get("Руст")),
               _num(row.get("Избыточное_давление")), level, _num(row.get("Рпл")), rho, j, k, how, row.get("Источник")]
        if has_file:
            rec += compare_row(rho, _num(row.get("Руст")), _num(row.get("Рпл")), _num(row.get("L")), _num(row.get("ES")),
                               _num(row.get("Плотность_воды")))
        out.append(rec)
    label = target_label(target)
    cols = ["Скважина", "Дата", "Горизонт", "Категория", "Руст, кгс/см2", "Избыточное давление, кгс/см2",
            "Уровень жидкости, м", "Рпл приведённое (на 670 м), кгс/см2", "Плотность воды, г/см3",
            "Рпл на %s, кгс/см2" % label, "Рпл на %s, бар" % label, "Способ расчёта", "Источник"]
    if has_file:
        cols += ["L из файла, м", "ES из файла", "ES расчёт (Рпл/Руст)", "Плотность из файла, г/см3", "Сверка с файлом"]
    res = pd.DataFrame(out, columns=cols)
    res.attrs["bar_col"] = "Рпл на %s, бар" % label
    return res


def compare_row(rho, p_wh, p_res, l, es, rho_file):
    """Значения L, ES, плотности из файла рядом с расчётом и пометка расхождений: [L, ES файла, ES расчёт, ρ файла, пометка]."""
    es_calc = p_res / p_wh if p_res is not None and p_wh not in (None, 0) else None
    marks = []
    if es is not None and es_calc is not None and abs(es - es_calc) > ES_TOLERANCE:
        marks.append("ES: в файле %.3f, по Рпл/Руст %.3f" % (es, es_calc))
    if rho_file is not None and rho is not None and abs(rho_file - rho) > RHO_TOLERANCE:
        marks.append("плотность: в файле %.3f, расчёт %.3f" % (rho_file, rho))
    if rho_file is not None and rho is None:
        marks.append("плотность есть в файле, расчёт её не даёт")
    return [l, es, es_calc, rho_file, "; ".join(marks)]


def normalize_db(df: pd.DataFrame) -> pd.DataFrame:
    """Таблица «Данные» базы замеров -> имена столбцов без пробелов, недостающие столбцы пустые."""
    df = df.rename(columns={"Избыточное давление": "Избыточное_давление", "Уровень жидкости": "Уровень_жидкости"})
    for col in ("Руст", "Избыточное_давление", "Рпл", "Уровень_жидкости"):
        if col not in df.columns:
            df[col] = None
    for col in ("Горизонт", "Категория", "Источник"):
        if col not in df.columns:
            df[col] = None
    for col in ("L", "ES", "Плотность воды"):
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    df["Дата"] = pd.to_datetime(df["Дата"], errors="coerce")
    return df.dropna(subset=["Скважина"])


def read_measurements(path: str) -> pd.DataFrame:
    """Лист «Данные» базы «База давлений по замерам» с именами столбцов без пробелов."""
    return normalize_db(pd.read_excel(path, sheet_name="Данные"))


def include_frames(res: pd.DataFrame, add_atm: bool = False, average: Optional[pd.DataFrame] = None) -> Dict[str, pd.DataFrame]:
    """Include-таблицы «скважина, дата, давление (бар)» как в Excel-файле.

    Среднее пластовое берётся готовым столбцом из `average` (как в Excel, без пересчёта единиц); без него считается как среднее Рпл
    эксплуатационных скважин по дате, в барах.
    """
    bar = res.attrs.get("bar_col", "Рпл на верх перфораций, бар")

    def triple(df, col):
        t = df.dropna(subset=["Дата", col])[["Скважина", "Дата", col]].sort_values(["Дата", "Скважина"])
        t.columns = ["Скважина", "Дата", "Давление, бар"]
        return t

    raw_bar = (res["Рпл приведённое (на 670 м), кгс/см2"] * KGF_TO_BAR + (ATM if add_atm else 0.0))
    res = res.assign(_raw=raw_bar)
    exploit = res[res["Категория"] == "Эксплуатационные"]
    obs = res[res["Категория"] != "Эксплуатационные"]
    if average is not None and len(average):
        mean = average.rename(columns={"Среднее пластовое, кгс/см2": "Среднее пластовое"})
    else:
        mean = (exploit.dropna(subset=["Дата", "_raw"]).groupby("Дата")["_raw"].mean().reset_index())
        mean.columns = ["Дата", "Среднее пластовое, бар (расчёт по эксплуатационным)"]
    return {
        "include_среднее_пластовое": mean,
        "include_эксплуатационные": triple(exploit, "_raw"),
        "include_набл_с_пересчётом": triple(obs, bar),
        "include_набл_без_пересчёта": triple(obs, "_raw"),
    }
