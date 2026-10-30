"""Свод скважин («Данные по скважинам», форма ПХГ-01) → шаблон конструкции для Газового атласа.

Книга: слева шапка в два яруса (группа и подпись столбца), скважина занимает блок строк с объединёнными ячейками; в блоке
несколько строк, если интервалов перфорации, колонн или пакеров больше одного. Столбцы ищутся по подписям шапки, а не по
буквам: добавленный или переставленный столбец не ломает разбор.

Результат (`export`):
  Конструкция_для_Атласа.xlsx — один лист «Конструкция» в формате «Шаблон_конструкции.xlsx» Атласа (одна строка = один
                               элемент: колонна, НКТ, пакер, фильтр, перфорация, интервал пласта);
  Свод_скважин_разбор.xlsx     — «Паспорт» (основные данные и наземная часть), «Работы» (проведённые работы с датами),
                               «Конструкция а.о.» (те же элементы в абсолютных отметках), «Сверка альтитуд» (свод и Атлас),
                               «Соответствие столбцов» (что куда попало), «Проверка» (замечания к своду).
Глубины свода — по стволу (MD) от стола ротора. Абсолютная отметка: а.о. = альтитуда + высота стола ротора − глубина
(флажок `rotor`; без него а.о. = альтитуда − глубина, как в Атласе по умолчанию). Считается по альтитуде Атласа и по альтитуде
свода рядом: какую взять, решает инженер. Нет высоты стола у скважины при включённом флажке — а.о. пуста и есть замечание.
Проверка только сообщает: спорные значения не правятся молча, а попадают в «Проверку» с причиной; решение за инженером.
"""
from __future__ import annotations

import datetime as dt
import os
import re
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

ATLAS_COLUMNS = ["Скважина", "Дата конструкции", "Элемент", "Верх интервала, м", "Низ интервала, м",
                 "Наружный диаметр, мм", "Внутренний диаметр, мм", "Комментарий"]
ATLAS_FILE = "Конструкция_для_Атласа.xlsx"
REPORT_FILE = "Свод_скважин_разбор.xlsx"

# (ключ в названии группы шапки, элемент в Атласе, вид строки)
#   string   — колонна с башмаком: верх 0, низ = «Глубина спуска»;
#   interval — «от»–«до»; point — одна глубина установки (верх = низ).
ELEMENTS: List[Tuple[str, str, str]] = [
    ("(направление)", "Направление", "string"),
    ("(кондуктор)", "Кондуктор", "string"),
    ("(тех. колонна)", "Техническая колонна", "string"),
    ("(экспл. колонна)", "Эксплуатационная колонна", "string"),
    ("(потайная колонна)", "Потайная колонна", "interval"),
    ("(нкт)", "НКТ", "string"),
    ("(пакер)", "Пакер", "point"),
    ("(клапан-отсекатель)", "Клапан-отсекатель", "point"),
    ("(ингибиторный клапан)", "Ингибиторный клапан", "point"),
    ("(фильтр)", "Фильтр", "interval"),
    ("интервал перфорации", "Интервал перфорации", "interval"),
    ("интервал пласта", "Интервал пласта", "interval"),
]
SURFACE = {"фонтанная арматура": "ФА", "колонная головка": "Колонная головка"}
CASING = ("Направление", "Кондуктор", "Техническая колонна", "Эксплуатационная колонна")
LEVELS = ("ошибка", "предупреждение", "заметка")
DATE_RE = re.compile(r"(\d{1,2}\.\d{1,2}\.\d{4})\s*[-–—]\s*(\d{1,2}\.\d{1,2}\.\d{4})")
HEADER_DATE_RE = re.compile(r"состоянию\s+на:?\s*(\d{1,2}\.\d{1,2}\.\d{4})", re.I)
WORKS_SUB = "проведенные на скважине работы"
ALT_LABEL, TABLE_LABEL = "Альтитуда, м", "Высота стола ротора, м"
ALT_TOLERANCE = 1.0      # м: как допуск сверки альтитуд в Атласе (`atlas/base_model.py: TOLERANCE`)


@dataclass
class Issue:
    level: str
    code: str
    message: str
    well: str = ""
    row: str = ""
    hint: str = ""


@dataclass
class Well:
    number: str
    section: str
    rows: List[int]                      # номера строк книги, занятых скважиной
    main: Dict[str, object] = field(default_factory=dict)
    elements: List[dict] = field(default_factory=list)
    works: List[Tuple[dt.date, dt.date]] = field(default_factory=list)


@dataclass
class Book:
    path: str
    sheet: str
    snapshot: Optional[dt.date]
    snapshot_from_header: bool
    wells: List[Well] = field(default_factory=list)
    issues: List[Issue] = field(default_factory=list)
    mapping: List[dict] = field(default_factory=list)

    def add(self, level: str, code: str, message: str, well: str = "", row="", hint: str = "") -> None:
        self.issues.append(Issue(level, code, message, well, str(row or ""), hint))


def _norm(value) -> str:
    return re.sub(r"\s+", " ", str(value if value is not None else "").replace("ё", "е").replace("Ё", "Е")).strip().casefold()


def _num(value) -> Tuple[Optional[float], bool]:
    """(число, было_ли_нечисловое): пусто → (None, False), текст, не похожий на число → (None, True)."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, False
    if isinstance(value, bool):
        return None, True
    if isinstance(value, (int, float)):
        return float(value), False
    try:
        return float(str(value).strip().replace(" ", "").replace(",", ".")), False
    except ValueError:
        return None, True


def _well_name(value) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def parse_date(text) -> Optional[dt.date]:
    if isinstance(text, dt.datetime):
        return text.date()
    if isinstance(text, dt.date):
        return text
    s = str(text or "").strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d.%m.%y", "%Y-%m-%d %H:%M:%S"):
        try:
            return dt.datetime.strptime(s[:19] if "%H" in fmt else s, fmt).date()
        except ValueError:
            continue
    return None


def _find_header(ws) -> Tuple[int, int]:
    for r in range(1, min(ws.max_row, 60) + 1):
        for c in range(1, min(ws.max_column, 12) + 1):
            if _norm(ws.cell(r, c).value) == "№ скважины":
                return r, c
    raise ValueError("В книге нет шапки с подписью «№ скважины»: это не свод скважин (форма ПХГ-01).")


def _header(ws, hr: int) -> List[Tuple[int, str, str]]:
    """[(номер столбца, нормированная группа, нормированная подпись)]: группы шапки протягиваются вправо."""
    out, group = [], ""
    for c in range(1, ws.max_column + 1):
        g = ws.cell(hr - 1, c).value
        if g is not None and str(g).strip():
            group = _norm(g)
        sub = ws.cell(hr, c).value
        if sub is not None and str(sub).strip():
            out.append((c, group, _norm(sub)))
    return out


def _snapshot(ws, hr: int, option: str) -> Tuple[Optional[dt.date], bool]:
    if option:
        return parse_date(option), False
    for r in range(1, hr):
        for c in range(1, 6):
            m = HEADER_DATE_RE.search(str(ws.cell(r, c).value or ""))
            if m:
                return parse_date(m.group(1)), True
    return None, False


def _pick(cols: List[Tuple[int, str, str]], group_key: str, *needles: str, nth: int = 0) -> Optional[int]:
    found = [c for c, g, s in cols if group_key in g and all(n in s for n in needles)]
    return found[nth] if len(found) > nth else None


def read(path: str, date: str = "", sheet: str = "") -> Book:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.active
    hr, wc = _find_header(ws)
    cols = _header(ws, hr)
    snap, from_header = _snapshot(ws, hr, date.strip())
    book = Book(path, ws.title, snap, from_header)
    if date.strip() and snap is None:
        book.add("ошибка", "DATE", "Дата конструкции «%s» не распознана" % date, hint="Формат ДД.ММ.ГГГГ")
    if snap is None and not date.strip():
        book.add("ошибка", "DATE", "В шапке нет «По состоянию на», дата не задана", hint="Укажите дату конструкции параметром")
    elif from_header and snap and snap.year <= 2000:
        book.add("предупреждение", "DATE", "Дата в шапке свода %s похожа на шаблонную" % snap.strftime("%d.%m.%Y"),
                 hint="Атлас берёт последнюю дату конструкции скважины: задайте реальную дату снимка")
    # блоки скважин
    section, current = "", None
    for r in range(hr + 1, ws.max_row + 1):
        raw = ws.cell(r, wc).value
        if raw is not None and str(raw).strip():
            current = Well(_well_name(raw), section, [r])
            book.wells.append(current)
        else:
            a = ws.cell(r, 1).value
            if a is not None and str(a).strip():
                if current is not None and all(ws.cell(r, c).value is None for c in range(2, ws.max_column + 1)):
                    current = None           # подвал или заголовок следующего раздела
                section = str(a).strip()
            elif current is not None:
                current.rows.append(r)
    seen: Dict[str, int] = {}
    for w in book.wells:
        seen[w.number] = seen.get(w.number, 0) + 1
    for number, n in seen.items():
        if n > 1:
            book.add("предупреждение", "DUP", "Номер скважины встречается %d раз(а)" % n, number,
                     hint="Строки обеих скважин попадут в шаблон под одним номером: различите их вручную")
    _fill_main(book, ws, cols, hr)
    for w in book.wells:
        _fill_elements(book, ws, cols, w)
        _fill_works(book, ws, cols, w)
        _check_well(book, w)
    book.mapping = _mapping(ws, cols)
    return book


def _first(ws, rows: List[int], c: int):
    for r in rows:
        v = ws.cell(r, c).value
        if v is not None and not (isinstance(v, str) and not v.strip()):
            return v
    return None


def _fill_main(book: Book, ws, cols, hr: int) -> None:
    element_keys = [k for k, _, _ in ELEMENTS]
    names: Dict[str, int] = {}
    for c, g, s in cols:
        if any(k in g for k in element_keys):
            continue
        prefix = next((p for key, p in SURFACE.items() if key in g), "")
        sub = ws.cell(hr, c).value
        label = "%s: %s" % (prefix, sub) if prefix else str(sub)
        label = re.sub(r"\s+", " ", label).strip()
        names[label] = c
    for w in book.wells:
        for label, c in names.items():
            v = _first(ws, w.rows, c)
            if _norm(label).endswith(("дата окончания срока службы", "дата последнего диагностирования", "дата следующего диагностирования")) \
                    and v is not None:
                d = parse_date(v)
                v = d if d else v
            w.main[label] = v


def _fill_elements(book: Book, ws, cols, w: Well) -> None:
    for key, name, kind in ELEMENTS:
        if kind == "string":
            c_d, c_z = _pick(cols, key, ", мм"), _pick(cols, key, "глубина")
            c_cem, c_p, c_thr = _pick(cols, key, "цемент"), _pick(cols, key, "опрессовки"), _pick(cols, key, "резьбы")
            used = [r for r in w.rows if any(ws.cell(r, c).value is not None for c in (c_d, c_z) if c)]
            for n, r in enumerate(used, 1):
                d, bad_d = _num(ws.cell(r, c_d).value) if c_d else (None, False)
                z, bad_z = _num(ws.cell(r, c_z).value) if c_z else (None, False)
                _bad(book, w, r, name, bad_d, bad_z)
                if z is None:
                    if d is not None or bad_z:
                        book.add("предупреждение", "SKIP", "%s: нет глубины спуска, элемент не выгружен" % name, w.number, r,
                                 "Заполните «Глубина спуска» в своде или добавьте строку в шаблон вручную")
                    continue
                parts = []
                if len(used) > 1:
                    parts.append("колонна %d из %d, верх принят 0 м" % (n, len(used)))
                cem, _ = _num(ws.cell(r, c_cem).value) if c_cem else (None, False)
                if cem is not None:
                    parts.append("подъём цементного камня %g м (как в своде)" % cem)
                p, _ = _num(ws.cell(r, c_p).value) if c_p else (None, False)
                if p is not None:
                    parts.append("Pопрессовки %g МПа" % p)
                thr = ws.cell(r, c_thr).value if c_thr else None
                if thr:
                    parts.append("резьба %s" % str(thr).strip())
                _emit(w, name, 0.0, z, d, "; ".join(parts), r)
        elif kind == "interval":
            c_d = _pick(cols, key, ", мм")
            c_a = _pick(cols, key, "от, м")
            c_b = _pick(cols, key, "до, м")
            if c_a is None or c_b is None:
                continue
            for r in w.rows:
                a, bad_a = _num(ws.cell(r, c_a).value)
                b, bad_b = _num(ws.cell(r, c_b).value)
                d, bad_d = _num(ws.cell(r, c_d).value) if c_d else (None, False)
                if a is None and b is None and not bad_a and not bad_b:
                    continue
                _bad(book, w, r, name, bad_a, bad_b, bad_d)
                if a is None or b is None:
                    book.add("предупреждение", "SKIP", "%s: задана только одна граница интервала, строка не выгружена" % name,
                             w.number, r, "Допишите вторую границу в своде")
                    continue
                note = ""
                if a > b:
                    book.add("предупреждение", "SWAP", "%s: «от» %g больше «до» %g, границы поменяны местами" % (name, a, b),
                             w.number, r, "Проверьте глубины по первоисточнику")
                    a, b, note = b, a, "границы в своде перепутаны (%g–%g)" % (a, b)
                _emit(w, name, a, b, d, note, r)
        else:
            c_z = _pick(cols, key, "глубина")
            c_t = _pick(cols, key, "тип")
            c_d = _pick(cols, key, ", мм")
            if c_z is None:
                continue
            for r in w.rows:
                z, bad_z = _num(ws.cell(r, c_z).value)
                if z is None and not bad_z:
                    continue
                d, bad_d = _num(ws.cell(r, c_d).value) if c_d else (None, False)
                _bad(book, w, r, name, bad_z, bad_d)
                if z is None:
                    continue
                t = ws.cell(r, c_t).value if c_t else None
                _emit(w, name, z, z, d, str(t).strip() if t else "", r)


def _bad(book: Book, w: Well, row: int, name: str, *flags: bool) -> None:
    if any(flags):
        book.add("предупреждение", "NUM", "%s: в числовой ячейке текст, значение пропущено" % name, w.number, row,
                 "Исправьте ячейку в своде")


def _emit(w: Well, name: str, top: float, bottom: float, d: Optional[float], comment: str, row: int) -> None:
    if d is not None and d <= 0:
        d = None
    w.elements.append({"Элемент": name, "Верх": top, "Низ": bottom, "Диаметр": d, "Комментарий": comment, "Строка": row})


def _fill_works(book: Book, ws, cols, w: Well) -> None:
    c = next((c for c, g, s in cols if WORKS_SUB in s), None)
    if c is None:
        return
    text = str(_first(ws, w.rows, c) or "")
    if not text:
        return
    seen = set()
    for a, b in DATE_RE.findall(text):
        da, db = parse_date(a), parse_date(b)
        if da is None or db is None:
            book.add("предупреждение", "DATE", "Работы: дата «%s – %s» не распознана" % (a, b), w.number, w.rows[0])
            continue
        if (da, db) in seen:
            book.add("заметка", "DUP", "Работы: период %s – %s указан дважды, взят один раз" % (a, b), w.number, w.rows[0])
            continue
        seen.add((da, db))
        if db < da:
            book.add("предупреждение", "DATE", "Работы: окончание %s раньше начала %s" % (b, a), w.number, w.rows[0])
        w.works.append((da, db))
    rest = DATE_RE.sub("", text).replace(";", "").strip()
    if rest:
        book.add("заметка", "TEXT", "В графе работ есть текст без дат: «%s»" % rest[:80], w.number, w.rows[0])


def _by(w: Well, name: str) -> List[dict]:
    return [e for e in w.elements if e["Элемент"] == name]


def _check_well(book: Book, w: Well) -> None:
    prod = _by(w, "Эксплуатационная колонна")
    if not w.elements:
        book.add("заметка", "EMPTY", "В своде нет ни одного элемента конструкции", w.number, w.rows[0])
        return
    shoe = max((e["Низ"] for e in prod), default=None)
    if shoe is None:
        book.add("предупреждение", "NOPROD", "Нет эксплуатационной колонны", w.number, w.rows[0],
                 "Для скважины без эксплуатационной колонны проверка глубин перфорации и НКТ не делается")
    else:
        for e in _by(w, "НКТ"):
            if e["Низ"] > shoe:
                book.add("предупреждение", "DEPTH", "НКТ спущены на %g м, глубже башмака эксплуатационной колонны %g м" % (e["Низ"], shoe),
                         w.number, e["Строка"], "Проверьте глубину спуска НКТ")
        for e in _by(w, "Интервал перфорации"):
            if e["Низ"] > shoe:
                book.add("предупреждение", "DEPTH", "Перфорация до %g м ниже башмака эксплуатационной колонны %g м" % (e["Низ"], shoe),
                         w.number, e["Строка"], "Возможно, открытый ствол или ошибка в глубине")
    # вложенность обсадных колонн: диаметр уменьшается, башмак уходит глубже
    chain = [(n, e) for n in CASING for e in _by(w, n)]
    singles = [(n, e) for n, e in chain if len(_by(w, n)) == 1]
    for (n1, e1), (n2, e2) in zip(singles, singles[1:]):
        if e2["Низ"] < e1["Низ"]:
            book.add("предупреждение", "ORDER", "%s (%g м) спущена выше, чем предыдущая колонна «%s» (%g м)" % (n2, e2["Низ"], n1, e1["Низ"]),
                     w.number, e2["Строка"], "Проверьте глубины спуска колонн")
        if e1["Диаметр"] and e2["Диаметр"] and e2["Диаметр"] > e1["Диаметр"]:
            book.add("предупреждение", "ORDER", "%s Ø%g мм больше предыдущей колонны «%s» Ø%g мм" % (n2, e2["Диаметр"], n1, e1["Диаметр"]),
                     w.number, e2["Строка"], "Проверьте диаметры колонн")


def _mapping(ws, cols) -> List[dict]:
    hr, _ = _find_header(ws)
    rows = []
    for c, g, s in cols:
        group = re.sub(r"\s+", " ", str(next((ws.cell(hr - 1, k).value for k in range(c, 0, -1) if ws.cell(hr - 1, k).value), "") or "")).strip()
        sub = re.sub(r"\s+", " ", str(ws.cell(hr, c).value)).strip()
        target, rule = _target(g, s)
        rows.append({"Столбец": openpyxl.utils.get_column_letter(c), "Группа свода": group, "Подпись свода": sub,
                     "Куда": target, "Правило": rule})
    return rows


def _target(group: str, sub: str) -> Tuple[str, str]:
    for key, name, kind in ELEMENTS:
        if key not in group:
            continue
        if ", мм" in sub:
            return "«%s»: наружный диаметр" % name, "внутренний диаметр в своде не задан, остаётся пустым"
        if "глубина" in sub:
            return "«%s»: %s" % (name, "низ интервала" if kind == "string" else "верх и низ (точка)"), \
                "верх колонны принят 0 м" if kind == "string" else "одна глубина: верх = низ"
        if sub.endswith("от, м"):
            return "«%s»: верх интервала" % name, ""
        if sub.endswith("до, м"):
            return "«%s»: низ интервала" % name, "если «от» больше «до», границы меняются местами, замечание в «Проверке»"
        if "цемент" in sub:
            return "«%s»: комментарий" % name, "подъём цементного камня записан как в своде, без пересчёта"
        if "опрессовки" in sub or "резьбы" in sub or sub == "тип":
            return "«%s»: комментарий" % name, ""
        return "не используется", ""
    if any(k in group for k in SURFACE):
        return "Паспорт (наземная часть)", "в шаблон конструкции не идёт: Атлас не рисует устье"
    if WORKS_SUB in sub:
        return "лист «Работы»", "периоды «дд.мм.гггг - дд.мм.гггг» разобраны отдельными строками"
    return "Паспорт", "в шаблон конструкции не идёт"


# ---------------------------------------------------------------- альтитуды и а.о.

def _key(number) -> str:
    """Номер скважины как в Атласе: «31», «31.0» и 31 — одна скважина."""
    text = str(number).strip()
    try:
        return str(int(float(text)))
    except ValueError:
        return text


def atlas_altitudes(storage: str = "", project: str = "") -> Tuple[Dict[str, float], str]:
    """Альтитуды проекта Атласа (таблица `well_geo`): ({скважина: м}, подпись источника). Проект не задан — самый свежий с расходами."""
    from gsp_maps import atlas_source as src
    root = src.storage_root(storage)
    projects = src.list_projects(root)
    pid = project.strip() or src.default_project(projects)
    if root is None or not pid:
        raise ValueError("Проект Атласа не найден: укажите папку хранилища (storage) и проект.")
    data = src.AtlasData(root, pid)
    _, alt, _ = data.geo()
    return {_key(w): float(a) for w, a in alt.items()}, "проект Атласа «%s»" % data.name


def file_altitudes(path: str) -> Tuple[Dict[str, float], str]:
    """Альтитуды из файла Атласа/Карт «Глубины_перфораций.xlsx» (лист «альтитуды»)."""
    from gsp_maps.inputs import read_perforations
    _, alt = read_perforations(path)
    return {_key(w): float(a) for w, a in alt.items()}, "файл «%s»" % os.path.basename(path)


def _svod_alt(w: Well, label: str):
    v, _ = _num(w.main.get(label))
    return v


def compare_altitudes(book: Book, atlas: Optional[Dict[str, float]], source: str = "") -> pd.DataFrame:
    """Сверка альтитуд свода с Атласом; расхождение больше допуска — в «Проверку». Ничего не исправляет."""
    rows = []
    for w in book.wells:
        a_svod, a_atlas = _svod_alt(w, ALT_LABEL), (atlas or {}).get(_key(w.number))
        table_h = _svod_alt(w, TABLE_LABEL)
        if atlas is None:
            status = "Атлас не задан"
        elif a_svod is None and a_atlas is None:
            status = "нет нигде"
        elif a_svod is None:
            status = "нет в своде"
        elif a_atlas is None:
            status = "нет в Атласе"
        else:
            status = "совпадает" if abs(a_svod - a_atlas) <= ALT_TOLERANCE else "расходится"
        diff = round(a_svod - a_atlas, 2) if a_svod is not None and a_atlas is not None else None
        rows.append([w.number, a_svod, a_atlas, diff, table_h, status])
        if status == "расходится":
            book.add("предупреждение", "ALT", "Альтитуда свода %g м, в Атласе %g м (разница %+g м)" % (a_svod, a_atlas, diff), w.number,
                     w.rows[0], "а.о. в таблице «Конструкция а.о.» посчитана по обеим; решение за инженером")
        elif status in ("нет в Атласе", "нет в своде", "нет нигде"):
            book.add("заметка", "ALT", "Альтитуда: %s" % status, w.number, w.rows[0])
    return pd.DataFrame(rows, columns=["Скважина", "Альтитуда свода, м", "Альтитуда Атласа, м", "Свод − Атлас, м",
                                       "Высота стола ротора (свод), м", "Статус"])


def ao_table(book: Book, atlas: Optional[Dict[str, float]] = None, rotor: bool = True) -> pd.DataFrame:
    """Элементы конструкции в абсолютных отметках: а.о. = альтитуда [+ высота стола ротора] − глубина; две альтитуды рядом."""
    rows = []
    for w in book.wells:
        a_svod, a_atlas = _svod_alt(w, ALT_LABEL), (atlas or {}).get(_key(w.number))
        table_h = _svod_alt(w, TABLE_LABEL)
        if rotor and table_h is None and w.elements:
            book.add("предупреждение", "ALT", "Нет высоты стола ротора: а.о. не посчитана (флажок «учитывать высоту стола ротора»)",
                     w.number, w.rows[0], "Заполните «Высота стола ротора» в своде или снимите флажок")

        def ao(depth, alt):
            if alt is None or (rotor and table_h is None):
                return None
            return round(alt + (table_h if rotor else 0.0) - depth, 2)
        for e in w.elements:
            rows.append([w.number, e["Элемент"], e["Верх"], e["Низ"], a_atlas, ao(e["Верх"], a_atlas), ao(e["Низ"], a_atlas),
                         a_svod, ao(e["Верх"], a_svod), ao(e["Низ"], a_svod)])
    return pd.DataFrame(rows, columns=["Скважина", "Элемент", "Верх (MD), м", "Низ (MD), м", "Альтитуда Атласа, м",
                                       "Верх а.о. (Атлас), м", "Низ а.о. (Атлас), м", "Альтитуда свода, м",
                                       "Верх а.о. (свод), м", "Низ а.о. (свод), м"])


# ---------------------------------------------------------------- выгрузка

def atlas_table(book: Book) -> pd.DataFrame:
    rows = []
    for w in book.wells:
        for e in w.elements:
            rows.append([w.number, book.snapshot, e["Элемент"], e["Верх"], e["Низ"], e["Диаметр"], None, e["Комментарий"]])
    return pd.DataFrame(rows, columns=ATLAS_COLUMNS)


def passport_table(book: Book) -> pd.DataFrame:
    rows = []
    for w in book.wells:
        row = {"Раздел свода": w.section}
        row.update(w.main)
        row["Работ в графе"] = len(w.works)
        rows.append(row)
    table = pd.DataFrame(rows)
    for name in table.columns:      # годы и целые числа без «.0» после выравнивания пустых ячеек
        col = pd.to_numeric(table[name], errors="coerce")
        if table[name].notna().sum() and col.notna().sum() == table[name].notna().sum() and (col.dropna() % 1 == 0).all():
            table[name] = col.astype("Int64")
    return table


def works_table(book: Book) -> pd.DataFrame:
    rows = [[w.number, a, b, (b - a).days + 1] for w in book.wells for a, b in w.works]
    return pd.DataFrame(rows, columns=["Скважина", "Начало", "Окончание", "Дней"])


def issues_table(book: Book) -> pd.DataFrame:
    order = {lv: i for i, lv in enumerate(LEVELS)}
    rows = sorted(book.issues, key=lambda i: (order[i.level], i.code, i.well))
    return pd.DataFrame([[i.level, i.code, i.message, i.well, i.row, i.hint] for i in rows],
                        columns=["Уровень", "Вид", "Что не так", "Скважина", "Строка свода", "Что сделать"])


def _style(ws, widths: Optional[Dict[str, int]] = None) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="16746C")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 34
    for column in ws.columns:
        name = column[0].column_letter
        width = max([len(str(c.value)) for c in column[:60] if c.value is not None] or [10])
        ws.column_dimensions[name].width = (widths or {}).get(name, max(12, min(45, width + 2)))
        for cell in column[1:]:
            if isinstance(cell.value, (dt.date, dt.datetime)):
                cell.number_format = "DD.MM.YYYY"


def export(book: Book, out_dir: str, atlas_alt: Optional[Dict[str, float]] = None, atlas_source: str = "", rotor: bool = True) -> List[str]:
    os.makedirs(out_dir, exist_ok=True)
    atlas = os.path.join(out_dir, ATLAS_FILE)
    report = os.path.join(out_dir, REPORT_FILE)
    with pd.ExcelWriter(atlas, engine="openpyxl", date_format="DD.MM.YYYY") as xw:
        atlas_table(book).to_excel(xw, sheet_name="Конструкция", index=False)
        _style(xw.sheets["Конструкция"])
    with pd.ExcelWriter(report, engine="openpyxl", date_format="DD.MM.YYYY") as xw:
        altitudes = compare_altitudes(book, atlas_alt, atlas_source)     # раньше «Проверки»: добавляет замечания о расхождении
        for name, table in (("Паспорт", passport_table(book)), ("Работы", works_table(book)),
                            ("Конструкция а.о.", ao_table(book, atlas_alt, rotor)), ("Сверка альтитуд", altitudes),
                            ("Соответствие столбцов", pd.DataFrame(book.mapping)), ("Проверка", issues_table(book))):
            table.to_excel(xw, sheet_name=name, index=False)
            _style(xw.sheets[name])
    return [atlas, report]
