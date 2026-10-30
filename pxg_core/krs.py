"""Таблица «Анализ ревизии КРУ»: проведённые КР, песок по ревизиям, спец. ГДИ, потребность и приоритет КР.

Шапка ищется по тексту («№ скв», «Выполнение КР 2023», «Ревизия КРУ после отбора 2024-2025» …), поэтому новые сезоны, добавленные
справа, читаются без правки кода. Типы ремонтов берутся из легенды самой таблицы (заливка ячейки → подпись в соседнем столбце),
а где цвет не найден в легенде, тип определяется по кодам работ (КР 1.094 — замена фильтра, КР 2.013 — закрепление ПЗП).

Приоритет КР (1 — срочно) выведен из «Потребности» самой таблицы, а не задан формулой: первая ревизия КРУ с песком после
ремонта даёт 3, каждая следующая — на единицу ниже (3 → 2 → 1); год без нового песка число не меняет; промывка и СКО счёт
не сбрасывают, замена фильтра и закрепление ПЗП сбрасывают. Это подсказка с видимой причиной: инженер переписывает её вручную
(столбец «Приоритет вручную» в результате не затирается при пересчёте).
"""
from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

import pandas as pd

RGB = Tuple[int, int, int]

KIND_FILTER = "Замена фильтра"
KIND_CT = "Колтюбинг без СКО"
KIND_CTS = "Колтюбинг с СКО / СКО"
KIND_ARM = "Армирование ПЗП"
KIND_OTHER = "Тип не определён"
RESETTING = (KIND_FILTER, KIND_ARM)
KINDS = (KIND_FILTER, KIND_CT, KIND_CTS, KIND_ARM, KIND_OTHER)

# заливки легенды в файле, если легенда не найдена (палитра Excel 97–2003)
DEFAULT_LEGEND: Dict[RGB, str] = {
    (255, 204, 0): KIND_FILTER, (0, 204, 255): KIND_CT, (255, 255, 0): KIND_CTS, (0, 102, 204): KIND_ARM,
}
SAND_AN_FILL: RGB = (159, 159, 159)      # «Х» серым: песок по отчётам авторского надзора
SHLEIF_FILL: RGB = (255, 128, 128)       # «х» розовым: вероятно вынос остатков песка из шлейфа

RELAPSE_NEW = "new"       # после замены фильтра счёт идёт заново с 3 (по умолчанию)
RELAPSE_TWO = "two"       # после замены фильтра повтор песка сразу даёт 2 (так записана скв. 56)

CODE_RE = re.compile(r"(КР|ИС)\s*(\d+(?:\.\d+)?)\.?\s*[-–]?\s*(ГТ|ПА|ЗР|КТ)?")


@dataclass
class Grid:
    values: List[List[object]]
    fills: List[List[Optional[RGB]]]
    sheet: str = ""

    @property
    def nrows(self) -> int:
        return len(self.values)

    @property
    def ncols(self) -> int:
        return max((len(r) for r in self.values), default=0)

    def value(self, r: int, c: int):
        row = self.values[r]
        return row[c] if c < len(row) else ""

    def fill(self, r: int, c: int) -> Optional[RGB]:
        row = self.fills[r]
        return row[c] if c < len(row) else None


def read_grid(path: str, sheet: Optional[str] = None) -> Grid:
    """Значения и заливки ячеек листа из .xls (xlrd) или .xlsx (openpyxl). Лист по умолчанию — первый."""
    if path.lower().endswith(".xls"):
        import xlrd
        book = xlrd.open_workbook(path, formatting_info=True)
        sh = book.sheet_by_name(sheet) if sheet else book.sheet_by_index(0)
        values, fills = [], []
        for r in range(sh.nrows):
            vr, fr = [], []
            for c in range(sh.ncols):
                vr.append(sh.cell_value(r, c))
                xf = book.xf_list[sh.cell_xf_index(r, c)]
                color = book.colour_map.get(xf.background.pattern_colour_index) if xf.background.fill_pattern == 1 else None
                fr.append(tuple(color) if color else None)
            values.append(vr)
            fills.append(fr)
        return Grid(values, fills, sh.name)
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    values, fills = [], []
    for row in ws.iter_rows():
        vr, fr = [], []
        for cell in row:
            vr.append("" if cell.value is None else cell.value)
            rgb = None
            if cell.fill is not None and cell.fill.fill_type == "solid":
                raw = cell.fill.fgColor.rgb
                if isinstance(raw, str) and len(raw) >= 6 and not cell.fill.fgColor.type == "theme":
                    h = raw[-6:]
                    try:
                        rgb = (int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))
                    except ValueError:
                        rgb = None
            fr.append(rgb)
        values.append(vr)
        fills.append(fr)
    return Grid(values, fills, ws.title)


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def _is_mark(v) -> bool:
    return _text(v).lower() in ("х", "x")


def well_number(v) -> Optional[int]:
    m = re.match(r"\s*(\d+)", _text(v))
    return int(m.group(1)) if m else None


def group_label(v) -> str:
    t = _text(v)
    return "ГСП-%s" % t if re.fullmatch(r"\d+", t) else t


@dataclass
class Columns:
    header_row: int
    well: int
    gsp: Optional[int] = None
    works: Dict[int, int] = field(default_factory=dict)       # год → столбец «Выполнение КР»
    rev: Dict[str, int] = field(default_factory=dict)         # сезон «2024-2025» → столбец «Ревизия КРУ»
    sand: Dict[str, int] = field(default_factory=dict)        # сезон → столбец «Песок после отбора» (отчёты АН)
    need: Dict[int, int] = field(default_factory=dict)        # год → «Потребность в КР на конец года»
    gdi: Dict[str, int] = field(default_factory=dict)         # сезон → «Спец ГДИ в отбор»
    plan: Dict[int, int] = field(default_factory=dict)        # год → «План КР»
    required: Optional[int] = None                            # «Требуется КР (значимость 1-3)»


def find_columns(g: Grid) -> Columns:
    hdr = None
    for r in range(min(g.nrows, 30)):
        for c in range(g.ncols):
            if re.fullmatch(r"№\s*скв\.?", _text(g.value(r, c)).lower()):
                hdr = (r, c)
                break
        if hdr:
            break
    if hdr is None:
        raise ValueError("Не найден столбец «№ скв» в первых 30 строках листа «%s»" % g.sheet)
    cols = Columns(header_row=hdr[0], well=hdr[1])
    for c in range(g.ncols):
        h = re.sub(r"\s+", " ", _text(g.value(hdr[0], c)))
        low = h.lower()
        y = re.findall(r"20\d\d", h)
        season = "%s-%s" % (y[0], y[1]) if len(y) > 1 else ""
        if re.fullmatch(r"№ гсп", low) and cols.gsp is None:
            cols.gsp = c
        elif low.startswith("выполнение кр") and y:
            cols.works[int(y[0])] = c
        elif low.startswith("ревизия кру") and season:
            cols.rev[season] = c
        elif low.startswith("песок после отбора") and season:
            cols.sand[season] = c
        elif low.startswith("потребность в кр") and y:
            cols.need[int(y[0])] = c
        elif low.startswith("спец гди") and season:
            cols.gdi[season] = c
        elif low.startswith("план кр") and y:
            cols.plan[int(y[0])] = c
        elif low.startswith("требуется кр"):
            cols.required = c
    if not cols.works and not cols.rev:
        raise ValueError("В шапке нет столбцов «Выполнение КР …» и «Ревизия КРУ после отбора …»")
    return cols


def read_legend(g: Grid, cols: Columns) -> Tuple[Dict[RGB, str], Dict[RGB, str]]:
    """Легенда таблицы: ячейка с заливкой справа от данных + подпись в соседнем столбце. → (цвет → тип ремонта, цвет → подпись)."""
    start = max([c for c in (cols.required, *cols.plan.values()) if c is not None] + [0]) + 1
    legend: Dict[RGB, str] = {}
    captions: Dict[RGB, str] = {}
    for r in range(cols.header_row + 1, g.nrows):
        for c in range(start, g.ncols - 1):
            color = g.fill(r, c)
            cap = _text(g.value(r, c + 1))
            if not color or not cap or _text(g.value(r, c)):
                continue
            low = cap.lower()
            kind = None
            if "фильтр" in low:
                kind = KIND_FILTER
            elif "армирован" in low:
                kind = KIND_ARM
            elif "колтюбинг" in low or "ско" in low:
                kind = KIND_CT if ("без ско" in low and "с ско" not in low.replace("без ско", "")) else KIND_CTS
            if kind:
                legend[color] = kind
                captions[color] = cap
    return legend, captions


def split_works(text: str) -> List[Tuple[str, str]]:
    """Все коды работ в ячейке (разделитель в таблице «;» или «.», поэтому ищутся по коду): [(код, фрагмент текста до следующего кода)]."""
    text = text or ""
    hits = list(CODE_RE.finditer(text))
    out = []
    for i, m in enumerate(hits):
        end = hits[i + 1].start() if i + 1 < len(hits) else len(text)
        code = "%s %s" % (m.group(1), m.group(2)) + (" " + m.group(3) if m.group(3) else "")
        out.append((code, text[m.start():end].strip(" ;.\n")))
    return out


def kind_by_codes(codes: Iterable[str]) -> str:
    nums = {c.split(" ")[1] for c in codes if " " in c}
    if "1.094" in nums:
        return KIND_FILTER
    if "2.013" in nums:
        return KIND_ARM
    return KIND_OTHER


@dataclass
class KrsData:
    works: pd.DataFrame
    revisions: pd.DataFrame
    gdi: pd.DataFrame
    wells: pd.DataFrame
    table_need: pd.DataFrame
    notes: List[Tuple[str, str]]          # (уровень, текст)
    legend: Dict[RGB, str]
    columns: Columns


def parse(path: str, sheet: Optional[str] = None) -> Tuple[KrsData, Grid]:
    g = read_grid(path, sheet)
    cols = find_columns(g)
    legend, _ = read_legend(g, cols)
    notes: List[Tuple[str, str]] = []
    if not legend:
        legend = dict(DEFAULT_LEGEND)
        notes.append(("заметка", "Легенда цветов в таблице не найдена, взяты цвета по умолчанию (оранжевая — замена фильтра, голубая — колтюбинг без СКО, жёлтая — колтюбинг с СКО, синяя — армирование ПЗП)"))
    works_rows, rev_rows, gdi_rows, well_rows, need_rows = [], [], [], [], []
    seen: Dict[int, int] = {}
    for r in range(cols.header_row + 1, g.nrows):
        w = well_number(g.value(r, cols.well))
        if w is None:
            continue
        if w in seen:
            notes.append(("предупреждение", "Скважина %d повторена в строках %d и %d: берётся первая" % (w, seen[w] + 1, r + 1)))
            continue
        if any(isinstance(g.value(r, c), (int, float)) and not isinstance(g.value(r, c), bool) for c in (*cols.works.values(), *cols.rev.values())):
            continue      # строка итогов под шапкой (числа в столбцах ремонтов и ревизий)
        seen[w] = r
        gsp = group_label(g.value(r, cols.gsp)) if cols.gsp is not None else ""
        well_rows.append({"Скважина": w, "ГСП": gsp, "Строка": r + 1,
                          "Требуется КР (значимость)": g.value(r, cols.required) if cols.required is not None else ""})
        for year, c in cols.works.items():
            text = _text(g.value(r, c))
            if not text:
                continue
            items = split_works(text)
            codes = [code for code, _ in items if code]
            color = g.fill(r, c)
            kind, how = None, ""
            if color in legend:
                kind, how = legend[color], "цвет ячейки"
            else:
                kind, how = kind_by_codes(codes), "коды работ"
                if kind == KIND_OTHER:
                    notes.append(("предупреждение", "Скв. %d, %d год: тип ремонта не определён (цвет %s не в легенде, кодов типа нет)" % (w, year, color)))
                else:
                    notes.append(("заметка", "Скв. %d, %d год: цвет %s не в легенде, тип по кодам работ — «%s»" % (w, year, color, kind)))
            by_codes = kind_by_codes(codes)
            if how == "цвет ячейки" and by_codes not in (KIND_OTHER, kind) and kind in (KIND_CT, KIND_CTS):
                notes.append(("предупреждение", "Скв. %d, %d год: цвет «%s», но в кодах есть работа типа «%s»" % (w, year, kind, by_codes)))
            works_rows.append({"Скважина": w, "ГСП": gsp, "Год": year, "Тип ремонта": kind, "Тип определён по": how,
                               "Коды работ": "; ".join(codes), "Описание работ": text, "Работ": len(items)})
        for season, c in cols.rev.items():
            v = g.value(r, c)
            if not _text(v):
                continue
            shleif = g.fill(r, c) == SHLEIF_FILL
            sand = _is_mark(v)      # розовая «х» тоже даёт песок: в таблице она входит в счёт потребности (скв. 64), в примечании — причина
            rev_rows.append({"Скважина": w, "ГСП": gsp, "Сезон": season, "Источник": "Ревизия КРУ",
                             "Песок": "Да" if sand else "Нет",
                             "Примечание": "вероятно вынос остатков песка из шлейфа (в счёт входит)" if shleif else (_text(v) if not _is_mark(v) else "")})
        for season, c in cols.sand.items():
            v = g.value(r, c)
            if _is_mark(v):
                rev_rows.append({"Скважина": w, "ГСП": gsp, "Сезон": season, "Источник": "Отчёты АН", "Песок": "Да", "Примечание": ""})
        for season, c in cols.gdi.items():
            v = _text(g.value(r, c))
            if v:
                gdi_rows.append({"Скважина": w, "ГСП": gsp, "Сезон": season, "Результат": v})
        for year, c in cols.need.items():
            v = g.value(r, c)
            if _text(v):
                need_rows.append({"Скважина": w, "Год": year, "Потребность (таблица)": int(float(v)) if re.fullmatch(r"\d+(\.0)?", _text(v)) else _text(v)})
    plan_by_well = {}
    for year, c in cols.plan.items():
        for w, r in seen.items():
            if _text(g.value(r, c)):
                plan_by_well[w] = year
    wells = pd.DataFrame(well_rows)
    if not wells.empty:
        wells["В плане КР"] = wells["Скважина"].map(plan_by_well)
    return KrsData(
        works=pd.DataFrame(works_rows, columns=["Скважина", "ГСП", "Год", "Тип ремонта", "Тип определён по", "Коды работ", "Описание работ", "Работ"]),
        revisions=pd.DataFrame(rev_rows, columns=["Скважина", "ГСП", "Сезон", "Источник", "Песок", "Примечание"]),
        gdi=pd.DataFrame(gdi_rows, columns=["Скважина", "ГСП", "Сезон", "Результат"]),
        wells=wells,
        table_need=pd.DataFrame(need_rows, columns=["Скважина", "Год", "Потребность (таблица)"]),
        notes=notes, legend=legend, columns=cols), g


def season_end_year(season: str) -> int:
    return int(season.split("-")[1])


def add_revisions(data: KrsData, new: pd.DataFrame, season: str) -> List[Tuple[str, str]]:
    """Результаты вскрытия КРУ нового сезона (скважина, песок) → в тот же формат; повтор сезона заменяет прежнюю запись."""
    notes: List[Tuple[str, str]] = []
    if new.empty:
        return notes
    known = set(data.wells["Скважина"])
    gsp = dict(zip(data.wells["Скважина"], data.wells["ГСП"]))
    rows = []
    for _, row in new.iterrows():
        w = int(row["Скважина"])
        if w not in known:
            notes.append(("предупреждение", "Скв. %d из вскрытий КРУ %s нет в таблице анализа: добавлена без ГСП" % (w, season)))
        rows.append({"Скважина": w, "ГСП": gsp.get(w, ""), "Сезон": season, "Источник": "Ревизия КРУ",
                     "Песок": row["Песок"], "Примечание": row.get("Примечание", "")})
    add = pd.DataFrame(rows)
    old = data.revisions
    mask = (old["Сезон"] == season) & (old["Источник"] == "Ревизия КРУ") & old["Скважина"].isin(add["Скважина"])
    if mask.any():
        notes.append(("заметка", "Сезон %s: прежних записей по тем же скважинам заменено %d" % (season, int(mask.sum()))))
    data.revisions = pd.concat([old[~mask], add], ignore_index=True)
    return notes


def read_new_revisions(path: str, sand_words: Sequence[str] = ("да", "есть", "х", "x", "1", "песок")) -> pd.DataFrame:
    """Таблица вскрытий КРУ нового сезона: столбец с номером скважины и столбец «песок» (да/нет, х, есть/нет)."""
    df = pd.read_csv(path, sep=None, engine="python") if path.lower().endswith((".csv", ".txt")) else pd.read_excel(path)
    well_col = next((c for c in df.columns if re.search(r"скв", str(c), re.I)), None)
    sand_col = next((c for c in df.columns if re.search(r"песок|вынос|пескопроявл", str(c), re.I)), None)
    if well_col is None or sand_col is None:
        raise ValueError("Во вскрытиях КРУ нужны столбцы с номером скважины («№ скв») и песком («Песок»); найдены: %s" % ", ".join(map(str, df.columns)))
    note_col = next((c for c in df.columns if re.search(r"примеч|коммент", str(c), re.I)), None)
    rows = []
    for _, r in df.iterrows():
        w = well_number(r[well_col])
        if w is None:
            continue
        v = _text(r[sand_col]).lower()
        rows.append({"Скважина": w, "Песок": "Да" if v in sand_words else "Нет",
                     "Примечание": _text(r[note_col]) if note_col else ""})
    return pd.DataFrame(rows, columns=["Скважина", "Песок", "Примечание"]).drop_duplicates("Скважина", keep="last")


def level_from_count(count: int) -> Optional[int]:
    return None if count <= 0 else max(1, 4 - count)


def priority_history(sand_years: Sequence[int], repairs: Dict[int, List[str]], years: Sequence[int],
                     relapse: str = RELAPSE_NEW) -> List[Tuple[int, Optional[int], str]]:
    """Приоритет на конец каждого года: [(год, уровень или None, причина)].

    В году сначала ревизия (по ней песок или нет), потом ремонт — как в таблице: «Ревизия после отбора», затем «Выполнение КР».
    """
    out = []
    count = 0
    base = 0
    sand_set = set(sand_years)
    for y in years:
        reasons = []
        if y in sand_set:
            count += 1
            reasons.append("песок в ревизии %d-%d" % (y - 1, y))
        for kind in repairs.get(y, []):
            if kind in RESETTING:
                count = 0
                base = 1 if (relapse == RELAPSE_TWO and kind == KIND_FILTER) else 0
                count = base
                reasons.append("«%s» %d сбросила счёт" % (kind, y))
            else:
                reasons.append("«%s» %d счёт не сбросила" % (kind, y))
        sand_count = count - base
        out.append((y, level_from_count(count) if sand_count > 0 else None, "; ".join(reasons)))
    return out


def compute_priorities(data: KrsData, start_year: int = 2022, end_year: Optional[int] = None,
                       relapse: str = RELAPSE_NEW) -> pd.DataFrame:
    """Потребность по годам (расчёт) и сверка с «Потребностью» из таблицы."""
    ends = [season_end_year(s) for s in data.revisions["Сезон"]] if not data.revisions.empty else []
    end = end_year or max(ends + [start_year])
    years = list(range(start_year, end + 1))
    rev = data.revisions[(data.revisions["Источник"] == "Ревизия КРУ") & (data.revisions["Песок"] == "Да")]
    sand_by_well = rev.groupby("Скважина")["Сезон"].apply(lambda s: sorted({season_end_year(x) for x in s})).to_dict()
    reps: Dict[int, Dict[int, List[str]]] = {}
    for _, r in data.works.iterrows():
        reps.setdefault(int(r["Скважина"]), {}).setdefault(int(r["Год"]), []).append(r["Тип ремонта"])
    tbl = {(int(r["Скважина"]), int(r["Год"])): r["Потребность (таблица)"] for _, r in data.table_need.iterrows()}
    rows = []
    for w in data.wells["Скважина"]:
        for y, level, why in priority_history(sand_by_well.get(w, []), reps.get(w, {}), years, relapse):
            t = tbl.get((w, y))
            rows.append({"Скважина": w, "Год": y, "Потребность (расчёт)": level, "Потребность (таблица)": t,
                         "Совпало": "" if t is None and level is None else ("да" if t == level else "нет"), "Причина": why})
    return pd.DataFrame(rows)


def final_priorities(data: KrsData, calc: pd.DataFrame, manual: Optional[Dict[int, int]] = None) -> pd.DataFrame:
    """Итог по скважинам: приоритет на последний год, причина, ручное значение, итог (ручное главнее расчёта)."""
    manual = manual or {}
    last_year = int(calc["Год"].max()) if not calc.empty else 0
    cur = calc[calc["Год"] == last_year].set_index("Скважина")
    works = data.works
    rows = []
    for _, w in data.wells.iterrows():
        n = int(w["Скважина"])
        rec = cur.loc[n] if n in cur.index else None
        level = None if rec is None or pd.isna(rec["Потребность (расчёт)"]) else int(rec["Потребность (расчёт)"])
        mine = works[works["Скважина"] == n]
        last = mine.sort_values("Год").iloc[-1] if not mine.empty else None
        revs = data.revisions[(data.revisions["Скважина"] == n) & (data.revisions["Источник"] == "Ревизия КРУ") & (data.revisions["Песок"] == "Да")]
        hist = [calc[(calc["Скважина"] == n) & (calc["Год"] == y)]["Причина"].iloc[0] for y in sorted(calc["Год"].unique())] if not calc.empty else []
        reason = "; ".join(h for h in hist if h)
        rows.append({
            "Скважина": n, "ГСП": w["ГСП"], "На конец года": last_year,
            "Приоритет (расчёт)": level, "Приоритет вручную": manual.get(n),
            "Приоритет итог": manual.get(n) if manual.get(n) is not None else level,
            "Песок в ревизиях КРУ": ", ".join(sorted(revs["Сезон"])),
            "Последний КР": int(last["Год"]) if last is not None else None,
            "Тип последнего КР": last["Тип ремонта"] if last is not None else "",
            "Всего КР": len(mine),
            "В плане КР": w.get("В плане КР"),
            "Причина": reason,
        })
    return pd.DataFrame(rows)


def stats_tables(data: KrsData, prio: pd.DataFrame, calc: pd.DataFrame, per_season: Optional[int] = None,
                 horizon: int = 5) -> Dict[str, pd.DataFrame]:
    """Статистика: проведённые КР по годам и типам, по ГСП; потребность по годам и ГСП; очередь на ближайшие сезоны."""
    out: Dict[str, pd.DataFrame] = {}
    w = data.works
    if not w.empty:
        by = w.pivot_table(index="Год", columns="Тип ремонта", values="Скважина", aggfunc="count", fill_value=0)
        by["Всего КР"] = by.sum(axis=1)
        by["Скважин"] = w.groupby("Год")["Скважина"].nunique()
        out["КР по годам"] = by.reset_index()
        g = w.pivot_table(index="ГСП", columns="Тип ремонта", values="Скважина", aggfunc="count", fill_value=0)
        g["Всего КР"] = g.sum(axis=1)
        g["Скважин с КР"] = w.groupby("ГСП")["Скважина"].nunique()
        out["КР по ГСП"] = g.reset_index()
        rep = w.groupby("Скважина")["Год"].apply(lambda s: sorted(set(s)))
        gaps = [b - a for yrs in rep for a, b in zip(yrs, yrs[1:])]
        out["Интервалы между КР"] = pd.DataFrame([
            ["Скважин с КР", len(rep)], ["Скважин с двумя и более КР", int(sum(len(y) > 1 for y in rep))],
            ["Средний интервал между КР, лет", round(sum(gaps) / len(gaps), 1) if gaps else None],
            ["Скважин без КР за весь период", int(len(data.wells) - len(rep))]], columns=["Показатель", "Значение"])
    if not calc.empty:
        need = calc.dropna(subset=["Потребность (расчёт)"]).pivot_table(index="Год", columns="Потребность (расчёт)", values="Скважина",
                                                                         aggfunc="count", fill_value=0)
        need.columns = ["Приоритет %d" % c for c in need.columns]
        need["Всего нуждается"] = need.sum(axis=1)
        out["Потребность по годам"] = need.reindex(sorted(calc["Год"].unique()), fill_value=0).reset_index()
    if not prio.empty:
        p = prio.dropna(subset=["Приоритет итог"])
        if not p.empty:
            t = p.pivot_table(index="ГСП", columns="Приоритет итог", values="Скважина", aggfunc="count", fill_value=0)
            t.columns = ["Приоритет %d" % c for c in t.columns]
            t["Всего"] = t.sum(axis=1)
            out["Потребность по ГСП"] = t.reset_index()
        # очередь: порядок — приоритет (1 первым), затем больше ревизий с песком, затем дольше без КР; в плане КР — отдельно
        queue = p[p["В плане КР"].isna()].copy()
        queue["Ревизий с песком"] = queue["Песок в ревизиях КРУ"].map(lambda s: len([x for x in s.split(",") if x.strip()]))
        queue["Последний КР"] = queue["Последний КР"].fillna(0)
        queue = queue.sort_values(["Приоритет итог", "Ревизий с песком", "Последний КР"], ascending=[True, False, True])
        if per_season is None:
            by_year = w[w["Тип ремонта"] == KIND_FILTER].groupby("Год")["Скважина"].nunique() if not w.empty else pd.Series(dtype=int)
            recent = by_year[by_year.index >= (int(prio["На конец года"].max()) - 2)]
            per_season = max(1, int(round(recent.mean()))) if len(recent) else max(1, len(queue))
        base = int(prio["На конец года"].max())
        rows = []
        for i, (_, r) in enumerate(queue.iterrows()):
            season = base + i // per_season
            rows.append({"Очередь": i + 1, "Сезон ремонта": "%d-%d" % (season, season + 1) if season < base + horizon else "позже",
                         "Скважина": r["Скважина"], "ГСП": r["ГСП"], "Приоритет": int(r["Приоритет итог"]), "Причина": r["Причина"]})
        out["Очередь КР (подсказка)"] = pd.DataFrame(rows)
        out["Очередь КР (подсказка)"].attrs["per_season"] = per_season
    return out


def read_manual(path: str) -> Dict[int, int]:
    """Ручные приоритеты из прежнего результата (лист «Приоритет», столбец «Приоритет вручную»)."""
    try:
        df = pd.read_excel(path, sheet_name="Приоритет")
    except Exception:
        return {}
    if "Приоритет вручную" not in df.columns:
        return {}
    df = df.dropna(subset=["Приоритет вручную"])
    out = {}
    for _, r in df.iterrows():
        try:
            out[int(r["Скважина"])] = int(r["Приоритет вручную"])
        except (TypeError, ValueError):
            continue
    return out


def write_updated_table(src_grid: Grid, cols: Columns, data: KrsData, prio: pd.DataFrame, season: str, dst: str) -> None:
    """Копия исходной таблицы (значения и заливки) с вскрытиями КРУ нового сезона и приоритетом в двух столбцах справа."""
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = (src_grid.sheet or "Анализ")[:31]
    for r in range(src_grid.nrows):
        for c in range(src_grid.ncols):
            v = src_grid.value(r, c)
            if v != "":
                ws.cell(row=r + 1, column=c + 1, value=v)
            f = src_grid.fill(r, c)
            if f:
                ws.cell(row=r + 1, column=c + 1).fill = PatternFill("solid", fgColor="%02X%02X%02X" % f)
    rev_col = cols.rev.get(season)
    rows = dict(zip(data.wells["Скважина"], data.wells["Строка"]))
    if rev_col is not None:
        new = data.revisions[(data.revisions["Сезон"] == season) & (data.revisions["Источник"] == "Ревизия КРУ")]
        for _, r in new.iterrows():
            row = rows.get(int(r["Скважина"]))
            if row:
                ws.cell(row=row, column=rev_col + 1, value="х" if r["Песок"] == "Да" else None)
    c0 = src_grid.ncols + 2
    ws.cell(row=cols.header_row + 1, column=c0, value="Приоритет КР (расчёт, подсказка)")
    ws.cell(row=cols.header_row + 1, column=c0 + 1, value="Причина")
    for _, r in prio.iterrows():
        row = rows.get(int(r["Скважина"]))
        if row:
            v = r["Приоритет итог"]
            ws.cell(row=row, column=c0, value=None if pd.isna(v) else int(v))
            ws.cell(row=row, column=c0 + 1, value=r["Причина"])
    wb.save(dst)


def revision_notes(data: KrsData) -> pd.DataFrame:
    return pd.DataFrame(data.notes, columns=["Уровень", "Сообщение"])
