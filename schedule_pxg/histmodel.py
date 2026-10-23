"""Сопоставление «история / модель» по дебиту и приёмистости газа: по сезонам и по скважинам.

История — база расходов проекта (м³/сут на скважину и день), модель — вектора WGPR (отбор) и WGIR (закачка) сводки tNavigator
(SMSPEC/UNSMRY). Эталон — суточный объём с узла учёта (`sources.load_daily_total`): им проверяется, что собранный по истории schedule
даёт в модели тот же объём, что в БД (ошибка сборки schedule-секции).

Сезоны — подсказка из данных: непрерывные периоды работы между остановками (эталон, а без него история: ≥ `min_pause` суток подряд без
объёма), вид сезона — по преобладающему виду истории; осень/весна — это границы между сезонами. Список сезонов можно задать вручную.
Объём модели за сезон — интеграл дебита шага сводки по пересечению шага с сезоном; сравнивается тот же период, что покрыт моделью.
Единицы: объём — тыс. м³, расход — тыс. м³/сут. Отклонение = модель − история; недобор (отбор) и недозакачка (закачка) — модель меньше
истории больше допуска. Допуски настраиваются, пометки — только подсказка с видимой причиной; строки можно исключить вручную
(ключ `сезон|скважина` или `*|скважина`).
Python 3.8+.
"""
from __future__ import annotations

import datetime as dt
import math
from typing import Dict, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd

from tnav_results import Summary

INJ, PROD = "закачка", "отбор"
VECTOR = {PROD: "WGPR", INJ: "WGIR"}
FIELD_VECTOR = {PROD: "FGPR", INJ: "FGIR"}
SHORT, OVER, OK, NO_MODEL = "недобор", "превышение", "норма", "нет в модели"
SHORT_LABEL = {PROD: "недобор", INJ: "недозакачка"}
MIN_PAUSE = 5
DEFAULTS = {"rel_pct": 5.0, "abs_k": 1.0, "min_pause": MIN_PAUSE}     # допуск: % от истории и тыс. м³/сут; остановка — суток

_UNIT_K = {"SM3": 1.0, "M3": 1.0, "KSM3": 1e3, "KM3": 1e3, "MSM3": 1e6, "MM3": 1e6}


def unit_factor(unit: str) -> float:
    """Множитель сводки к м³/сут по единице вектора («SM3/DAY» → 1); неизвестная единица — понятная ошибка."""
    head = str(unit or "").strip().upper().split("/")[0].replace(" ", "").replace("-", "")
    if head in _UNIT_K:
        return _UNIT_K[head]
    raise ValueError("Незнакомая единица расхода в сводке модели: «%s». Задайте множитель к м³/сут в настройках" % unit)


# ───────────────────────── сезоны

def _daily_kind(hist: pd.DataFrame) -> Tuple[pd.Series, pd.Series]:
    """По суткам: объём истории (м³/сут, вся) и вид преобладающего объёма."""
    h = hist[(hist["well"].astype(str).str.strip() != "") & (hist["rate"] > 0) & hist["kind"].isin([INJ, PROD])]
    if not len(h):
        return pd.Series(dtype=float), pd.Series(dtype=object)
    h = h.assign(d=pd.to_datetime(h["date"]).dt.normalize())
    piv = h.pivot_table(index="d", columns="kind", values="rate", aggfunc="sum", fill_value=0.0)
    for k in (INJ, PROD):
        if k not in piv:
            piv[k] = 0.0
    vol = piv[INJ] + piv[PROD]
    kind = pd.Series(np.where(piv[INJ] >= piv[PROD], INJ, PROD), index=piv.index)
    return vol, kind


def detect_seasons(hist: pd.DataFrame, daily: Optional[pd.DataFrame] = None, min_pause: int = MIN_PAUSE) -> List[dict]:
    """Сезоны закачки и отбора: [{name, kind, start, end}] (даты ISO), по порядку. Подсказка — правится вручную."""
    hvol, hkind = _daily_kind(hist)
    if not len(hvol):
        return []
    if daily is not None and len(daily):
        act = daily.assign(d=pd.to_datetime(daily["Дата"]).dt.normalize()).groupby("d")["Объем"].sum()
    else:
        act = hvol
    idx = pd.date_range(min(act.index.min(), hvol.index.min()), max(act.index.max(), hvol.index.max()), freq="D")
    on = (act.reindex(idx, fill_value=0.0) > 0).to_numpy()
    kind = hkind.reindex(idx).ffill().bfill().to_numpy()
    # непрерывные периоды работы: остановка короче min_pause суток период не разрывает
    runs: List[List[int]] = []
    i, n = 0, len(idx)
    while i < n:
        if not on[i]:
            i += 1
            continue
        j = i
        gap = 0
        last = i
        while j < n:
            if on[j]:
                gap, last = 0, j
            else:
                gap += 1
                if gap >= min_pause:
                    break
            j += 1
        runs.append([i, last])
        i = last + 1
    # внутри периода вид меняется без остановки — делим; короткие (< min_pause) куски присоединяем к соседу
    segs: List[List] = []
    for a, b in runs:
        cur: List = []
        for k in range(a, b + 1):
            if cur and cur[2] == kind[k]:
                cur[1] = k
            else:
                if cur:
                    segs.append(cur)
                cur = [k, k, kind[k]]
        if cur:
            segs.append(cur)
    merged: List[List] = []
    for s in segs:
        if merged and merged[-1][1] + 1 >= s[0] - min_pause and (s[1] - s[0] + 1 < min_pause or merged[-1][2] == s[2]):
            merged[-1][1] = s[1]
        else:
            merged.append(list(s))
    out: List[dict] = []
    for a, b, k in merged:
        start, end = idx[a].date(), idx[b].date()
        label = str(start.year) if start.year == end.year else "%d/%02d" % (start.year, end.year % 100)
        out.append({"name": "%s %s" % (k.capitalize(), label), "kind": k, "start": start.isoformat(), "end": end.isoformat()})
    names: Dict[str, int] = {}
    for s in out:                                   # два сезона одного вида в год — различаем номером
        names[s["name"]] = names.get(s["name"], 0) + 1
        if names[s["name"]] > 1:
            s["name"] += " (%d)" % names[s["name"]]
    return out


# ───────────────────────── объёмы

def _model_series(s: Summary, kind: str, factor: Optional[float] = None) -> Tuple[Dict[str, np.ndarray], Optional[np.ndarray], List[dt.datetime]]:
    """Дебиты модели по скважинам (м³/сут, по шагам сводки) и вектор по всему объекту, если он есть."""
    kw, fkw = VECTOR[kind], FIELD_VECTOR[kind]
    data = np.asarray(s.data, dtype=float) if len(s.data) else np.zeros((0, len(s.columns)))
    wells: Dict[str, np.ndarray] = {}
    total = None
    for ci, (k, name, _n, unit) in enumerate(s.columns):
        if k == kw and name and not name.startswith(":+:"):
            wells.setdefault(name, np.abs(data[:, ci]) * (factor if factor else unit_factor(unit)))
        elif k == fkw and total is None:
            total = np.abs(data[:, ci]) * (factor if factor else unit_factor(unit))
    return wells, total, list(s.dates)


def _integrate(rate: np.ndarray, edges: np.ndarray, a: float, b: float) -> float:
    """Объём (м³) за [a, b] суток от начала расчёта: дебит шага i действует на (edges[i-1], edges[i]]."""
    lo = np.concatenate(([0.0], edges[:-1]))
    over = np.clip(np.minimum(edges, b) - np.maximum(lo, a), 0.0, None)
    return float(np.dot(rate, over))


def _key(project, name: str) -> str:
    n = str(name).strip()
    r = project.resolve(n) if project is not None else None
    return r or n


def compare(hist: pd.DataFrame, summary: Summary, project=None, daily: Optional[pd.DataFrame] = None,
            seasons: Optional[Sequence[dict]] = None, settings: Optional[dict] = None, excluded: Sequence[str] = (),
            unit_k: Optional[float] = None) -> dict:
    """Сравнение история/модель/эталон: {seasons: [...], notes, settings}. Сезон: итог по объекту, группы, строки по скважинам."""
    st = dict(DEFAULTS, **{k: v for k, v in (settings or {}).items() if v not in (None, "")})
    rel, abs_k = float(st["rel_pct"]), float(st["abs_k"])
    notes: List[str] = []
    hist = hist[hist["kind"].isin([INJ, PROD])]
    seasons = list(seasons) if seasons else detect_seasons(hist, daily, int(st["min_pause"]))
    if not seasons:
        return {"seasons": [], "notes": ["В истории нет расходов: сезонов нет"], "settings": st}
    excl = set(excluded)
    if not summary.dates:
        raise ValueError("В сводке модели нет шагов")
    t0 = summary.start
    edges = np.array([(d - t0).total_seconds() / 86400.0 for d in summary.dates])
    m_from, m_to = summary.start, summary.dates[-1]          # расчёт покрывает [старт, последний шаг]
    cov_a = m_from.date() if m_from.time() == dt.time() else m_from.date() + dt.timedelta(days=1)
    cov_b = m_to.date() - dt.timedelta(days=1) if m_to.time() == dt.time() else m_to.date() - dt.timedelta(days=1)
    series = {k: _model_series(summary, k, unit_k) for k in (PROD, INJ)}
    h = hist[hist["well"].astype(str).str.strip() != ""].assign(d=lambda x: pd.to_datetime(x["date"]).dt.normalize(),
                                                                key=lambda x: x["well"].map(lambda w: _key(project, w)))
    ref = None
    if daily is not None and len(daily):
        ref = daily.assign(d=pd.to_datetime(daily["Дата"]).dt.normalize()).groupby("d")["Объем"].sum()
    gof = (project.well_group if project is not None else {})
    out = []
    for se in seasons:
        kind = se["kind"]
        a, b = dt.date.fromisoformat(se["start"]), dt.date.fromisoformat(se["end"])
        a_m, b_m = max(a, cov_a), min(b, cov_b)
        row = {"name": se["name"], "kind": kind, "start": se["start"], "end": se["end"], "rows": [], "groups": [], "notes": []}
        if b_m < a_m:
            row["notes"].append("Модель не покрывает сезон: сравнивать нечего")
            row["compared"] = None
            out.append(row)
            continue
        days = (b_m - a_m).days + 1
        row["compared"] = {"start": a_m.isoformat(), "end": b_m.isoformat(), "days": days, "full": (a_m, b_m) == (a, b)}
        if not row["compared"]["full"]:
            row["notes"].append("Модель покрывает сезон частично: сравнивается %s — %s (%d сут)" % (a_m, b_m, days))
        ta, tb = pd.Timestamp(a_m), pd.Timestamp(b_m)
        ma = (dt.datetime.combine(a_m, dt.time()) - t0).total_seconds() / 86400.0
        mb = ma + days
        hs = h[(h["kind"] == kind) & (h["d"] >= ta) & (h["d"] <= tb) & (h["rate"] > 0)]
        hv = hs.groupby("key")["rate"].sum() / 1000.0
        wells_m, total_m, _dates = series[kind]
        mv: Dict[str, float] = {}
        for name, r in wells_m.items():
            k = _key(project, name)
            mv[k] = mv.get(k, 0.0) + _integrate(r, edges, ma, mb) / 1000.0
        for key in sorted(set(hv.index) | set(mv), key=lambda s: (len(s), s)):
            hh, mm = float(hv.get(key, 0.0)), float(mv.get(key, 0.0))
            if hh == 0 and mm == 0:
                continue
            diff = mm - hh
            tol = max(abs_k, hh / days * rel / 100.0)
            if key not in mv:
                status, why = NO_MODEL, "скважины нет в сводке модели"
            elif (mm - hh) / days < -tol:
                status, why = SHORT, "модель ниже истории на %.1f тыс. м³/сут при допуске %.1f" % (-diff / days, tol)
            elif (mm - hh) / days > tol:
                status, why = OVER, "модель выше истории на %.1f тыс. м³/сут при допуске %.1f" % (diff / days, tol)
            else:
                status, why = OK, ""
            row["rows"].append({"well": key, "group": gof.get(key, ""), "hist": hh, "model": mm, "diff": diff,
                                "hist_avg": hh / days, "model_avg": mm / days, "diff_avg": diff / days,
                                "pct": (diff / hh * 100.0) if hh else None, "status": status, "why": why,
                                "label": SHORT_LABEL[kind] if status == SHORT else status,
                                "excluded": ("%s|%s" % (se["name"], key) in excl) or ("*|%s" % key in excl),
                                "key": "%s|%s" % (se["name"], key)})
        use = [r for r in row["rows"] if not r["excluded"]]
        h_sum, m_sum = sum(r["hist"] for r in use), sum(r["model"] for r in use)
        e_sum = None
        if ref is not None:
            e_sum = float(ref[(ref.index >= ta) & (ref.index <= tb)].sum()) / 1000.0
            if not len(ref[(ref.index >= ta) & (ref.index <= tb)]):
                e_sum = None
        tot = {"hist": h_sum, "model": m_sum, "ref": e_sum, "hist_avg": h_sum / days, "model_avg": m_sum / days,
               "ref_avg": e_sum / days if e_sum is not None else None, "model_field": None}
        if total_m is not None:
            tot["model_field"] = _integrate(total_m, edges, ma, mb) / 1000.0
        base = e_sum if e_sum is not None else h_sum
        tol_t = max(abs_k * days, base * rel / 100.0)
        reasons = []
        if e_sum is None:
            reasons.append("эталон в этом сезоне не задан: сравнение с историей")
        if m_sum - base < -tol_t:
            tot["status"] = SHORT
            reasons.append("модель ниже %s на %.1f тыс. м³ (%.1f%%), допуск %.1f" %
                           ("эталона" if e_sum is not None else "истории", base - m_sum, (base - m_sum) / base * 100.0 if base else 0.0, tol_t))
        elif m_sum - base > tol_t:
            tot["status"] = OVER
            reasons.append("модель выше %s на %.1f тыс. м³, допуск %.1f" % ("эталона" if e_sum is not None else "истории", m_sum - base, tol_t))
        else:
            tot["status"] = OK
        tot["differs_from_db"] = bool(e_sum is not None and abs(m_sum - e_sum) > tol_t)
        tot["hist_vs_ref"] = None if e_sum is None else h_sum - e_sum
        if e_sum is not None and abs(h_sum - e_sum) > tol_t:
            reasons.append("история по скважинам расходится с эталоном на %.1f тыс. м³ (проверьте БД и узел учёта)" % (h_sum - e_sum))
        tot["label"] = SHORT_LABEL[kind] if tot["status"] == SHORT else tot["status"]
        tot["why"] = "; ".join(reasons)
        row["total"] = tot
        groups: Dict[str, List[float]] = {}
        for r in use:
            g = groups.setdefault(r["group"] or "без группы", [0.0, 0.0])
            g[0] += r["hist"]
            g[1] += r["model"]
        row["groups"] = [{"group": g, "hist": x, "model": y, "diff": y - x, "hist_avg": x / days, "model_avg": y / days,
                          "pct": ((y - x) / x * 100.0) if x else None} for g, (x, y) in sorted(groups.items())]
        out.append(row)
    return {"seasons": out, "notes": notes, "settings": st}


def find_schedule(model_path: str) -> List[str]:
    """Файлы schedule с историей (WCONHIST/WCONINJH) рядом с моделью: папка модели и подпапки до двух уровней вглубь
    (INCLUDE и т.п.). Нет файлов — пустой список."""
    import os
    from . import history
    folder = model_path if os.path.isdir(model_path) else os.path.dirname(model_path)
    if not folder or not os.path.isdir(folder):
        return []
    found: List[str] = []
    base = folder.rstrip("\\/").count(os.sep)
    for root, dirs, files in os.walk(folder):
        if root.rstrip("\\/").count(os.sep) - base >= 2:
            dirs[:] = []
        for n in sorted(files):
            if os.path.splitext(n)[1].lower() in (".inc", ".sch", ".txt", ".data"):
                path = os.path.join(root, n)
                try:
                    if os.path.getsize(path) < 3e8 and history.detect_format(path) == history.FORMAT_SCHEDULE:
                        found.append(path)
                except OSError:
                    continue
    return found


def schedule_history(paths: Sequence[str]) -> pd.DataFrame:
    """История по файлам schedule (как её читает ГДМ при импорте); одинаковые скважина+сутки+вид — последняя запись."""
    from . import history
    frames = [history.read_schedule(p) for p in paths]
    df = pd.concat([f for f in frames if len(f)], ignore_index=True) if any(len(f) for f in frames) else history._empty()
    return df.drop_duplicates(["well", "date", "kind"], keep="last").reset_index(drop=True)


def view(project, model_path: str, body: Optional[dict] = None) -> dict:
    """Сравнение для экрана: история и эталон проекта против сводки модели. Присланные настройки, сезоны и исключения
    запоминаются в проекте (`project.histmodel`), без них берутся прежние; вернуть автоподсказку сезонов — `seasons: []`.
    История: файл schedule, выбранный вручную (`schedule_path`), иначе база расходов проекта, иначе schedule из папки модели."""
    import os
    from . import results, sources
    body = body or {}
    saved = project.histmodel
    for k in ("settings", "seasons", "excluded", "unit_k", "schedule_path"):
        if k in body:
            saved[k] = body[k]
    notes: List[str] = []
    manual = str(saved.get("schedule_path") or "").strip().strip('"')
    if manual:
        if not os.path.isfile(manual):
            raise ValueError("Файл schedule не найден: %s" % manual)
        hist, source = schedule_history([manual]), "schedule: %s" % os.path.basename(manual)
    elif sources.flow_paths(project):
        hist, source = sources.load_project_history(project), "база расходов проекта"
    else:
        files = find_schedule(model_path)
        if not files:
            raise ValueError("Историю взять неоткуда: в папке модели нет файла schedule с WCONHIST/WCONINJH, "
                             "а база расходов в проекте не задана. Укажите файл schedule или базу расходов")
        hist, source = schedule_history(files), "schedule из папки модели: %s" % ", ".join(os.path.basename(f) for f in files[:5])
    if not len(hist):
        raise ValueError("В источнике истории нет расходов по скважинам (%s)" % source)
    daily = sources.load_daily_total(project)
    out = compare(hist, results.load_summary(model_path), project, daily, saved.get("seasons") or None, saved.get("settings"),
                  saved.get("excluded") or [], float(saved["unit_k"]) if saved.get("unit_k") else None)
    out["notes"] = notes + out["notes"]
    out["history_source"] = source
    out["saved"] = {"seasons": saved.get("seasons") or [], "excluded": saved.get("excluded") or [], "unit_k": saved.get("unit_k"),
                    "schedule_path": manual}
    if daily is None or not len(daily):
        out["notes"].append("Эталонный суточный объём не задан: итог по сезону сравнивается с историей по скважинам")
    return out


def save_xlsx(path: str, res: dict) -> None:
    """Выгрузка в Excel: лист «Скважины» (история/модель по сезонам, разности — формулами) и «Сезоны» (итоги по объекту:
    суммы по скважинам без исключённых — формулами SUMIFS, отличие от эталона — формулами). Пометки статуса — текст-подсказка."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Сезоны"
    wl = wb.create_sheet("Скважины")
    wl.append(["Сезон", "Скважина", "Группа", "Вид", "Суток", "История, тыс. м³", "Модель, тыс. м³", "Модель − история, тыс. м³",
               "Отклонение, %", "История, тыс. м³/сут", "Модель, тыс. м³/сут", "Модель − история, тыс. м³/сут", "Пометка", "Причина", "Исключена"])
    ws.append(["Сезон", "Вид", "С", "По", "Сравнено суток", "История, тыс. м³", "Модель, тыс. м³", "Эталон (узел учёта), тыс. м³",
               "Модель − эталон, тыс. м³", "Модель − эталон, %", "История − эталон, тыс. м³", "Пометка", "Причина"])
    n = 1
    for se in res["seasons"]:
        for r in se["rows"]:
            n += 1
            wl.append([se["name"], r["well"], r["group"], se["kind"], se["compared"]["days"], r["hist"], r["model"],
                       "=G%d-F%d" % (n, n), '=IF(F%d=0,"",H%d/F%d)' % (n, n, n), "=F%d/E%d" % (n, n), "=G%d/E%d" % (n, n),
                       "=K%d-J%d" % (n, n), r["label"], r["why"], "да" if r["excluded"] else "нет"])
            wl.cell(row=n, column=9).number_format = "0.0%"
    last = max(n, 2)
    m = 1
    for se in res["seasons"]:
        m += 1
        t, c = se.get("total"), se.get("compared")
        if not t:
            ws.append([se["name"], se["kind"], se["start"], se["end"], 0, None, None, None, None, None, None, "", "; ".join(se["notes"])])
            continue
        sums = ['=SUMIFS(Скважины!$%s$2:$%s$%d,Скважины!$A$2:$A$%d,$A%d,Скважины!$O$2:$O$%d,"нет")' % (col, col, last, last, m, last) for col in "FG"]
        ws.append([se["name"], se["kind"], c["start"], c["end"], c["days"], sums[0], sums[1], t["ref"],
                   '=IF(H%d="","",G%d-H%d)' % (m, m, m), '=IF(OR(H%d="",H%d=0),"",I%d/H%d)' % (m, m, m, m),
                   '=IF(H%d="","",F%d-H%d)' % (m, m, m), t["label"], t["why"] or "; ".join(se["notes"])])
        ws.cell(row=m, column=10).number_format = "0.0%"
    for sheet in (ws, wl):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for i in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(i)].width = 18
        sheet.freeze_panes = "A2"
    wb.save(path)
