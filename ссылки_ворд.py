import zipfile
import os
import shutil
import xml.etree.ElementTree as ET
from pathlib import Path
import subprocess
import sys

class DocxRepair:
    def __init__(self):
        self.temp_dir = "temp_docx_repair"
    
    def repair_docx_structure(self, input_path, output_path=None):
        """
        Восстанавливает структуру DOCX файла, удаляя поврежденные элементы
        """
        if output_path is None:
            output_path = str(Path(input_path).parent / f"{Path(input_path).stem}_repaired.docx")
        
        print(f"🔧 Восстанавливаю структуру DOCX: {input_path}")
        
        try:
            # Создаем временную директорию
            if os.path.exists(self.temp_dir):
                shutil.rmtree(self.temp_dir)
            os.makedirs(self.temp_dir)
            
            # Распаковываем DOCX как ZIP
            with zipfile.ZipFile(input_path, 'r') as zip_ref:
                zip_ref.extractall(self.temp_dir)
            
            print("✅ DOCX распакован")
            
            # Удаляем проблемные файлы если они существуют
            problematic_files = [
                'word/embeddings/maskFile.bin',
                'word/embeddings/oleObject1.bin',
                'word/embeddings/',
            ]
            
            for file_path in problematic_files:
                full_path = os.path.join(self.temp_dir, file_path)
                if os.path.exists(full_path):
                    if os.path.isfile(full_path):
                        os.remove(full_path)
                        print(f"✅ Удален проблемный файл: {file_path}")
                    elif os.path.isdir(full_path):
                        shutil.rmtree(full_path)
                        print(f"✅ Удалена проблемная директория: {file_path}")
            
            # Чистим ссылки на отсутствующие файлы из [Content_Types].xml
            self.clean_content_types()
            
            # Чистим ссылки из .rels файлов
            self.clean_relationships()
            
            # Пересоздаем DOCX архив
            with zipfile.ZipFile(output_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
                for root, dirs, files in os.walk(self.temp_dir):
                    for file in files:
                        file_path = os.path.join(root, file)
                        arcname = os.path.relpath(file_path, self.temp_dir)
                        zipf.write(file_path, arcname)
            
            print(f"✅ Восстановленный DOCX сохранен: {output_path}")
            
            # Очищаем временные файлы
            shutil.rmtree(self.temp_dir)
            
            return output_path
            
        except Exception as e:
            print(f"❌ Ошибка восстановления: {e}")
            # Пытаемся очистить временные файлы даже при ошибке
            if os.path.exists(self.temp_dir):
                shutil.rmtree(self.temp_dir)
            return None
    
    def clean_content_types(self):
        """Очищает [Content_Types].xml от ссылок на отсутствующие файлы"""
        try:
            content_types_path = os.path.join(self.temp_dir, '[Content_Types].xml')
            if not os.path.exists(content_types_path):
                return
            
            tree = ET.parse(content_types_path)
            root = tree.getroot()
            
            # Удаляем Override для embeddings
            for override in root.findall('{http://schemas.openxmlformats.org/package/2006/content-types}Override'):
                part_name = override.get('PartName')
                if part_name and 'embeddings' in part_name:
                    # Проверяем существует ли файл
                    file_path = os.path.join(self.temp_dir, part_name.lstrip('/'))
                    if not os.path.exists(file_path):
                        root.remove(override)
                        print(f"✅ Удалена ссылка на отсутствующий файл: {part_name}")
            
            tree.write(content_types_path, encoding='utf-8', xml_declaration=True)
            
        except Exception as e:
            print(f"⚠️  Ошибка очистки Content_Types: {e}")
    
    def clean_relationships(self):
        """Очищает .rels файлы от ссылок на отсутствующие ресурсы"""
        try:
            # Проверяем основные .rels файлы
            rels_files = [
                'word/_rels/document.xml.rels',
                '_rels/.rels'
            ]
            
            for rels_file in rels_files:
                rels_path = os.path.join(self.temp_dir, rels_file)
                if os.path.exists(rels_path):
                    self.clean_single_rels_file(rels_path)
                    
        except Exception as e:
            print(f"⚠️  Ошибка очистки relationships: {e}")
    
    def clean_single_rels_file(self, rels_path):
        """Очищает один .rels файл"""
        try:
            tree = ET.parse(rels_path)
            root = tree.getroot()
            namespace = '{http://schemas.openxmlformats.org/package/2006/relationships}'
            
            removed_count = 0
            for relationship in root.findall(f'{namespace}Relationship'):
                target = relationship.get('Target')
                if target and 'embeddings' in target:
                    # Проверяем существует ли файл
                    target_path = os.path.join(self.temp_dir, rels_path.replace('_rels/', '').replace('.rels', ''), target)
                    target_path = os.path.normpath(target_path)
                    if not os.path.exists(target_path):
                        root.remove(relationship)
                        removed_count += 1
            
            if removed_count > 0:
                tree.write(rels_path, encoding='utf-8', xml_declaration=True)
                print(f"✅ Удалено {removed_count} ссылок из {os.path.basename(rels_path)}")
                
        except Exception as e:
            print(f"⚠️  Ошибка очистки {rels_path}: {e}")


class DocumentNumberingFixer:
    def __init__(self):
        self.repair = DocxRepair()
    
    def safe_open_document(self, doc_path):
        """
        Безопасно открывает документ, предварительно восстанавливая его при необходимости
        """
        try:
            from docx import Document
            return Document(doc_path)
        except Exception as e:
            if "maskFile.bin" in str(e) or "embeddings" in str(e):
                print("⚠️  Обнаружена поврежденная структура DOCX, восстанавливаю...")
                repaired_path = self.repair.repair_docx_structure(doc_path)
                if repaired_path:
                    return Document(repaired_path)
                else:
                    raise Exception("Не удалось восстановить структуру DOCX")
            else:
                raise
    
    def fix_numbering(self, input_path, output_path=None):
        """
        Основная функция исправления нумерации с обработкой поврежденных файлов
        """
        if output_path is None:
            output_path = str(Path(input_path).parent / f"{Path(input_path).stem}_numbered.docx")
        
        try:
            print("🔧 Начинаю исправление нумерации...")
            
            # Пытаемся открыть документ с восстановлением при необходимости
            doc = self.safe_open_document(input_path)
            
            # Счетчики и карта ссылок
            figure_count = 0
            table_count = 0
            reference_map = {}
            
            print("📝 Первый проход: поиск и нумерация подписей...")
            
            # Первый проход: находим все подписи
            for i, paragraph in enumerate(doc.paragraphs):
                text = paragraph.text.strip()
                lower_text = text.lower()
                
                # Ищем подписи рисунков
                if any(keyword in lower_text for keyword in ['рисунок', 'figure']):
                    # Ищем номер в подписи
                    import re
                    num_match = re.search(r'(\d+)', text)
                    old_num = num_match.group(1) if num_match else str(figure_count + 1)
                    
                    figure_count += 1
                    reference_map[f'figure_{old_num}'] = figure_count
                    
                    # Обновляем текст подписи
                    if 'рисунок' in lower_text:
                        new_text = re.sub(r'Рисунок\s+\d+', f'Рисунок {figure_count}', text)
                    else:
                        new_text = re.sub(r'Figure\s+\d+', f'Figure {figure_count}', text)
                    
                    paragraph.text = new_text
                    print(f"  📊 Рисунок {old_num} → {figure_count}")
                
                # Ищем подписи таблиц
                elif any(keyword in lower_text for keyword in ['таблица', 'table']):
                    num_match = re.search(r'(\d+)', text)
                    old_num = num_match.group(1) if num_match else str(table_count + 1)
                    
                    table_count += 1
                    reference_map[f'table_{old_num}'] = table_count
                    
                    if 'таблица' in lower_text:
                        new_text = re.sub(r'Таблица\s+\d+', f'Таблица {table_count}', text)
                    else:
                        new_text = re.sub(r'Table\s+\d+', f'Table {table_count}', text)
                    
                    paragraph.text = new_text
                    print(f"  📋 Таблица {old_num} → {table_count}")
            
            print("🔗 Второй проход: обновление ссылок в тексте...")
            
            # Второй проход: обновляем ссылки в тексте
            for paragraph in doc.paragraphs:
                original_text = paragraph.text
                text = original_text
                
                # Обновляем ссылки на рисунки
                import re
                figure_refs = re.finditer(r'(\bрис\.?\s*|рисунок\s*|figure\s*)(\d+)', text, re.IGNORECASE)
                for match in figure_refs:
                    prefix = match.group(1).strip()
                    old_num = match.group(2)
                    ref_key = f'figure_{old_num}'
                    
                    if ref_key in reference_map:
                        new_num = reference_map[ref_key]
                        old_ref = match.group(0)
                        new_ref = f"{prefix} {new_num}"
                        text = text.replace(old_ref, new_ref)
                        print(f"  🔄 Ссылка на рисунок: {old_ref} → {new_ref}")
                
                # Обновляем ссылки на таблицы
                table_refs = re.finditer(r'(\bтабл\.?\s*|таблица\s*|table\s*)(\d+)', text, re.IGNORECASE)
                for match in table_refs:
                    prefix = match.group(1).strip()
                    old_num = match.group(2)
                    ref_key = f'table_{old_num}'
                    
                    if ref_key in reference_map:
                        new_num = reference_map[ref_key]
                        old_ref = match.group(0)
                        new_ref = f"{prefix} {new_num}"
                        text = text.replace(old_ref, new_ref)
                        print(f"  🔄 Ссылка на таблицу: {old_ref} → {new_ref}")
                
                if text != original_text:
                    paragraph.text = text
            
            print("💾 Сохраняю результат...")
            doc.save(output_path)
            
            print(f"✅ Нумерация исправлена! Файл сохранен: {output_path}")
            print(f"📊 Статистика: {figure_count} рисунков, {table_count} таблиц")
            
            return output_path
            
        except ImportError:
            print("❌ python-docx не установлен. Установите: pip install python-docx")
            return None
        except Exception as e:
            print(f"❌ Ошибка при исправлении нумерации: {e}")
            return None


def main():
    """Основная функция"""
    print("=" * 60)
    print("🔧 DOCX Repair & Numbering Fixer")
    print("=" * 60)
    
    # Получаем путь к файлу
    if len(sys.argv) > 1:
        input_file = sys.argv[1]
    else:
        input_file = input("Введите путь к поврежденному DOCX файлу: ").strip().strip('"')
    
    if not input_file or not os.path.exists(input_file):
        print("❌ Файл не существует!")
        return
    
    # Создаем фиксер
    fixer = DocumentNumberingFixer()
    
    # Обрабатываем документ
    output_file = str(Path(input_file).parent / f"{Path(input_file).stem}_FIXED.docx")
    
    print(f"📁 Входной файл: {input_file}")
    print(f"💾 Выходной файл: {output_file}")
    print()
    
    result = fixer.fix_numbering(input_file, output_file)
    
    if result:
        print("\n🎉 Обработка завершена успешно!")
        print(f"✅ Исправленный файл: {result}")
    else:
        print("\n❌ Обработка завершена с ошибками!")
        
        # Пробуем просто восстановить структуру без нумерации
        print("\n🔄 Пробую просто восстановить структуру DOCX...")
        repair = DocxRepair()
        repaired = repair.repair_docx_structure(input_file)
        if repaired:
            print(f"✅ Структура восстановлена: {repaired}")
        else:
            print("❌ Не удалось восстановить структуру")


def install_dependencies():
    """Устанавливает необходимые зависимости"""
    print("Устанавливаю зависимости...")
    
    try:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
        print("✅ python-docx установлен")
    except subprocess.CalledProcessError:
        print("❌ Не удалось установить python-docx")
    
    print("\nЗависимости установлены. Запустите скрипт снова.")


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--install":
        install_dependencies()
    else:
        main()
