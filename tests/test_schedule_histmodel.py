"""Сопоставление история/модель по газу: сезоны, объёмы, недоборы, эталон."""
import datetime as dt
import os
import sys

import pandas as pd
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from schedule_pxg import histmodel as hm  # noqa: E402
from schedule_pxg.project import Project  # noqa: E402
from tnav_results import Summary  # noqa: E402


def _hist(rows):
    return pd.DataFrame(rows, columns=["well", "date", "rate", "hours", "kind"]).assign(date=lambda d: pd.to_datetime(d["date"]))


def _days(start, n):
    return [pd.Timestamp(start) + pd.Timedelta(days=i) for i in range(n)]


def _case():
    """Закачка 10 сут (W1=100, W2=50 тыс. м³/сут), пауза 6 сут, отбор 10 сут (W1=80, W2=40)."""
    rows = []
    for d in _days("2020-01-01", 10):
        rows += [("W1", d, 100000.0, 24, hm.INJ), ("W2", d, 50000.0, 24, hm.INJ)]
    for d in _days("2020-01-17", 10):
        rows += [("W1", d, 80000.0, 24, hm.PROD), ("W2", d, 40000.0, 24, hm.PROD)]
    return _hist(rows)


def _daily(h):
    t = h.groupby("date")["rate"].sum().reset_index()
    return t.rename(columns={"date": "Дата", "rate": "Объем"})


def _summary(inj=(100.0, 50.0), prod=(80.0, 40.0), unit="SM3/DAY"):
    """Шаги по 1 сут с 2020-01-01: закачка дни 0–9, пауза, отбор 16–25 (дебит в тыс. → в сводке м³/сут)."""
    cols = [("TIME", "", 0, "DAYS"), ("WGIR", "W1", 0, unit), ("WGIR", "W2", 0, unit), ("WGPR", "W1", 0, unit), ("WGPR", "W2", 0, unit),
            ("FGPR", ":+:+:+:+", 0, unit)]
    dates, data = [], []
    for i in range(1, 31):                         # шаг i покрывает сутки i-1 (с 00:00)
        day = i - 1
        gi = inj if day < 10 else (0.0, 0.0)
        gp = prod if 16 <= day < 26 else (0.0, 0.0)
        dates.append(dt.datetime(2020, 1, 1) + dt.timedelta(days=i))
        data.append([float(i), gi[0] * 1000, gi[1] * 1000, gp[0] * 1000, gp[1] * 1000, sum(gp) * 1000])
    return Summary(dt.datetime(2020, 1, 1), cols, dates, data)


def _write_model(mdl, sm):
    from tests.test_tnav_results import _blk
    (mdl / "M.SMSPEC").write_bytes(
        _blk("STARTDAT", "INTE", [1, 1, 2020, 0, 0, 0]) + _blk("KEYWORDS", "CHAR", [c[0] for c in sm.columns]) +
        _blk("WGNAMES", "CHAR", [c[1] or ":+:+:+:+" for c in sm.columns]) + _blk("NUMS", "INTE", [0] * len(sm.columns)) +
        _blk("UNITS", "CHAR", [c[3] for c in sm.columns]))
    (mdl / "M.UNSMRY").write_bytes(b"".join(_blk("MINISTEP", "INTE", [i]) + _blk("PARAMS", "REAL", r) for i, r in enumerate(sm.data)))


def test_unit_factor():
    assert hm.unit_factor("SM3/DAY") == 1.0 and hm.unit_factor("KSM3/DAY") == 1000.0
    with pytest.raises(ValueError):
        hm.unit_factor("ЧТО-ТО/DAY")


def test_detect_seasons():
    h = _case()
    s = hm.detect_seasons(h, _daily(h))
    assert [(x["kind"], x["start"], x["end"]) for x in s] == [(hm.INJ, "2020-01-01", "2020-01-10"), (hm.PROD, "2020-01-17", "2020-01-26")]
    assert s[0]["name"] == "Закачка 2020"
    assert len(hm.detect_seasons(h)) == 2                       # без эталона — по истории


def test_compare_ok_and_short():
    h = _case()
    r = hm.compare(h, _summary(), None, _daily(h))
    inj, prod = r["seasons"]
    assert inj["total"]["hist"] == 1500.0 and inj["total"]["model"] == 1500.0 and inj["total"]["status"] == hm.OK
    assert inj["compared"]["days"] == 10 and prod["total"]["ref"] == 1200.0
    # недобор отбора по W2: модель 20 вместо 40
    r = hm.compare(h, _summary(prod=(80.0, 20.0)), None, _daily(h))
    prod = r["seasons"][1]
    w2 = next(x for x in prod["rows"] if x["well"] == "W2")
    assert w2["status"] == hm.SHORT and w2["label"] == "недобор" and w2["hist"] == 400.0 and w2["model"] == 200.0 and w2["diff_avg"] == -20.0
    assert prod["total"]["status"] == hm.SHORT and prod["total"]["differs_from_db"] and "ниже эталона" in prod["total"]["why"]
    assert r["seasons"][0]["total"]["status"] == hm.OK
    # закачка называется недозакачкой
    r = hm.compare(h, _summary(inj=(60.0, 50.0)), None, _daily(h))
    assert next(x for x in r["seasons"][0]["rows"] if x["well"] == "W1")["label"] == "недозакачка"


def test_tolerance_exclude_and_reference_mismatch():
    h = _case()
    sm = _summary(prod=(79.0, 40.0))                           # -1 тыс./сут на W1: на границе допуска 1.0 и 5% — норма
    assert hm.compare(h, sm, None, _daily(h))["seasons"][1]["total"]["status"] == hm.OK
    r = hm.compare(h, sm, None, _daily(h), settings={"abs_k": 0.1, "rel_pct": 0.1})
    assert r["seasons"][1]["total"]["status"] == hm.SHORT
    # исключили скважину вручную — из итога уходит
    key = r["seasons"][1]["rows"][0]["key"]
    r2 = hm.compare(h, sm, None, _daily(h), settings={"abs_k": 0.1, "rel_pct": 0.1}, excluded=[key])
    assert r2["seasons"][1]["rows"][0]["excluded"] and r2["seasons"][1]["total"]["hist"] == 400.0
    # эталон отличается от истории (ошибка в БД): подсказка в причине
    d = _daily(h)
    d.loc[d["Дата"] == pd.Timestamp("2020-01-17"), "Объем"] += 200000.0
    r3 = hm.compare(h, _summary(), None, d)
    assert r3["seasons"][1]["total"]["hist_vs_ref"] == pytest.approx(-200.0) and "расходится с эталоном" in r3["seasons"][1]["total"]["why"]


def test_partial_coverage_synonyms_and_manual_seasons():
    h = _case()
    p = Project()
    p.add_well("Well-1", synonyms=["W1"])
    sm = _summary()
    cut = Summary(sm.start, sm.columns, sm.dates[:20], sm.data[:20])        # модель до 21.01
    r = hm.compare(h, cut, p, _daily(h))
    prod = r["seasons"][1]
    assert prod["compared"] == {"start": "2020-01-17", "end": "2020-01-20", "days": 4, "full": False}
    assert prod["total"]["hist"] == 480.0 and prod["total"]["model"] == 480.0 and any(x["well"] == "Well-1" for x in prod["rows"])
    r = hm.compare(h, cut, None, None, seasons=[{"name": "Весна", "kind": hm.PROD, "start": "2020-03-01", "end": "2020-03-31"}])
    assert r["seasons"][0]["compared"] is None
    miss = hm.compare(h, _summary(), None, _daily(h))["seasons"][0]
    assert miss["total"]["model_field"] is None
    assert hm.compare(h, _summary(), None, _daily(h))["seasons"][1]["total"]["model_field"] == 1200.0


def test_api_view_saves_settings(tmp_path, monkeypatch):
    pytest.importorskip("starlette")
    from starlette.testclient import TestClient
    from schedule_pxg import api, sources
    mdl = tmp_path / "m"
    mdl.mkdir()
    _write_model(mdl, _summary(prod=(80.0, 20.0)))
    h = _case()
    monkeypatch.setattr(sources, "flow_paths", lambda p: ["x.xlsx"])
    monkeypatch.setattr(sources, "load_project_history", lambda p, *a, **k: h)
    monkeypatch.setattr(sources, "load_daily_total", lambda p, *a, **k: _daily(h))
    folder = str(tmp_path / "proj")
    monkeypatch.setattr(api, "FOLDER", folder)
    Project().save(folder)
    c = TestClient(api.build_app())
    assert c.post("/api/results/histmodel", json={}).status_code == 404
    r = c.post("/api/results/histmodel", json={"model": str(mdl), "settings": {"rel_pct": 3}, "excluded": ["*|W1"]}).json()
    prod = r["seasons"][1]
    assert prod["total"]["status"] == hm.SHORT and next(x for x in prod["rows"] if x["well"] == "W1")["excluded"]
    assert Project.load(folder).histmodel["settings"] == {"rel_pct": 3}
    again = c.post("/api/results/histmodel", json={"model": str(mdl)}).json()          # настройки взяты из проекта
    assert again["settings"]["rel_pct"] == 3 and again["saved"]["excluded"] == ["*|W1"]
    x = c.post("/api/results/histmodel/xlsx", json={"model": str(mdl)})
    assert x.status_code == 200 and x.content[:2] == b"PK"


def test_xlsx(tmp_path):
    from openpyxl import load_workbook
    h = _case()
    res = hm.compare(h, _summary(prod=(80.0, 20.0)), None, _daily(h), excluded=["Закачка 2020|W2"])
    out = str(tmp_path / "x.xlsx")
    hm.save_xlsx(out, res)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Сезоны", "Скважины"]
    wl, ws = wb["Скважины"], wb["Сезоны"]
    assert wl["H2"].value == "=G2-F2" and wl["F2"].value == 1000.0 and wl["O3"].value == "да"
    assert ws["F2"].value.startswith("=SUMIFS(Скважины!$F$2") and ws["H3"].value == 1200.0 and ws["I3"].value == "=IF(H3=\"\",\"\",G3-H3)"
    assert ws["L3"].value == "недобор"


def _sched(h):
    """schedule tNav: запись под DATES d относится к суткам d+1."""
    lines = []
    for d, g in h.groupby("date"):
        e = d - pd.Timedelta(days=1)
        lines.append("DATES\n %d %s %d /\n/\n" % (e.day, e.strftime("%b").upper(), e.year))
        for kind, kw in ((hm.PROD, "WCONHIST"), (hm.INJ, "WCONINJH")):
            r = g[g["kind"] == kind]
            if len(r):
                lines.append(kw + "\n" + "".join(("'%s' OPEN GRAT 1* 1* %g /\n" if kw == "WCONHIST" else "'%s' GAS OPEN %g /\n") % (x.well, x.rate)
                                                   for x in r.itertuples()) + "/\n")
    return "".join(lines)


def test_history_from_model_folder_schedule(tmp_path, monkeypatch):
    pytest.importorskip("starlette")
    from starlette.testclient import TestClient
    from schedule_pxg import api, sources
    h = _case()
    mdl = tmp_path / "m"
    (mdl / "INCLUDE").mkdir(parents=True)
    _write_model(mdl, _summary(prod=(80.0, 20.0)))
    (mdl / "INCLUDE" / "sched.inc").write_text(_sched(h), encoding="utf-8")
    (mdl / "readme.txt").write_text("не schedule", encoding="utf-8")
    assert [os.path.basename(x) for x in hm.find_schedule(str(mdl))] == ["sched.inc"]
    got = hm.schedule_history(hm.find_schedule(str(mdl)))
    assert got["rate"].sum() == h["rate"].sum() and set(got["kind"]) == {hm.INJ, hm.PROD}
    monkeypatch.setattr(sources, "flow_paths", lambda p: [])
    monkeypatch.setattr(sources, "load_daily_total", lambda p, *a, **k: _daily(h))
    p = Project()
    out = hm.view(p, str(mdl), {})
    assert "schedule из папки модели" in out["history_source"]
    inj, prod = out["seasons"]
    assert inj["total"]["hist"] == 1500.0 and prod["total"]["status"] == hm.SHORT
    own = tmp_path / "own.inc"                                       # файл, выбранный вручную, главнее
    own.write_text(_sched(h[h["kind"] == hm.INJ]), encoding="utf-8")
    out = hm.view(p, str(mdl), {"schedule_path": str(own)})
    assert out["history_source"] == "schedule: own.inc" and out["seasons"][0]["total"]["hist"] == 1500.0
    bare = tmp_path / "bare"
    bare.mkdir()
    _write_model(bare, _summary())
    with pytest.raises(ValueError, match="Историю взять неоткуда"):
        hm.view(Project(), str(bare), {})
