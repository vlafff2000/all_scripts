"""Анализ технического состояния скважин (КРС): разбор таблицы по шапке и цветам, счёт приоритета, вскрытия КРУ, ручной приоритет."""
import importlib.util
import os
import sys

import pytest

pd = pytest.importorskip("pandas")
openpyxl = pytest.importorskip("openpyxl")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl.styles import PatternFill  # noqa: E402

from pxg_core import krs  # noqa: E402

FILTER, CT, CTS = "FFCC00", "00CCFF", "FFFF00"


def fill(cell, rgb):
    cell.fill = PatternFill("solid", fgColor=rgb)


def make_table(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    head = ["№ скв", "№ ГСП", "Ревизия КРУ после отбора 2021-2022", "Выполнение КР 2022", "Ревизия КРУ после отбора 2022-2023",
            "Выполнение КР 2023", "Потребность в КР на конец 2023 года", "Ревизия КРУ после отбора 2023-2024", "Выполнение КР 2024",
            "Потребность в КР на конец 2024 года", "Ревизия КРУ после отбора 2025-2026", "Требуется КР (значимость 1-3)"]
    ws.append(head)
    ws.append([3, "", "", "", "", "", "", "", "", "", "", ""])      # строка итогов под шапкой
    ws.cell(2, 3).value = 5.0
    ws.append([11, 1, "х", "КР 2.022 ГТ Ликвидация пробки; КР 4.033 Кислотная обработка", "х", "КР 1.094 ПА Замена противопесочного фильтра", "", "х", "", 3, "", ""])
    ws.append([12, 1, "х", "", "х", "", 2, "х", "КР 17.3 Промывка песчаных пробок", 1, "", ""])
    ws.append([13, 2, "", "", "", "", "", "", "", "", "", ""])
    fill(ws.cell(3, 4), CTS)
    fill(ws.cell(3, 6), FILTER)
    fill(ws.cell(4, 9), CT)
    ws.cell(1, 20).value = "легенда"
    ws.cell(6, 14).value = None
    fill(ws.cell(6, 14), FILTER)
    ws.cell(6, 15).value = "Капитальный ремонт с заменой фильтра"
    fill(ws.cell(7, 14), CT)
    ws.cell(7, 15).value = "Промывка скв колтюбингом (без СКО)"
    fill(ws.cell(8, 14), CTS)
    ws.cell(8, 15).value = "Промывка скв колтюбингом (с СКО) или просто СКО"
    wb.save(path)


def test_priority_rules():
    f, ct = krs.KIND_FILTER, krs.KIND_CTS
    res = krs.priority_history([2022, 2023, 2025], {2024: [ct]}, [2022, 2023, 2024, 2025])
    assert [r[1] for r in res] == [3, 2, 2, 1]          # промывка счёт не сбрасывает, год без песка число не меняет
    res = krs.priority_history([2022, 2023], {2023: [f]}, [2022, 2023, 2024])
    assert [r[1] for r in res] == [3, None, None]       # замена фильтра сбрасывает после ревизии того же года
    res = krs.priority_history([2022, 2024, 2025], {2023: [f]}, [2022, 2023, 2024, 2025], relapse=krs.RELAPSE_TWO)
    assert [r[1] for r in res] == [3, None, 2, 1]       # как скв. 56: повтор после фильтра сразу 2


def test_parse_and_calc(tmp_path):
    src = tmp_path / "t.xlsx"
    make_table(src)
    data, _ = krs.parse(str(src))
    assert sorted(data.wells["Скважина"]) == [11, 12, 13]       # строка итогов не скважина
    w = data.works
    assert set(w[w["Скважина"] == 11]["Тип ремонта"]) == {krs.KIND_CTS, krs.KIND_FILTER}
    assert w[w["Скважина"] == 11]["Коды работ"].iloc[0] == "КР 2.022 ГТ; КР 4.033"
    assert w[w["Скважина"] == 12]["Тип ремонта"].iloc[0] == krs.KIND_CT
    calc = krs.compute_priorities(data)
    c = calc.set_index(["Скважина", "Год"])
    assert c.loc[(12, 2023), "Потребность (расчёт)"] == 2 and c.loc[(12, 2024), "Потребность (расчёт)"] == 1
    assert pd.isna(c.loc[(11, 2023), "Потребность (расчёт)"])    # фильтр 2023 после ревизии 2022-23
    assert c.loc[(12, 2023), "Совпало"] == "да"


def test_new_revisions_and_manual(tmp_path):
    src = tmp_path / "t.xlsx"
    make_table(src)
    data, grid = krs.parse(str(src))
    new = pd.DataFrame({"Скважина": [11, 13, 99], "Песок": ["Да", "Да", "Нет"], "Примечание": ["", "", ""]})
    notes = krs.add_revisions(data, new, "2025-2026")
    assert any("99" in n[1] for n in notes)
    calc = krs.compute_priorities(data)
    prio = krs.final_priorities(data, calc, {13: 1})
    p = prio.set_index("Скважина")
    assert p.loc[11, "Приоритет итог"] == 2 and p.loc[13, "Приоритет (расчёт)"] == 3
    assert p.loc[13, "Приоритет итог"] == 1                       # ручное главнее расчёта
    assert p.loc[12, "Приоритет итог"] == 1
    out = tmp_path / "upd.xlsx"
    krs.write_updated_table(grid, data.columns, data, prio, "2025-2026", str(out))
    ws = openpyxl.load_workbook(out).active
    assert ws.cell(5, 11).value == "х" and ws.cell(3, 11).value == "х"


def test_module_runs(tmp_path, monkeypatch):
    src = tmp_path / "t.xlsx"
    make_table(src)
    kru = tmp_path / "kru.csv"
    kru.write_text("№ скв;Песок\n11;да\n12;нет\n", encoding="utf-8")
    spec = importlib.util.spec_from_file_location("krs_mod", os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "pxg_base", "modules", "Анализ_технического_состояния_скважин_КРС.py"))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    monkeypatch.chdir(tmp_path)
    assert mod.main(["--table", str(src), "--kru", str(kru), "--season", "2025-2026"]) == 0
    x = pd.read_excel(tmp_path / "КРС_для_Атласа.xlsx", sheet_name=None)
    assert {"Ремонты", "Ревизии КРУ", "Приоритет", "Потребность", "КР по годам", "Проверка"} <= set(x)
    p = x["Приоритет"].set_index("Скважина")
    p.loc[11, "Приоритет вручную"] = 2
    with pd.ExcelWriter(tmp_path / "prev.xlsx") as xw:
        p.reset_index().to_excel(xw, sheet_name="Приоритет", index=False)
    assert krs.read_manual(str(tmp_path / "prev.xlsx")) == {11: 2}
    assert (tmp_path / "Таблица_анализа_обновлённая.xlsx").exists()


def make_word_revision(path, season="2025/26"):
    """Минимальный .docx: абзацы и таблица «№ скв.» с меткой ГСП над ней — как ревизия КРУ в Газовом Атласе."""
    import zipfile
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    p = lambda t: '<w:p><w:r><w:t>%s</w:t></w:r></w:p>' % t
    row = lambda *c: "<w:tr>%s</w:tr>" % "".join("<w:tc>%s</w:tc>" % p(x) for x in c)
    body = (p("Ревизия КРУ, сезон %s" % season) + p("ГСП-1") +
            "<w:tbl>" + row("№ скв.", "Описание ревизии") + row("11", "разрушение фильтра, песок") + row("13", "износ") + "</w:tbl>" +
            p("ГСП 2") + "<w:tbl>" + row("№ скв.", "Описание ревизии") + row("99", "износ") + "</w:tbl>")
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("word/document.xml", '<w:document xmlns:w="%s"><w:body>%s</w:body></w:document>' % (W, body))


def test_word_revision_like_atlas(tmp_path):
    doc = tmp_path / "rev.docx"
    make_word_revision(doc)
    new = krs.read_new_revisions(str(doc))
    assert sorted(new["Скважина"]) == [11, 13, 99] and set(new["Песок"]) == {"Да"}
    assert new.attrs["season_hints"] == ["2025-2026"]
    assert set(krs.read_new_revisions(str(doc), filter_regex="песок")["Скважина"]) == {11}
    with pytest.raises(ValueError):
        krs.read_new_revisions(str(doc), tables=[7])


def test_module_word_revision(tmp_path, monkeypatch):
    src = tmp_path / "t.xlsx"
    make_table(src)
    doc = tmp_path / "rev.docx"
    make_word_revision(doc)
    spec = importlib.util.spec_from_file_location("krs_mod2", os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "pxg_base", "modules", "Анализ_технического_состояния_скважин_КРС.py"))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    monkeypatch.chdir(tmp_path)
    assert mod.main(["--table", str(src), "--kru", str(doc), "--season", "2025-2026"]) == 0
    x = pd.read_excel(tmp_path / "КРС_для_Атласа.xlsx", sheet_name=None)
    rev = x["Ревизии КРУ"]
    assert set(rev[rev["Сезон"] == "2025-2026"]["Скважина"]) == {11, 13, 99}
    p = x["Приоритет"].set_index("Скважина")
    assert p.loc[13, "Приоритет (расчёт)"] == 3 and p.loc[11, "Приоритет (расчёт)"] == 2
    from pxg_base import checks
    rep = checks.run("Анализ_технического_состояния_скважин_КРС", {"table": str(src), "kru": str(doc), "season": "2025-2026"})
    assert not any(i.level == "ошибка" for i in rep.issues)


def test_check_reports_notes(tmp_path):
    from pxg_base import checks
    src = tmp_path / "t.xlsx"
    make_table(src)
    rep = checks.run("Анализ_технического_состояния_скважин_КРС", {"table": str(src), "season": "2025/26"})
    assert any(i.code == "SEASON" for i in rep.issues)
    rep = checks.run("Анализ_технического_состояния_скважин_КРС", {"table": str(src), "season": "2025-2026"})
    assert not any(i.level == "ошибка" for i in rep.issues)
