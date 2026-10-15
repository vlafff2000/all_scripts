"""Запуск модулей из формы (как в веб-интерфейсе) на небольших таблицах: без tkinter, без вопросов в консоли."""
import os
import sys
import time
from datetime import datetime

import pytest

pytest.importorskip("starlette")
pd = pytest.importorskip("pandas")
pytest.importorskip("openpyxl")
pytest.importorskip("xlsxwriter")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from starlette.testclient import TestClient  # noqa: E402

from pxg_base.api import app  # noqa: E402
from pxg_base.registry import MODULES  # noqa: E402
from pxg_base.testdata.flows import make_flow_tree  # noqa: E402
from pxg_base.webspec import SPECS, get  # noqa: E402


def run_module(module, params, out_dir, timeout=180):
    client = TestClient(app)
    job = client.post("/api/jobs", json={"module": module, "params": dict(params, out_dir=str(out_dir))}).json()
    assert "error" not in job, job
    deadline = time.time() + timeout
    while time.time() < deadline:
        job = client.get("/api/jobs/" + job["id"]).json()
        if job["status"] != "running":
            assert job["status"] == "done", "\n".join(job["log"][-20:])
            return job
        time.sleep(0.4)
    raise AssertionError("модуль не завершился")


def test_every_module_has_a_form():
    assert {s.module for s in SPECS} == {m[1] for m in MODULES}
    for spec in SPECS:
        ids = {p.id for p in spec.params}
        used = set(spec.answers) | {p for p, _ in spec.args} | set(spec.dialogs) | set(spec.strings) | {p for p, _ in spec.env}
        assert used <= ids, (spec.module, used - ids)


def test_vlookup_single_and_composite_key(tmp_path):
    main = tmp_path / "main.xlsx"
    look = tmp_path / "look.xlsx"
    pd.DataFrame({"Скв": ["1", "2", "3"], "Год": [2020, 2020, 2021], "A": [1, 2, 3]}).to_excel(main, index=False)
    pd.DataFrame({"Номер": ["1", "2", "2"], "Y": [2020, 2020, 2021], "B": [10, 20, 99]}).to_excel(look, index=False)
    look1 = tmp_path / "look1.xlsx"
    pd.DataFrame({"Номер": ["1", "2"], "B": [10, 20]}).to_excel(look1, index=False)
    job = run_module("Сопоставление_таблиц_ВПР", {"main_file": str(main), "lookup_file": str(look1), "mode": "1", "main_key": "Скв",
                             "lookup_key": "Номер", "columns": "B", "output": "r1.xlsx"}, tmp_path / "o1")
    df = pd.read_excel(tmp_path / "o1" / "r1.xlsx")
    assert "r1.xlsx" in job["files"] and "B" in df.columns and df["B"].notna().sum() == 2
    run_module("Сопоставление_таблиц_ВПР", {"main_file": str(main), "lookup_file": str(look), "mode": "2", "main_key": "Скв", "lookup_key": "Номер",
                       "pairs": "Год = Y", "columns": "B", "output": "r2.xlsx"}, tmp_path / "o2")
    df = pd.read_excel(tmp_path / "o2" / "r2.xlsx")
    assert df["B"].notna().sum() == 2 and df.loc[df["Скв"] == 2, "B"].iloc[0] == 20


def test_pressure_table(tmp_path):
    rows = [[None, "Скв. №56", None, "Скв. №83", None, None], [None, "Ру", "Рпл", "Ру", "Рпл", "Рср"],
            ["2020-01-01", 1, 2, 3, 4, 5], ["2020-01-02", 6, 7, 8, 9, 10]]
    src = tmp_path / "p.xlsx"
    pd.DataFrame(rows).to_excel(src, header=False, index=False)
    run_module("Итоговая_таблица_давлений_2006_2025", {"file": str(src)}, tmp_path / "o")
    df = pd.read_excel(tmp_path / "o" / "БД_давления_2006-2025.xlsx")
    assert len(df) == 4 and set(df["Номер ГСП"]) == {1, 2}


def test_average_flow_table(tmp_path):
    days = pd.date_range("2024-04-01", "2024-10-31", freq="7D")
    db = pd.DataFrame([{"Дата": d, "Скважина": w, "Источник": g, "Суточный расход газа": q * 1000}
                       for d in days for w, g, q in (("1", "ГСП-1", 120), ("2", "ГСП-1", 260), ("3", "ГСП-2", 40))])
    src = tmp_path / "db.xlsx"
    with pd.ExcelWriter(src) as w:
        db.to_excel(w, sheet_name="Закачка", index=False)
    job = run_module("Таблица_среднесуточных_расходов_из_базы", {"year": "2024", "file": str(src), "sheet": "Закачка"}, tmp_path / "o")
    assert "результат_приемистость_по_скважинам_2024.xlsx" in job["files"]


def test_interannular_analysis_and_collection(tmp_path):
    db = pd.DataFrame([{"дата": "2024-01-31", "год": 2024, "месяц": 1, "сезон": s, "номер_скважины": w, "расход_газа_МК_сут": q, "расход_газа_МК_мес": q * 30, "давление_МК": p}
                       for s in ("2023-2024", "2024-2025") for w, q, p in ((1, 0, 0), (2, 15, 10), (3, 40, 31))])
    src = tmp_path / "mk.xlsx"
    db.to_excel(src, index=False)
    job = run_module("Анализ_межколонных_давлений_для_авторского_надзора", {"db": str(src), "seasons": "2024-2025"}, tmp_path / "o")
    assert any(f.startswith("анализ_МКД_МКП_2024-2025") for f in job["files"]), job["log"][-8:]


def test_update_flow_db_with_new_files(tmp_path):
    tree = make_flow_tree(tmp_path / "data")
    built = run_module("Создание_базы_данных_расходов", {"root": str(tree.root), "periods": str(tree.periods_file)}, tmp_path / "db")
    db = tmp_path / "db" / "Сводка_закачка_отбор_обновленный_скрипт.xlsx"
    assert db.name in built["files"]
    folder = next(p for p in (tree.root / "Отбор").rglob("*") if p.is_dir() and any(p.glob("*.xlsx")))
    job = run_module("Дополнение_базы_данных_расходов", {"db": str(db), "max_date": "01.12.2023", "periods": str(tree.periods_file),
                                             "kind": "отбор", "folder": str(folder)}, tmp_path / "upd")
    assert any("ШАГ 5" in line for line in job["log"]) and any("ГОТОВО" in line for line in job["log"])


def test_extract_schedule_keywords(tmp_path):
    src = tmp_path / "sched.inc"
    src.write_text("DATES\n 01 JAN 2024 /\n/\n\nWELSPECS\n'56' 'G' 1 1 /\n/\n\nWRONG\n/\n", encoding="utf-8")
    out = tmp_path / "o"
    job = run_module("Извлечение_ключевых_слов_из_schedule", {"file": str(src), "output": "kw.inc"}, out)
    assert "WELSPECS" in (out / "kw.inc").read_text(encoding="utf-8"), "\n".join(job["log"][-15:])


@pytest.mark.parametrize("module", ["Дополнение_базы_данных_расходов", "Сбор_данных_по_межколонным_давлениям",
                                    "Анализ_межколонных_давлений_для_авторского_надзора"])
def test_no_tkinter_windows(module):
    from pathlib import Path
    src = (Path(__file__).resolve().parents[1] / "pxg_base" / "modules" / (module + ".py")).read_text(encoding="utf-8")
    assert "tkinter" not in src


def test_interannular_analysis_console_args(tmp_path):
    """Консольный запуск без формы: параметры аргументами, без окон."""
    import subprocess
    db = pd.DataFrame([{"дата": "2024-01-31", "год": 2024, "месяц": 1, "сезон": "2024-2025", "номер_скважины": w,
                        "расход_газа_МК_сут": q, "расход_газа_МК_мес": q * 30, "давление_МК": p} for w, q, p in ((2, 15, 10), (3, 40, 31))])
    src = tmp_path / "mk.xlsx"
    db.to_excel(src, index=False)
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    mod = os.path.join(root, "pxg_base", "modules", "Анализ_межколонных_давлений_для_авторского_надзора.py")
    r = subprocess.run([sys.executable, mod, "--db=" + str(src), "--seasons=2024-2025", "--save=да"], cwd=str(tmp_path),
                       env=dict(os.environ, PYTHONPATH=root, PYTHONIOENCODING="utf-8"), stdin=subprocess.DEVNULL,
                       stdout=subprocess.PIPE, stderr=subprocess.STDOUT, universal_newlines=True, encoding="utf-8")
    assert r.returncode == 0, r.stdout
    assert (tmp_path / "анализ_МКД_МКП_2024-2025.xlsx").exists()


def test_trim_unfilled_days_future_templates():
    from pxg_core.расходы_файлы import trim_unfilled_days
    days = pd.to_datetime(["2026-10-01", "2026-10-02", "2026-10-03", "2026-10-04", "2026-10-05"])
    df = pd.DataFrame({"Скважина": ["1"] * 5, "Дата": days, "Часовой расход газа": [5.0, 0.0, 7.0, 0.0, 0.0],
                       "Время работы": [24.0, 0.0, 24.0, 0.0, 0.0]})
    out = trim_unfilled_days(df)
    assert list(out["Дата"].dt.day) == [1, 2, 3]            # хвост-шаблон убран, простой внутри остался
    assert list(trim_unfilled_days(df, "02.10.2026")["Дата"].dt.day) == [1]   # граница 02.10, дальше данных нет
    assert len(trim_unfilled_days(df.assign(**{"Часовой расход газа": 0.0, "Время работы": 0.0}))) == 5


def _plant_bad_dates(tree):
    """В первой книге: в заголовке листа «Январь» одна дата → 1970 (между соседними), другая → 2203 (с краю соседей нет)."""
    import openpyxl
    book = next(p for p in sorted((tree.root / "Отбор").rglob("*.xlsx")))
    wb = openpyxl.load_workbook(str(book))
    ws = wb["Январь"]
    hdr = next(r for r in range(1, 20) if "скв" in str(ws.cell(r, 1).value or ws.cell(r, 2).value or "").lower())
    col0 = 1 if "скв" in str(ws.cell(hdr, 1).value or "").lower() else 2
    ws.cell(hdr, col0 + 5).value = datetime(1970, 1, 1)
    ws.cell(hdr, col0 + 20).value = datetime(2203, 1, 21)
    wb.save(str(book))
    return book, hdr, col0


def test_bad_header_dates_are_fixed_or_excluded_and_reported(tmp_path):
    tree = make_flow_tree(tmp_path / "data")
    book, hdr, col0 = _plant_bad_dates(tree)
    job = run_module("Создание_базы_данных_расходов", {"root": str(tree.root), "periods": str(tree.periods_file)}, tmp_path / "db")
    assert "Даты_к_проверке.xlsx" in job["files"], job["log"][-10:]
    db = pd.read_excel(tmp_path / "db" / "Сводка_закачка_отбор_обновленный_скрипт.xlsx", sheet_name="Отборы")
    years = pd.to_datetime(db["Дата"], errors="coerce").dt.year
    assert years.notna().all() and years.min() >= 2023 and years.max() <= 2024
    rep = pd.read_excel(tmp_path / "db" / "Даты_к_проверке.xlsx", sheet_name="Замечания")
    assert set(rep["Что сделано"]) <= {"исправлено автоматически", "исключено"} and len(rep) == 2
    assert book.stem in set(rep["Файл"].str.replace(".xlsx", "", regex=False))
    # решение пользователя: вторую дату задать вручную
    row = rep[rep["Значение"].str.endswith("2203")].index[0]
    rep["Решение"] = rep["Решение"].astype(object)
    rep.loc[row, "Решение"] = "21.01.2024"
    fixes = tmp_path / "fixes.xlsx"
    with pd.ExcelWriter(fixes) as xw:
        rep.to_excel(xw, sheet_name="Замечания", index=False)
    run_module("Создание_базы_данных_расходов", {"root": str(tree.root), "periods": str(tree.periods_file),
                                             "date_fixes": str(fixes)}, tmp_path / "db2")
    rep2 = pd.read_excel(tmp_path / "db2" / "Даты_к_проверке.xlsx", sheet_name="Замечания")
    assert "исправлено вручную" in set(rep2["Что сделано"])


def test_date_issues_api_and_fixes_file(tmp_path):
    tree = make_flow_tree(tmp_path / "data")
    _plant_bad_dates(tree)
    job = run_module("Создание_базы_данных_расходов", {"root": str(tree.root), "periods": str(tree.periods_file)}, tmp_path / "db")
    client = TestClient(app)
    issues = client.get("/api/jobs/%s/date-issues" % job["id"]).json()["issues"]
    assert len(issues) == 2 and issues[0]["fragment"]
    i = issues[0]
    saved = client.post("/api/jobs/%s/date-fixes" % job["id"], json={"fixes": [
        {"file": i["Файл"], "sheet": i["Лист"], "cell": i["Ячейка"], "decision": "исключить"}]}).json()
    assert saved["count"] == 1 and os.path.isfile(saved["path"])
    job2 = run_module("Создание_базы_данных_расходов", {"root": str(tree.root), "periods": str(tree.periods_file),
                                                      "date_fixes": saved["path"]}, tmp_path / "db2")
    rep = pd.read_excel(tmp_path / "db2" / "Даты_к_проверке.xlsx", sheet_name="Замечания")
    assert "исключено" in set(rep["Что сделано"])


def test_zamery_old_template_latin_headers_and_section_in_column_a(tmp_path):
    """Старые шаблоны (2016–2017): «Pуст», «Hст», «Pпл» латиницей, подпись раздела в столбце A, раздел действует на обе таблицы."""
    from pxg_core import замеры
    head = ["№№ п/п", "№№ скв.", "L", "Дата замера", "Pуст., кг/см2", "Hст., м", "ES", "A", "плот   ность", "Pпл., кг/см2"]
    rows = [[None] * 17, ["Результаты замеров за апрель 2017г"] + [None] * 16, head + head,
            ["Эксплуатационные скважины"] + [None] * 16,
            [1, 56, 769, datetime(2017, 4, 28), 52.0, None, None, None, None, 55.0, 5, 237, 765, datetime(2017, 4, 28), 52.5, None, None, None, None, 55.6, None],
            ["Наблюдательные скважины"] + [None] * 20,
            [1, 116, 7, datetime(2017, 4, 19), None, 47.0, 1.06, None, 1.11, 82.3, 2, 2, 7, datetime(2017, 4, 18), None, 136.4, None, None, None, 73.5, None]]
    path = tmp_path / "old.xlsx"
    pd.DataFrame(rows).to_excel(path, header=False, index=False)
    df = замеры.to_frame(замеры.parse_file(str(path)).rows).set_index("Скважина")
    assert df.loc[237, "Категория"] == "Эксплуатационные" and df.loc[237, "Руст"] == 52.5
    assert df.loc[2, "Категория"] == "Наблюдательные" and df.loc[116, "Уровень жидкости"] == 47.0
    assert df.loc[237, "L"] == 765 and df.loc[116, "ES"] == 1.06 and df.loc[116, "Плотность воды"] == 1.11


def test_office_converter_batch_and_fallback(tmp_path):
    import pytest
    from pxg_core import office_convert as oc
    assert oc.pair_supported("xls", "xlsx")[0] and not oc.pair_supported("xlsx", "docx")[0] and oc.pair_supported("docx", "pdf")[0]
    assert not oc.pair_supported("pdf", "xlsx")[0]
    soffice = oc.find_soffice()
    if not soffice:
        pytest.skip("нет LibreOffice: для подготовки .xls")
    src = tmp_path / "in" / "sub"
    src.mkdir(parents=True)
    pd.DataFrame({"Скв": ["1", "2"], "P": [1.5, 2.5]}).to_excel(tmp_path / "a.xlsx", index=False)
    xls = src / "a.xls"
    tmp = tmp_path / "tmp"
    tmp.mkdir()
    import shutil
    shutil.move(oc._soffice_convert(soffice, str(tmp_path / "a.xlsx"), "xls", str(tmp)), xls)
    # без LibreOffice: запасной путь на Python, подпапки и «~$»-файлы
    (src / "~$a.xls").write_bytes(b"x")
    files = oc.collect([str(tmp_path / "in")])
    assert [os.path.basename(f) for f, _ in files] == ["a.xls"]
    f, base = files[0]
    dst = oc.dest_path(f, base, "xlsx", str(tmp_path / "out"))
    res = oc.convert_file(f, "xlsx", dst, None)
    assert res["status"] == "ok" and res["engine"] == "xlrd+openpyxl" and dst.endswith(os.path.join("sub", "a.xlsx"))
    assert pd.read_excel(dst)["P"].tolist() == [1.5, 2.5]
    assert oc.convert_file(f, "xlsx", dst, None)["status"] == "skip"
    assert oc.convert_file(f, "docx", dst + "x", None)["status"] == "error"
    # LibreOffice: Excel → pdf, лист → csv
    res = oc.convert_file(f, "pdf", str(tmp_path / "a.pdf"), soffice)
    assert res["status"] == "ok" and (tmp_path / "a.pdf").stat().st_size > 0
