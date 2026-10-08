"""Режим «история» «Скедул ПХГ» (А13): шаги по суткам и по датам замеров, поправка по ПЗРГ, сшивка; паритет со старыми скриптами."""
import contextlib
import importlib.util
import io
import os
import sys
from datetime import date, timedelta

import pandas as pd
import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from schedule_pxg import forecast as fc  # noqa: E402
from schedule_pxg import historymode as hm  # noqa: E402
from schedule_pxg.project import Project  # noqa: E402

OLD_DIR = os.path.join(ROOT, "apps", "schedule_tr")
OLD_DAILY = os.path.join(OLD_DIR, "Schedule_по_пропорциональным_коэффициентам_шаг_1_сутки.py")
OLD_DATES = os.path.join(OLD_DIR, "Schedule_по_датам_замеров_давлений.py")
PROD, INJ = "отбор", "закачка"


def old(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    try:
        spec.loader.exec_module(mod)
    except Exception as e:
        pytest.skip("старый скрипт не импортируется: %s" % e)
    return mod


def frame(rows):
    """rows: (скважина, дата, расход, вид)."""
    return pd.DataFrame({"well": [r[0] for r in rows], "date": pd.to_datetime([r[1] for r in rows]),
                         "rate": [float(r[2]) for r in rows], "hours": 24.0, "kind": [r[3] for r in rows]})


def old_periods():
    """PERIODS в виде старых скриптов: словари с datetime и типом prod/inj/none, конец — за сутки до начала следующего."""
    kinds = {INJ: "inj", PROD: "prod", "нейтральный": "none"}
    ends = [PERIODS[1][0] - timedelta(days=1), PERIODS[2][0] - timedelta(days=1), date(2026, 12, 31)]
    return [{"start": pd.Timestamp(a).to_pydatetime(), "end": pd.Timestamp(e).to_pydatetime(), "type": kinds[k]}
            for (a, k), e in zip(PERIODS, ends)]


def old_frames(df):
    """Таблицы листов «Отборы» и «Закачка» в виде старых скриптов."""
    def one(kind):
        d = df[df["kind"] == kind]
        return pd.DataFrame({"Скважина": d["well"].values, "Дата": d["date"].values, "Суточный_расход_газа": d["rate"].values})
    return one(PROD), one(INJ)


def blocks(text):
    """Схедул → [(дата, секция, отсортированные строки)] — без учёта порядка скважин и разделителей."""
    out, cur, sect = [], None, None
    for ln in text.splitlines():
        s = ln.strip()
        if not s or s.startswith("---") or s == "/":
            sect = None if s == "/" else sect
            continue
        if s in ("DATES", "WELOPEN", "WCONHIST", "WCONINJH", "WEFAC"):
            sect = s
            if s == "DATES":
                cur = None
            continue
        if sect == "DATES" and cur is None:
            cur = s.replace("/", "").split()
            out.append([tuple(cur), {}])
            sect = None
            continue
        if cur is not None and sect:
            out[-1][1].setdefault(sect, []).append(s)
    return [(d, {k: sorted(v) for k, v in secs.items()}) for d, secs in out]


def project():
    p = Project()
    p.add_group("ГСП 1")
    for w in ("1", "2", "3", "54", "80"):
        p.add_well(w, group="ГСП 1")
    return p


# ---------------------------------------------------------------- читалки

def test_parse_dates_and_periods_report_bad_lines():
    ds, n = hm.parse_dates("# даты\n01.05.2026\n15.05.2026\nпривет\n01.05.2026\n")
    assert ds == [date(2026, 5, 1), date(2026, 5, 15)] and len(n) == 1 and "привет" in n[0]
    ps, n = hm.parse_periods("01.05.2026 inj\n01.10.2026 prod\n01.04.2027 none\nплохо\n")
    assert ps == [(date(2026, 5, 1), INJ), (date(2026, 10, 1), PROD), (date(2027, 4, 1), "нейтральный")] and len(n) == 1
    assert hm.kind_at(date(2026, 4, 1), ps) == "нейтральный" and hm.kind_at(date(2026, 10, 1), ps) == PROD


def test_periods_from_history_merge_and_gaps():
    h = frame([("1", "2026-05-01", 100, INJ), ("1", "2026-05-02", 100, INJ), ("1", "2026-05-03", 0, INJ),
               ("1", "2026-05-04", 90, INJ), ("1", "2026-10-01", 50, PROD)])
    assert hm.periods_from_history(h) == [(date(2026, 5, 1), INJ), (date(2026, 5, 3), "нейтральный"), (date(2026, 5, 4), INJ),
                                          (date(2026, 10, 1), PROD)]


def test_read_pzrg_abs_and_sorted(tmp_path):
    f = tmp_path / "pzrg.xlsx"
    pd.DataFrame([["2026-05-02", "x", -200.0], ["2026-05-01", "x", 100.0], ["не дата", "x", 5]]).to_excel(f, header=False, index=False)
    z = hm.read_pzrg(str(f))
    assert list(z["rate"]) == [100.0, 200.0] and z["date"].is_monotonic_increasing


def test_split_wells_halves_rate():
    h = frame([("54/80", "2026-05-01", 100, INJ), ("1", "2026-05-01", 60, INJ)])
    s = hm.split_wells(h)
    assert sorted(zip(s["well"], s["rate"])) == [("1", 60.0), ("54", 50.0), ("80", 50.0)]
    assert len(hm.split_wells(h, {})) == 2


# ---------------------------------------------------------------- шаг сутки: паритет

def sample():
    rows = []
    for k in range(6):
        d = date(2026, 5, 1) + timedelta(days=k)
        rows += [("1", d, 100000 + 1000 * k, INJ), ("2", d, 50000 + k, INJ), ("3", d, 0 if k == 2 else 150000, INJ), ("54/80", d, 200000, INJ)]
    for k in range(4):
        d = date(2026, 10, 1) + timedelta(days=k)
        rows += [("1", d, 300000, PROD), ("2", d, 20000 + 7 * k, PROD)]
    return frame(rows)


PERIODS = [(date(2026, 5, 1), INJ), (date(2026, 5, 5), "нейтральный"), (date(2026, 10, 1), PROD)]


def test_daily_text_parity_with_old_script(tmp_path):
    o = old(OLD_DAILY, "old_daily")
    h = sample()
    res = hm.build_history(h, mode="daily", periods=PERIODS)
    prod, inj = old_frames(hm.split_wells(h))
    out = str(tmp_path / "old.inc")
    with contextlib.redirect_stdout(io.StringIO()):
        o.create_include_file(prod, inj, old_periods(), out)
    new = hm.render(res)
    assert blocks(new) == blocks(open(out, encoding="utf-8").read())


def test_daily_pzrg_parity_with_old_correction():
    o = old(OLD_DAILY, "old_daily2")
    h = hm.split_wells(sample())
    inj = h[h["kind"] == INJ]
    pz = pd.DataFrame({"date": pd.to_datetime(["2026-05-01", "2026-05-02", "2026-05-03", "2026-05-04"]),
                       "rate": [450000.0, 200000.0, 0.0, 600000.0]})
    got, log, notes = hm.correct_pzrg(inj, pz)
    old_df = old_frames(h)[1]
    pz_old = pz.rename(columns={"date": "Дата", "rate": "Суточный_расход_ПЗРГ"})
    with contextlib.redirect_stdout(io.StringIO()):
        fixed = o.calculate_correction_coefficients(old_df, pz_old, None)
    assert list(got["rate_corr"]) == pytest.approx(list(fixed["Суточный_расход_газа_скорректированный"]))
    methods = {r["start"]: r["method"] for r in log}
    assert methods["2026-05-03"] == "SKIP" and methods["2026-05-01"] in ("IN_RANGE_ONLY", "ALL_WELLS")
    assert hm.log_summary(log)["skipped"] == 3  # ПЗРГ = 0 и два дня без ПЗРГ


def test_correct_set_branches():
    new, info = hm._correct_set([100000.0, 200000.0, 10000.0], 400000.0, 80000, 600000)
    assert info["method"] == "IN_RANGE_ONLY" and new[2] == 10000.0 and sum(new) == pytest.approx(400000.0)
    new, info = hm._correct_set([100000.0, 10000.0], 5000.0, 80000, 600000)       # после вычета ≤ 0 — правятся все
    assert info["method"] == "ALL_WELLS" and sum(new) == pytest.approx(5000.0)
    new, info = hm._correct_set([1000.0, 3000.0], 8000.0, 80000, 600000)           # нет рабочих — правятся все
    assert info["method"] == "ALL_WELLS" and new == pytest.approx([2000.0, 6000.0])


# ---------------------------------------------------------------- по датам замеров: паритет

DATES = [date(2026, 5, 3), date(2026, 5, 5), date(2026, 10, 2)]


def test_dates_text_parity_with_old_script(tmp_path):
    o = old(OLD_DATES, "old_dates")
    h = sample()
    res = hm.build_history(h, mode="dates", model_dates=DATES, periods=PERIODS)
    prod, inj = old_frames(hm.split_wells(h))
    allv = sorted(set(pd.to_datetime(prod["Дата"]).tolist() + pd.to_datetime(inj["Дата"]).tolist()))
    to_dt = lambda d: pd.Timestamp(d).to_pydatetime()
    per = old_periods()
    steps = o.get_model_dates_and_periods([to_dt(d) for d in DATES], [v.to_pydatetime() for v in allv], per)
    out = str(tmp_path / "old.inc")
    with contextlib.redirect_stdout(io.StringIO()):
        o.create_include_file(prod, inj, steps, out)
    assert blocks(hm.render(res)) == blocks(open(out, encoding="utf-8").read())
    assert [i["from_file"] for i in res.info][:3] == [True, True, True] and not res.info[-1]["from_file"]


def test_dates_pzrg_single_period_matches_old_function():
    o = old(OLD_DATES, "old_dates2")
    h = hm.split_wells(sample())
    inj = h[h["kind"] == INJ]
    pz = pd.DataFrame({"date": pd.date_range("2026-05-01", "2026-05-06"), "rate": [450000.0] * 6})
    a, b = date(2026, 5, 1), date(2026, 5, 4)
    got, log, _ = hm.correct_pzrg(inj, pz, {INJ: [(a, b)]})
    od = old_frames(h)[1]
    with contextlib.redirect_stdout(io.StringIO()):
        fixed = o.calculate_correction_coefficients_for_period(
            od, pz.rename(columns={"date": "Дата", "rate": "Суточный_расход_ПЗРГ"}), pd.Timestamp(a), pd.Timestamp(b), None)
    assert list(got["rate_corr"]) == pytest.approx(list(fixed["Суточный_расход_газа_скорректированный"]))


def test_correction_windows_do_not_share_boundary_day():
    ms = hm.model_steps(DATES[:2], [date(2026, 5, 1) + timedelta(days=k) for k in range(5)], [(date(2026, 5, 1), INJ)])
    w = hm.correction_windows(ms)
    assert w[INJ][0] == (date(2026, 5, 1), date(2026, 5, 3)) and w[INJ][1] == (date(2026, 5, 4), date(2026, 5, 5))
    # поправка каждого периода сохраняется (в старом скрипте оставалась только последняя)
    h = frame([("1", date(2026, 5, 1) + timedelta(days=k), 100000, INJ) for k in range(5)])
    pz = pd.DataFrame({"date": pd.date_range("2026-05-01", "2026-05-05"), "rate": [200000.0, 200000.0, 200000.0, 150000.0, 150000.0]})
    got, log, _ = hm.correct_pzrg(h, pz, w)
    assert list(got["rate_corr"]) == pytest.approx([200000, 200000, 200000, 150000, 150000]) and len(log) == 2


def test_build_history_with_pzrg_notes_and_names():
    p = project()
    h = frame([("1", "2026-05-01", 100000, INJ), ("1", "2026-05-02", 100000, INJ), ("77", "2026-05-01", 90000, INJ)])
    pz = pd.DataFrame({"date": pd.to_datetime(["2026-05-01"]), "rate": [380000.0]})
    r = hm.build_history(h, p, "daily", pzrg=pz)
    txt = " ".join(r.notes)
    assert "77" in txt and "из самой истории" in txt and "пропущена" in txt and len(r.steps) == 2
    assert r.steps[0].rates["1"] == pytest.approx(r.steps[0].rates["1"]) and sum(r.steps[0].rates.values()) == pytest.approx(380000.0)
    with pytest.raises(ValueError):
        hm.build_history(h, p, "dates")
    assert hm.build_history(h.iloc[0:0], p).notes == ["История пуста — шагов нет"]


# ---------------------------------------------------------------- сшивка

def step(a, b, kind="закачка", rates=None):
    return fc.Step(a, b, (b - a).days + 1, kind, rates or {"1": 10.0})


def test_stitch_gap_overlap_and_text():
    hist = [step(date(2026, 5, 1), date(2026, 5, 3)), step(date(2026, 5, 4), date(2026, 5, 9))]
    fore = [step(date(2026, 5, 20), date(2026, 5, 31))]
    s = hm.stitch(hist, fore)
    assert s.gap == (date(2026, 5, 10), date(2026, 5, 19)) and s.steps[-2].kind == "нейтральный" and hm.stitch_issues(s.steps) == []
    ov = hm.stitch(hist, [step(date(2026, 5, 6), date(2026, 5, 31))])
    assert ov.cut == 0 and ov.steps[1].end == date(2026, 5, 5) and hm.stitch_issues(ov.steps) == [] and "обрезан" in " ".join(ov.notes)
    ov2 = hm.stitch(hist, [step(date(2026, 5, 2), date(2026, 5, 31))])
    assert ov2.cut == 1 and ov2.steps[0].end == date(2026, 5, 1)
    txt = fc.render_schedule(s.steps)
    assert txt.count("DATES") == len(s.steps) + 1 and "WELOPEN" in txt
    assert len(hm.stitch([], fore).notes) == 1


def test_dates_steps_are_contiguous_for_stitching():
    r = hm.build_history(sample(), mode="dates", model_dates=DATES, periods=PERIODS)
    assert hm.stitch_issues(r.steps) == []


# ---------------------------------------------------------------- API

@pytest.fixture
def client(tmp_path, monkeypatch):
    from starlette.testclient import TestClient
    from schedule_pxg import api
    from schedule_pxg import history as hist
    from tests.test_schedule_scenarios import _lib
    folder = str(tmp_path / "proj")
    monkeypatch.setattr(api, "FOLDER", folder)
    h = sample()
    xl = str(tmp_path / "hist.xlsx")
    pd.DataFrame({"Скважина": h["well"], "Дата": h["date"], "Расход": h["rate"], "Вид": h["kind"]}).to_excel(xl, index=False)
    p = project()
    p.techmaps = _lib()
    p.templates = {"т": hist.Template(name="т", well="Скважина", date="Дата", rate="Расход", kind="Вид", unit="м3/сут").to_dict()}
    p.save(folder)
    return TestClient(api.build_app()), xl, tmp_path


def test_api_history_summary_and_schedule(client):
    c, xl, tmp = client
    assert c.post("/api/history", json={}).status_code == 400
    assert c.post("/api/history", json={"files": [xl], "mode": "dates"}).status_code == 400
    (tmp / "dates.txt").write_text("03.05.2026\n05.05.2026\n02.10.2026\n", encoding="utf-8")
    (tmp / "per.txt").write_text("01.05.2026 inj\n05.05.2026 none\n01.10.2026 prod\n", encoding="utf-8")
    r = c.post("/api/history", json={"files": [xl], "mode": "dates", "dates_file": str(tmp / "dates.txt"), "periods_file": str(tmp / "per.txt")}).json()
    assert r["mode"] == "dates" and r["steps"] > 3 and r["correction"] is None
    r = c.post("/api/history", json={"files": [xl], "mode": "daily"}).json()
    assert r["steps"] == 10 and any("из самой истории" in n for n in r["notes"])
    t = c.post("/api/history/schedule", json={"files": [xl], "mode": "daily"})
    assert t.status_code == 200 and "WCONINJH" in t.text and "WCONHIST" in t.text
    assert c.post("/api/history", json={"files": [xl], "stitch": "нет"}).status_code == 404


def test_api_history_stitch_with_scenario(client, monkeypatch):
    from schedule_pxg import api, scenarios, techmap as tmod
    c, xl, _ = client
    p = api._project()
    tm = tmod.TechMap("Закачка", "закачка", ["Июнь"])
    tm.days = {"Июнь": 30}
    tm.volumes = {"ГСП 1": {"Июнь": 300.0}}
    p.techmaps = {"Закачка": tm.to_dict()}
    sc = scenarios.Scenarios()
    sc.add("Базовый", {"calendar": [scenarios.season_entry(2026, "Закачка")]})
    p.scenarios = sc.to_dict()
    p.save(api.FOLDER)
    r = c.post("/api/history", json={"files": [xl], "stitch": "Базовый"}).json()
    assert r["to"] == "2026-06-30" and r["stitch"] == [] and any("нейтральным шагом" in n for n in r["notes"])
    t = c.post("/api/history/schedule", json={"files": [xl], "stitch": "Базовый"}).text
    assert t.count("DATES") > 30 and "JUN" in t and t.rstrip().endswith("-" * 80)


def test_correction_xlsx_parity_with_old_report(tmp_path):
    """Excel-отчёт по поправкам: детали по скважинам и статистика совпадают с `create_correction_report` старого скрипта."""
    from openpyxl import load_workbook
    o = old(OLD_DAILY, "old_daily_report")
    h = hm.split_wells(sample())
    d = h[h["kind"] == INJ]
    pz = pd.DataFrame({"date": pd.to_datetime(["2026-05-01", "2026-05-02", "2026-05-03", "2026-05-04"]),
                       "rate": [450000.0, 200000.0, 0.0, 600000.0]})
    inj_only = sample()[sample()["kind"] == INJ]
    res = hm.build_history(inj_only, mode="daily", periods=PERIODS, pzrg=pz)
    assert res.data is not None and "rate_corr" in res.data
    new_path = hm.write_correction_xlsx(str(tmp_path / "new.xlsx"), res)
    # старый скрипт: тот же журнал в его формате и таблицы закачки/отбора после его поправки
    cols = ["Дата", "Метод_коррекции", "ПЗРГ_сут", "Скважин_в_диапазоне", "Скважин_вне_диапазона", "Сумма_в_диапазоне",
            "Сумма_вне_диапазона", "Коэффициент", "Итоговая_сумма_сут", "Расхождение_%", "Категория", "Комментарий"]
    rows = [[r["start"], r["method"], r["pzrg"], r["n_in"], r["n_out"], r["sum_in"], r["sum_out"], r["coef"], r["total_after"],
             r["discrepancy"], r["category"], r["comment"]] for r in res.log]
    logcsv = str(tmp_path / "log.csv")
    pd.DataFrame(rows, columns=cols).to_csv(logcsv, sep=";", index=False, encoding="utf-8")
    pz_old = pz.rename(columns={"date": "Дата", "rate": "Суточный_расход_ПЗРГ"})
    prod_old, inj_old = old_frames(h)
    with contextlib.redirect_stdout(io.StringIO()):
        inj_old = o.calculate_correction_coefficients(inj_old, pz_old, None)
        o.create_correction_report(logcsv, prod_old, inj_old, str(tmp_path / "old.xlsx"))
    a = load_workbook(new_path)
    b = load_workbook(str(tmp_path / "old.xlsx"))
    assert a.sheetnames[:2] == b.sheetnames[:2]

    def body(ws, first):
        return [[c for c in r] for r in ws.iter_rows(min_row=first, values_only=True) if any(v is not None for v in r)]
    na, nb = body(a["Детали по скважинам"], 4), body(b["Детали по скважинам"], 4)
    assert len(na) == len(nb) > 0
    key = lambda r: (r[0], r[1], r[2])  # noqa: E731
    for x, y in zip(sorted(na, key=key), sorted(nb, key=key)):
        assert x[:3] == y[:3] and x[6] == y[6]
        assert list(x[3:6]) + list(x[7:9]) == pytest.approx(list(y[3:6]) + list(y[7:9]))
    sa, sb = body(a["Статистика по периодам"], 4), body(b["Статистика по дням"], 4)
    assert len(sa) == len(sb) == 3
    for x, y in zip(sa, sb):      # ours: начало, конец, всего, в диапазоне, вне, ПЗРГ, до, после, расхождение, метод, коэффициент, категория
        assert x[0] == y[0] and x[2:5] == y[1:4] and x[9] == y[8] and x[11] == y[10]
        assert list(x[5:9]) + [x[10]] == pytest.approx(list(y[4:8]) + [y[9]])
    assert [r[2] for r in body(a["Сводная информация"], 7)] == [r[1] for r in body(b["Сводная информация"], 7)]  # методы по периодам


def test_correction_xlsx_empty_and_api(tmp_path):
    res = hm.build_history(sample(), mode="daily", periods=PERIODS)
    assert res.log == []
    p = hm.write_correction_xlsx(str(tmp_path / "e.xlsx"), res)
    assert os.path.isfile(p)
