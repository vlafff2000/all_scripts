"""Модуль «База_шахматка_Лена»: шахматка отбора, пересчёт Qгаз.пром, проверка исходников."""
from __future__ import annotations

import math
import os
import sys

import pytest

pd = pytest.importorskip("pandas")
openpyxl = pytest.importorskip("openpyxl")

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from pxg_base.checks import shakhmatka as check  # noqa: E402
from pxg_core import qc, шахматка  # noqa: E402
from tests.test_pxg_modules import run_module  # noqa: E402

Q_CALC = 0.0141 * 10 ** 2 * 50 * math.sqrt(1 + 0.002253 * 50)   # t = 24 ч: /24*t = 1


def make_otbor(path, bad_label=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Шахматка-газ. Журнал учёта"])
    ws.append([])
    months = ["Ноябрь", "Ноябрь", "Декабрь" if bad_label else "Ноябрь"]
    ws.append(["№№", "Показатели", "Ед.изм.", None] + months)
    ws.append(["скв.", None, None, None, 1, 2, 3])
    ws.append(["Всего по ПХГ:"])
    ws.append([None, "Qгаза отч", "т.м3", None, 100, 100, 100])
    for well, q_prom in (("1", 50.0), ("2", 40.0)):      # у скважины 2 в Qгаз.пром стоит неверное число
        ws.append([well, "Состояние", None, None, "Рб", "Рб", "Рб"])
        ws.append(["бл.1", "t работы", "час", None, 24, 24, 24])
        ws.append([None, "Диам. штуцера", "мм", None, 10, 10, 10])
        ws.append([None, "Ршт. блок", "кгс/см2", None, 50, 50, 50])
        ws.append([None, "qгаза (расч)", "т.м3/сут", None] + [Q_CALC] * 3)
        ws.append([None, "q газа пром", "т.м3/сут", None] + [q_prom] * 3)
    wb.save(path)


def test_recompute_k_and_difference(tmp_path):
    f = tmp_path / "otbor 2025.xlsx"
    make_otbor(f)
    sh = шахматка.Shakh(str(f), "отбор", 2025, str(tmp_path))
    assert [d.isoformat() for _c, d, _m in sh.day_cols] == ["2025-11-01", "2025-11-02", "2025-11-03"]
    df = шахматка.recompute(sh, sh.wells_table())
    day = df[df["Дата"] == df["Дата"].min()]
    assert day["Qгаз.пром (пересчёт)"].sum() == pytest.approx(100.0)       # сумма по скважинам = Qотч
    bad = df[df["Расхождение > допуска"]]
    assert set(bad["Скважина"]) == {"2"} and len(bad) == 3
    assert set(bad["Причина расхождения"]) == {"в ячейке число, а не формула"}


def test_month_label_does_not_move_dates(tmp_path):
    f = tmp_path / "o.xlsx"
    make_otbor(f, bad_label=True)
    sh = шахматка.Shakh(str(f), "отбор", 2025, str(tmp_path))
    assert sh.day_cols[-1][1].isoformat() == "2025-11-03"
    assert any(i["Вид"] == "Подпись месяца" for i in sh.issues)


def test_run_module_two_seasons_in_one_window(tmp_path):
    src = tmp_path / "src"
    src.mkdir()
    make_otbor(src / "Шахматка отбор 2025.xlsx")
    make_otbor(src / "Шахматка отбор 2026.xlsx")
    (src / "посторонний.txt").write_text("не Excel")
    files = "\n".join([str(src / "Шахматка отбор 2025.xlsx"), str(src / "Шахматка отбор 2026.xlsx"), str(src / "посторонний.txt")])
    assert шахматка.classify(str(src / "Шахматка отбор 2025.xlsx")) == "отбор"
    out = tmp_path / "out"
    run_module("База_шахматка_Лена", {"files": files, "output": "db.xlsx"}, out)
    book = pd.read_excel(out / "db.xlsx", sheet_name=None)
    assert {"Сводка", "Отбор_по_скважинам", "Отбор_ПХГ_по_суткам", "Скважины", "Контроль_формул"} <= set(book)
    wells = book["Отбор_по_скважинам"]
    assert len(wells) == 12                 # по 6 строк (2 скважины × 3 суток) из каждого сезона
    assert set(wells["Файл"]) >= {"Шахматка отбор 2026.xlsx"}
    assert (out / "db.sqlite").exists()
    rep = check.check_shakhmatka({"files": files})
    assert not [i for i in rep.issues if i.level == qc.ERROR and i.code != "FILE"]


def test_check_without_files():
    rep = check.check_shakhmatka({})
    assert any(i.level == qc.ERROR for i in rep.issues)


def test_calamine_panic_falls_back(tmp_path, monkeypatch):
    """Паника calamine (BaseException) не должна ронять чтение: берём openpyxl/xlrd."""
    import types
    from pxg_core import шахматка as sh

    class Panic(BaseException):
        pass

    class Boom:
        @staticmethod
        def from_path(_p):
            raise Panic("range start index 7 out of range")

    mod = types.ModuleType("python_calamine")
    mod.CalamineWorkbook = Boom
    monkeypatch.setitem(__import__("sys").modules, "python_calamine", mod)
    monkeypatch.delenv("PXG_EXCEL_ENGINE", raising=False)
    import openpyxl
    p = tmp_path / "a.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = 5
    wb.save(str(p))
    V, name = sh._values(str(p), 10, 10)
    assert V[0][0] == 5


def test_broken_autofilter_does_not_crash(tmp_path):
    """После LibreOffice у листа бывает автофильтр с битым ref: openpyxl падает, мы читаем копию без него."""
    import shutil
    import zipfile
    import openpyxl
    from pxg_core import шахматка as sh

    src = tmp_path / "ok.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = 7
    wb.active["B1"] = "=A1*2"
    wb.save(str(src))
    bad = tmp_path / "bad.xlsx"
    with zipfile.ZipFile(str(src)) as zi, zipfile.ZipFile(str(bad), "w") as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"</sheetData>", b'</sheetData><autoFilter ref="!!bad!!"/>')
            zo.writestr(it, data)
    # убеждаемся, что без обхода openpyxl действительно падает
    try:
        list(openpyxl.load_workbook(str(bad), read_only=True)["Sheet"].iter_rows(values_only=True))
        failed = False
    except ValueError:
        failed = True
    assert failed
    F = sh._formulas(str(bad), 10, 10)
    assert F[0][1] == "=A1*2"
    V, _ = sh._values_plain(str(bad), 10, 10)
    assert V[0][0] == 7


@pytest.mark.parametrize("name,year", [
    ("Шахм-газ 11 - 12.xlsx", 2011), ("Шахм-газ 13-14.xlsx", 2013), ("Шахм-газ копия 16-17.xlsx", 2016),
    ("Шахматка отбор 23-24 копия.xlsx", 2023), ("Шахматка газ 2017-18.xlsx", 2017), ("Шахматка газ 2019-2020 копия.xlsx", 2019),
    ("Шахматка отбор 2025-2026_06042026.xlsx", 2025), ("Шахматка закачка 2018г копия.xlsx", 2018), ("Шахматка закачка  2013.xlsx", 2013),
    ("Шахматка_отбор_2024-2025.xlsx", 2024), ("Шахматка 10-15.xlsx", None), ("отбор.xlsx", None)])
def test_infer_year(name, year):
    assert шахматка._infer_year(name) == year


def test_year_from_header_manual_and_unknown(tmp_path):
    src = tmp_path / "src"
    src.mkdir()
    make_otbor(src / "отбор без года.xlsx")                       # года нет ни в имени, ни в шапке
    wb = openpyxl.load_workbook(str(src / "отбор без года.xlsx"))
    wb.active["A1"] = "Журнал за сезон 2019"                      # год в шапке листа
    wb.save(str(src / "отбор в шапке.xlsx"))
    sh = шахматка.Shakh(str(src / "отбор в шапке.xlsx"), "отбор", None, str(tmp_path))
    assert sh.day_cols[0][1].year == 2019
    log = []
    res = шахматка.build([str(src / "отбор без года.xlsx")], str(tmp_path / "a.xlsx"), log=log.append)
    assert any("Год не определён" == v for v in res["issues"]["Вид"])
    assert any("год первых суток не определён" in m for m in log)
    res = шахматка.build([str(src / "отбор без года.xlsx")], str(tmp_path / "b.xlsx"), years="без года = 2014")
    assert not len(res["issues"][res["issues"]["Вид"] == "Год не определён"])
    book = pd.read_excel(str(tmp_path / "b.xlsx"), sheet_name="Отбор_по_скважинам")
    assert str(book["Дата"].min())[:4] == "2014"
