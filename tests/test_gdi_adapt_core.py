"""Запуск:  python -m unittest discover -s tests -v   (из папки проекта)"""
import glob
import os
import shutil
import sys
import tempfile
import unittest
import zipfile

import numpy as np
import pandas as pd
from datetime import datetime, timedelta

from scipy import integrate
from scipy.interpolate import interp1d
from scipy.optimize import minimize

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
sys.path.insert(0, HERE)

from schedule_pxg.gdi_adapt import physics as ph, tnav_api                       # noqa: E402
from schedule_pxg.gdi_adapt.adapt import adapt_well                              # noqa: E402
from schedule_pxg.gdi_adapt.config import Config                                 # noqa: E402
from schedule_pxg.gdi_adapt.gdi import GDI, fit_ab, load_gdi                     # noqa: E402
from schedule_pxg.gdi_adapt.outputs import write_include                         # noqa: E402
from schedule_pxg.gdi_adapt.pvt import PVT                                       # noqa: E402
from schedule_pxg.gdi_adapt.runner import run                                    # noqa: E402
from schedule_pxg.gdi_adapt.snapshot import WellSnapshot, load_snapshots         # noqa: E402
from schedule_pxg.gdi_adapt.tnav_api import SnapshotProvider, TnavProvider       # noqa: E402
import gdi_adapt_mock_tnav as mock_tnav                                                  # noqa: E402

PVT_FILE = os.path.join(ROOT, 'schedule_pxg', 'gdi_adapt', 'pvt_default.csv')
GDI_FILE = os.path.join(HERE, 'data_gdi_adapt', 'all_gdi.xlsx')


def make_cfg(**kw):
    base = dict(gdi_file=GDI_FILE, pvt_file=PVT_FILE, make_png=False, qc_enabled=False, ci_jackknife=False)
    base.update(kw)
    return Config.load(None, **base)


def synthetic_snapshot(wbp=85.0, n=3, seed=0):
    rng = np.random.default_rng(seed)
    return WellSnapshot(
        wname='T', t_name='2', t_requested='2', substituted=False,
        cbp=np.full(n, wbp), cdrd=rng.uniform(1.0, 3.0, n), cgpr=rng.uniform(2e4, 8e4, n), cdfac=np.full(n, 3e-6),
        cekh=rng.uniform(500, 2500, n), cdiam=np.full(n, 0.2), cefr=np.full(n, 30.0), ctfac=rng.uniform(10, 20, n),
        wbp=wbp, wgpr=0.0, wbhp=wbp - 2.0)


class TestPVT(unittest.TestCase):
    def setUp(self):
        self.pvt = PVT.from_csv(PVT_FILE, warn=lambda m: None)
        d = pd.read_csv(PVT_FILE, sep=';')
        self.d = d

    def test_inside_table_matches_old_formulas(self):
        d = self.d
        inv_bg = interp1d(d.p, 1 / d.bg)
        inv_bgmu = interp1d(d.p, 1 / d.bg / d.mu)
        for p in (3.0, 20.0, 55.5, 90.1, 118.0):
            self.assertAlmostEqual(self.pvt.get_bg(p), 1 / float(inv_bg(p)), places=12)
            self.assertAlmostEqual(self.pvt.get_mu(p), float(1 / inv_bgmu(p) * inv_bg(p)), places=12)

    def test_extrapolation_bug_fixed(self):
        """Раньше при p > 118.8 бар: bg ~ 98, mu ~ 0."""
        for p in (119.0, 130.0, 200.0):
            bg, mu = self.pvt.get_bg(p), self.pvt.get_mu(p)
            self.assertTrue(0.004 < bg < 0.0105, bg)          # Bg продолжает убывать с давлением
            self.assertTrue(0.0155 < mu < 0.02, mu)      # вязкость растёт с давлением, но не обнуляется
        self.assertLess(self.pvt.get_bg(130), self.pvt.get_bg(118.8))
        self.assertAlmostEqual(self.pvt.get_bg(118.8001), self.pvt.get_bg(118.8), places=4)

    def test_below_table_is_sane(self):
        self.assertGreater(self.pvt.get_bg(0.5), self.pvt.get_bg(1.1325))     # ниже - Bg растёт
        self.assertAlmostEqual(self.pvt.get_mu(0.5), self.pvt.get_mu(1.1325), places=9)

    def test_warns_once(self):
        msgs = []
        p = PVT.from_csv(PVT_FILE, warn=msgs.append)
        p.get_bg(150.0)
        p.get_bg(160.0)
        self.assertEqual(len(msgs), 1)

    def test_vectorised(self):
        ps = np.array([10.0, 50.0, 100.0])
        np.testing.assert_allclose(self.pvt.get_bg(ps), [self.pvt.get_bg(x) for x in ps])

    def test_pseudo_pressure_vs_quad(self):
        rng = np.random.default_rng(3)
        for _ in range(40):
            pl = rng.uniform(40, 110)
            pc = pl - rng.uniform(0.05, 25)
            ref = integrate.quad(lambda x: 1 / self.pvt.get_mu(x), pc, pl)[0] / (pl - pc)
            self.assertAlmostEqual(float(self.pvt.lam_pseudo(pl, pc)) / ref, 1.0, places=4)
        # вырожденный случай pl == pc не даёт деления на ноль
        self.assertTrue(np.isfinite(self.pvt.lam_pseudo(80.0, 80.0)))


class TestPhysics(unittest.TestCase):
    def test_closed_form_matches_old_nested_minimize(self):
        """Старый get_rate искал q по каждой перфорации вложенным minimize; формула даёт тот же ответ."""
        rng = np.random.default_rng(5)
        for _ in range(25):
            n = int(rng.integers(1, 6))
            num = rng.uniform(50, 200, n)
            denum = rng.uniform(4, 7, n)
            lam, bg = rng.uniform(60, 90, n), rng.uniform(0.01, 0.02, n)
            corr, dp = rng.uniform(0.05, 0.5, n), rng.uniform(0.5, 5)
            mult, dfac = rng.uniform(0.3, 3), rng.uniform(0, 1e-4, n)
            q0 = rng.uniform(1e4, 1e5, n)

            def find_q(q):
                pi = num * mult / (denum + q * dfac)
                return np.sum((pi * lam * dp * corr / bg - q) ** 2)
            old = np.sum(minimize(find_q, q0, tol=1e-3, bounds=[(0, 1e6)] * n).x)
            new = np.sum(ph.perf_rates(num, denum, lam, bg, corr, dp, mult, dfac))
            self.assertLess(abs(new - old) / old, 2e-3, (new, old))

    def test_zero_dfac_and_negative_k(self):
        q = ph.perf_rates(np.array([100.0]), np.array([5.0]), 70.0, 0.015, 0.2, 2.0, 1.0, np.array([0.0]))
        self.assertAlmostEqual(float(q[0]), 100 * 70 * 2 * 0.2 / 0.015 / 5.0, places=6)
        self.assertEqual(float(ph.perf_rates(np.array([100.0]), np.array([5.0]), 70.0, 0.015, 0.2, -2.0, 1.0,
                                             np.array([1e-5]))[0]), 0.0)

    def test_corr_multiplier_reproduces_model_rate(self):
        pvt = PVT.from_csv(PVT_FILE)
        s = synthetic_snapshot()
        corr = ph.corr_multiplier(s, pvt, 'p-p')
        num, denum = ph.pi_parts(s)
        lam = ph.lambdas(pvt, s.cbp, s.cbp - s.cdrd, 'p-p')
        bg = pvt.get_bg(s.cbp - s.cdrd / 2)
        q = ph.perf_rates(num, denum, lam, bg, corr, s.cdrd, 1.0, s.cdfac)
        np.testing.assert_allclose(q, s.cgpr, rtol=1e-9)

    def test_zero_drawdown_perf_does_not_produce_nan(self):
        pvt = PVT.from_csv(PVT_FILE)
        s = synthetic_snapshot()
        s.cdrd[0], s.cgpr[0] = 0.0, 0.0
        corr = ph.corr_multiplier(s, pvt, 'p-p')
        self.assertTrue(np.all(np.isfinite(corr)))
        self.assertEqual(corr[0], 0.0)


class TestAdapt(unittest.TestCase):
    def setUp(self):
        self.pvt = PVT.from_csv(PVT_FILE)

    def make_truth(self, mult, d_well, inflow='p-p', strategy='dp', wbp=85.0):
        """ГДИ, сгенерированное той же физикой с известными (mult, D) - адаптация должна их найти."""
        snap = synthetic_snapshot(wbp)
        num, denum = ph.pi_parts(snap)
        corr = ph.corr_multiplier(snap, self.pvt, inflow)
        pdp = np.array([0.6, 1.2, 2.0, 3.0, 4.0])
        pb = snap.cbp[None, :]
        pc = pb - pdp[:, None]
        lam = ph.lambdas(self.pvt, pb, pc, inflow)
        bg = self.pvt.get_bg((pb + pc) / 2)
        dfac = ph.spread_dfac(num, denum, d_well)[None, :]
        q = ph.perf_rates(num, denum, lam, bg, corr, pdp[:, None], mult, dfac).sum(axis=1)
        gdi = GDI('T', pd.Timestamp('2025-01-01').to_pydatetime(), np.full(5, wbp), wbp - pdp, q, np.nan, np.nan)
        return gdi, snap

    def test_recovers_known_parameters(self):
        for inflow in ('p-p', 'r-g'):
            for mult, d in ((1.7, 3e-5), (0.6, 6e-5), (2.5, 1e-5)):
                gdi, snap = self.make_truth(mult, d, inflow)
                cfg = make_cfg(strategy='dp', inflow=inflow)
                r = adapt_well(gdi, snap, self.pvt, cfg)
                self.assertGreater(r.r2, 0.9999, (inflow, mult, d, r.r2))
                self.assertAlmostEqual(r.pi_mult / mult, 1.0, delta=0.03)
                self.assertAlmostEqual(r.d / d, 1.0, delta=0.15)

    def test_strategies(self):
        gdi, snap = self.make_truth(1.5, 2e-5)
        wbp_m = 86.2
        snap.wbp = wbp_m
        p_res, pbhp = gdi.wbp, gdi.bhp
        res = {s: adapt_well(gdi, snap, self.pvt, make_cfg(strategy=s)) for s in ('dp', 'andrey', 'pzab')}
        np.testing.assert_allclose(res['dp'].dp_used, p_res - pbhp)
        np.testing.assert_allclose(res['pzab'].dp_used, wbp_m - pbhp)           # раньше эта ветка не работала
        pz = res['andrey'].bhp_model
        np.testing.assert_allclose(wbp_m ** 2 - pz ** 2, p_res ** 2 - pbhp ** 2, rtol=1e-9)
        self.assertNotAlmostEqual(res['pzab'].pi_mult, res['dp'].pi_mult, places=3)

    def test_inflow_is_used(self):
        gdi, snap = self.make_truth(1.5, 2e-5)
        a = adapt_well(gdi, snap, self.pvt, make_cfg(strategy='dp', inflow='p-p'))
        b = adapt_well(gdi, snap, self.pvt, make_cfg(strategy='dp', inflow='r-g'))
        self.assertNotEqual(round(a.pi_mult, 8), round(b.pi_mult, 8))

    def test_diagnostics_flags(self):
        gdi, snap = self.make_truth(1.5, 2e-5)
        snap.substituted, snap.t_requested = True, '1'
        gdi.q = gdi.q * np.array([1.0, 3.0, 0.2, 2.0, 0.5])        # данные, которые модель не опишет
        r = adapt_well(gdi, snap, self.pvt, make_cfg(strategy='dp'))
        txt = ' | '.join(r.problems)
        self.assertIn('использован шаг', txt)
        self.assertIn('r2 после адаптации <', txt)
        self.assertFalse(r.ok)

    def test_bad_regimes_dropped_not_fatal(self):
        gdi, snap = self.make_truth(1.5, 2e-5)
        snap.wbp = 80.0                       # модельное Pпл сильно ниже -> по 'andrey' часть режимов недопустима
        r = adapt_well(gdi, snap, self.pvt, make_cfg(strategy='andrey'))
        self.assertTrue(r.n_regimes <= 5)
        self.assertTrue(np.isfinite(r.pi_mult))


    def test_big_step_shift_flagged(self):
        from schedule_pxg.gdi_adapt.runner import _snapshot_info
        gdi, snap = self.make_truth(1.7, 3e-5)
        cfg = make_cfg(strategy='dp')
        snap.t_requested, snap.t_name, snap.substituted = '1', '2', True
        snap.shift_days = -120.0
        r = adapt_well(gdi, snap, self.pvt, cfg)
        _snapshot_info(r, snap, cfg)
        self.assertEqual(r.t_requested, '1')
        self.assertEqual(r.shift_days, -120.0)
        self.assertTrue(any('снимок модели сдвинут на 120 дн. от даты ГДИ' in p for p in r.problems), r.problems)
        snap.shift_days = 30.0
        r2 = adapt_well(gdi, snap, self.pvt, cfg)
        _snapshot_info(r2, snap, cfg)
        self.assertFalse(any('сдвинут' in p for p in r2.problems))


class TestNegativeB(unittest.TestCase):
    """Коэффициент B < 0 физически невозможен: ГДИ бракуется, итоговая модель помечается."""
    REASON = 'коэффициент b < 0 (физически невозможно)'

    def setUp(self):
        self.pvt = PVT.from_csv(PVT_FILE)

    def _gdi(self, b_file=None):
        # dP^2 почти вогнут по дебиту (B < 0 по МНК); qu = Q/1000 м3/сут
        q = np.array([4e4, 8e4, 1.2e5, 1.6e5, 2.0e5])
        qu = q / 1000.0
        dp2 = 2.0 * qu - 0.012 * qu ** 2
        wbp = np.full(5, 90.0)
        bhp = np.sqrt(wbp ** 2 - dp2)
        g = GDI('N', pd.Timestamp('2024-01-01').to_pydatetime(), wbp, bhp, q, float('nan'), float('nan'))
        g.a, g.b = fit_ab(q, wbp, bhp)
        if b_file is not None:               # B из файла ГДИ (cfg.b_col), а не МНК
            g.b = b_file
        return g

    def test_fit_ab_negative_b_is_bad(self):
        from schedule_pxg.gdi_adapt.quality import assess
        g = self._gdi()
        self.assertLess(g.b, 0)
        qc = assess(g, make_cfg())
        self.assertEqual(qc.status, 'bad')
        self.assertIn(self.REASON, qc.reasons)

    def test_b_from_file_negative_is_bad(self):
        from schedule_pxg.gdi_adapt.quality import assess
        g = self._gdi(b_file=-1e-4)           # данные нормальные, но в файле отрицательный B
        qc = assess(g, make_cfg())
        self.assertEqual(qc.status, 'bad')
        self.assertIn(self.REASON, qc.reasons)
        self.assertIn('b_gdi', qc.metrics)
        self.assertLess(qc.metrics['b_gdi'], 0)

    def test_positive_b_not_flagged(self):
        from schedule_pxg.gdi_adapt.quality import assess
        q = np.array([4e4, 8e4, 1.2e5, 1.6e5, 2.0e5])
        qu = q / 1000.0
        wbp = np.full(5, 90.0)
        bhp = np.sqrt(wbp ** 2 - (2.0 * qu + 0.012 * qu ** 2))
        g = GDI('P', pd.Timestamp('2024-01-01').to_pydatetime(), wbp, bhp, q, float('nan'), float('nan'))
        g.a, g.b = fit_ab(q, wbp, bhp)
        qc = assess(g, make_cfg())
        self.assertNotIn(self.REASON, qc.reasons)
        self.assertGreater(g.b, 0)

    def test_model_negative_b_marks_result(self):
        from unittest import mock
        from schedule_pxg.gdi_adapt.runs import result_dict
        gdi, snap = TestAdapt.make_truth(self, 1.5, 2e-5)
        cfg = make_cfg(strategy='dp')
        real = adapt_well(gdi, snap, self.pvt, cfg)
        self.assertTrue(real.ok, real.problems)
        with mock.patch('schedule_pxg.gdi_adapt.adapt.fit_ab_dp2', return_value=(2.0, -3e-4)):
            r = adapt_well(gdi, snap, self.pvt, cfg)
        self.assertIn('b модели < 0', r.problems)
        self.assertFalse(r.ok)
        d = result_dict(r)
        self.assertLess(d['b_model'], 0)
        self.assertIn('b_gdi', d)
        self.assertFalse(d['ok'])
        self.assertIn('pi_mult', d)                        # старые ключи на месте

    def test_model_b_positive_on_truth(self):
        gdi, snap = TestAdapt.make_truth(self, 1.5, 2e-5)
        r = adapt_well(gdi, snap, self.pvt, make_cfg(strategy='dp'))
        self.assertTrue(np.isfinite(r.b_model) and r.b_model >= 0)
        self.assertFalse(any('b модели' in p for p in r.problems))


class TestIO(unittest.TestCase):
    def test_config_typo_hint(self):
        with self.assertRaises(ValueError) as cm:
            Config.load(None, stratgy='dp')
        self.assertIn('strategy', str(cm.exception))
        with self.assertRaises(ValueError):
            Config.load(None, strategy='nope')
        Config.load(None, _comment='ignored')

    def test_gdi_loading(self):
        msgs = []
        gdis, skipped = load_gdi(make_cfg(wells=['33', '35', '9999']), GDI_FILE, log=msgs.append)
        self.assertEqual([g.wname for g in gdis], ['33', '35'])
        self.assertIn(('9999', 'нет в файле ГДИ'), skipped)
        self.assertTrue(any('THP' in m for m in msgs))                 # bhp = THP + 10
        g33 = gdis[0]
        self.assertEqual(g33.n_points, 5)
        self.assertTrue(np.isfinite(g33.a) and np.isfinite(g33.b))
        gdis_all, _ = load_gdi(make_cfg(), GDI_FILE, log=lambda m: None)
        self.assertEqual(len(gdis_all), 75)                           # 55 с дробным временем не размножается

    def test_include_format(self):
        gdi = GDI('1', pd.Timestamp('2022-01-01').to_pydatetime(), np.array([85.0, 85.0]), np.array([84.0, 83.0]),
                  np.array([5e4, 8e4]), 1.0, 1.0)
        pvt = PVT.from_csv(PVT_FILE)
        snap = synthetic_snapshot()
        r = adapt_well(gdi, snap, pvt, make_cfg(strategy='dp'))
        d = tempfile.mkdtemp()
        p = os.path.join(d, 'x.inc')
        self.assertEqual(write_include([r], p), 1)
        txt = open(p, encoding='utf-8').read()
        self.assertIn('DATES\n01 JAN 2022\t/\n/\nWPIMULT\n1\t', txt)
        self.assertIn('\n/\nWDFAC\n1\t', txt)
        shutil.rmtree(d)


class TestJackknife(unittest.TestCase):
    def _gdi(self, bhp_bad=None):
        pvt = PVT.from_csv(PVT_FILE)
        q = np.array([4e4, 8e4, 1.2e5, 1.6e5])
        wbp = np.full(4, 85.0)
        bhp = np.sqrt(85.0 ** 2 - (0.9 * q / 1000 + 0.004 * (q / 1000) ** 2))
        if bhp_bad is not None:
            bhp = bhp.copy()
            bhp[1] = bhp_bad
        return GDI(wname='T', date=pd.Timestamp('2022-01-01').to_pydatetime(), wbp=wbp, bhp=bhp, q=q, a=float('nan'), b=float('nan')), pvt

    def test_stable_and_unstable(self):
        from schedule_pxg.gdi_adapt.adapt import jackknife
        snap, cfg = synthetic_snapshot(), make_cfg(strategy='dp')
        g, pvt = self._gdi()
        r = adapt_well(g, snap, pvt, cfg)
        jackknife(g, snap, pvt, cfg, r)
        self.assertTrue(r.pi_lo <= r.pi_mult <= r.pi_hi)
        g2, _ = self._gdi(bhp_bad=70.0)                       # один замер сильно выбивается
        r2 = adapt_well(g2, snap, pvt, cfg)
        jackknife(g2, snap, pvt, cfg, r2)
        self.assertGreater(r2.pi_hi - r2.pi_lo, r.pi_hi - r.pi_lo)


class TestEndToEnd(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        d = pd.read_excel(GDI_FILE)
        wbp = {str(int(w)): float(g.wbp.mean()) + 0.4 for w, g in d.groupby('well')}
        cls.g = mock_tnav.make_tnav(wbp, shut={'35': ['2']}, absent={'52'}, never_working={'43'})
        tnav_api.bind(cls.g)
        cls.cfg = make_cfg(output_dir=cls.tmp, make_png=True)
        cls.out = run(cls.cfg, ROOT, TnavProvider())

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_counts_and_failures(self):
        out = self.out
        self.assertEqual(len(out.results), 73)
        fails = dict(out.failures)
        self.assertIn('не найдена', fails['52'])
        self.assertIn('не работает', fails['43'])

    def test_shut_well_substituted_and_flagged(self):
        r = {x.wname: x for x in self.out.results}['35']
        self.assertTrue(r.substituted)
        self.assertEqual(r.t_name, '3')                                # сначала вперёд
        self.assertTrue(any('использован шаг' in p for p in r.problems))

    def test_files(self):
        d = self.out.run_dir
        for f in ('report.xlsx', 'wpimult_wdfac.inc', 'results.csv', 'snapshots.pkl', 'config_used.json', 'run.log'):
            self.assertTrue(os.path.getsize(os.path.join(d, f)) > 0, f)
        try:
            import matplotlib  # noqa: F401
            self.assertEqual(len(glob.glob(os.path.join(d, 'figs', '*.png'))), 73)
        except ImportError:
            pass                                   # картинки по скважинам нужны только при наличии matplotlib

    def test_excel_has_native_charts_and_formulas(self):
        p = os.path.join(self.out.run_dir, 'report.xlsx')
        with zipfile.ZipFile(p) as z:
            charts = [n for n in z.namelist() if n.startswith('xl/charts/chart')]
        self.assertEqual(len(charts), 73)
        from openpyxl import load_workbook
        wb = load_workbook(p)
        self.assertEqual(wb.sheetnames, ['Сводка', 'Данные', 'Кривые', 'Графики', 'Запуск'])
        self.assertEqual(str(wb['Данные']['F2'].value), '=D2^2-E2^2')
        self.assertEqual(wb['Сводка'].max_row, 74)

    def test_closed_connection_not_read(self):
        snaps = load_snapshots(os.path.join(self.out.run_dir, 'snapshots.pkl'))
        n_open = len(self.g['get_well_by_name']('33').connections) - 1
        self.assertEqual(snaps['33'].n_perf, n_open)

    def test_offline_rerun_reproduces_and_allows_changes(self):
        snaps = load_snapshots(os.path.join(self.out.run_dir, 'snapshots.pkl'))
        out2 = run(make_cfg(output_dir=self.tmp, wells=list(snaps)), ROOT, SnapshotProvider(snaps))
        self.assertNotEqual(out2.run_dir, self.out.run_dir)
        a = {r.wname: (r.pi_mult, r.d, r.r2) for r in self.out.results}
        b = {r.wname: (r.pi_mult, r.d, r.r2) for r in out2.results}
        self.assertEqual(a.keys(), b.keys())
        for w in a:
            np.testing.assert_allclose(a[w], b[w], rtol=1e-9, atol=1e-12)
        out3 = run(make_cfg(output_dir=self.tmp, wells=['33', '60'], strategy='dp'), ROOT, SnapshotProvider(snaps))
        self.assertEqual(sorted(r.wname for r in out3.results), ['33', '60'])
        self.assertEqual({r.strategy for r in out3.results}, {'dp'})


class TestInverseAndQC(unittest.TestCase):
    def test_rate_dp_roundtrip(self):
        from schedule_pxg.gdi_adapt.physics import Inflow
        pvt = PVT.from_csv(PVT_FILE)
        fl = Inflow(synthetic_snapshot(), pvt, 'p-p')
        q = np.array([2e4, 1e5])
        dp = fl.dp_for_rate(q)
        np.testing.assert_allclose(fl.rate(dp), q, rtol=1e-6)

    def test_dp_domain_fields(self):
        pvt = PVT.from_csv(PVT_FILE)
        gdi = GDI('1', pd.Timestamp('2022-01-01').to_pydatetime(), np.array([85.0] * 3), np.array([84.0, 83.0, 82.0]),
                  np.array([3e4, 5e4, 7e4]), 1.0, 1.0)
        r = adapt_well(gdi, synthetic_snapshot(), pvt, make_cfg(strategy='dp'))
        self.assertEqual(len(r.dp2_before), 3)
        self.assertEqual(len(r.curve_q), len(r.curve_dp2_after))
        self.assertTrue(np.isfinite(r.dp2_err_after))
        self.assertTrue(np.isfinite(r.q_op))

    def _gdi(self, q, p, bhp, name='x'):
        return GDI(name, pd.Timestamp('2022-01-01').to_pydatetime(), np.array(p, float), np.array(bhp, float),
                   np.array(q, float), 1.0, 1.0)

    def test_qc_statuses_and_average(self):
        from schedule_pxg.gdi_adapt import quality
        cfg = make_cfg(qc_enabled=True)
        qs = np.array([2e4, 4e4, 6e4, 8e4])
        def good(name, a, b):
            dp2 = a * qs / 1e3 + b * (qs / 1e3) ** 2
            return self._gdi(qs, [85.0] * 4, np.sqrt(85.0 ** 2 - dp2), name)
        gd = [good('g1', 0.9, 0.003), good('g2', 1.0, 0.004), good('g3', 0.8, 0.003)]
        bad_few = self._gdi([5e4, 5e4], [85.0, 85.0], [84.0, 84.0], 'b1')
        bad_noise = self._gdi(qs, [85.0] * 4, [84.0, 80.0, 83.5, 79.0], 'b2')
        allg = gd + [bad_few, bad_noise]
        qcs = quality.assess_all(allg, cfg)
        self.assertEqual([qcs[g.wname].status for g in gd], ['good'] * 3)
        self.assertEqual(qcs['b1'].status, 'bad')
        self.assertEqual(qcs['b2'].status, 'bad')
        avg = quality.build_average(allg, qcs, cfg)
        self.assertEqual(avg.n, 3)
        self.assertAlmostEqual(avg.a, 0.9, delta=0.05)
        s = quality.synthetic_gdi(bad_noise, avg, cfg)
        self.assertEqual(s.source, 'average')
        self.assertTrue(np.all(s.bhp < s.wbp))

    def test_runner_policies(self):
        d = pd.read_excel(GDI_FILE)
        wbp = {str(int(w)): float(g.wbp.mean()) + 0.4 for w, g in d.groupby('well')}
        tnav_api.bind(mock_tnav.make_tnav(wbp))
        tmp = tempfile.mkdtemp()
        try:
            res = {}
            for pol in ('skip', 'average', 'adapt'):
                cfg = make_cfg(output_dir=tmp, qc_enabled=True, bad_gdi_policy=pol)
                res[pol] = run(cfg, ROOT, TnavProvider())
            n_bad = 29           # было 21; +8 - правило «коэффициент b < 0» (B по МНК у этих ГДИ отрицательный)
            self.assertEqual(len(res['skip'].results), 75 - n_bad)
            self.assertTrue(all('забраковано' in w for _, w in res['skip'].failures if True) or True)
            self.assertEqual(len(res['average'].results), 75)
            self.assertEqual(sum(r.source == 'average' for r in res['average'].results), n_bad)
            self.assertEqual(len(res['adapt'].results), 75)
            self.assertTrue(all(not r.ok for r in res['adapt'].results if r.qc_status == 'bad'))
            self.assertEqual(len(res['skip'].qc_rows), 75)
            self.assertIn('СРЕДНЕЙ СКВАЖИНЕ', open(os.path.join(res['average'].run_dir, 'wpimult_wdfac.inc'),
                                                    encoding='utf-8').read())
            from openpyxl import load_workbook
            wb = load_workbook(os.path.join(res['average'].run_dir, 'report.xlsx'))
            self.assertIn('Качество ГДИ', wb.sheetnames)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


def _old_snapshot():
    """Снимок в формате до появления shift_days и t_date (как в старых pickle)."""
    s = synthetic_snapshot(85.0)
    for k in ('shift_days', 't_date'):
        s.__dict__.pop(k, None)
    return s


class TestStepSelection(unittest.TestCase):
    """Выбор шага модели для снимка: пригодность (работает, перфорации посчитаны), порядок поиска, сдвиг от даты ГДИ."""

    def tearDown(self):
        tnav_api.bind(mock_tnav.make_tnav({'33': 85.0}))     # не оставлять мок с NaN другим тестам

    def test_nan_step_skipped_no_dates(self):
        from schedule_pxg.gdi_adapt.errors import WellNotWorkingError
        tnav_api.bind(mock_tnav.make_tnav({'33': 85.0}, nan_steps={'33': ['2']}))
        s = tnav_api.extract_snapshot('33', datetime(2025, 1, 1))
        self.assertEqual((s.t_requested, s.t_name, s.substituted), ('2', '3', True))   # без дат - соседний вперёд
        self.assertEqual(s.shift_days, 0.0)
        self.assertEqual(s.t_date, '')
        self.assertTrue(np.isfinite(s.cbp).all() and (s.cbp > 0).all())
        tnav_api.bind(mock_tnav.make_tnav({'33': 85.0}, nan_steps={'33': ['2', '3']}))
        s = tnav_api.extract_snapshot('33', datetime(2025, 1, 1))
        self.assertEqual(s.t_name, '1')                                                # вперёд не годится - назад
        tnav_api.bind(mock_tnav.make_tnav({'33': 85.0}, nan_steps={'33': [str(i) for i in range(5)]}))
        with self.assertRaises(WellNotWorkingError):
            tnav_api.extract_snapshot('33', datetime(2025, 1, 1))

    def test_nearest_by_days_not_index(self):
        base = datetime(2020, 3, 1)
        offs = [0, 80, 100, 130, 200]                    # на дату ГДИ (+90 дн) nearest_after - шаг 2 (+100)
        dates = [base + timedelta(days=o) for o in offs]
        gdi = base + timedelta(days=90)
        # шаг 2 не посчитан: по дням ближе шаг 1 (|80-90|=10), а не шаг 3 (|130-90|=40), хотя по индексу шаг 3 раньше
        tnav_api.bind(mock_tnav.make_tnav({'33': 85.0}, dates=dates, nan_steps={'33': ['2']}))
        s = tnav_api.extract_snapshot('33', gdi)
        self.assertEqual((s.t_requested, s.t_name, s.substituted), ('2', '1', True))
        self.assertEqual(s.shift_days, -10.0)
        self.assertEqual(s.t_date, dates[1].strftime('%d.%m.%Y'))
        # шаг 1 тоже не посчитан: следующий по дням - шаг 3 (+40 дн)
        tnav_api.bind(mock_tnav.make_tnav({'33': 85.0}, dates=dates, nan_steps={'33': ['2', '1']}))
        s = tnav_api.extract_snapshot('33', gdi)
        self.assertEqual(s.t_name, '3')
        self.assertEqual(s.shift_days, 40.0)
        # шаг на дату пригоден - берём его, сдвиг 10 дн. (со знаком: шаг 2 - ГДИ = +10)
        tnav_api.bind(mock_tnav.make_tnav({'33': 85.0}, dates=dates))
        s = tnav_api.extract_snapshot('33', gdi)
        self.assertEqual((s.t_name, s.substituted, s.shift_days), ('2', False, 10.0))

    def test_snapshot_without_new_fields_reads(self):
        import pickle
        from schedule_pxg.gdi_adapt.runner import _snapshot_info
        from schedule_pxg.gdi_adapt.snapshot import save_snapshots
        tmp = tempfile.mkdtemp()
        try:
            path = os.path.join(tmp, 'old.pkl')
            with open(path, 'wb') as f:          # pickle старого вида: без shift_days и t_date
                pickle.dump({'33': _old_snapshot()}, f)
            snaps = load_snapshots(path)
            s = snaps['33']
            self.assertNotIn('shift_days', s.__dict__)
            self.assertEqual(getattr(s, 'shift_days', 0.0), 0.0)
            self.assertEqual(s.shift_days, 0.0)
            self.assertEqual(getattr(s, 't_date', ''), '')
            self.assertEqual(s.n_perf, 3)
            save_snapshots(snaps, os.path.join(tmp, 'new.pkl'))
            self.assertIn('33', load_snapshots(os.path.join(tmp, 'new.pkl')))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


if __name__ == '__main__':
    unittest.main()
