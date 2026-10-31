import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from starlette.applications import Starlette  # noqa: E402
from starlette.testclient import TestClient  # noqa: E402

from pxg_core import table_export  # noqa: E402

TABLE = {"id": "t", "title": "Итого", "count": 2, "columns": [
    {"key": "well", "label": "Скважина", "unit": "", "decimals": None, "kind": "text"},
    {"key": "q", "label": "Расход", "unit": "тыс. м³/сут", "decimals": 1, "kind": "number"},
    {"key": "d", "label": "Дата", "unit": "", "decimals": None, "kind": "date"}],
    "rows": [["101", "102"], [12.34, None], ["2026-01-05", "2026-01-06T00:00:00"]]}


def _client():
    return TestClient(Starlette(routes=table_export.routes()))


def test_xlsx_column_major_rows():
    from openpyxl import load_workbook
    r = _client().post("/api/table/export", json={"table": TABLE, "format": "xlsx"})
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    assert [c.value for c in ws[1]] == ["Скважина", "Расход, тыс. м³/сут", "Дата"]
    assert [c.value for c in ws[2]][:2] == ["101", 12.34]
    assert ws.cell(3, 2).value is None and ws.cell(3, 3).value.day == 6


def test_csv_russian_excel():
    r = _client().post("/api/table/export", json={"table": TABLE, "format": "csv"})
    assert r.content.startswith("﻿".encode())
    assert "101;12,34;05.01.2026" in r.content.decode("utf-8")


def test_bad_request():
    assert _client().post("/api/table/export", json={"format": "xlsx"}).status_code == 400
