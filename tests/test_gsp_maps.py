"""«Карты ГСП»: паритет со старыми функциями скрипта карт (достаются через ast, сам скрипт не импортируется), XY, API."""
import json
import ast
import os
import re
import sys
from datetime import datetime
from pathlib import Path

import pytest

np = pytest.importorskip("numpy")
pd = pytest.importorskip("pandas")
pytest.importorskip("openpyxl")
pytest.importorskip("xlsxwriter")
pytest.importorskip("starlette")
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from gsp_maps import inputs, wellmap  # noqa: E402
from gsp_maps.store import KIND_NAMES, OTBOR, ZAKACHKA, Store, season_of  # noqa: E402

OLD = next((ROOT / "apps" / "map_dashboards").glob("Секторные_диаграммы*.py"))
_NEEDED = {"get_season_from_month", "process_main_data", "load_water_files", "load_map_file", "load_pressure_file",
           "load_seasons_file", "analyze_trends", "load_perforation_depths", "integrate_water_data",
           "create_well_share_analysis", "create_pressure_flow_sheet", "_format_header", "_color_season_groups"}


@pytest.fixture(scope="module")
def old():
    """Старые функции в своём пространстве имён (остальной модуль с побочными эффектами не выполняется)."""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    tree = ast.parse(OLD.read_text(encoding="utf-8"))
    ns = {"pd": pd, "np": np, "re": re, "os": os, "datetime": datetime, "load_workbook": load_workbook,
          "Alignment": Alignment, "Border": Border, "Font": Font, "PatternFill": PatternFill, "Side": Side, "get_column_letter": get_column_letter}
    for node in tree.body:
        if isinstance(node, ast.FunctionDef) and node.name in _NEEDED:
            exec(compile(ast.Module([node], []), str(OLD), "exec"), ns)
    return ns


def make_db(rng, gsps=("ГСП 1", "ГСП 2"), wells=(11, 12, 13, 14), runtime=True):
    rows_o, rows_z = [], []
    for gsp in gsps:
        for w in wells:
            for day in pd.date_range("2022-10-01", "2023-04-30"):
                f = float(rng.choice([0, 0, 15000, 22000.5, 31000]))
                rt = float(rng.choice([0, 12, 24]))
                rows_o.append({"Скважина": w, "Дата": day, "Месяц": day.month, "Часовой расход газа": f / 24, "Время работы": rt,
                               "Суточный расход газа": f, "Тип данных": "отбор", "Источник": gsp, "Год": day.year,
                               "Сезон": season_of(day.month, day.year)[0]})
            for day in pd.date_range("2023-05-10", "2023-09-20"):
                f = float(rng.choice([0, 40000, 52000]))
                rows_z.append({"Скважина": w, "Дата": day, "Месяц": day.month, "Часовой расход газа": f / 24, "Время работы": 24.0,
                               "Суточный расход газа": f, "Тип данных": "закачка", "Источник": gsp, "Год": day.year})
    o, z = pd.DataFrame(rows_o), pd.DataFrame(rows_z)
    if not runtime:
        o, z = o.drop(columns="Время работы"), z.drop(columns="Время работы")
    return o, z


@pytest.fixture(scope="module")
def db(tmp_path_factory):
    o, z = make_db(np.random.RandomState(7))
    path = tmp_path_factory.mktemp("db") / "БД_расходы.xlsx"
    with pd.ExcelWriter(str(path), engine="xlsxwriter") as xw:
        o.to_excel(xw, sheet_name="Отборы", index=False)
        z.to_excel(xw, sheet_name="Закачка", index=False)
    return path, o, z


def test_season_of_matches_old(old):
    for y in (2019, 2024):
        for m in range(1, 13):
            assert season_of(m, y) == old["get_season_from_month"](m, y)


def test_aggregates_match_process_main_data(old, db, tmp_path, monkeypatch):
    monkeypatch.setenv("GSP_MAPS_CACHE", str(tmp_path / "cache"))
    path, o, z = db
    store = Store.open(str(path))
    old_o, old_z = old["process_main_data"](o.copy(), z.copy(), "ГСП 2")
    for kind, ref in ((OTBOR, old_o), (ZAKACHKA, old_z)):
        for season, part in ref.groupby("Сезон"):
            new = store.aggregate("ГСП 2", kind, str(season)).set_index("Скважина")
            part = part.set_index("Скважина")
            for col in part.columns:
                if col in ("Сезон", "Тип"):
                    continue
                assert np.allclose(new.loc[part.index, col].to_numpy(dtype=float), part[col].to_numpy(dtype=float), equal_nan=True), (kind, season, col)
    assert sorted(store.seasons("ГСП 2", OTBOR)) == sorted(str(s) for s in old_o["Сезон"].unique())
    assert sorted(store.seasons("ГСП 2", ZAKACHKA)) == sorted(str(s) for s in old_z["Сезон"].unique())


def test_cache_roundtrip_and_matrix(db, tmp_path, monkeypatch):
    monkeypatch.setenv("GSP_MAPS_CACHE", str(tmp_path / "cache"))
    path, o, _z = db
    first = Store.open(str(path))
    log = []
    second = Store.open(str(path), log.append)
    assert any("кэша" in m for m in log)
    a, b = first.matrix("ГСП 1", OTBOR, "2022-2023"), second.matrix("ГСП 1", OTBOR, "2022-2023")
    assert np.array_equal(a["flow"], b["flow"]) and a["flow"].shape == (4, len(pd.date_range("2022-10-01", "2023-04-30")))
    sub = o[(o["Источник"] == "ГСП 1") & (o["Скважина"] == 12)].sort_values("Дата")
    assert np.allclose(a["flow"][list(a["wells"]).index(12)], sub["Суточный расход газа"].to_numpy())
    assert first.gsp_list() == ["ГСП 1", "ГСП 2"]


def test_without_runtime_column(old, tmp_path, monkeypatch):
    monkeypatch.setenv("GSP_MAPS_CACHE", str(tmp_path / "cache"))
    o, z = make_db(np.random.RandomState(1), gsps=("ГСП 5",), runtime=False)
    path = tmp_path / "bd.xlsx"
    with pd.ExcelWriter(str(path), engine="xlsxwriter") as xw:
        o.to_excel(xw, sheet_name="Отборы", index=False)
        z.to_excel(xw, sheet_name="Закачка", index=False)
    store = Store.open(str(path))
    df = store.aggregate("ГСП 5", OTBOR, "2022-2023")
    assert "Суммарное_время_работы" not in df.columns and "Накопленный_расход_газа" in df.columns


# ---------------- вход: периоды, давление, вода, карта ----------------
def test_periods_and_pressure_match_old(old, tmp_path):
    p = tmp_path / "per.txt"
    p.write_text("08.04.2012 inj\n29.10.2012 prod\n\nмусор\n14.04.2013 NONE\n", encoding="utf-8")
    new, _ = inputs.load_periods(str(p))
    ref = old["load_seasons_file"](str(p))
    assert [(inputs.iso(d), t) for d, t in new] == [(r["date"].strftime("%Y-%m-%d"), r["type"]) for r in ref]
    g = tmp_path / "gsp.txt"
    g.write_text("520\t02.08.2006\t\n520\t04.08.2006\t94,9\n520\t22.11.2024\t95,42\n", encoding="utf-8")
    o = tmp_path / "obj.txt"
    o.write_text("02.08.2006\t94,7\n04.08.2006\t94,9\n", encoding="utf-8")
    for path, is_gsp in ((g, True), (o, False)):
        new_df, _ = inputs.load_pressure(str(path), is_gsp)
        ref_df = old["load_pressure_file"](str(path), is_gsp=is_gsp, gsp_filter="ГСП 9").reset_index(drop=True)
        assert list(new_df["Дата"]) == list(ref_df["Дата"]) and list(new_df["Давление_бар"]) == list(ref_df["Давление_бар"])


def test_pressure_tolerates_typos(tmp_path):
    p = tmp_path / "o.txt"
    p.write_text("27.10.23.\t90,5\n11.30.22\t80\n01.02.2024\t70\n", encoding="utf-8")
    df, warn = inputs.load_pressure(str(p), False)
    assert len(df) == 2 and warn  # дата с точкой в конце прочитана, невозможная — в предупреждении


def test_water_old_format_matches_old(old, tmp_path):
    rows = [["№ скв.", "Вынос воды", "Расход газа", "Вода, л/ч"], [500, "46/1000", 26000, 1196], [501, "1000м³/79л.", 31000, 2449], [502, "нет воды", 20000, 0],
            [503, "песок в ремонт", 0, None], [504, "12,5", 24000, 300]]
    f = tmp_path / "ГСП 9 замеры ноябрь 2025.xlsx"
    pd.DataFrame(rows).to_excel(str(f), index=False, header=False)
    new, _ = inputs.load_water([str(f)])
    ref = old["load_water_files"]([str(f)])
    key = ["Скважина", "Месяц", "Год", "Метка_замера", "Примечание"]
    assert new[key].reset_index(drop=True).equals(ref[key].reset_index(drop=True))
    for col in ("Водный_фактор", "Расход_воды_лч"):
        assert np.allclose(new[col].to_numpy(dtype=float), ref[col].to_numpy(dtype=float), equal_nan=True), col


def test_water_litres_format(tmp_path):
    rows = [["ГСП-9 Замеры на воду март 2016", None, None, None, None, None],
            ["№скв", "литров/за час", "литров/за сутки", "расход", None, "ВФ"],
            [500, "338л.", "8112л.", 26000, 338.0, 0.013], [503, "2900л.", "69600л.", 20000, 2900.0, 0.145]]
    f = tmp_path / "ГСП-9 вода.xlsx"  # месяца в имени файла нет — берётся из заголовка листа
    pd.DataFrame(rows).to_excel(str(f), index=False, header=False)
    df, warn = inputs.load_water([str(f)])
    assert not warn and list(df["Метка_замера"]) == ["03.2016", "03.2016"]
    assert list(df["Расход_воды_лч"]) == [338.0, 2900.0] and np.allclose(df["Водный_фактор"], [13.0, 145.0])


def make_map(path, cells, sheet="ГСП 9", rows=40, cols=14):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.cell(rows, cols).value = None
    for w, (r, c) in cells.items():
        ws.cell(r, c).value = w
    ws.cell(rows, cols).number_format = "0"  # размеры листа как у реального файла: больше, чем занято
    wb.save(str(path))


CELLS = {500: (3, 9), 501: (4, 7), 502: (3, 3), 503: (30, 4), 504: (35, 12), 505: (20, 7), 506: (7, 14)}


def test_grid_matches_old(old, tmp_path):
    path = tmp_path / "карта.xlsx"
    make_map(path, CELLS)
    cells, dims, name = wellmap.read_grid(str(path), "ГСП 9")
    ref_cells, ref_dirs = old["load_map_file"](str(path), "ГСП 9")
    assert cells == ref_cells and name == "ГСП 9"
    assert wellmap.grid_directions(cells, dims) == ref_dirs
    lay = wellmap.layout(list(CELLS), cells, dims, {}, "auto")
    assert lay["mode"] == "grid" and lay["wells"][500]["y"] == 35 - 3 + 1 and lay["wells"][500]["x"] == 9.0
    assert {w: v["dir"] for w, v in lay["wells"].items()} == ref_dirs
    with pytest.raises(KeyError):
        wellmap.read_grid(str(path), "ГСП 3")


# ---------------- XY из tNavigator ----------------
def test_xy_parsing_variants():
    xy, w = wellmap.parse_xy_lines(["Well\tX\tY\tZ", "W500\t1234,5\t999.1\t10", "501 100 200", "'Скв_502'; 5; 6", "S-503;7,5;8,5;0", ""])
    assert xy == {500: (1234.5, 999.1), 501: (100.0, 200.0), 502: (5.0, 6.0), 503: (7.5, 8.5)} and not w
    xy, _ = wellmap.parse_xy_lines(["WELLTRACK 'W500'", "10 20 0 0", "11 21 5 5 /", "WELLTRACK 'W501'", "1.5 2.5 0 0 /"])
    assert xy == {500: (10.0, 20.0), 501: (1.5, 2.5)}
    xy, w = wellmap.parse_xy_lines(["ничего полезного"])
    assert not xy and w


def test_xy_excel_and_auto_layout(tmp_path):
    path = tmp_path / "xy.xlsx"
    pd.DataFrame({"Скважина": [500, 501, 502, 503], "X": [1000, 1100, 1200, 1300], "Y": [2000, 2100, 2050, 2300]}).to_excel(str(path), index=False)
    xy, _ = wellmap.read_xy(str(path))
    assert xy[502] == (1200.0, 2050.0)
    # сетка = XY с масштабом 100 и сдвигом: недостающая скважина должна встать точно
    cells = {500: (10, 2), 501: (9, 3), 502: (9, 4), 503: (7, 5), 504: (8, 3)}
    gx = {w: (c, 11 - r) for w, (r, c) in cells.items()}
    true = {w: (1000 + (x - 2) * 100, 2000 + (y - 1) * 100) for w, (x, y) in gx.items()}
    xy = {w: true[w] for w in (500, 501, 502, 503)}
    lay = wellmap.layout(list(cells), cells, (12, 6), xy, "auto")
    assert lay["mode"] == "xy" and lay["wells"][504]["src"] == "fit" and lay["wells"][500]["src"] == "xy"
    assert lay["wells"][504]["x"] == pytest.approx(true[504][0]) and lay["wells"][504]["y"] == pytest.approx(true[504][1])
    only = wellmap.layout(list(cells), cells, (12, 6), xy, "xy")
    assert 504 in only["missing"] and only["wells"][500]["src"] == "xy"
    assert wellmap.layout(list(cells), cells, (12, 6), {}, "xy")["mode"] == "grid"  # XY пуст — запасной вариант, а не пустая карта


def test_angle_directions():
    assert wellmap.angle_direction(0, 5) == "Север" and wellmap.angle_direction(5, 5) == "Северо-Восток"
    assert wellmap.angle_direction(-5, 0) == "Запад" and wellmap.angle_direction(-1, -5) == "Юг"


# ---------------- тренды и API ----------------
def test_trends_flow_slope_matches_old(old, db, tmp_path, monkeypatch):
    monkeypatch.setenv("GSP_MAPS_CACHE", str(tmp_path / "cache"))
    monkeypatch.setenv("GSP_MAPS_HOME", str(tmp_path / "home"))
    from gsp_maps.project import Project
    path, o, z = db
    # два сезона отбора, чтобы тренд существовал
    o2 = o.copy()
    o2["Дата"] = o2["Дата"] + pd.DateOffset(years=1)
    o2["Сезон"] = o2["Сезон"].map(lambda s: "%d-%d" % tuple(int(p) + 1 for p in s.split("-")))
    o2["Суточный расход газа"] = o2["Суточный расход газа"] * 0.7
    both = pd.concat([o, o2], ignore_index=True)
    p2 = tmp_path / "bd2.xlsx"
    with pd.ExcelWriter(str(p2), engine="xlsxwriter") as xw:
        both.to_excel(xw, sheet_name="Отборы", index=False)
        z.to_excel(xw, sheet_name="Закачка", index=False)
    pr = Project()
    pr.set_paths({"db": str(p2)}, str(tmp_path / "out"))
    pr.store = Store.open(str(p2))
    frames = [pr.season_frame("ГСП 1", OTBOR, s, {}) for s in pr.store.seasons("ГСП 1", OTBOR)]
    new = pr.trends(frames).set_index("Скважина")
    old_o, _ = old["process_main_data"](both.copy(), z.copy(), "ГСП 1")
    ref = old["analyze_trends"](old_o).set_index("Скважина")
    assert np.allclose(new["Тренд_расхода"], ref.loc[new.index, "Тренд_расхода"]) and np.allclose(new["Средний_расход"], ref.loc[new.index, "Средний_расход"])


def test_api_flow(db, tmp_path, monkeypatch):
    monkeypatch.setenv("GSP_MAPS_CACHE", str(tmp_path / "cache"))
    monkeypatch.setenv("GSP_MAPS_HOME", str(tmp_path / "home"))
    import importlib
    import time
    from starlette.testclient import TestClient
    from gsp_maps import api, project
    importlib.reload(project)
    importlib.reload(api)
    path, _o, _z = db
    data = tmp_path / "data"
    data.mkdir()
    make_map(data / "Карта расположения скважин на ГСП.xlsx", {11: (3, 3), 12: (3, 8), 13: (9, 3), 14: (9, 8)}, sheet="ГСП 1")
    (data / "xy_tnav.txt").write_text("W11 100 900\nW12 300 900\nW13 100 700\nW14 300 700\n", encoding="utf-8")
    (data / "давление_по_ГСП.txt").write_text("11\t10.11.2022\t95,4\n12\t10.11.2022\t95,6\n11\t10.01.2023\t90,1\n99\t10.01.2023\t10,0\n", encoding="utf-8")
    c = TestClient(api.app)
    found = c.post("/api/scan", json={"folder": str(data)})
    assert found.status_code == 200 and found.json()["paths"]["xy"].endswith("xy_tnav.txt")
    assert c.post("/api/config", json={"paths": {"db": str(path)}, "resultsDir": str(tmp_path / "out")}).status_code == 200
    for _ in range(100):
        st = c.get("/api/state").json()
        if st["state"] != "loading":
            break
        time.sleep(0.1)
    assert st["state"] == "ready" and st["gsps"] == ["ГСП 1", "ГСП 2"]
    assert st["gspMeta"]["ГСП 1"] == {"grid": True, "xy": True} and st["gspMeta"]["ГСП 2"]["grid"] is False
    g = c.get("/api/gsp", params={"name": "ГСП 1", "mode": "auto"}).json()
    assert g["layout"]["mode"] == "xy" and g["layout"]["wells"]["11"] == {"x": 100.0, "y": 900.0, "src": "xy", "dir": "Северо-Запад"}
    assert [s["key"] for s in g["seasons"]["Отбор"]] == ["2022-2023"] and g["seasonPressure"]["gsp"]["Отбор|2022-2023"] == pytest.approx(92.8)
    assert g["pressure"]["gsp"]["bar"] == [95.5, 90.1]  # замеры скважин 11 и 12 за день усреднены, скважина 99 чужая
    gg = c.get("/api/gsp", params={"name": "ГСП 1", "mode": "grid"}).json()
    assert gg["layout"]["mode"] == "grid"
    s = c.get("/api/season", params={"gsp": "ГСП 1", "kind": "Отбор", "season": "2022-2023"}).json()
    assert s["wells"] == [11, 12, 13, 14] and len(s["days"]) == 212 and len(s["flow"][0]) == 212
    assert c.get("/api/gsp", params={"name": "нет"}).status_code == 404
    ex = c.post("/api/export", json={"gsp": "ГСП 1", "mode": "auto"}).json()
    book = pd.ExcelFile(ex["path"])
    assert {"Отбор по сезонам", "Закачка по сезонам", "Скважины", "Работа по сезонам", "Работа по месяцам", "Очерёдность ввода"} <= set(book.sheet_names)
    wk = c.get("/api/work", params={"name": "ГСП 1"}).json()
    assert wk["wells"] == [11, 12, 13, 14] and len(wk["seasons"]) == 2 and len(wk["monthFlow"]) == 4
    assert wk["months"][0] == 2022 * 12 + 9 and wk["months"][-1] == 2023 * 12 + 8
    # дни работы по месяцам складываются в дни сезона
    o = next(s for s in wk["seasons"] if s["kind"] == "Отбор")
    assert sum(wk["monthDays"][0][:7]) == o["days"][0]
    assert c.get("/api/work", params={"name": "нет"}).status_code == 404
    assert c.get("/api/files/" + ex["name"]).status_code == 200
    assert c.get("/api/files/../x").status_code == 404


@pytest.mark.parametrize("with_water", [False, True])
def test_share_tables_match_old(old, db, tmp_path, monkeypatch, with_water):
    monkeypatch.setenv("GSP_MAPS_CACHE", str(tmp_path / "cache"))
    monkeypatch.setenv("GSP_MAPS_HOME", str(tmp_path / "home"))
    from gsp_maps.project import Project
    path, o, z = db
    paths = {"db": str(path)}
    df_water = None
    if with_water:
        wd = tmp_path / "вода"
        wd.mkdir()
        for name, k in (("ноябрь 2022", 1), ("январь 2023", 2)):
            rows = [["№", "Вынос", "Расход газа", "Вода, л/ч"]] + [[w, "%d/1000" % (10 * k + w % 5), 20000, 100.0 * k] for w in (11, 12, 13)]
            pd.DataFrame(rows).to_excel(str(wd / ("ГСП замеры %s.xlsx" % name)), index=False, header=False)
        paths["water"] = str(wd)
        df_water = old["load_water_files"]([str(f) for f in sorted(wd.iterdir())])
    pr = Project()
    pr.set_paths(paths, str(tmp_path / "out"))
    pr.store = Store.open(str(path))
    dirs = {11: "Север", 12: "Юг", 13: "Север", 14: "Восток"}
    mine = pr.share_tables("ГСП 1", dirs)
    old_o, old_z = old["process_main_data"](o.copy(), z.copy(), "ГСП 1")
    if df_water is not None:
        old_o = old["integrate_water_data"](old_o, df_water)
    book = tmp_path / "old.xlsx"
    with pd.ExcelWriter(str(book), engine="openpyxl") as writer:
        old["create_well_share_analysis"](old_o, old_z, dirs, writer)
    ref = pd.read_excel(str(book), sheet_name=None)
    assert set(mine) == set(ref)
    for name, df in mine.items():
        r = ref[name]
        assert list(df.columns) == list(r.columns), name
        key = ["Сезон", "Направление"] if name == "Направления_по_сезонам" else ["Скважина"]
        a, b = df.sort_values(key).reset_index(drop=True), r.sort_values(key).reset_index(drop=True)
        for col in df.columns:
            if not pd.api.types.is_numeric_dtype(a[col]):
                assert list(a[col]) == list(b[col]), (name, col)
            else:
                assert np.allclose(a[col].to_numpy(dtype=float), b[col].to_numpy(dtype=float), equal_nan=True), (name, col)


def _project(tmp_path, monkeypatch, path):
    monkeypatch.setenv("GSP_MAPS_CACHE", str(tmp_path / "cache"))
    monkeypatch.setenv("GSP_MAPS_HOME", str(tmp_path / "home"))
    from gsp_maps.project import Project
    pr = Project()
    pr.set_paths({"db": str(path)}, str(tmp_path / "out"))
    pr.store = Store.open(str(path))
    return pr


def test_pressure_flow_table_matches_old(old, db, tmp_path, monkeypatch):
    path, o, z = db
    pr = _project(tmp_path, monkeypatch, path)
    days = pd.date_range("2022-09-25", "2023-02-10", freq="5D")
    pg = pd.DataFrame({"Дата": days, "Давление_бар": np.linspace(80, 60, len(days)), "Скважина": 0})
    po = pd.DataFrame({"Дата": days[::2] + pd.Timedelta(days=1), "Давление_бар": np.linspace(70, 50, len(days[::2]))})
    pr.pressure = lambda which: ((pg if which == "gsp" else po), [])  # type: ignore[assignment]
    mine = pr.pressure_flow_table("ГСП 1")
    book = tmp_path / "old.xlsx"
    with pd.ExcelWriter(str(book), engine="openpyxl") as writer:
        # старая функция берёт первый столбец «…расход…газ…», то есть часовой («Часовой расход газа»), и подписывает его суточным:
        # это ошибка старого листа, здесь расход суточный; для сверки убираем часовой столбец
        old["create_pressure_flow_sheet"](writer, o.drop(columns="Часовой расход газа"), z.drop(columns="Часовой расход газа"), pg, po, "ГСП 1")
    ref = pd.read_excel(str(book), sheet_name="Данные_для_графика_давления", header=3)
    assert list(mine.columns) == list(ref.columns)
    assert len(mine) == len(ref)
    assert list(mine["Дата"]) == [str(v) for v in ref["Дата"]]
    for col in mine.columns[1:]:
        a = pd.to_numeric(mine[col].replace("", np.nan)).to_numpy(dtype=float)
        b = pd.to_numeric(ref[col].replace("", np.nan)).to_numpy(dtype=float)
        assert np.allclose(a, b, equal_nan=True), col


def test_summary_tables_and_summary_payload(old, db, tmp_path, monkeypatch):
    path, o, z = db
    pr = _project(tmp_path, monkeypatch, path)
    dirs = {11: "Север", 12: "Юг", 13: "Север", 14: "Восток"}
    frames = {k: [pr.season_frame("ГСП 1", k, s, dirs) for s in pr.store.seasons("ГСП 1", k)] for k in (OTBOR, ZAKACHKA)}
    t = pr.summary_tables(frames)
    old_o, old_z = old["process_main_data"](o.copy(), z.copy(), "ГСП 1")
    ref = pd.concat([old_o, old_z], ignore_index=True, sort=False).sort_values(["Тип", "Сезон", "Скважина"])
    mine = t["Сводка_все_сезоны"]
    assert len(mine) == len(ref)
    for col in ("Накопленный_расход_газа", "Средний_суточный_расход", "Количество_дней", "Суммарное_время_работы", "Накопленный_расход_газа_закачка"):
        assert np.allclose(mine[col].to_numpy(dtype=float), ref[col].to_numpy(dtype=float), equal_nan=True), col
    assert "Вода_по_направлениям" not in t  # воды нет — листа нет, как и в старом скрипте
    sp = pr.summary_payload("ГСП 1", "Отбор", "2022-2023")
    assert "Суммарное_время_работы" in sp["columns"] and len(sp["rows"]) == 4
    allp = pr.summary_payload("ГСП 1", "Отбор", "*")
    assert len(allp["rows"]) == 4 * len(pr.store.seasons("ГСП 1", OTBOR))
    json.dumps(sp), json.dumps(allp)
    gp = pr.gsp_payload("ГСП 1", "grid")
    assert len(gp["gspFlow"]["days"]) == len(gp["gspFlow"]["bar"]) > 0
