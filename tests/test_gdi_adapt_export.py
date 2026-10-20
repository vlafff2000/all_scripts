"""Выгрузка индикаторных диаграмм адаптации по ГДИ: PNG / SVG / PDF / XLSX, пакет ZIP, маршруты API."""
import io
import json
import os
import sys
import zipfile

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from schedule_pxg.gdi_adapt import export_charts as ec  # noqa: E402


def _row(well="33", date="2024-05-01", n=5, source="gdi"):
    q = [100.0 * (i + 1) for i in range(n)]
    dp2 = [0.5 * x + 0.002 * x * x for x in q]
    return {"well": well, "date": date, "source": source, "q_gdi": q, "dp2_gdi": dp2,
            "curve_q": [50.0 * i for i in range(1, 13)],
            "curve_before": [0.5 * 50 * i + 0.002 * (50 * i) ** 2 + 30 for i in range(1, 13)],
            "curve_after": [0.5 * 50 * i + 0.002 * (50 * i) ** 2 for i in range(1, 13)],
            "dp2_err_before": 35.0, "dp2_err_after": 4.0}


def _run_dir(tmp_path, rows):
    d = tmp_path / "runs" / "2026-01-01_10-00"
    d.mkdir(parents=True)
    (d / "results.json").write_text(json.dumps({"rows": rows}, ensure_ascii=False), encoding="utf-8")
    return str(d)


@pytest.mark.parametrize("y", ["dp2", "dp2q"])
@pytest.mark.parametrize("fmt", ["png", "svg", "pdf"])
def test_render_each_format_nonempty(y, fmt):
    data = ec.render(_row(), ec.Style(y=y, dpi=300, width_mm=160), fmt)
    assert len(data) > 1000
    magic = {"png": b"\x89PNG", "svg": b"<?xml", "pdf": b"%PDF"}[fmt]
    assert data.lstrip().startswith(magic) or (fmt == "svg" and b"<svg" in data[:2000])


def test_render_variants_and_excel_look():
    for variant in ec.VARIANTS:
        for look in ec.LOOKS:
            data = ec.render(_row(), ec.Style(variant=variant, look=look), "png")
            assert data.startswith(b"\x89PNG")


def test_png_dpi_changes_size():
    small = ec.render(_row(), ec.Style(dpi=300), "png")
    big = ec.render(_row(), ec.Style(dpi=600), "png")
    assert len(big) > len(small)


def test_render_without_gdi_points_still_draws_curves():
    row = _row()
    row["q_gdi"] = [None, None]
    row["dp2_gdi"] = [None, None]
    assert ec.render(row, ec.Style(), "png").startswith(b"\x89PNG")


def test_xlsx_is_readable_and_has_chart():
    from openpyxl import load_workbook
    data = ec.xlsx_bytes(_row(), ec.Style(y="dp2q"))
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws["A2"].value == 100.0 and ws["B2"].value is not None
    assert ws.max_row == 13                                # заголовок + 12 точек кривой
    assert len(ws._charts) == 1


def test_zip_counts_files_per_row_and_format(tmp_path):
    rows = [_row("33", "2024-05-01"), _row("35", "2023-02-01"), _row("36", "2020-02-01")]
    data, count = ec.zip_bytes(rows, ec.Style(), ["png", "svg"])
    assert count == 6
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = z.namelist()
        assert len(names) == 6 and all(n.startswith("charts/") for n in names)
        assert all(z.getinfo(n).file_size > 0 for n in names)
        assert sum(n.endswith(".png") for n in names) == 3


def test_load_rows_filters_by_id_and_rejects_empty(tmp_path):
    run = _run_dir(tmp_path, [_row("33", "2024-05-01"), _row("35", "2023-02-01")])
    assert len(ec.load_rows(run)) == 2
    picked = ec.load_rows(run, ["35|2023-02-01"])
    assert [r["well"] for r in picked] == ["35"]
    with pytest.raises(ValueError):
        ec.load_rows(run, ["99|2020-01-01"])
    with pytest.raises(ValueError):
        ec.load_rows(str(tmp_path / "нет"))


def test_style_validation_messages():
    assert ec.style_from({}).dpi == 300
    for bad in ({"dpi": 150}, {"width": 50}, {"y": "x"}, {"look": "neon"}, {"variant": "mid"}, {"height": 10}, {"width": "много"}):
        with pytest.raises(ValueError):
            ec.style_from(bad)
    with pytest.raises(ValueError):
        ec.check_format("docx")


def test_api_chart_and_zip(tmp_path, monkeypatch):
    from starlette.testclient import TestClient
    from schedule_pxg import api
    run = _run_dir(tmp_path, [_row("33", "2024-05-01"), _row("35", "2023-02-01"), {**_row("36", "2020-02-01"), "curve_q": []}])
    c = TestClient(api.build_app())
    r = c.post("/api/adapt/export/chart", json={"run_dir": run, "id": "33|2024-05-01", "format": "png", "dpi": 300, "y": "dp2q"})
    assert r.status_code == 200 and r.content.startswith(b"\x89PNG")
    assert "attachment" in r.headers["content-disposition"]
    r = c.post("/api/adapt/export/chart", json={"run_dir": run, "id": "33|2024-05-01", "format": "pdf"})
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    r = c.post("/api/adapt/export/chart", json={"run_dir": run, "id": "33|2024-05-01", "format": "xlsx"})
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert c.post("/api/adapt/export/chart", json={"run_dir": run, "id": "33|2024-05-01", "format": "png", "dpi": 77}).status_code == 400
    assert c.post("/api/adapt/export/chart", json={"run_dir": run, "id": "36|2020-02-01", "format": "png"}).status_code == 400
    r = c.post("/api/adapt/export/zip", json={"run_dir": run, "ids": None, "formats": ["png", "svg"], "y": "dp2"})
    assert r.status_code == 200 and r.headers["x-charts-count"] == "4"             # 36 без кривых пропущена
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        assert len(z.namelist()) == 4
    r = c.post("/api/adapt/export/zip", json={"run_dir": run, "ids": ["35|2023-02-01"], "formats": ["pdf"]})
    assert r.status_code == 200 and r.headers["x-charts-count"] == "1"
    assert c.post("/api/adapt/export/zip", json={"run_dir": run, "ids": ["36|2020-02-01"], "formats": ["png"]}).status_code == 400
    assert c.post("/api/adapt/export/zip", json={"run_dir": run, "ids": None, "formats": ["gif"]}).status_code == 400
    assert c.post("/api/adapt/export/zip", data="не json".encode("utf-8")).status_code == 400
