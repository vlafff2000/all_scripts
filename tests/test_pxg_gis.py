"""Разбор «Базы ГИС-контроль» (pxg_core.gis): даты, блоки, формулы, пачки по заливке, ГВК, сверка с Excel."""
import datetime as dt
import json
import os

import openpyxl
import pandas as pd
import pytest
from openpyxl.styles import PatternFill

from pxg_core import gis

ETALON = "/mnt/project-files/gis-gvk/База_ГИС-контроль.xlsx"


def fill(rgb):
    return PatternFill("solid", fgColor=rgb)


def make_book(path, rows, alt=100.0, title="7", headings=("Верхняя песчаная пачка", "Нижняя песчаная пачка")):
    """Лист в формате файла Егора: шапка в строке 3–4, данные с 5-й строки. rows: (A, B, C, D, Кп, Кг, заливка C)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    ws["E3"] = f"С к в а ж и н а №{title}"
    ws["L3"], ws["M3"] = "Алт.=", alt
    for i, h in enumerate(headings):
        ws.cell(3, 16 + 2 * i, h)
    for c, v in enumerate(["Дата", "ГВК", "Газонасыщ.интервал (м)", None, "hi", "Кпi", "Кгi"], start=1):
        ws.cell(4, c, v)
    for r, (a, b, c, d, kp, kg, color) in enumerate(rows, start=5):
        for col, v in zip((1, 2, 3, 4, 6, 7), (a, b, c, d, kp, kg)):
            ws.cell(r, col, v)
        if color:
            ws.cell(r, 3).fill = fill(color)
    wb.save(path)
    return path


@pytest.mark.parametrize("text,ymd,season,tail", [
    ("Декабрь 1997г", (1997, 12, 1), "осень", ""),
    ("Апрель   1998г", (1998, 4, 1), "весна", ""),
    ("Дек.1982г", (1982, 12, 1), "осень", ""),
    ("Март1983г", (1983, 3, 1), "весна", ""),
    ("Октябрь 2004г ", (2004, 10, 1), "осень", ""),
    ("Сент.1983г", (1983, 9, 1), "осень", ""),
    ("Аперль 2021г", (2021, 4, 1), "весна", ""),
    ("Апрель 2007гг", (2007, 4, 1), "весна", ""),
    ("Июль   2009г спец. исследования", (2009, 7, 1), "без сезона", "спец. исследования"),
    ("Январь 1998г", (1998, 1, 1), "без сезона", ""),
    ("Май 2002г", (2002, 5, 1), "весна", ""),
    ("Июнь 2003г", (2003, 6, 1), "без сезона", ""),
])
def test_parse_date_words(text, ymd, season, tail):
    p = gis.parse_date(text)
    assert p.date == dt.date(*ymd) and p.season == season and p.note == tail and p.quality == ""


def test_parse_date_real_date_year_only_and_text():
    p = gis.parse_date(dt.datetime(2019, 5, 30))
    assert p.date == dt.date(2019, 5, 1) and p.exact == dt.date(2019, 5, 30) and p.season == "весна"
    assert gis.parse_date("1997г").quality == "только год"
    for text in ("При бурении", "(переаттестация)", "Исследования не были выполнены"):
        p = gis.parse_date(text)
        assert p.date is None and p.quality == "нет даты"


def test_blocks_formulas_and_gvk(tmp_path):
    book = make_book(tmp_path / "g.xlsx", [
        ("Октябрь 2020г", "кровля 1", 100, 102, 0.3, 0.5, None),
        (None, None, 102, 104, 0.3, 0.2, None),                 # Кг ≤ 0,25: hэф = 0
        (None, "ГВК", 104, 105, 0.3, 0.6, None),
        ("Март 2021г", "ГВК", 100, 104, 0.3, 0.25, None),        # ровно порог: не эффективный
        ("Апрель 2021г", None, 100, 104, 0.3, 0.4, None),
    ])
    res = gis.load_book(str(book))["7"]
    assert res.well == "7" and res.alt == 100.0 and len(res.blocks) == 3
    b1 = res.blocks[0]
    assert (b1.date, b1.season) == (dt.date(2020, 10, 1), "осень")
    p = gis.to_frames({"7": res})["Пачки"]
    first = p.iloc[0]
    assert first["hг эф, м"] == pytest.approx(3.0)                       # 2 + 0 + 1
    assert first["Кг ср.взв"] == pytest.approx((2 * 0.5 + 1 * 0.6) / 3)
    assert first["Этаж, м"] == pytest.approx(5.0)                         # до последней − от первой
    assert first["ГВК глубина, м"] == 105 and first["ГВК а.о., м"] == pytest.approx(-5.0)   # −(105 − 100)
    assert p.iloc[1]["hг эф, м"] == 0 and p.iloc[1]["Кг ср.взв"] != p.iloc[1]["Кг ср.взв"]  # NaN: нет эффективной толщины
    assert p.iloc[2]["hг эф, м"] == pytest.approx(4.0)


def test_text_notes_statuses_and_text_values(tmp_path):
    book = make_book(tmp_path / "g.xlsx", [
        ("Март 2010г", None, 100, 102, 0.3, 0.6, None),
        ("При проведении ГИС прихват прибора на глубине 777м.", None, 102, 104, 0.3, 0.6, None),
        ("Октябрь 2010г", None, "Работы проведены на скважине №103", None, None, None, None),
        ("Апрель 2011г", None, 100, "811,5*", 0.3, 0, None),
        ("Октябрь 2011г", None, 100, 102, 0.3, "не ясно", None),
    ])
    res = gis.load_book(str(book))["7"]
    assert len(res.blocks) == 4                                            # примечание в A не открывает новый блок
    assert "прихват" in res.blocks[0].note
    assert res.blocks[1].rows[0].status.startswith("Работы проведены")
    assert res.blocks[2].rows[0].to == 811.5 and "811,5*" in res.blocks[2].rows[0].note
    bad = res.blocks[3].rows[0]
    assert bad.kg is None and bad.kg_text == "не ясно" and bad.hef == 0
    codes = {i.code for i in res.issues}
    assert {"interval_text", "kg_text"} <= codes


def test_legend_by_fill_and_manual_override(tmp_path):
    book = make_book(tmp_path / "g.xlsx", [
        ("Октябрь 2020г", "ГВК-1", 100, 103, 0.3, 0.6, None),
        (None, "ГВК-2", 110, 114, 0.3, 0.7, "FF99CCFF"),
        (None, None, 120, 121, 0.3, 0.7, "FF00FF00"),                      # нет в легенде
    ], headings=("Верхняя песчаная пачка", "Нижняя песчаная пачка"))
    res = gis.load_book(str(book))["7"]
    packs = {r.xl_row: r.pack for r in res.rows}
    assert packs == {5: gis.VERH, 6: gis.NIZ, 7: ""}
    unk = [e for e in res.legend if e.color == "00FF00"][0]
    assert unk.status == "не распознан" and any(i.code == "legend_unknown" for i in res.issues)
    # ручной выбор инженера: цвет → пачка, ошибка уходит
    res2 = gis.load_book(str(book), legends={"7": {"00FF00": gis.PERE}})["7"]
    assert {r.xl_row: r.pack for r in res2.rows}[7] == gis.PERE
    assert not any(i.code == "legend_unknown" for i in res2.issues)
    assert [e for e in res2.legend if e.color == "00FF00"][0].source == "вручную"


def test_legend_label_hint_overrides_default_with_check(tmp_path):
    """Персиковая заливка при метке «ГВК*» (покрышка): подсказка по меткам, статус «проверить» (скв. 447)."""
    book = make_book(tmp_path / "g.xlsx", [
        ("Октябрь 2020г", "ГВК*", 100, 103, 0.3, 0.6, "FFFFCC99"),
    ], headings=("Песчаный прослой в глинистой покрышке", "Верхняя песчаная пачка"))
    res = gis.load_book(str(book))["7"]
    e = res.legend[0]
    assert e.pack == gis.POKR and e.status == "проверить" and "ГВК*" in e.reason
    assert res.rows[0].pack == gis.POKR


def test_one_gvk_per_pack_second_goes_to_errors(tmp_path):
    book = make_book(tmp_path / "g.xlsx", [
        ("Октябрь 2020г", "ГВК-1", 100, 103, 0.3, 0.6, None),
        (None, "ГВК-1", 103, 104, 0.3, 0.6, None),                          # второй в той же пачке
        (None, "ГВК-2", 110, 114, 0.3, 0.7, "FF99CCFF"),                    # другая пачка: можно
    ])
    res = gis.load_book(str(book))["7"]
    assert [r.gvk_extra for r in res.rows] == [False, True, False]
    err = [i for i in res.issues if i.code == "gvk_twice"]
    assert len(err) == 1 and err[0].level == "error" and err[0].row == 6
    p = gis.to_frames({"7": res})["Пачки"]
    assert list(p["ГВК глубина, м"]) == [103, 114] and list(p["Лишних ГВК"]) == [1, 0]


def test_alt_fallback_from_formula_and_missing(tmp_path):
    book = make_book(tmp_path / "g.xlsx", [("Октябрь 2020г", "ГВК", 100, 103, 0.3, 0.6, None)])
    wb = openpyxl.load_workbook(book)
    ws = wb["7"]
    ws["L3"] = None                      # подписи «Алт.=» нет, как на листе 23
    ws["M13"] = 126
    ws["M5"] = '=IF(B5="ГВК",(D5-$M$13)*(-1)," ")'
    wb.save(book)
    # у openpyxl нет кэша значений → значение в M13 прочитается, формула укажет на неё
    res = gis.load_book(str(book))["7"]
    assert res.alt == 126 and "формулы" in res.alt_source
    wb = openpyxl.load_workbook(book)
    wb["7"]["M13"] = None
    wb["7"]["M3"] = None
    wb["7"]["M5"] = None
    wb.save(book)
    res = gis.load_book(str(book))["7"]
    assert res.alt is None and any(i.code == "no_alt" and i.level == "error" for i in res.issues)
    assert gis.load_book(str(book), alts={"7": 90.0})["7"].alt == 90.0


def test_parity_with_stored_values(tmp_path):
    """Формулы Excel (диапазон из самой формулы) против сохранённых значений; порог листа ≠ 0,25 помечается отдельно."""
    book = make_book(tmp_path / "g.xlsx", [
        ("Октябрь 2020г", None, 100, 102, 0.3, 0.5, None),
        (None, "ГВК", 102, 104, 0.3, 0.26, None),
    ])
    res = gis.load_book(str(book))["7"]
    r1, r2 = res.rows
    r1.xf.update(J="=IF(G5<=0.25,0,E5)", K="=SUM(J5:J6)", H="=SUM(I5:I6)/K5", L="=D6-C5")
    r1.xl.update(J=2.0, K=4.0, H=(2 * 0.5 + 2 * 0.26) / 4, L=4.0)
    r2.xf.update(J="=IF(G6<=0.27,0,E6)", M='=IF(B6="ГВК",(D6-$M$3)*(-1)," ")')
    r2.xl.update(J=0.0, M=-4.0)
    par = gis.parity({"7": res})
    verdict = {(r["Строка"], r["Столбец"]): r["Вердикт"] for _, r in par.iterrows()}
    assert verdict[(5, "J")] == "совпало" and verdict[(5, "K")] == "совпало" and verdict[(5, "L")] == "совпало"
    assert verdict[(6, "J")] == "порог листа"        # в формуле 0,27, у нас 0,25: Excel дал 0, мы дали бы 2
    assert verdict[(6, "M")] == "совпало"


@pytest.mark.skipif(not os.path.exists(ETALON), reason="эталонный файл Егора лежит в файлах проекта")
def test_etalon_book_parity_and_legends():
    res = gis.load_book(ETALON)
    assert len(res) == 36
    par = gis.parity(res)
    ok = (par["Вердикт"] == "совпало").sum()
    assert ok / len(par) > 0.99                                      # всё остальное — правки формул в самом файле
    bad = par[par["Вердикт"] == "расхождение"]
    assert (bad["Пояснение"] != "").all()                            # каждое расхождение с причиной
    legend = {(e.sheet, e.color): e for r in res.values() for e in r.legend}
    # решения Егора: «ГВК**» = перемычка, «ГВК*» = покрышка; скв. 115 как в файле
    assert legend[("115", "FFCC99")].pack == gis.PERE and legend[("115", "FFCC99")].status == "ок"
    assert legend[("18", "FFFF99")].pack == gis.POKR and legend[("18", "99CCFF")].pack == gis.NIZ
    assert legend[("447", "FFCC99")].pack == gis.POKR and legend[("447", "FFCC99")].status == "проверить"
    assert legend[("13", "FF0000")].status == "не распознан"
    assert legend[("110", "FFCC99")].status == "проверить"
    assert all(b.date for r in res.values() for b in r.blocks if r.sheet != "447" or b.block > 1)
    assert res["18"].alt == 135 and res["23"].alt == 126 and res["124"].alt is None


def test_altitude_file_beats_sheet_and_manual_beats_file(tmp_path):
    book = make_book(tmp_path / "g.xlsx", [("Октябрь 2020г", "ГВК", 100, 103, 0.3, 0.6, None)], alt=100.0)
    base = gis.load_book(str(book))["7"]
    assert base.alt == 100.0 and base.rows[0].gvk_abs == pytest.approx(-3.0)
    res = gis.load_book(str(book), file_alts={"7": 95.0}, perfs={"7": (-20.0, -10.0)})["7"]
    assert res.alt == 95.0 and res.alt_sheet == 100.0 and res.alt_source == "файл альтитуд"
    assert res.rows[0].gvk_abs == pytest.approx(-8.0)                       # −(103 − 95)
    assert [i for i in res.issues if i.code == "alt_mismatch"][0].value == "100 / 95"
    row = gis.to_frames({"7": res})["Пачки"].iloc[0]
    assert (row["Перфорация верх, м"], row["Перфорация низ, м"], row["Альтитуда по листу, м"]) == (-20.0, -10.0, 100.0)
    assert gis.load_book(str(book), alts={"7": 90.0}, file_alts={"7": 95.0})["7"].alt == 90.0
    same = gis.load_book(str(book), file_alts={"7": 100.0})["7"]
    assert not any(i.code == "alt_mismatch" for i in same.issues)


def test_statuses_instead_of_gaps(tmp_path):
    book = make_book(tmp_path / "g.xlsx", [
        ("Март 2010г", None, 100, 102, 0.3, 0.6, None),
        ("Октябрь 2010г", None, "Исследования не были выполнены, т.к. не были выполнены ремонтные работы", None, None, None, None),
        ("Май 2011г", None, "Пласт-коллетор исследованиями не охвачен", None, None, None, None),
        ("Ноябрь 2011г", None, 100, 102, 0.3, "обводнен", None),
        ("Апрель 2012г", "Остановка", None, None, None, None, None),
        ("Май 2012г", None, "газ в пласте-коллекторе не отмечается", None, None, None, None),
        ("Март 2013г", None, "что-то непонятное", None, None, None, None),
        ("Октябрь 2013г", None, 100, 102, 0.3, 0.5, None),
        ("При проведении ГИС произошел прихват прибора на глубине 777м.", None, 102, 104, 0.3, 0.5, None),
    ])
    res = gis.load_book(str(book))["7"]
    blocks = gis.to_frames({"7": res})["Замеры"]
    st = list(blocks["Статус"])
    assert st[0] == "выполнено"
    assert st[1] == "не выполнено" and st[2] == "пласт не охвачен"
    assert "обводнена" in st[3] and st[3].startswith("выполнено")
    assert st[4] == "остановка" and st[5] == "газ не отмечается"
    assert st[6] == "нет данных" and "что-то непонятное" in blocks["Текст статуса"][6]
    assert "прихват" in st[7]
    # текст не потерян, а непонятное — в «Проверке»
    assert [i for i in res.issues if i.code == "block_empty"][0].value == "что-то непонятное"


def test_tab_color_and_current_bottom(tmp_path):
    book = make_book(tmp_path / "g.xlsx", [("Октябрь 2020г", None, 100, 103, 0.3, 0.6, None)])
    wb = openpyxl.load_workbook(book)
    wb["7"].sheet_properties.tabColor = "FF00FF00"
    ws = wb.create_sheet("Тек.забой")
    ws["B1"] = "Максимальная глубина прохождения геофизического прибора (м)"
    for c, v in zip("ABC", ["№ скв.", 7, 8]):
        ws[f"{c}2"] = v
    ws["B3"], ws["C3"] = 850, 821
    ws["A4"], ws["B4"] = dt.datetime(1997, 9, 1), 800
    wb.save(book)
    res = gis.load_book(str(book))
    assert list(res) == ["7"]                                   # служебный лист не разбирается как скважина
    assert res["7"].tab_color == "00FF00"
    row = gis.to_frames(res)["Листы"].iloc[0]
    assert "перемычка" in row["Смысл цвета"]
    b = gis.read_bottom(str(book))
    assert sorted(b["Скважина"] + ":" + b["Вид"] + ":" + b["Глубина, м"].astype(str)) == [
        "7:прохождение прибора:800.0", "7:текущий забой:850.0", "8:текущий забой:821.0"]


def test_export_excel_and_atlas_tables(tmp_path):
    from pxg_core import gis_export
    book = make_book(tmp_path / "g.xlsx", [
        ("Октябрь 2020г", "ГВК-1", 100, 103, 0.3, 0.6, None),
        (None, "ГВК-2", 110, 114, 0.3, 0.7, "FF99CCFF"),
        ("Апрель 2021г", None, "Пласт-коллетор исследованиями не охвачен", None, None, None, None),
    ])
    res = gis.load_book(str(book))
    out = tmp_path / "out"
    files = gis_export.export(res, str(out), bottom=pd.DataFrame(
        [["7", None, 850.0, "текущий забой"]], columns=["Скважина", "Дата", "Глубина, м", "Вид"]), source="g.xlsx")
    names = {os.path.basename(f) for f in files}
    assert {"ГИС_свод.xlsx", "gis_observations.parquet", "gis_packs.csv", "gis_intervals.parquet", "gis_db_meta.json",
            "gis_bottom.parquet"} <= names
    sheets = openpyxl.load_workbook(out / "ГИС_свод.xlsx").sheetnames
    assert {"Свод", "Верхняя пачка", "Нижняя пачка", "Замеры", "Легенда", "Проверка", "Тек.забой"} <= set(sheets)
    intervals = pd.read_excel(out / "ГИС_свод.xlsx", sheet_name="Интервалы")
    assert len(intervals) == 3 and intervals["Замер"].notna().all()     # 2 интервала замера 1 + пустая строка замера 2 («не охвачен»)
    assert intervals["От, м"].notna().sum() == 2                         # замер без интервалов не пропал, но интервалов у него нет
    obs = pd.read_parquet(out / "gis_observations.parquet")
    assert list(obs["status"]) == ["выполнено", "пласт не охвачен"]       # замер без данных не пропал
    packs = pd.read_parquet(out / "gis_packs.parquet")
    assert list(packs["pack"]) == [gis.VERH, gis.NIZ] and packs["gvk_abs_m"].tolist() == [-3.0, -14.0]
    meta = gis_export.read_meta(str(out))
    assert meta["version"] == gis_export.FORMAT_VERSION and meta["tables"]["gis_packs"]["rows"] == 2
    data = json.loads((out / "gis_db_meta.json").read_text(encoding="utf8"))
    data["version"] = 99
    (out / "gis_db_meta.json").write_text(json.dumps(data), encoding="utf8")
    with pytest.raises(ValueError):
        gis_export.read_meta(str(out))


def test_date_below_first_interval_is_moved_to_its_block(tmp_path):
    """Дата стоит строкой ниже первого интервала, формула листа считает от строки выше (скв. 14, строка 24–25)."""
    book = make_book(tmp_path / "g.xlsx", [
        ("Сент.1984г", None, 770, 771, 0.3, 0.6, None),
        (None, None, 771, 772, 0.3, 0.6, None),
        (None, None, 770, 771.2, 0.3, 0.6, None),          # строка 7: принадлежит следующему замеру
        ("Окт.1984г", None, 771.2, 772, 0.3, 0.6, None),
    ])
    wb = openpyxl.load_workbook(book)
    wb["7"]["K8"] = "=SUM(J7:J8)"
    wb.save(book)
    res = gis.load_book(str(book))["7"]
    assert [len(b.rows) for b in res.blocks] == [2, 2]
    assert any(i.code == "date_shifted" for i in res.issues)
