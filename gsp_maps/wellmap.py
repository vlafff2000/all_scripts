"""Положение скважин: карта-сетка из Excel (номер скважины в ячейке) и координаты X, Y из tNavigator.

Координаты на выходе — «картографические»: x растёт на восток, y — на север.
Режимы: «grid» — только сетка, «xy» — только XY, «auto» — XY там, где они есть; скважины, которых нет в XY,
ставятся по сетке через преобразование «ячейки → XY», подобранное по общим скважинам.
"""
from __future__ import annotations

import math
import os
import re
import threading
from typing import Dict, List, Optional, Tuple

import numpy as np

DIRECTIONS8 = ["Север", "Северо-Восток", "Восток", "Юго-Восток", "Юг", "Юго-Запад", "Запад", "Северо-Запад"]


# ---------------- сетка из Excel ----------------
def list_sheets(path: str) -> List[str]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _sheet_key(name: str) -> str:
    """«ГСП-7», «ГСП_7» и «ГСП 7» — одно имя."""
    return re.sub(r"[\s_\-]+", "", str(name)).lower()


_GRIDS: Dict[str, Tuple[Tuple[int, int], Dict[str, tuple]]] = {}
_GRIDS_VERSION = 1
_GRIDS_LOCK = threading.Lock()   # группы запрашиваются параллельно: книгу карты читает один поток, остальные ждут


def _parse_sheet(ws) -> tuple:
    cells: Dict[int, Tuple[int, int]] = {}
    # только существующие ячейки в порядке «строка, столбец» (iter_rows по всему листу создаёт каждую пустую ячейку)
    for (r, c), cell in sorted(ws._cells.items()):
        if cell.value is not None:
            try:
                cells[int(float(str(cell.value).strip()))] = (r, c)
            except ValueError:
                pass
    return cells, (ws.max_row, ws.max_column), ws.title


def read_grids(path: str) -> Dict[str, tuple]:
    """Сетки всех листов книги за один проход: {лист: (скважина → (строка, столбец), размеры, имя)}. Книга читается один раз
    и запоминается (в памяти и в кэше на диске), а не заново для каждой группы."""
    st = os.stat(path)
    sig = (st.st_size, st.st_mtime_ns)
    hit = _GRIDS.get(path)
    if hit and hit[0] == sig:
        return hit[1]
    with _GRIDS_LOCK:
        hit = _GRIDS.get(path)
        if hit and hit[0] == sig:
            return hit[1]
        return _read_grids(path, sig)


def _read_grids(path: str, sig: Tuple[int, int]) -> Dict[str, tuple]:
    from openpyxl import load_workbook
    cache = None
    try:
        import hashlib
        import json
        from .store import cache_dir
        key = "%s|%d|%d|%d" % (os.path.abspath(path), sig[0], sig[1], _GRIDS_VERSION)
        cache = cache_dir() / ("grid_" + hashlib.sha1(key.encode("utf-8")).hexdigest()[:20] + ".json")
        raw = json.loads(cache.read_text(encoding="utf-8"))
        grids = {n: ({int(w): (r, c) for w, r, c in v["cells"]}, (v["dims"][0], v["dims"][1]), v["title"]) for n, v in raw.items()}
    except Exception:
        grids = None
    if grids is None:
        wb = load_workbook(path, data_only=True)
        grids = {name: _parse_sheet(wb[name]) for name in wb.sheetnames}
        if cache is not None:
            try:
                cache.parent.mkdir(parents=True, exist_ok=True)
                tmp = cache.with_name(cache.name + ".tmp")
                tmp.write_text(json.dumps({n: {"cells": [[w, r, c] for w, (r, c) in g[0].items()], "dims": list(g[1]), "title": g[2]}
                                           for n, g in grids.items()}), encoding="utf-8")
                os.replace(str(tmp), str(cache))
            except Exception:
                pass  # кэш не обязателен
    _GRIDS[path] = (sig, grids)
    return grids


def read_grid(path: str, gsp: str) -> Tuple[Dict[int, Tuple[int, int]], Tuple[int, int], str]:
    """{скважина: (строка, столбец)}, размеры листа (строк, столбцов), имя листа. Лист — «ГСП 9» или содержащий это имя
    (дефис, подчёркивание и пробел не различаются). Книга читается один раз на все группы."""
    grids = read_grids(path)
    key = _sheet_key(gsp)
    name = next((n for n in grids if _sheet_key(n) == key), None) or next((n for n in grids if key in _sheet_key(n)), None)
    if name is None:
        raise KeyError("На листах карты нет «%s» (есть: %s)" % (gsp, ", ".join(grids)))
    cells, dims, title = grids[name]
    return dict(cells), dims, title


def grid_directions(cells: Dict[int, Tuple[int, int]], dims: Tuple[int, int]) -> Dict[int, str]:
    """Направление по четвертям листа — точно как в старом скрипте."""
    center_row, center_col = dims[0] / 2, dims[1] / 2
    out = {}
    for well, (row, col) in cells.items():
        if row < center_row and col < center_col:
            d = "Северо-Запад"
        elif row < center_row and col > center_col:
            d = "Северо-Восток"
        elif row > center_row and col < center_col:
            d = "Юго-Запад"
        elif row > center_row and col > center_col:
            d = "Юго-Восток"
        elif row < center_row:
            d = "Север"
        elif row > center_row:
            d = "Юг"
        elif col < center_col:
            d = "Запад"
        else:
            d = "Восток"
        out[well] = d
    return out


def grid_xy(cells: Dict[int, Tuple[int, int]]) -> Dict[int, Tuple[float, float]]:
    """Как у старого скрипта: x — столбец, y — перевёрнутая строка (север сверху)."""
    if not cells:
        return {}
    max_row = max(r for r, _ in cells.values())
    return {w: (float(c), float(max_row - r + 1)) for w, (r, c) in cells.items()}


# ---------------- XY из tNavigator ----------------
_NUM = re.compile(r"^[+-]?(?:\d+(?:[.,]\d*)?|[.,]\d+)(?:[eE][+-]?\d+)?$")


def _float(tok: str) -> Optional[float]:
    tok = tok.strip().strip("'\"")
    if _NUM.match(tok):
        return float(tok.replace(",", "."))
    return None


def _well_id(name: str) -> Optional[int]:
    name = name.strip().strip("'\"")
    m = re.match(r"^\D*(\d+)\D*$", name) or re.search(r"(\d+)", name)
    return int(m.group(1)) if m else None


def _lines(path: str) -> List[str]:
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm", ".xls", ".ods"):
        import pandas as pd
        df = pd.read_excel(path, header=None)
        return ["\t".join("" if v != v else str(v) for v in row) for row in df.itertuples(index=False)]
    for enc in ("utf-8-sig", "cp1251"):
        try:
            with open(path, "r", encoding=enc) as f:
                return f.read().splitlines()
        except UnicodeDecodeError:
            continue
    return []


def parse_xy_lines(lines: List[str]) -> Tuple[Dict[int, Tuple[float, float]], List[str]]:
    """Таблица «скважина X Y …» (табуляция, пробелы, «;» или запятая; заголовок и пустые строки пропускаются)
    либо траектории WELLTRACK 'имя' (берётся первая точка — устье). Возвращает (координаты, предупреждения)."""
    xy: Dict[int, Tuple[float, float]] = {}
    warn: List[str] = []
    text = "\n".join(lines)
    if re.search(r"\bWELLTRACK\b", text, re.I):
        name = None
        for line in lines:
            m = re.match(r"\s*WELLTRACK\s+(\S+)", line, re.I)
            if m:
                name = m.group(1)
                continue
            if name is not None:
                nums = [_float(t) for t in line.replace("/", " ").split()]
                if len(nums) >= 2 and nums[0] is not None and nums[1] is not None:
                    wid = _well_id(name)
                    if wid is not None:
                        xy[wid] = (nums[0], nums[1])
                    name = None
        if xy:
            return xy, warn
    skipped = 0
    for raw in lines:
        line = raw.strip()
        if not line or line.startswith(("#", "--", "//")):
            continue
        if "\t" in line:
            toks = line.split("\t")
        elif ";" in line:
            toks = line.split(";")
        elif line.count(",") >= 2 and not re.search(r"\d,\d+\s", line):
            toks = line.split(",")
        else:
            toks = line.split()
        toks = [t.strip() for t in toks if t.strip() != ""]
        if len(toks) < 3:
            skipped += 1 if toks and _float(toks[0]) is not None else 0
            continue
        vals = [_float(t) for t in toks]
        if all(v is None for v in vals):  # заголовок таблицы
            continue
        if vals[0] is None:  # первый столбец — имя скважины (W500, 500-1, 'Скв_500')
            wid = _well_id(toks[0])
            nums = [v for v in vals[1:] if v is not None]
        else:  # первый столбец — число: номер скважины, дальше X и Y
            wid = int(vals[0]) if float(vals[0]).is_integer() else None
            nums = [v for v in vals[1:] if v is not None]
        if wid is None or len(nums) < 2:
            skipped += 1
            continue
        xy[wid] = (nums[0], nums[1])
    if skipped:
        warn.append("XY: не разобрано строк — %d" % skipped)
    if not xy:
        warn.append("XY: в файле не нашлось строк «скважина X Y»")
    return xy, warn


def read_xy(path: str) -> Tuple[Dict[int, Tuple[float, float]], List[str]]:
    return parse_xy_lines(_lines(path))


# ---------------- совмещение ----------------
def fit_affine(src: np.ndarray, dst: np.ndarray) -> Optional[np.ndarray]:
    """Аффинное преобразование (3×2) по общим точкам; нужно ≥ 3 неколлинеарных."""
    if len(src) < 3:
        return None
    a = np.c_[src, np.ones(len(src))]
    if np.linalg.matrix_rank(a) < 3:
        return None
    coef, *_ = np.linalg.lstsq(a, dst, rcond=None)
    return coef


def center_of(points: Dict[int, Tuple[float, float]]) -> Tuple[float, float]:
    xs = [p[0] for p in points.values()]
    ys = [p[1] for p in points.values()]
    return (min(xs) + max(xs)) / 2, (min(ys) + max(ys)) / 2


def angle_direction(dx: float, dy: float) -> str:
    """Направление по азимуту (8 румбов); dy вверх — север."""
    az = (math.degrees(math.atan2(dx, dy)) + 360) % 360
    return DIRECTIONS8[int((az + 22.5) // 45) % 8]


def layout(wells: List[int], cells: Dict[int, Tuple[int, int]], dims: Tuple[int, int],
           xy: Dict[int, Tuple[float, float]], mode: str = "auto") -> dict:
    """Положения скважин ГСП. Результат: {"mode", "wells": {номер: {x, y, src, dir}}, "missing": [...], "notes": [...]}"""
    want = set(wells)
    notes: List[str] = []
    gxy = grid_xy(cells)
    use_xy = mode in ("auto", "xy") and any(w in xy for w in want)
    if mode == "xy" and not use_xy:
        notes.append("В файле XY нет скважин этого ГСП — показана сетка.")
    out: Dict[int, dict] = {}
    if not use_xy:
        gd = grid_directions(cells, dims)
        for w in wells:
            if w in gxy:
                out[w] = {"x": gxy[w][0], "y": gxy[w][1], "src": "grid", "dir": gd.get(w, "")}
        eff = "grid"
    else:
        pos = {w: xy[w] for w in wells if w in xy}
        src = {w: "xy" for w in pos}
        if mode == "auto":
            common = [w for w in pos if w in gxy]
            coef = fit_affine(np.array([gxy[w] for w in common]), np.array([pos[w] for w in common])) if common else None
            lacking = [w for w in wells if w not in pos and w in gxy]
            if lacking and coef is not None:
                for w in lacking:
                    p = np.r_[gxy[w], 1.0] @ coef
                    pos[w], src[w] = (float(p[0]), float(p[1])), "fit"
                notes.append("Скважин без XY, поставленных по сетке через преобразование: %d" % len(lacking))
            elif lacking:
                notes.append("Скважин без XY: %d (по сетке не поставить — мало общих скважин)" % len(lacking))
        cx, cy = center_of(pos)
        for w, (x, y) in pos.items():
            out[w] = {"x": x, "y": y, "src": src[w], "dir": angle_direction(x - cx, y - cy)}
        eff = "xy"
    missing = [w for w in wells if w not in out]
    return {"mode": eff, "wells": out, "missing": missing, "notes": notes}
