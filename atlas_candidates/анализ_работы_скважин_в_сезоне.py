import pandas as pd
import numpy as np
from datetime import datetime
import tkinter as tk
from tkinter import filedialog, messagebox
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.chart import BarChart, LineChart, Reference
import warnings

warnings.filterwarnings('ignore')


def select_file():
    """Диалог выбора файла"""
    root = tk.Tk()
    root.withdraw()
    file_path = filedialog.askopenfilename(
        title="Выберите файл с базой данных",
        filetypes=[("Excel files", "*.xlsx *.xls"), ("All files", "*.*")]
    )
    root.destroy()
    return file_path


def get_available_seasons(df):
    """Получение списка доступных сезонов из данных"""
    production_seasons = []
    injection_seasons = []

    for season in df['Сезон'].unique():
        if pd.isna(season):
            continue
        season_str = str(season)
        season_data = df[df['Сезон'] == season_str]
        season_type = season_data['Тип данных'].iloc[0] if len(season_data) > 0 else None

        if season_type == 'отбор':
            production_seasons.append(season_str)
        elif season_type == 'закачка':
            injection_seasons.append(season_str)

    production_seasons = sorted(production_seasons, reverse=True)
    injection_seasons = sorted(injection_seasons, reverse=True)

    return production_seasons, injection_seasons


def create_season_selection_window(production_seasons, injection_seasons):
    """Создание окна выбора сезона"""
    root = tk.Tk()
    root.title("Выбор сезона для анализа")
    root.geometry("650x450")
    root.resizable(False, False)

    # Переменные для хранения результата
    result = {"type": None, "season": None}

    # Заголовок
    title_label = tk.Label(root, text="ВЫБОР СЕЗОНА ДЛЯ АНАЛИЗА",
                           font=("Arial", 14, "bold"))
    title_label.pack(pady=20)

    # Фрейм для выбора типа сезона
    type_frame = tk.LabelFrame(root, text="Тип сезона", padx=15, pady=15)
    type_frame.pack(pady=10, padx=20, fill="x")

    season_type_var = tk.StringVar(value="отбор")

    def on_type_selected():
        """Функция, вызываемая при выборе типа сезона"""
        selected_type = season_type_var.get()
        update_season_list(selected_type)

    prod_radio = tk.Radiobutton(type_frame, text="ОТБОР газа",
                                variable=season_type_var, value="отбор",
                                font=("Arial", 12, "bold"),
                                command=on_type_selected)
    prod_radio.pack(side="left", padx=30, pady=5)

    inj_radio = tk.Radiobutton(type_frame, text="ЗАКАЧКА газа",
                               variable=season_type_var, value="закачка",
                               font=("Arial", 12, "bold"),
                               command=on_type_selected)
    inj_radio.pack(side="left", padx=30, pady=5)

    # Фрейм для выбора сезона
    season_frame = tk.LabelFrame(root, text="Доступные сезоны", padx=15, pady=15)
    season_frame.pack(pady=10, padx=20, fill="both", expand=True)

    # Информационная метка
    info_label = tk.Label(season_frame, text="", font=("Arial", 11), fg="blue")
    info_label.pack(pady=5)

    # Список сезонов с прокруткой
    list_frame = tk.Frame(season_frame)
    list_frame.pack(fill="both", expand=True)

    scrollbar = tk.Scrollbar(list_frame)
    scrollbar.pack(side="right", fill="y")

    listbox = tk.Listbox(list_frame, yscrollcommand=scrollbar.set,
                         font=("Arial", 11), height=8,
                         selectmode="single",
                         activestyle="dotbox")
    listbox.pack(side="left", fill="both", expand=True)
    scrollbar.config(command=listbox.yview)

    # Словарь для соответствия индекса и значения
    idx_to_value = {}

    def update_season_list(season_type):
        """Обновление списка сезонов"""
        listbox.delete(0, tk.END)
        idx_to_value.clear()

        if season_type == "отбор":
            seasons = production_seasons
            info_label.config(text=f"📊 СЕЗОНЫ ОТБОРА (доступно: {len(seasons)})")
        else:
            seasons = injection_seasons
            info_label.config(text=f"📊 СЕЗОНЫ ЗАКАЧКИ (доступно: {len(seasons)})")

        for i, season in enumerate(seasons):
            display_text = f"  {season}  "
            listbox.insert(tk.END, display_text)
            idx_to_value[i] = season

        if seasons:
            listbox.selection_set(0)
            listbox.activate(0)
            listbox.focus_set()

    def on_analyze():
        """Обработчик кнопки Анализировать"""
        selection = listbox.curselection()
        if not selection:
            messagebox.showwarning("Предупреждение", "⚠️  Выберите сезон из списка!")
            return

        selected_idx = selection[0]
        result["type"] = season_type_var.get()
        result["season"] = idx_to_value[selected_idx]
        root.destroy()

    def on_cancel():
        """Обработчик кнопки Отмена"""
        root.destroy()

    # Кнопки
    button_frame = tk.Frame(root)
    button_frame.pack(pady=20)

    analyze_btn = tk.Button(button_frame, text="▶  АНАЛИЗИРОВАТЬ",
                            command=on_analyze,
                            width=20, height=2,
                            bg="#366092", fg="white",
                            font=("Arial", 11, "bold"))
    analyze_btn.pack(side="left", padx=10)

    cancel_btn = tk.Button(button_frame, text="❌  Отмена",
                           command=on_cancel,
                           width=15, height=2,
                           font=("Arial", 11))
    cancel_btn.pack(side="left", padx=10)

    # Двойной клик для быстрого выбора
    listbox.bind('<Double-Button-1>', lambda e: on_analyze())

    # Инициализация списка
    update_season_list("отбор")

    # Центрирование окна
    root.update_idletasks()
    w = root.winfo_width()
    h = root.winfo_height()
    x = (root.winfo_screenwidth() // 2) - (w // 2)
    y = (root.winfo_screenheight() // 2) - (h // 2)
    root.geometry(f'{w}x{h}+{x}+{y}')

    root.mainloop()

    return result["type"], result["season"]


def get_well_weight(well_id):
    """Определяет вес скважины"""
    well_str = str(well_id)
    if well_str == '54/80':
        return 2
    return 1


def count_unique_wells_with_weights(well_list):
    """Подсчитывает количество уникальных скважин с учетом весов"""
    well_weights = {}
    for well in well_list:
        well_str = str(well)
        if well_str not in well_weights:
            well_weights[well_str] = get_well_weight(well_str)
    return sum(well_weights.values())


def load_data(file_path):
    """Загрузка и подготовка данных"""
    print("\n📂 Загрузка данных...")

    df_production = pd.read_excel(file_path, sheet_name='Отборы')
    df_injection = pd.read_excel(file_path, sheet_name='Закачка')

    print(f"✓ Загружено записей по отбору: {len(df_production)}")
    print(f"✓ Загружено записей по закачке: {len(df_injection)}")

    # Проверяем наличие столбца "Сезон" в данных закачки
    if 'Сезон' not in df_injection.columns and 'Год' in df_injection.columns:
        print("⚠️  В данных закачки отсутствует столбец 'Сезон', использую столбец 'Год'")
        df_injection['Сезон'] = df_injection['Год'].astype(str)

    # Объединяем данные
    df = pd.concat([df_production, df_injection], ignore_index=True)

    # Фильтруем нейтральный период
    df = df[df['Тип данных'] != 'нейтральный период'].copy()

    # Преобразуем дату
    df['Дата'] = pd.to_datetime(df['Дата'], dayfirst=True)

    return df


def filter_target_season(df, season_type, target_season):
    """Фильтрация данных по выбранному сезону"""
    df_target = df[(df['Сезон'] == target_season) & (df['Тип данных'] == season_type)].copy()

    # Получаем предыдущие 5 сезонов того же типа
    previous_seasons = []
    previous_seasons_data = {}

    all_seasons = sorted(df[df['Тип данных'] == season_type]['Сезон'].unique(), reverse=True)

    try:
        current_idx = all_seasons.index(target_season)
    except ValueError:
        current_idx = -1

    for i in range(current_idx + 1, min(current_idx + 6, len(all_seasons))):
        prev_season = all_seasons[i]
        df_prev = df[(df['Сезон'] == prev_season) & (df['Тип данных'] == season_type)]
        if len(df_prev) > 0:
            previous_seasons.append(prev_season)
            previous_seasons_data[prev_season] = df_prev.copy()

    return df_target, previous_seasons, previous_seasons_data


def get_gsp_info():
    """Получение информации о количестве скважин по ГСП"""
    gsp_info = {
        'ГСП 1': 40, 'ГСП 2': 50, 'ГСП 3': 27, 'ГСП 4': 48, 'ГСП 5': 30,
        'ГСП 6': 28, 'ГСП 7': 40, 'ГСП 8': 24, 'ГСП 9': 44
    }
    return gsp_info


def check_units_and_adjust_intervals(df_target):
    """Проверка единиц измерения"""
    print("\n=== ПРОВЕРКА ЕДИНИЦ ИЗМЕРЕНИЯ ===")

    non_zero = df_target[df_target['Суточный расход газа'] > 0]

    if len(non_zero) == 0:
        print("Нет данных с ненулевым расходом!")
        return None, None

    max_val = non_zero['Суточный расход газа'].max()

    if max_val > 1000000:
        print("⚠️  Данные в м3/сут, пересчет в тыс. м3/сут")
        scale_factor = 1000
    else:
        scale_factor = 1

    # Интервалы
    if scale_factor == 1000:
        intervals = [
            ('0', 0, 0, 'нулевой расход'),
            ('Менее 100', 1, 99999, 'менее 100 тыс.'),
            ('От 100 до 150', 100000, 149999, '100-150 тыс.'),
            ('От 150 до 200', 150000, 199999, '150-200 тыс.'),
            ('От 200 до 250', 200000, 249999, '200-250 тыс.'),
            ('От 250 до 300', 250000, 299999, '250-300 тыс.'),
            ('Более 300', 300000, float('inf'), 'более 300 тыс.')
        ]
    else:
        intervals = [
            ('0', 0, 0, 'нулевой расход'),
            ('Менее 100', 0.01, 99.99, 'менее 100'),
            ('От 100 до 150', 100, 149.99, '100-150'),
            ('От 150 до 200', 150, 199.99, '150-200'),
            ('От 200 до 250', 200, 249.99, '200-250'),
            ('От 250 до 300', 250, 299.99, '250-300'),
            ('Более 300', 300, float('inf'), 'более 300')
        ]

    return intervals, scale_factor


def get_season_months_from_data(df_target):
    """Определение месяцев сезона на основе фактических данных"""
    months_data = df_target[['Дата', 'Месяц']].drop_duplicates()
    months_data = months_data.sort_values('Дата')

    unique_months = months_data['Месяц'].unique()

    month_abbr = {
        'Январь': 'Янв', 'Февраль': 'Фев', 'Март': 'Мар',
        'Апрель': 'Апр', 'Май': 'Май', 'Июнь': 'Июн',
        'Июль': 'Июл', 'Август': 'Авг', 'Сентябрь': 'Сен',
        'Октябрь': 'Окт', 'Ноябрь': 'Ноя', 'Декабрь': 'Дек'
    }

    months_order = {}
    for month in unique_months:
        month_data = months_data[months_data['Месяц'] == month]
        sample_date = month_data['Дата'].iloc[0]
        year_short = str(sample_date.year)[-2:]
        month_short = month_abbr.get(month, month[:3])
        months_order[month] = f"{month_short}.{year_short}"

    print(f"✓ Месяцы сезона: {', '.join(months_order.values())}")

    return months_order


def create_table_1(df_target, gsp_info):
    """Создание таблицы активности по месяцам"""
    gsp_list = sorted(df_target['Источник'].unique())
    months_order = get_season_months_from_data(df_target)

    result_data = []

    for gsp in gsp_list:
        gsp_data = df_target[df_target['Источник'] == gsp]
        total_wells = gsp_info.get(gsp, 0)

        row = {'№ ГСП': gsp, 'Всего скв.': total_wells}

        for month_ru, month_short in months_order.items():
            month_data = gsp_data[gsp_data['Месяц'] == month_ru]
            if len(month_data) > 0:
                active_wells_list = month_data[month_data['Суточный расход газа'] > 0]['Скважина'].unique()
                active_wells = count_unique_wells_with_weights(active_wells_list)
                row[month_short] = active_wells
            else:
                row[month_short] = 0

        result_data.append(row)

    df_table1 = pd.DataFrame(result_data)
    month_cols = [v for k, v in months_order.items() if v in df_table1.columns]
    df_table1 = df_table1[['№ ГСП', 'Всего скв.'] + month_cols]

    return df_table1, months_order


def create_table_2(df_target, gsp_info, intervals):
    """Создание таблицы производительности"""
    gsp_list = sorted(df_target['Источник'].unique())

    result_data = []
    total_counts = {interval[0]: 0 for interval in intervals}

    for gsp in gsp_list:
        gsp_data = df_target[df_target['Источник'] == gsp]
        wells = gsp_data['Скважина'].unique()

        total_wells_in_gsp = gsp_info.get(gsp, 0)

        row = {'Номер ГСП': gsp, 'Количество скважин': total_wells_in_gsp}

        interval_counts = {interval[0]: 0 for interval in intervals}
        zero_wells_list = []

        for well in wells:
            well_data = gsp_data[gsp_data['Скважина'] == well]
            non_zero_days = well_data[well_data['Суточный расход газа'] > 0]
            well_weight = get_well_weight(well)

            if len(non_zero_days) == 0:
                interval_counts['0'] += well_weight
                suffix = f" (x{well_weight})" if well_weight > 1 else ""
                zero_wells_list.append(f"{well}{suffix}")
            else:
                avg_rate = non_zero_days['Суточный расход газа'].mean()

                interval_found = False
                for interval_name, min_val, max_val, _ in intervals:
                    if interval_name != '0' and min_val <= avg_rate <= max_val:
                        interval_counts[interval_name] += well_weight
                        interval_found = True
                        break

                if not interval_found and avg_rate >= intervals[-1][1]:
                    interval_counts[intervals[-1][0]] += well_weight

        for interval_name, _, _, _ in intervals:
            if interval_name == '0':
                if interval_counts['0'] > 0:
                    row[interval_name] = f"{interval_counts['0']} (№ {', '.join(zero_wells_list)})"
                else:
                    row[interval_name] = "0"
            else:
                row[interval_name] = interval_counts[interval_name]
            total_counts[interval_name] += interval_counts[interval_name]

        result_data.append(row)

    # ИТОГО
    total_wells_all = sum(gsp_info.values())
    total_row = {'Номер ГСП': 'ИТОГО', 'Количество скважин': total_wells_all}
    for interval_name, _, _, _ in intervals:
        if interval_name == '0':
            total_row[interval_name] = str(total_counts[interval_name])
        else:
            total_row[interval_name] = total_counts[interval_name]

    result_data.append(total_row)

    df_table2 = pd.DataFrame(result_data)
    column_order = ['Номер ГСП', 'Количество скважин'] + [interval[0] for interval in intervals]
    df_table2 = df_table2[column_order]

    return df_table2


def analyze_well_performance(df_target, scale_factor, season_type):
    """Анализ производительности скважин"""
    wells_analysis = []

    for well in df_target['Скважина'].unique():
        well_data = df_target[df_target['Скважина'] == well]
        well_weight = get_well_weight(well)

        total_gas = well_data['Суточный расход газа'].sum()
        working_hours = well_data['Время работы'].sum()

        non_zero_days = well_data[well_data['Суточный расход газа'] > 0]
        working_days = len(non_zero_days)
        total_days_in_season = len(well_data)

        max_daily_rate = well_data['Суточный расход газа'].max() / scale_factor

        zero_gas_with_hours = well_data[(well_data['Суточный расход газа'] == 0) & (well_data['Время работы'] > 0)]
        days_zero_gas_with_hours = len(zero_gas_with_hours)
        hours_zero_gas_with_hours = zero_gas_with_hours['Время работы'].sum()

        avg_rate = non_zero_days['Суточный расход газа'].mean() / scale_factor if working_days > 0 else 0
        total_gas_adjusted = total_gas / scale_factor

        working_ratio = working_days / total_days_in_season if total_days_in_season > 0 else 0

        action_word = "отбора" if season_type == "отбор" else "закачки"

        if working_days == 0:
            category = "Не работала"
        elif working_days < 5:
            category = f"Подозрительно (менее 5 дней {action_word})"
        elif days_zero_gas_with_hours > 10:
            category = "Подозрительно (много дней с часами, но нулевым расходом)"
        elif working_days > 150 and avg_rate > 200:
            category = "Отлично работающая"
        elif working_days > 150 and avg_rate < 100:
            category = f"Подозрительно (много дней {action_word}, низкий расход)"
        elif working_ratio > 0.8 and avg_rate > 150:
            category = "Хорошо работающая"
        else:
            category = "Нормально работающая"

        well_display = f"{well} (x{well_weight})" if well_weight > 1 else str(well)

        wells_analysis.append({
            'Скважина': well_display,
            'Вес': well_weight,
            'ГСП': well_data['Источник'].iloc[0],
            'Накопленный объем, тыс. м3': round(total_gas_adjusted, 2),
            'Макс. суточный расход, тыс. м3/сут': round(max_daily_rate, 2),
            'Средний расход, тыс. м3/сут': round(avg_rate, 2),
            'Время работы, часы': working_hours,
            'Кол-во рабочих дней': working_days,
            'Всего дней в сезоне': total_days_in_season,
            'Доля рабочих дней, %': round(working_ratio * 100, 1),
            'Дней с часами работы, но нулевым расходом': days_zero_gas_with_hours,
            'Часов работы при нулевом расходе': hours_zero_gas_with_hours,
            'Категория': category
        })

    df_analysis = pd.DataFrame(wells_analysis)
    df_analysis = df_analysis.sort_values('Накопленный объем, тыс. м3', ascending=False)

    return df_analysis


def analyze_season_comparison(df_target, previous_seasons_data, scale_factor):
    """Сравнительный анализ по сезонам"""
    print("\n=== СРАВНИТЕЛЬНЫЙ АНАЛИЗ ПО СЕЗОНАМ ===")

    current_stats = calculate_season_stats(df_target, scale_factor, "Текущий")
    print(f"✓ Текущий сезон проанализирован")

    historical_stats = []
    for season_name, df_season in previous_seasons_data.items():
        stats = calculate_season_stats(df_season, scale_factor, season_name)
        historical_stats.append(stats)
        print(f"✓ Сезон {season_name} проанализирован")

    all_stats = [current_stats] + historical_stats

    # Сводная таблица по сезонам
    summary_data = []
    for stats in all_stats:
        summary_data.append({
            'Сезон': stats['season_name'],
            'Всего скважин': stats['total_wells'],
            'Активных скважин': stats['active_wells'],
            'Доля активных, %': round(stats['active_wells'] / stats['total_wells'] * 100, 1) if stats[
                                                                                                    'total_wells'] > 0 else 0,
            'Общий объем, тыс. м3': round(stats['total_gas'], 2),
            'Средний расход, тыс. м3/сут': round(stats['avg_rate'], 2),
            'Медианный расход, тыс. м3/сут': round(stats['median_rate'], 2),
            'Макс. суточный расход, тыс. м3/сут': round(stats['max_rate'], 2),
            'Общее время работы, часы': stats['total_hours'],
            'Среднее время работы на скв., часы': round(stats['avg_hours_per_well'], 1),
            'Среднее кол-во рабочих дней': round(stats['avg_working_days'], 1),
            'Дней с нулевым расходом при работе': stats['zero_gas_with_hours_days'],
            'Подозрительных скважин': stats['suspicious_wells'],
            'Отличных скважин': stats['excellent_wells'],
            'Не работало скважин': stats['inactive_wells']
        })

    df_season_comparison = pd.DataFrame(summary_data)

    # Детальное сравнение по скважинам
    well_comparison = compare_wells_across_seasons(df_target, previous_seasons_data, scale_factor)

    return df_season_comparison, well_comparison


def calculate_season_stats(df_season, scale_factor, season_name):
    """Расчет статистики для одного сезона"""
    total_wells_unique = df_season['Скважина'].unique()
    total_wells = count_unique_wells_with_weights(total_wells_unique)

    active_wells_list = df_season[df_season['Суточный расход газа'] > 0]['Скважина'].unique()
    active_wells = count_unique_wells_with_weights(active_wells_list)

    total_gas = df_season['Суточный расход газа'].sum() / scale_factor
    total_hours = df_season['Время работы'].sum()
    max_rate = df_season['Суточный расход газа'].max() / scale_factor

    non_zero = df_season[df_season['Суточный расход газа'] > 0]
    avg_rate = non_zero['Суточный расход газа'].mean() / scale_factor if len(non_zero) > 0 else 0
    median_rate = non_zero['Суточный расход газа'].median() / scale_factor if len(non_zero) > 0 else 0

    zero_gas_with_hours = df_season[(df_season['Суточный расход газа'] == 0) & (df_season['Время работы'] > 0)]
    zero_gas_with_hours_days = len(zero_gas_with_hours)

    suspicious_wells = 0
    excellent_wells = 0
    inactive_wells = 0
    total_working_days = 0
    working_wells_count = 0

    for well in total_wells_unique:
        well_data = df_season[df_season['Скважина'] == well]
        well_weight = get_well_weight(well)
        working_days = len(well_data[well_data['Суточный расход газа'] > 0])
        total_days = len(well_data)

        if working_days == 0:
            inactive_wells += well_weight
        else:
            working_wells_count += well_weight
            total_working_days += working_days * well_weight

            if working_days > 0:
                well_avg = well_data[well_data['Суточный расход газа'] > 0][
                               'Суточный расход газа'].mean() / scale_factor
                working_ratio = working_days / total_days if total_days > 0 else 0

                if working_days < 5 or (working_ratio < 0.3 and well_avg > 200):
                    suspicious_wells += well_weight
                elif working_days > 150 and well_avg > 200:
                    excellent_wells += well_weight

    avg_working_days = total_working_days / working_wells_count if working_wells_count > 0 else 0
    avg_hours_per_well = total_hours / active_wells if active_wells > 0 else 0

    return {
        'season_name': season_name,
        'total_wells': total_wells,
        'active_wells': active_wells,
        'total_gas': total_gas,
        'avg_rate': avg_rate,
        'median_rate': median_rate,
        'max_rate': max_rate,
        'total_hours': total_hours,
        'avg_hours_per_well': avg_hours_per_well,
        'avg_working_days': avg_working_days,
        'zero_gas_with_hours_days': zero_gas_with_hours_days,
        'suspicious_wells': suspicious_wells,
        'excellent_wells': excellent_wells,
        'inactive_wells': inactive_wells
    }


def compare_wells_across_seasons(df_target, previous_seasons_data, scale_factor):
    """Сравнение производительности скважин по сезонам"""
    all_wells = set(df_target['Скважина'].unique())
    for df_season in previous_seasons_data.values():
        all_wells.update(df_season['Скважина'].unique())

    comparison_data = []

    for well in sorted(all_wells):
        well_weight = get_well_weight(well)

        current_data = df_target[df_target['Скважина'] == well]
        current_stats = get_well_season_stats(current_data, scale_factor)

        row = {
            'Скважина': well,
            'Вес': well_weight,
            'ГСП': current_data['Источник'].iloc[0] if len(current_data) > 0 else 'Н/Д'
        }

        row['Текущий_объем'] = current_stats['total_gas']
        row['Текущий_расход'] = current_stats['avg_rate']
        row['Текущий_макс_расход'] = current_stats['max_rate']
        row['Текущий_дни'] = current_stats['working_days']
        row['Текущий_нулевой_расход_часы'] = current_stats['zero_gas_with_hours']

        for i, (season_name, df_season) in enumerate(previous_seasons_data.items(), 1):
            season_data = df_season[df_season['Скважина'] == well]
            stats = get_well_season_stats(season_data, scale_factor)

            row[f'Сезон{i}_объем'] = stats['total_gas']
            row[f'Сезон{i}_расход'] = stats['avg_rate']
            row[f'Сезон{i}_макс_расход'] = stats['max_rate']
            row[f'Сезон{i}_дни'] = stats['working_days']
            row[f'Сезон{i}_нулевой_расход_часы'] = stats['zero_gas_with_hours']

        # Тренды
        historical_rates = [row[f'Сезон{i}_расход'] for i in range(1, len(previous_seasons_data) + 1)
                            if f'Сезон{i}_расход' in row and row[f'Сезон{i}_расход'] > 0]
        historical_days = [row[f'Сезон{i}_дни'] for i in range(1, len(previous_seasons_data) + 1)
                           if f'Сезон{i}_дни' in row and row[f'Сезон{i}_дни'] > 0]

        if len(historical_rates) > 0 and current_stats['avg_rate'] > 0:
            avg_hist_rate = np.mean(historical_rates)
            row['Тренд_расход_%'] = round((current_stats['avg_rate'] / avg_hist_rate - 1) * 100, 1)
        else:
            row['Тренд_расход_%'] = 0

        if len(historical_days) > 0 and current_stats['working_days'] > 0:
            avg_hist_days = np.mean(historical_days)
            row['Тренд_дни_%'] = round((current_stats['working_days'] / avg_hist_days - 1) * 100, 1)
        else:
            row['Тренд_дни_%'] = 0

        # Аномалии
        if current_stats['working_days'] > 0 and len(historical_days) >= 3:
            if current_stats['working_days'] < avg_hist_days * 0.3:
                row['Аномалия'] = 'Резкое снижение дней работы'
            elif current_stats['avg_rate'] > 0 and len(historical_rates) >= 3:
                if current_stats['avg_rate'] > avg_hist_rate * 1.5:
                    row['Аномалия'] = 'Резкий рост расхода'
                elif current_stats['avg_rate'] < avg_hist_rate * 0.5:
                    row['Аномалия'] = 'Резкое падение расхода'
                else:
                    row['Аномалия'] = 'Норма'
            else:
                row['Аномалия'] = 'Норма'
        else:
            row['Аномалия'] = 'Недостаточно данных'

        if current_stats['zero_gas_with_hours'] > 5:
            if row['Аномалия'] == 'Норма':
                row['Аномалия'] = 'Много дней с часами, но нулевым расходом'
            else:
                row['Аномалия'] += ' + нулевой расход при работе'

        comparison_data.append(row)

    df_comparison = pd.DataFrame(comparison_data)
    df_comparison = df_comparison.sort_values('Аномалия')

    return df_comparison


def get_well_season_stats(well_data, scale_factor):
    """Статистика по скважине за сезон"""
    if len(well_data) == 0:
        return {'total_gas': 0, 'avg_rate': 0, 'max_rate': 0, 'working_days': 0, 'zero_gas_with_hours': 0}

    total_gas = well_data['Суточный расход газа'].sum() / scale_factor
    max_rate = well_data['Суточный расход газа'].max() / scale_factor
    working_days = len(well_data[well_data['Суточный расход газа'] > 0])
    zero_gas_with_hours = len(well_data[(well_data['Суточный расход газа'] == 0) & (well_data['Время работы'] > 0)])

    avg_rate = well_data[well_data['Суточный расход газа'] > 0][
                   'Суточный расход газа'].mean() / scale_factor if working_days > 0 else 0

    return {
        'total_gas': round(total_gas, 2),
        'avg_rate': round(avg_rate, 2),
        'max_rate': round(max_rate, 2),
        'working_days': working_days,
        'zero_gas_with_hours': zero_gas_with_hours
    }


def format_excel_sheet(writer, sheet_name, df):
    """Форматирование листа Excel"""
    workbook = writer.book
    worksheet = writer.sheets[sheet_name]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    for col in range(1, len(df.columns) + 1):
        cell = worksheet.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for column in worksheet.columns:
        max_length = 0
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        worksheet.column_dimensions[column[0].column_letter].width = min(max_length + 2, 45)


def apply_color_formatting(worksheet, df, sheet_name):
    """Цветовое форматирование"""
    if 'Категория' in df.columns:
        category_col = df.columns.get_loc('Категория') + 1
        for row in range(2, len(df) + 2):
            cell = worksheet.cell(row=row, column=category_col)
            val = str(cell.value)
            if 'Подозрительно' in val:
                cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                cell.font = Font(color="9C0006", bold=True)
            elif 'Отлично' in val:
                cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                cell.font = Font(color="006100", bold=True)
            elif 'Хорошо' in val:
                cell.fill = PatternFill(start_color="D4F4D4", end_color="D4F4D4", fill_type="solid")
            elif 'Не работала' in val:
                cell.fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")

    if sheet_name == 'Сравнение скважин' and 'Аномалия' in df.columns:
        anomaly_col = df.columns.get_loc('Аномалия') + 1
        for row in range(2, len(df) + 2):
            cell = worksheet.cell(row=row, column=anomaly_col)
            val = str(cell.value)
            if 'Резкое' in val or 'Много' in val:
                cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                cell.font = Font(color="9C0006", bold=True)


def create_chart_all_months(worksheet, df, title, position):
    """Создание гистограммы"""
    if 'ИТОГО' in df.iloc[:, 0].values:
        df = df[df.iloc[:, 0] != 'ИТОГО']

    numeric_cols = []
    for col in df.columns:
        if col not in ['№ ГСП', 'Всего скв.', 'Номер ГСП', 'Количество скважин']:
            try:
                pd.to_numeric(df[col])
                numeric_cols.append(col)
            except:
                pass

    if len(numeric_cols) == 0:
        return

    chart = BarChart()
    chart.type = "col"
    chart.title = title
    chart.y_axis.title = 'Количество скважин'
    chart.x_axis.title = 'ГСП'

    first_data_col = df.columns.get_loc(numeric_cols[0]) + 1
    last_data_col = df.columns.get_loc(numeric_cols[-1]) + 1

    data = Reference(worksheet, min_col=first_data_col, min_row=1, max_row=len(df) + 1, max_col=last_data_col)
    cats = Reference(worksheet, min_col=1, min_row=2, max_row=len(df) + 1)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.width = 22
    chart.height = 12

    worksheet.add_chart(chart, position)


def create_trend_chart(worksheet, df_seasons, title, position):
    """Создание графика тренда по сезонам"""
    if len(df_seasons) <= 1:
        return

    chart = LineChart()
    chart.title = title
    chart.y_axis.title = 'Значение'
    chart.x_axis.title = 'Сезон'
    chart.width = 20
    chart.height = 10

    data = Reference(worksheet, min_col=2, min_row=1, max_row=len(df_seasons) + 1, max_col=6)
    cats = Reference(worksheet, min_col=1, min_row=2, max_row=len(df_seasons) + 1)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)

    worksheet.add_chart(chart, position)


def main():
    print("=" * 70)
    print("           АНАЛИТИКА РАБОТЫ СКВАЖИН ПХГ")
    print("=" * 70)
    print("\n⚠️  Скважина '54/80' учитывается как ДВЕ скважины")

    file_path = select_file()
    if not file_path:
        print("❌ Файл не выбран.")
        return

    df = load_data(file_path)

    production_seasons, injection_seasons = get_available_seasons(df)

    print(f"\n📊 Доступно сезонов отбора: {len(production_seasons)}")
    if production_seasons:
        print(f"   {', '.join(production_seasons[:5])}{'...' if len(production_seasons) > 5 else ''}")
    print(f"📊 Доступно сезонов закачки: {len(injection_seasons)}")
    if injection_seasons:
        print(f"   {', '.join(injection_seasons[:5])}{'...' if len(injection_seasons) > 5 else ''}")

    if len(production_seasons) == 0 and len(injection_seasons) == 0:
        print("❌ Нет доступных сезонов!")
        return

    print("\n🔍 Открываю окно выбора сезона...")

    season_type, target_season = create_season_selection_window(production_seasons, injection_seasons)

    if season_type is None or target_season is None:
        print("❌ Сезон не выбран.")
        return

    print(f"\n✅ Выбран тип: {season_type.upper()}")
    print(f"✅ Выбран сезон: {target_season}")

    df_target, previous_seasons, previous_seasons_data = filter_target_season(df, season_type, target_season)

    if len(df_target) == 0:
        print(f"❌ Нет данных за сезон {target_season}")
        return

    print(f"✓ Записей: {len(df_target)}")
    print(f"✓ Даты: {df_target['Дата'].min().strftime('%d.%m.%Y')} - {df_target['Дата'].max().strftime('%d.%m.%Y')}")

    if previous_seasons:
        print(f"✓ Предыдущие сезоны для сравнения: {', '.join(previous_seasons)}")
    else:
        print("⚠️  Нет предыдущих сезонов для сравнения")

    intervals, scale_factor = check_units_and_adjust_intervals(df_target)
    if intervals is None:
        return

    gsp_info = get_gsp_info()

    print("\n" + "=" * 70)
    print("ФОРМИРОВАНИЕ ТАБЛИЦ")
    print("=" * 70)

    print("\n📊 Таблица 1: Активность по месяцам")
    df_table1, months_order = create_table_1(df_target, gsp_info)
    print(df_table1.to_string(index=False))

    print("\n📊 Таблица 2: Производительность")
    df_table2 = create_table_2(df_target, gsp_info, intervals)
    print(df_table2.to_string(index=False))

    print("\n📊 Анализ скважин")
    df_analysis = analyze_well_performance(df_target, scale_factor, season_type)
    print(f"Проанализировано: {len(df_analysis)} скважин")

    # Статистика по категориям
    category_counts = df_analysis['Категория'].value_counts()
    for cat, count in category_counts.items():
        print(f"  • {cat}: {count}")

    print("\n📊 Сравнение сезонов")
    df_season_comparison, df_well_comparison = analyze_season_comparison(
        df_target, previous_seasons_data, scale_factor
    )

    # Сохранение
    season_display = target_season.replace('/', '-')
    output_file = f"Аналитика_{season_type}_{season_display}.xlsx"

    print(f"\n💾 Сохранение в: {output_file}")

    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        df_table1.to_excel(writer, sheet_name='Активность по месяцам', index=False)
        df_table2.to_excel(writer, sheet_name='Производительность', index=False)
        df_analysis.to_excel(writer, sheet_name='Анализ скважин', index=False)
        df_season_comparison.to_excel(writer, sheet_name='Сравнение сезонов', index=False)
        df_well_comparison.to_excel(writer, sheet_name='Сравнение скважин', index=False)

        for name, dff in [('Активность по месяцам', df_table1),
                          ('Производительность', df_table2),
                          ('Анализ скважин', df_analysis),
                          ('Сравнение сезонов', df_season_comparison),
                          ('Сравнение скважин', df_well_comparison)]:
            format_excel_sheet(writer, name, dff)

        apply_color_formatting(writer.sheets['Анализ скважин'], df_analysis, 'Анализ скважин')
        apply_color_formatting(writer.sheets['Сравнение скважин'], df_well_comparison, 'Сравнение скважин')

        # Графики
        ws1 = writer.sheets['Активность по месяцам']
        create_chart_all_months(ws1, df_table1, f'Активность ({season_type} {target_season})', 'A15')

        ws2 = writer.sheets['Производительность']
        chart_data2 = df_table2[df_table2['Номер ГСП'] != 'ИТОГО'].copy()
        for col in chart_data2.columns[2:]:
            chart_data2[col] = chart_data2[col].apply(
                lambda x: int(str(x).split()[0]) if isinstance(x, str) and '(' in str(x) else x
            )
        create_chart_all_months(ws2, chart_data2, f'Производительность ({season_type})', 'A15')

        if len(df_season_comparison) > 1:
            ws3 = writer.sheets['Сравнение сезонов']
            create_trend_chart(ws3, df_season_comparison, 'Динамика показателей по сезонам', 'A18')

    print("\n" + "=" * 70)
    print("✅ АНАЛИТИКА ЗАВЕРШЕНА!")
    print(f"📁 Файл: {output_file}")
    print("📋 Листы:")
    print("   • Активность по месяцам")
    print("   • Производительность")
    print("   • Анализ скважин")
    print("   • Сравнение сезонов")
    print("   • Сравнение скважин")


if __name__ == "__main__":
    main()